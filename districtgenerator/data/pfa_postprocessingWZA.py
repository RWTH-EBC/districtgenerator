from pathlib import Path
from collections import defaultdict
import hashlib
import math
import re
import unicodedata

import numpy as np
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation


# ============================================================
# 1. KONFIGURATION
# ============================================================

BASE_DIR = Path("WZ2025_Vorbereitung")


# ------------------------------------------------------------
# Referenzdatensatz der anderen Forscher
# ------------------------------------------------------------

REFERENCE_FILE = (
        BASE_DIR
        / "Final_Dataset_070825.xlsx"
)

REFERENCE_SHEET = "Final_Dataset_070825"

# =============================================================================
# DATENQUELLE – WISSENSCHAFTLICHE VERGLEICHSKLASSIFIKATION
# =============================================================================
#
# Unveröffentlichter Datensatz:
# Angstmann, M. (o. J.): Unveröffentlichte Vergleichsdaten zur
# Klassifizierung industrieller Abwärmepotenziale der Plattform
# für Abwärme (PfA). Institut Arbeit und Technik (IAT),
# Westfälische Hochschule, Gelsenkirchen.
# Persönlich zur Verfügung gestellter Forschungsdatensatz,
# erhalten am 29.09.2026. Nicht öffentlich zugänglich.
#
# Zugehörige wissenschaftliche Publikation:
# Angstmann, M., Meyer, K. & Gärtner, S. (2025):
# Exploring the Potential of Urban Manufacturer's Waste Heat
# for the Residential Heating Transition in Germany:
# A Spatial Analysis Across Four Federal States.
# IAT Discussion Papers, No. 25/03.
# Institut Arbeit und Technik (IAT), Westfälische Hochschule,
# Gelsenkirchen.
# DOI: https://doi.org/10.53190/dp/202503
#
# Verwendung:
# Die persönlich bereitgestellten Vergleichsdaten dienen
# der wissenschaftlichen Gegenüberstellung und Validierung
# der eigenen WZ-Klassifizierung der PFA-Daten.


# ------------------------------------------------------------
# Alle 68 klassifizierten WZ-Batches
# ------------------------------------------------------------

BATCH_DIR = Path(
    "WZ2025_Vorbereitung/WZBatches_klassifiziert"
)

BATCH_PATTERN = (
    "WZ2025_Batch_*_klassifiziert*.xlsx"
)

BATCH_SHEET = "Standorte"


# ------------------------------------------------------------
# Neue Validierungs- UND Arbeitsdatei
#
# Diese Datei kannst du anschließend manuell bearbeiten.
#
# Wenn du den Code später erneut ausführst, versucht er,
# vorhandene manuelle Entscheidungen aus dieser Datei
# wieder einzulesen und zu erhalten.
# ------------------------------------------------------------

OUTPUT_FILE = (
        BASE_DIR
        / "WZ2025_Validierung_v4.xlsx"
)


# ------------------------------------------------------------
# Optionaler Radius für räumlich nahe Standorte desselben
# Unternehmens, FALLS Koordinaten in den Batchdateien
# vorhanden sind.
#
# Auch solche Fälle werden nur vorgeschlagen,
# niemals automatisch zusammengeführt.
# ------------------------------------------------------------

NEARBY_DISTANCE_METERS = 500


# ============================================================
# 2. MANUELLE ENTSCHEIDUNGSWERTE
# ============================================================

# ------------------------------------------------------------
# Standortverbünde
# ------------------------------------------------------------

VALID_SITE_DECISIONS = {
    "",
    "OFFEN",
    "ZUSAMMENFASSEN",
    "GETRENNT_LASSEN",
}


# ------------------------------------------------------------
# WZ-Prüfung
# ------------------------------------------------------------

VALID_WZ_DECISIONS = {
    "",
    "OFFEN",
    "KI_BESTAETIGT",
    "WZ_KORRIGIEREN",
    "NUR_WZ2_SICHER",
    "NICHT_SICHER_KLASSIFIZIERBAR",
}


VALID_WZ_LEVELS = {
    "",
    "Gruppe",
    "Abteilung",
    "Abschnitt",
}


# ============================================================
# 3. USER_SECTOR -> WZ CROSSWALK
# ============================================================

SECTOR_CROSSWALK = {

    "Produktion": {
        "strict": (
            "01.* | 02.* | 03.* | 05.* | 06.* | 07.* | 08.* | "
            "09.* | 10.* | 11.* | 12.* | 13.* | 14.* | 15.* | "
            "16.* | 17.* | 18.* | 19.* | 20.* | 21.* | 22.* | "
            "23.* | 24.* | 25.* | 26.* | 27.* | 28.* | 29.* | "
            "30.* | 31.* | 32.* | 33.*"
        ),
        "plausible": "",
    },

    "Handel": {
        "strict":
            "45.* | 46.* | 47.*",

        "plausible":
            "52.*",
    },

    "Dienstleistung": {
        "strict": (
            "49.* | 50.* | 51.* | 52.* | 53.* | 54.* | 55.* | "
            "56.* | 57.* | 58.* | 59.* | 60.* | 61.* | 62.* | "
            "63.* | 64.* | 65.* | 66.* | 68.* | 69.* | 70.* | "
            "71.* | 72.* | 73.* | 74.* | 75.* | 77.* | 78.* | "
            "79.* | 80.* | 81.* | 82.* | 84.* | 85.* | 86.* | "
            "87.* | 88.* | 90.* | 91.* | 92.* | 93.* | 94.* | "
            "95.* | 96.*"
        ),
        "plausible":
            "33.* | 36.* | 37.*",
    },

    "Energie": {
        "strict":
            "35.*",

        "plausible":
            "19.2 | 06.*",
    },

    "Entsorgung": {
        "strict":
            "37.* | 38.* | 39.*",

        "plausible":
            "35.* | 36.*",
    },
}


# ============================================================
# 4. USER_GOODS -> WZ CROSSWALK
# ============================================================

GOODS_CROSSWALK = {

    "Abfall": (
        "38.*",
        "37.* | 39.* | 35.1",
    ),

    "Abwasser": (
        "37.*",
        "38.*",
    ),

    "Automotive": (
        "29.*",
        "22.1 | 22.2 | 24.* | 25.* | 26.* | 27.* | 28.* | 30.*",
    ),

    "Baustoffe und Bauelemente": (
        "16.2 | 22.2 | 23.* | 25.1",
        "25.9 | 31.0",
    ),

    "Bildung": (
        "85.*",
        "",
    ),

    "Biomasse": (
        "38.*",
        "35.1 | 35.3",
    ),

    "Bäderbetrieb": (
        "93.1",
        "93.*",
    ),

    "Büroausstattung": (
        "46.* | 47.*",
        "31.0",
    ),

    "Chemie": (
        "20.*",
        "",
    ),

    "Druckerei": (
        "18.1",
        "18.*",
    ),

    "Einzelhandel": (
        "47.*",
        "52.1",
    ),

    "Elektrotechnik, Elektronik & elektrische Komponenten": (
        "26.* | 27.*",
        "28.*",
    ),

    "Finanzen": (
        "64.* | 66.*",
        "",
    ),

    "Forschung": (
        "72.*",
        "",
    ),

    "Freizeit": (
        "93.*",
        "",
    ),

    "Getränke": (
        "11.0",
        "10.3 | 10.8 | 46.* | 47.* | 52.1",
    ),

    "Glas und Keramik": (
        "23.1 | 23.2 | 23.3 | 23.4",
        "23.9",
    ),

    "Großhandel": (
        "46.*",
        "47.* | 52.*",
    ),

    "Holzverarbeitung": (
        "16.*",
        "",
    ),

    "Hotelbetrieb": (
        "55.*",
        "",
    ),

    "Hygiene und Drogerieprodukte": (
        "20.4",
        "21.*",
    ),

    "IT": (
        "62.* | 63.1",
        "46.5 | 47.*",
    ),

    "Immobilien": (
        "68.*",
        "",
    ),

    "Kernkraft": (
        "35.1",
        "35.*",
    ),

    "Klinik": (
        "86.1",
        "86.*",
    ),

    "Klärschlamm": (
        "37.* | 38.*",
        "35.*",
    ),

    "Kraftstoffe": (
        "19.2",
        "06.* | 35.* | 46.* | 47.*",
    ),

    "Kultur": (
        "90.* | 91.*",
        "",
    ),

    "Kunststoffe und Kunststofferzeugnisse": (
        "22.2",
        "20.1",
    ),

    "Labor / Prüfstand": (
        "71.2",
        "72.*",
    ),

    "Lebensmittel": (
        "10.1 | 10.2 | 10.3 | 10.4 | 10.5 | 10.6 | 10.7 | 10.8",
        "46.* | 47.*",
    ),

    "Logistik": (
        "49.* | 50.* | 51.* | 52.*",
        "53.* | 46.*",
    ),

    "Luftfahrt": (
        "30.3",
        "33.1 | 25.4 | 26.* | 27.*",
    ),

    "Maschinen- und Anlagenbau": (
        "28.*",
        "",
    ),

    "Medizintechnik": (
        "32.5",
        "26.* | 27.*",
    ),

    "Messe": (
        "82.3",
        "82.*",
    ),

    "Metallverarbeitung": (
        "24.* | 25.*",
        "28.4",
    ),

    "Mineralische Produkte": (
        "23.*",
        "08.*",
    ),

    "Oberflächenveredelung": (
        "25.5",
        "25.*",
    ),

    "Ofenbau": (
        "28.2",
        "28.*",
    ),

    "Papier- und Pappe": (
        "17.*",
        "",
    ),

    "Pflege": (
        "87.* | 88.*",
        "86.*",
    ),

    "Pharma- und Medizinprodukte": (
        "21.* | 32.5",
        "26.*",
    ),

    "Rohrreinigung": (
        "81.2",
        "37.*",
    ),

    "Sanitärprodukte": (
        "22.2 | 23.4 | 25.*",
        "28.*",
    ),

    "Schleifmittel": (
        "23.9",
        "23.*",
    ),

    "Sozialeinrichtung": (
        "87.* | 88.*",
        "86.*",
    ),

    "Sparkasse / IT": (
        "64.*",
        "62.* | 63.1",
    ),

    "Tabakwaren": (
        "12.0",
        "",
    ),

    "Textilien": (
        "13.*",
        "14.*",
    ),

    "Tierkadaver": (
        "38.*",
        "10.9",
    ),

    "Tiernahrung": (
        "10.9",
        "",
    ),

    "Verkaufsstände": (
        "25.1 | 25.9 | 31.0",
        "16.2 | 22.2",
    ),

    "Verkehrsbetrieb": (
        "49.* | 50.* | 51.*",
        "52.*",
    ),

    "Verpackungsmaterial": (
        "16.2 | 17.2 | 22.2 | 23.1 | 25.9",
        "20.1",
    ),

    "Versicherung": (
        "65.*",
        "",
    ),

    "Verteidigung": (
        "25.4 | 30.3 | 30.4",
        "26.* | 27.* | 28.*",
    ),

    "Wartung": (
        "33.1 | 33.2",
        "95.*",
    ),

    "Wasserverband": (
        "36.* | 37.*",
        "42.9 | 84.*",
    ),

    "Werkzeuge": (
        "25.6 | 28.4",
        "25.* | 28.*",
    ),

    "Wäscherei": (
        "96.1",
        "96.*",
    ),
}


# ============================================================
# 5. ALLGEMEINE HILFSFUNKTIONEN
# ============================================================

def clean_string(value) -> str:

    if pd.isna(value):
        return ""

    text = str(value).strip()

    if text.lower() in {
        "nan",
        "none",
    }:
        return ""

    return text


def normalize_text(value) -> str:
    """
    Allgemeine Textnormalisierung.
    """

    text = clean_string(value)

    if not text:
        return ""

    text = unicodedata.normalize(
        "NFKC",
        text,
    )

    text = (
        text
        .replace("\u200b", "")
        .replace("\ufeff", "")
        .casefold()
    )

    text = (
        text
        .replace("ä", "ae")
        .replace("ö", "oe")
        .replace("ü", "ue")
        .replace("ß", "ss")
    )

    text = re.sub(
        r"[^a-z0-9]+",
        " ",
        text,
    )

    text = re.sub(
        r"\s+",
        " ",
        text,
    )

    return text.strip()


def normalize_company(value) -> str:
    """
    Firmenname normalisieren und häufige Rechtsformen entfernen.
    """

    text = normalize_text(
        value
    )

    if not text:
        return ""

    legal_terms = {
        "gmbh",
        "mbh",
        "ag",
        "kg",
        "kgaa",
        "ohg",
        "se",
        "ug",
        "eg",
        "ev",
        "co",
        "haftungsbeschraenkt",
    }

    parts = [
        part
        for part in text.split()
        if part not in legal_terms
    ]

    return " ".join(
        parts
    )


def normalize_street(value) -> str:

    text = normalize_text(
        value
    )

    text = re.sub(
        r"\bstr\b",
        "strasse",
        text,
    )

    return text


def street_without_house_number(value) -> str:
    """
    Entfernt Hausnummer am Ende.

    Industriestrasse 10 -> Industriestrasse
    Industriestrasse 10a -> Industriestrasse
    Industriestrasse 10-14 -> Industriestrasse
    """

    text = normalize_street(
        value
    )

    if not text:
        return ""

    text = re.sub(
        r"\s+\d+\s*[a-z]?"
        r"(?:\s*[-/]\s*\d+\s*[a-z]?)?"
        r"\s*$",
        "",
        text,
    )

    return text.strip()


def normalize_postal(value) -> str:

    return re.sub(
        r"\D",
        "",
        clean_string(value),
    )


def first_existing_column(
        df,
        candidates,
):

    for column in candidates:

        if column in df.columns:
            return column

    return None


def get_series(
        df,
        candidates,
):

    column = first_existing_column(
        df,
        candidates,
    )

    if column is None:

        return pd.Series(
            [""] * len(df),
            index=df.index,
            )

    return df[column]


def batch_number_from_filename(
        filename,
):

    match = re.search(
        r"Batch[_\s-]*(\d+)",
        filename,
        flags=re.IGNORECASE,
    )

    if not match:
        return ""

    return (
            "BATCH_"
            + str(
        int(
            match.group(1)
        )
    ).zfill(3)
    )


def make_address(
        street,
        postal,
        city,
):

    parts = [
        clean_string(street),
        clean_string(postal),
        clean_string(city),
    ]

    return ", ".join(
        part
        for part in parts
        if part
    )


def combine_reference_description(
        company,
        location,
        address,
):

    values = [
        clean_string(company),
        clean_string(location),
        clean_string(address),
    ]

    return " | ".join(
        value
        for value in values
        if value
    )


# ============================================================
# 6. WZ-PATTERN
# ============================================================

def split_patterns(text):

    if not clean_string(text):
        return []

    return [
        item.strip()
        for item in str(text).split("|")
        if item.strip()
    ]


def wz_matches_pattern(
        code,
        pattern,
):

    code = clean_string(
        code
    )

    pattern = clean_string(
        pattern
    )

    if not code or not pattern:
        return False

    if pattern.endswith(".*"):

        prefix = pattern[:-1]

        return code.startswith(
            prefix
        )

    return (
            code == pattern
    )


def matches_any_wz(
        code,
        patterns,
):

    return any(
        wz_matches_pattern(
            code,
            pattern,
        )
        for pattern in split_patterns(
            patterns
        )
    )


# ============================================================
# 7. WZ-BATCHES VEREINHEITLICHEN
# ============================================================

def canonicalize_batch(
        df,
        filename,
):

    result = pd.DataFrame(
        index=df.index
    )


    result[
        "Batch_ID"
    ] = batch_number_from_filename(
        filename
    )


    mapping = {

        "Klassifikations_ID": [
            "Klassifikations_ID"
        ],

        "Unternehmens_ID": [
            "Unternehmens_ID"
        ],

        "Firmenname": [
            "Firmenname"
        ],

        "Standortname": [
            "Standortname"
        ],

        "Strasse_Hausnummer": [
            "Strasse_Hausnummer",
            "Straße_Hausnummer",
            "Straße und Hausnummer",
        ],

        "PLZ": [
            "PLZ"
        ],

        "Ort": [
            "Ort"
        ],

        "Melde_IDs": [
            "Melde_IDs"
        ],

        "Originaldatensatz_IDs": [
            "Originaldatensatz_IDs"
        ],

        "Anzahl_Originalzeilen": [
            "Anzahl_Originalzeilen"
        ],

        "WZ_Gruppe_Code": [
            "WZ_Gruppe_Code"
        ],

        "WZ_Gruppe_Bezeichnung": [
            "WZ_Gruppe_Bezeichnung"
        ],

        "Standorttaetigkeit_kurz": [
            "Standorttaetigkeit_kurz"
        ],

        "WZ_Begruendung": [
            "WZ_Begruendung"
        ],

        "Recherchequelle_1": [
            "Recherchequelle_1"
        ],

        "Recherchequelle_2": [
            "Recherchequelle_2"
        ],

        "Konfidenz": [
            "Konfidenz"
        ],
    }


    for target, candidates in mapping.items():

        result[target] = (
            get_series(
                df,
                candidates,
            )
            .map(
                clean_string
            )
        )


    # ========================================================
    # SICHERE FALLBACK-EBENE
    # ========================================================

    result[
        "WZ_Sichere_Ebene"
    ] = (

        get_series(
            df,
            [
                "WZ_Sichere_Ebene",
                "WZ2025_Sichere_Ebene",
            ],
        )

        .map(
            clean_string
        )
    )


    direct_safe_code = (

        get_series(
            df,
            [
                "WZ_Sicherer_Code",
                "WZ2025_Tiefster_Code",
            ],
        )

        .map(
            clean_string
        )
    )


    direct_safe_label = (

        get_series(
            df,
            [
                "WZ_Sichere_Bezeichnung",
                "WZ2025_Tiefste_Bezeichnung",
            ],
        )

        .map(
            clean_string
        )
    )


    def infer_safe_code(index):

        existing = direct_safe_code.loc[
            index
        ]

        if existing:
            return existing

        level = (
            result.loc[
                index,
                "WZ_Sichere_Ebene",
            ]
            .casefold()
        )


        possible = {

            "unterklasse":
                "WZ_Unterklasse_Code",

            "klasse":
                "WZ_Klasse_Code",

            "gruppe":
                "WZ_Gruppe_Code",

            "abteilung":
                "WZ_Abteilung_Code",

            "abschnitt":
                "WZ_Abschnitt_Code",
        }


        for key, column in possible.items():

            if (
                    key in level
                    and
                    column in df.columns
            ):

                return clean_string(
                    df.loc[
                        index,
                        column,
                    ]
                )

        return ""


    def infer_safe_label(index):

        existing = direct_safe_label.loc[
            index
        ]

        if existing:
            return existing

        level = (
            result.loc[
                index,
                "WZ_Sichere_Ebene",
            ]
            .casefold()
        )


        possible = {

            "unterklasse":
                "WZ_Unterklasse_Bezeichnung",

            "klasse":
                "WZ_Klasse_Bezeichnung",

            "gruppe":
                "WZ_Gruppe_Bezeichnung",

            "abteilung":
                "WZ_Abteilung_Bezeichnung",

            "abschnitt":
                "WZ_Abschnitt_Bezeichnung",
        }


        for key, column in possible.items():

            if (
                    key in level
                    and
                    column in df.columns
            ):

                return clean_string(
                    df.loc[
                        index,
                        column,
                    ]
                )

        return ""


    result[
        "WZ_Sicherer_Code"
    ] = [
        infer_safe_code(index)
        for index in df.index
    ]


    result[
        "WZ_Sichere_Bezeichnung"
    ] = [
        infer_safe_label(index)
        for index in df.index
    ]


    # ========================================================
    # OPTIONALE KOORDINATEN
    # ========================================================

    latitude_column = first_existing_column(
        df,
        [
            "Breitengrad",
            "Latitude",
            "latitude",
            "LAT",
        ],
    )


    longitude_column = first_existing_column(
        df,
        [
            "Laengengrad",
            "Längengrad",
            "Longitude",
            "longitude",
            "LON",
        ],
    )


    result[
        "Latitude"
    ] = (
        df[latitude_column]
        if latitude_column
        else ""
    )


    result[
        "Longitude"
    ] = (
        df[longitude_column]
        if longitude_column
        else ""
    )


    result[
        "Quelldatei"
    ] = filename


    return result


def load_all_batches():

    batch_files = sorted(
        file
        for file in BATCH_DIR.glob(
            BATCH_PATTERN
        )
        if not file.name.startswith(
            "~$"
        )
    )


    if not batch_files:

        raise FileNotFoundError(
            "Keine klassifizierten WZ-Batches gefunden."
        )


    frames = []


    for file in batch_files:

        print(
            f"   Einlesen: {file.name}"
        )

        raw = pd.read_excel(
            file,
            sheet_name=BATCH_SHEET,
            dtype=str,
            keep_default_na=False,
        )


        frames.append(
            canonicalize_batch(
                raw,
                file.name,
            )
        )


    sites = pd.concat(
        frames,
        ignore_index=True,
    )


    if sites[
        "Klassifikations_ID"
    ].eq("").any():

        raise ValueError(
            "Mindestens eine Klassifikations_ID fehlt."
        )


    duplicated = (
        sites[
            "Klassifikations_ID"
        ]
        .duplicated(
            keep=False
        )
    )


    if duplicated.any():

        examples = (

            sites.loc[
                duplicated,
                [
                    "Klassifikations_ID",
                    "Quelldatei",
                ],
            ]

            .head(30)

            .to_string(
                index=False
            )
        )


        raise ValueError(
            "Klassifikations_ID kommt mehrfach vor.\n\n"
            + examples
        )


    return (
        sites,
        batch_files,
    )


# ============================================================
# 8. REFERENZDATEI VORBEREITEN
# ============================================================

def split_reference_address(value):

    parts = [
        item.strip()
        for item in clean_string(
            value
        ).split(";")
    ]

    return (
        parts[0]
        if len(parts) >= 1
        else "",

        parts[1]
        if len(parts) >= 2
        else "",

        parts[2]
        if len(parts) >= 3
        else "",
    )


def prepare_reference():

    reference = pd.read_excel(
        REFERENCE_FILE,
        sheet_name=REFERENCE_SHEET,
        dtype=str,
        keep_default_na=False,
    )


    required = {
        "USER_UniqueID",
        "USER_Address",
        "USER_CompanyName",
        "USER_LocationName",
        "USER_Sector",
        "USER_Goods",
    }


    missing = (
            required
            - set(reference.columns)
    )


    if missing:

        raise ValueError(
            "Referenzdatei unvollständig:\n"
            + "\n".join(
                sorted(missing)
            )
        )


    if not reference[
        "USER_UniqueID"
    ].is_unique:

        raise ValueError(
            "USER_UniqueID ist nicht eindeutig."
        )


    address_parts = (
        reference[
            "USER_Address"
        ]
        .map(
            split_reference_address
        )
    )


    reference[
        "_Street"
    ] = [
        item[0]
        for item in address_parts
    ]


    reference[
        "_Postal_Address"
    ] = [
        item[1]
        for item in address_parts
    ]


    reference[
        "_City"
    ] = [
        item[2]
        for item in address_parts
    ]


    if "Postal" in reference.columns:

        reference[
            "_Postal"
        ] = (

            reference[
                "Postal"
            ]

            .map(
                normalize_postal
            )
        )


        missing_postal = (
                reference[
                    "_Postal"
                ]
                == ""
        )


        reference.loc[
            missing_postal,
            "_Postal",
        ] = (

            reference.loc[
                missing_postal,
                "_Postal_Address",
            ]

            .map(
                normalize_postal
            )
        )

    else:

        reference[
            "_Postal"
        ] = (

            reference[
                "_Postal_Address"
            ]

            .map(
                normalize_postal
            )
        )


    reference[
        "_Street_Norm"
    ] = (

        reference[
            "_Street"
        ]

        .map(
            normalize_street
        )
    )


    reference[
        "_City_Norm"
    ] = (

        reference[
            "_City"
        ]

        .map(
            normalize_text
        )
    )


    reference[
        "_Company_Norm"
    ] = (

        reference[
            "USER_CompanyName"
        ]

        .map(
            normalize_company
        )
    )


    reference[
        "_Location_Norm"
    ] = (

        reference[
            "USER_LocationName"
        ]

        .map(
            normalize_text
        )
    )


    return reference


# ============================================================
# 9. MATCHING
# ============================================================

try:

    from rapidfuzz import fuzz

    def similarity(a, b):

        if not a or not b:
            return 0.0

        return float(
            fuzz.WRatio(
                a,
                b,
            )
        )

except ImportError:

    from difflib import SequenceMatcher

    def similarity(a, b):

        if not a or not b:
            return 0.0

        return (
                SequenceMatcher(
                    None,
                    a,
                    b,
                ).ratio()
                * 100
        )


def create_index(
        reference,
        columns,
):

    index = defaultdict(
        list
    )


    for row_index, row in reference.iterrows():

        key = tuple(
            row[column]
            for column in columns
        )


        # Sehr wichtig:
        # Keine exakten Matches mit leeren Schlüsselteilen.
        if all(key):

            index[
                key
            ].append(
                row_index
            )


    return index


def unique_candidate(
        index,
        key,
):

    if not all(key):
        return None

    candidates = index.get(
        tuple(key),
        [],
    )

    candidates = list(
        dict.fromkeys(
            candidates
        )
    )

    if len(candidates) == 1:

        return candidates[0]

    return None


def prepare_site_matching_fields(
        sites,
):

    sites = sites.copy()


    sites[
        "_Company"
    ] = (

        sites[
            "Firmenname"
        ]

        .map(
            normalize_company
        )
    )


    sites[
        "_Street"
    ] = (

        sites[
            "Strasse_Hausnummer"
        ]

        .map(
            normalize_street
        )
    )


    sites[
        "_Postal"
    ] = (

        sites[
            "PLZ"
        ]

        .map(
            normalize_postal
        )
    )


    sites[
        "_City"
    ] = (

        sites[
            "Ort"
        ]

        .map(
            normalize_text
        )
    )


    sites[
        "_Location"
    ] = (

        sites[
            "Standortname"
        ]

        .map(
            normalize_text
        )
    )


    return sites


def fuzzy_score(
        site,
        ref,
):

    company_score = similarity(
        site["_Company"],
        ref["_Company_Norm"],
    )


    street_score = similarity(
        site["_Street"],
        ref["_Street_Norm"],
    )


    city_score = similarity(
        site["_City"],
        ref["_City_Norm"],
    )


    location_score = similarity(
        site["_Location"],
        ref["_Location_Norm"],
    )


    score = (
            0.45 * company_score
            + 0.35 * street_score
            + 0.10 * city_score
            + 0.10 * location_score
    )


    if (
            site["_Postal"]
            and
            site["_Postal"]
            == ref["_Postal"]
    ):

        score += 3


    return (
        min(
            score,
            100,
        ),
        company_score,
    )


def find_fuzzy_candidate(
        site,
        reference,
):

    # Erst gleiche PLZ
    if site["_Postal"]:

        candidates = reference[
            reference[
                "_Postal"
            ]
            == site["_Postal"]
            ]

    else:

        candidates = pd.DataFrame()


    # Sonst gleicher Ort
    if candidates.empty and site["_City"]:

        candidates = reference[
            reference[
                "_City_Norm"
            ]
            == site["_City"]
            ]


    if candidates.empty:

        return (
            None,
            0.0,
            0.0,
        )


    scores = []


    for ref_index, ref_row in candidates.iterrows():

        score, company_score = (
            fuzzy_score(
                site,
                ref_row,
            )
        )


        scores.append(
            (
                ref_index,
                score,
                company_score,
            )
        )


    scores.sort(
        key=lambda item: item[1],
        reverse=True,
    )


    best = scores[0]


    second_score = (
        scores[1][1]
        if len(scores) > 1
        else 0
    )


    return (
        best,
        best[1],
        best[1] - second_score,
    )


def match_sites_to_reference(
        sites,
        reference,
):

    sites = prepare_site_matching_fields(
        sites
    )


    indices = {

        "company_address":
            create_index(
                reference,
                [
                    "_Company_Norm",
                    "_Street_Norm",
                    "_Postal",
                    "_City_Norm",
                ],
            ),

        "address":
            create_index(
                reference,
                [
                    "_Street_Norm",
                    "_Postal",
                    "_City_Norm",
                ],
            ),

        "street_postal":
            create_index(
                reference,
                [
                    "_Street_Norm",
                    "_Postal",
                ],
            ),

        "company_postal_city":
            create_index(
                reference,
                [
                    "_Company_Norm",
                    "_Postal",
                    "_City_Norm",
                ],
            ),

        "company_postal":
            create_index(
                reference,
                [
                    "_Company_Norm",
                    "_Postal",
                ],
            ),

        "company_city":
            create_index(
                reference,
                [
                    "_Company_Norm",
                    "_City_Norm",
                ],
            ),
    }


    rows = []


    for _, site in sites.iterrows():

        ref_index = None
        candidate_index = None

        match_status = ""
        match_method = ""
        match_score = 0.0


        # ====================================================
        # 1. FIRMA + KOMPLETTE ADRESSE
        # ====================================================

        ref_index = unique_candidate(
            indices[
                "company_address"
            ],
            (
                site["_Company"],
                site["_Street"],
                site["_Postal"],
                site["_City"],
            ),
        )


        if ref_index is not None:

            match_status = "MATCH"
            match_method = "EXAKT_FIRMA_ADRESSE"
            match_score = 100.0


        # ====================================================
        # 2. EINDEUTIGE KOMPLETTE ADRESSE
        # ====================================================

        if ref_index is None:

            ref_index = unique_candidate(
                indices[
                    "address"
                ],
                (
                    site["_Street"],
                    site["_Postal"],
                    site["_City"],
                ),
            )


            if ref_index is not None:

                match_status = "MATCH"
                match_method = "EXAKT_EINDEUTIGE_ADRESSE"
                match_score = 99.0


        # ====================================================
        # 3. STRASSE + PLZ
        # ====================================================

        if ref_index is None:

            ref_index = unique_candidate(
                indices[
                    "street_postal"
                ],
                (
                    site["_Street"],
                    site["_Postal"],
                ),
            )


            if ref_index is not None:

                match_status = "MATCH"
                match_method = "EXAKT_EINDEUTIGE_STRASSE_PLZ"
                match_score = 98.0


        # ====================================================
        # 4. FIRMA + PLZ + ORT
        # ====================================================

        if ref_index is None:

            ref_index = unique_candidate(
                indices[
                    "company_postal_city"
                ],
                (
                    site["_Company"],
                    site["_Postal"],
                    site["_City"],
                ),
            )


            if ref_index is not None:

                match_status = "MATCH"
                match_method = "EXAKT_FIRMA_PLZ_ORT"
                match_score = 97.0


        # ====================================================
        # 5. FIRMA + PLZ
        # ====================================================

        if ref_index is None:

            ref_index = unique_candidate(
                indices[
                    "company_postal"
                ],
                (
                    site["_Company"],
                    site["_Postal"],
                ),
            )


            if ref_index is not None:

                match_status = "MATCH"
                match_method = "EXAKT_FIRMA_PLZ"
                match_score = 95.0


        # ====================================================
        # 6. FIRMA + ORT
        # ====================================================

        if ref_index is None:

            ref_index = unique_candidate(
                indices[
                    "company_city"
                ],
                (
                    site["_Company"],
                    site["_City"],
                ),
            )


            if ref_index is not None:

                match_status = "MATCH"
                match_method = "EXAKT_FIRMA_ORT"
                match_score = 93.0


        # ====================================================
        # 7. FUZZY NUR ALS KONSERVATIVER FALLBACK
        # ====================================================

        if ref_index is None:

            (
                best,
                best_score,
                margin,
            ) = find_fuzzy_candidate(
                site,
                reference,
            )


            if best is None:

                match_status = "KEIN_MATCH"
                match_method = "KEIN_KANDIDAT"

            else:

                candidate_index = best[0]
                company_score = best[2]


                candidate_ref = reference.loc[
                    candidate_index
                ]


                same_postal = (
                        site["_Postal"]
                        and
                        site["_Postal"]
                        == candidate_ref["_Postal"]
                )


                same_city = (
                        site["_City"]
                        and
                        site["_City"]
                        == candidate_ref["_City_Norm"]
                )


                if (
                        best_score >= 88
                        and
                        margin >= 8
                        and
                        company_score >= 90
                        and
                        (
                                same_postal
                                or
                                same_city
                        )
                ):

                    ref_index = candidate_index

                    match_status = "MATCH"
                    match_method = "FUZZY"

                    match_score = round(
                        best_score,
                        1,
                    )

                else:

                    match_status = "UNSICHER"
                    match_method = "FUZZY_KANDIDAT"

                    match_score = round(
                        best_score,
                        1,
                    )


        output = site.to_dict()


        output[
            "Match_Status"
        ] = match_status


        output[
            "Match_Methode"
        ] = match_method


        output[
            "Match_Score"
        ] = match_score


        # ====================================================
        # UNSICHERER KANDIDAT
        # ====================================================

        if (
                match_status == "UNSICHER"
                and
                candidate_index is not None
        ):

            candidate = reference.loc[
                candidate_index
            ]


            output[
                "Kandidat_USER_UniqueID"
            ] = candidate[
                "USER_UniqueID"
            ]


            output[
                "Kandidat_CompanyName"
            ] = candidate[
                "USER_CompanyName"
            ]


            output[
                "Kandidat_LocationName"
            ] = candidate[
                "USER_LocationName"
            ]


            output[
                "Kandidat_Address"
            ] = candidate[
                "USER_Address"
            ]

        else:

            output[
                "Kandidat_USER_UniqueID"
            ] = ""

            output[
                "Kandidat_CompanyName"
            ] = ""

            output[
                "Kandidat_LocationName"
            ] = ""

            output[
                "Kandidat_Address"
            ] = ""


        # ====================================================
        # SICHERER REFERENZMATCH
        # ====================================================

        if ref_index is not None:

            ref = reference.loc[
                ref_index
            ]


            output[
                "Ref_USER_UniqueID"
            ] = ref[
                "USER_UniqueID"
            ]


            output[
                "Ref_CompanyName"
            ] = ref[
                "USER_CompanyName"
            ]


            output[
                "Ref_LocationName"
            ] = ref[
                "USER_LocationName"
            ]


            output[
                "Ref_Address"
            ] = ref[
                "USER_Address"
            ]


            output[
                "Alt_USER_Sector"
            ] = ref[
                "USER_Sector"
            ]


            output[
                "Alt_USER_Goods"
            ] = ref[
                "USER_Goods"
            ]

        else:

            for column in [
                "Ref_USER_UniqueID",
                "Ref_CompanyName",
                "Ref_LocationName",
                "Ref_Address",
                "Alt_USER_Sector",
                "Alt_USER_Goods",
            ]:

                output[column] = ""


        rows.append(
            output
        )


    results = pd.DataFrame(
        rows
    )


    results = results.drop(
        columns=[
            "_Company",
            "_Street",
            "_Postal",
            "_City",
            "_Location",
        ],
        errors="ignore",
    )


    return results


# ============================================================
# 10. WZ CROSSWALK-PRÜFUNG
# ============================================================

def determine_comparison_wz(
        row,
):

    group_code = clean_string(
        row[
            "WZ_Gruppe_Code"
        ]
    )


    if group_code:

        return (
            group_code,
            "Gruppe",
        )


    safe_code = clean_string(
        row[
            "WZ_Sicherer_Code"
        ]
    )


    safe_level = clean_string(
        row[
            "WZ_Sichere_Ebene"
        ]
    )


    if safe_code:

        return (
            safe_code,
            safe_level,
        )


    return (
        "",
        "Nicht klassifiziert",
    )


def check_sector(
        sector,
        wz_code,
):

    sector = clean_string(
        sector
    )

    wz_code = clean_string(
        wz_code
    )


    if not wz_code:

        return "KEIN_CODE"


    config = (
        SECTOR_CROSSWALK.get(
            sector
        )
    )


    if config is None:

        return "NEIN"


    if matches_any_wz(
            wz_code,
            config[
                "strict"
            ],
    ):

        return "STRIKT"


    if matches_any_wz(
            wz_code,
            config[
                "plausible"
            ],
    ):

        return "PLAUSIBEL"


    return "NEIN"


def check_goods(
        goods,
        wz_code,
        wz_level,
):

    goods = clean_string(
        goods
    )


    if not goods:

        return "KEIN_USER_GOODS"


    if (
            "gruppe"
            not in clean_string(
        wz_level
    ).casefold()
    ):

        return "WZ_ZU_GROB"


    config = (
        GOODS_CROSSWALK.get(
            goods
        )
    )


    if config is None:

        return "NEIN"


    strict, plausible = config


    if matches_any_wz(
            wz_code,
            strict,
    ):

        return "STRIKT"


    if matches_any_wz(
            wz_code,
            plausible,
    ):

        return "PLAUSIBEL"


    return "NEIN"


def determine_basis_check(
        row,
):

    if row[
        "Match_Status"
    ] != "MATCH":

        return "NICHT_VERGLEICHBAR"


    sector = row[
        "Sektor_Check"
    ]

    goods = row[
        "Goods_Check"
    ]


    if (
            sector == "KEIN_CODE"
            or
            goods == "WZ_ZU_GROB"
    ):

        return "WZ_ZU_GROB_ODER_FEHLT"


    if sector in {
        "STRIKT",
        "PLAUSIBEL",
    }:

        if goods == "STRIKT":

            return "OK_STRIKT"

        if goods == "PLAUSIBEL":

            return "OK_PLAUSIBEL"

        if goods == "KEIN_USER_GOODS":

            return "OK_NUR_SEKTOR"

        if goods == "NEIN":

            return "PRUEFEN_BRANCHE"


    if sector == "NEIN":

        return "WIDERSPRUCH"


    return "NICHT_VERGLEICHBAR"


# ============================================================
# 11. STANDORTFUNKTION
# ============================================================

def contains_any(
        text,
        keywords,
):

    return any(
        keyword in text
        for keyword in keywords
    )


def detect_site_function(
        row,
):

    wz = clean_string(
        row[
            "WZ_Gruppe_Code"
        ]
    )


    text = normalize_text(
        " ".join(
            [
                clean_string(
                    row[
                        "Firmenname"
                    ]
                ),

                clean_string(
                    row[
                        "Standortname"
                    ]
                ),

                clean_string(
                    row[
                        "Standorttaetigkeit_kurz"
                    ]
                ),

                clean_string(
                    row[
                        "WZ_Begruendung"
                    ]
                ),
            ]
        )
    )


    if (
            wz.startswith(
                (
                        "52.",
                        "53.",
                )
            )
            and
            contains_any(
                text,
                [
                    "lager",
                    "logistik",
                    "distribution",
                    "depot",
                    "terminal",
                    "zentrallager",
                ],
            )
    ):

        return "LAGER_LOGISTIK"


    if (
            wz.startswith(
                "47."
            )
            and
            contains_any(
                text,
                [
                    "filiale",
                    "supermarkt",
                    "markt",
                    "discounter",
                    "warenhaus",
                    "kaufhaus",
                ],
            )
    ):

        return "EINZELHANDEL_FILIALE"


    if (
            wz == "70.1"
            and
            contains_any(
                text,
                [
                    "verwaltung",
                    "hauptverwaltung",
                    "zentrale",
                    "headquarter",
                    "headquarters",
                ],
            )
    ):

        return "UNTERNEHMENSVERWALTUNG"


    if (
            wz.startswith(
                "72."
            )
            and
            contains_any(
                text,
                [
                    "forschung",
                    "entwicklung",
                    "research",
                    "development",
                    "forschungszentrum",
                    "entwicklungszentrum",
                ],
            )
    ):

        return "FORSCHUNG_ENTWICKLUNG"


    if (
            wz in {
        "35.1",
        "35.3",
    }
            and
            contains_any(
                text,
                [
                    "bhkw",
                    "blockheizkraftwerk",
                    "kesselhaus",
                    "heizzentrale",
                    "heizwerk",
                    "kraftwerk",
                    "energiezentrale",
                    "waermezentrale",
                    "dampfzentrale",
                ],
            )
    ):

        return "ENERGIE_TECHNISCHE_ANLAGE"


    if (
            wz in {
        "33.1",
        "33.2",
    }
            and
            contains_any(
                text,
                [
                    "instandhaltung",
                    "reparatur",
                    "werkstatt",
                    "servicewerkstatt",
                ],
            )
    ):

        return "INSTANDHALTUNG_REPARATUR"


    return ""


# ============================================================
# 12. VALIDIERUNGSERGEBNIS
# ============================================================

def add_validation_results(
        results,
):

    results = results.copy()


    comparisons = results.apply(
        determine_comparison_wz,
        axis=1,
    )


    results[
        "WZ_Vergleichscode"
    ] = [
        item[0]
        for item in comparisons
    ]


    results[
        "WZ_Vergleichsebene"
    ] = [
        item[1]
        for item in comparisons
    ]


    results[
        "Sektor_Check"
    ] = results.apply(
        lambda row:
        check_sector(
            row[
                "Alt_USER_Sector"
            ],
            row[
                "WZ_Vergleichscode"
            ],
        )
        if row[
               "Match_Status"
           ] == "MATCH"
        else "",
        axis=1,
    )


    results[
        "Goods_Check"
    ] = results.apply(
        lambda row:
        check_goods(
            row[
                "Alt_USER_Goods"
            ],
            row[
                "WZ_Vergleichscode"
            ],
            row[
                "WZ_Vergleichsebene"
            ],
        )
        if row[
               "Match_Status"
           ] == "MATCH"
        else "",
        axis=1,
    )


    results[
        "Basis_Check"
    ] = results.apply(
        determine_basis_check,
        axis=1,
    )


    results[
        "Standortfunktion"
    ] = results.apply(
        detect_site_function,
        axis=1,
    )


    results[
        "Gesamt_Check"
    ] = results[
        "Basis_Check"
    ]


    # Standortfunktion kann einen nominellen Konflikt erklären.
    function_mask = (

            results[
                "Standortfunktion"
            ].ne("")

            &

            results[
                "Basis_Check"
            ].isin(
                {
                    "WIDERSPRUCH",
                    "PRUEFEN_BRANCHE",
                }
            )
    )


    results.loc[
        function_mask,
        "Gesamt_Check",
    ] = "STANDORTFUNKTION_ABWEICHEND"


    # ========================================================
    # REFERENZ-MEHRFACHMATCH
    # ========================================================

    matched_mask = (
            results[
                "Match_Status"
            ]
            == "MATCH"
    )


    ref_counts = (

        results.loc[
            matched_mask,
            "Ref_USER_UniqueID",
        ]

        .value_counts()
    )


    results[
        "Anzahl_Neue_Standorte_pro_Referenz"
    ] = (

        results[
            "Ref_USER_UniqueID"
        ]

        .map(
            ref_counts
        )

        .fillna(0)

        .astype(int)
    )


    results[
        "Referenz_Mehrfach_Match"
    ] = (

            results[
                "Anzahl_Neue_Standorte_pro_Referenz"
            ]
            > 1
    )


    return results


# ============================================================
# 13. REFERENZABDECKUNG
# ============================================================

def build_reference_coverage(
        reference,
        results,
):

    matched = results[
        results[
            "Match_Status"
        ]
        == "MATCH"
        ]


    uncertain = results[
        results[
            "Match_Status"
        ]
        == "UNSICHER"
        ]


    rows = []


    for _, ref in reference.iterrows():

        ref_id = clean_string(
            ref[
                "USER_UniqueID"
            ]
        )


        safe_matches = matched[
            matched[
                "Ref_USER_UniqueID"
            ]
            == ref_id
            ]


        uncertain_matches = uncertain[
            uncertain[
                "Kandidat_USER_UniqueID"
            ]
            == ref_id
            ]


        rows.append(
            {

                "USER_UniqueID":
                    ref_id,

                "USER_CompanyName":
                    ref[
                        "USER_CompanyName"
                    ],

                "USER_LocationName":
                    ref[
                        "USER_LocationName"
                    ],

                "USER_Address":
                    ref[
                        "USER_Address"
                    ],

                "USER_Sector":
                    ref[
                        "USER_Sector"
                    ],

                "USER_Goods":
                    ref[
                        "USER_Goods"
                    ],

                "Sicher_wiedergefunden":
                    len(
                        safe_matches
                    ) > 0,

                "Anzahl_sichere_Matches":
                    len(
                        safe_matches
                    ),

                "Gematchte_Klassifikations_IDs":
                    " | ".join(
                        safe_matches[
                            "Klassifikations_ID"
                        ]
                        .astype(str)
                        .tolist()
                    ),

                "Anzahl_unsichere_Kandidaten":
                    len(
                        uncertain_matches
                    ),

                "Unsichere_Kandidaten_IDs":
                    " | ".join(
                        uncertain_matches[
                            "Klassifikations_ID"
                        ]
                        .astype(str)
                        .tolist()
                    ),
            }
        )


    return pd.DataFrame(
        rows
    )


# ============================================================
# 14. STANDORTVERBUND-KANDIDATEN
# ============================================================

class UnionFind:

    def __init__(
            self,
            items,
    ):

        self.parent = {
            item: item
            for item in items
        }


    def find(
            self,
            item,
    ):

        if self.parent[item] != item:

            self.parent[item] = (
                self.find(
                    self.parent[item]
                )
            )

        return self.parent[item]


    def union(
            self,
            a,
            b,
    ):

        root_a = self.find(
            a
        )

        root_b = self.find(
            b
        )


        if root_a != root_b:

            self.parent[
                root_b
            ] = root_a


def company_key(
        row,
):

    company_id = clean_string(
        row[
            "Unternehmens_ID"
        ]
    )


    if company_id:

        return (
                "ID::"
                + company_id
        )


    return (
            "NAME::"
            + normalize_company(
        row[
            "Firmenname"
        ]
    )
    )


def parse_coordinate(
        value,
):

    text = clean_string(
        value
    ).replace(
        ",",
        ".",
    )


    if not text:
        return np.nan


    try:

        return float(
            text
        )

    except ValueError:

        return np.nan


def haversine_distance_m(
        lat1,
        lon1,
        lat2,
        lon2,
):

    radius = 6371000


    phi1 = math.radians(
        lat1
    )

    phi2 = math.radians(
        lat2
    )


    dphi = math.radians(
        lat2 - lat1
    )

    dlambda = math.radians(
        lon2 - lon1
    )


    a = (
            math.sin(
                dphi / 2
            ) ** 2
            +
            math.cos(phi1)
            * math.cos(phi2)
            * math.sin(
        dlambda / 2
    ) ** 2
    )


    return (
            2
            * radius
            * math.asin(
        math.sqrt(a)
    )
    )


def connect_group(
        group,
        union_find,
        reason_map,
        reason,
):

    ids = (
        group[
            "Klassifikations_ID"
        ]
        .drop_duplicates()
        .tolist()
    )


    if len(ids) < 2:
        return


    first = ids[0]


    for other in ids[1:]:

        union_find.union(
            first,
            other,
        )


    for site_id in ids:

        reason_map[
            site_id
        ].add(
            reason
        )


def build_site_union_review(
        results,
):

    work = results.copy()


    work[
        "_Company_Key"
    ] = work.apply(
        company_key,
        axis=1,
    )


    work[
        "_Address"
    ] = (

        work[
            "Strasse_Hausnummer"
        ]

        .map(
            normalize_street
        )
    )


    work[
        "_Street_Base"
    ] = (

        work[
            "Strasse_Hausnummer"
        ]

        .map(
            street_without_house_number
        )
    )


    work[
        "_Postal"
    ] = (

        work[
            "PLZ"
        ]

        .map(
            normalize_postal
        )
    )


    work[
        "_City"
    ] = (

        work[
            "Ort"
        ]

        .map(
            normalize_text
        )
    )


    work[
        "_Location"
    ] = (

        work[
            "Standortname"
        ]

        .map(
            normalize_text
        )
    )


    ids = work[
        "Klassifikations_ID"
    ].tolist()


    uf = UnionFind(
        ids
    )


    reasons = defaultdict(
        set
    )


    # ========================================================
    # A) GLEICHE FIRMA + EXAKT GLEICHE ADRESSE
    # ========================================================

    same_address = work[
        work[
            "_Address"
        ].ne("")
    ]


    for _, group in same_address.groupby(
            [
                "_Company_Key",
                "_Address",
                "_Postal",
                "_City",
            ],
            sort=False,
    ):

        connect_group(
            group,
            uf,
            reasons,
            "GLEICHE_FIRMA_GLEICHE_ADRESSE",
        )


    # ========================================================
    # B) GLEICHE FIRMA + GLEICHE STRASSE,
    #    ABER UNTERSCHIEDLICHE HAUSNUMMER
    # ========================================================

    same_street = work[
        work[
            "_Street_Base"
        ].ne("")
    ]


    for _, group in same_street.groupby(
            [
                "_Company_Key",
                "_Street_Base",
                "_Postal",
                "_City",
            ],
            sort=False,
    ):

        if group[
            "_Address"
        ].nunique() >= 2:

            connect_group(
                group,
                uf,
                reasons,
                "GLEICHE_FIRMA_GLEICHE_STRASSE_ANDERE_HAUSNUMMER",
            )


    # ========================================================
    # C) GLEICHE FIRMA + GLEICHER STANDORTNAME
    # ========================================================

    meaningful = work[
        work[
            "_Location"
        ].str.len()
        >= 4
        ].copy()


    for _, group in meaningful.groupby(
            [
                "_Company_Key",
                "_Location",
                "_Postal",
                "_City",
            ],
            sort=False,
    ):

        if group[
            "_Address"
        ].nunique() >= 2:

            connect_group(
                group,
                uf,
                reasons,
                "GLEICHE_FIRMA_GLEICHER_STANDORTNAME",
            )


    # ========================================================
    # D) MEHRERE NEUE STANDORTE MATCHEN DENSELBEN
    #    ALTEN REFERENZSTANDORT
    #
    # Sehr hilfreicher zusätzlicher Hinweis.
    # ========================================================

    reference_matches = work[
        work[
            "Ref_USER_UniqueID"
        ].ne("")
    ]


    for _, group in reference_matches.groupby(
            "Ref_USER_UniqueID",
            sort=False,
    ):

        if len(group) >= 2:

            connect_group(
                group,
                uf,
                reasons,
                "GLEICHER_REFERENZSTANDORT",
            )


    # ========================================================
    # E) OPTIONAL KOORDINATEN
    # ========================================================

    work[
        "_Lat"
    ] = work[
        "Latitude"
    ].map(
        parse_coordinate
    )


    work[
        "_Lon"
    ] = work[
        "Longitude"
    ].map(
        parse_coordinate
    )


    geo = work[
        work[
            "_Lat"
        ].notna()
        &
        work[
            "_Lon"
        ].notna()
        ]


    if not geo.empty:

        for _, group in geo.groupby(
                "_Company_Key",
                sort=False,
        ):

            records = list(
                group.iterrows()
            )


            for i in range(
                    len(records)
            ):

                _, a = records[i]


                for j in range(
                        i + 1,
                        len(records),
                ):

                    _, b = records[j]


                    distance = (
                        haversine_distance_m(
                            a["_Lat"],
                            a["_Lon"],
                            b["_Lat"],
                            b["_Lon"],
                        )
                    )


                    if distance <= NEARBY_DISTANCE_METERS:

                        id_a = a[
                            "Klassifikations_ID"
                        ]

                        id_b = b[
                            "Klassifikations_ID"
                        ]


                        uf.union(
                            id_a,
                            id_b,
                        )


                        reason = (
                            "GLEICHE_FIRMA_RAEUMLICH_NAH_"
                            f"{round(distance)}m"
                        )


                        reasons[
                            id_a
                        ].add(
                            reason
                        )


                        reasons[
                            id_b
                        ].add(
                            reason
                        )


    # ========================================================
    # KOMPONENTEN BILDEN
    # ========================================================

    components = defaultdict(
        list
    )


    for site_id in ids:

        components[
            uf.find(
                site_id
            )
        ].append(
            site_id
        )


    candidate_id_map = {}
    candidate_size_map = {}


    for members in components.values():

        if len(members) < 2:

            continue


        members = sorted(
            members
        )


        digest = (

            hashlib.sha1(
                "|".join(
                    members
                ).encode(
                    "utf-8"
                )
            )

            .hexdigest()[:10]

            .upper()
        )


        candidate_id = (
                "KAND_"
                + digest
        )


        for site_id in members:

            candidate_id_map[
                site_id
            ] = candidate_id

            candidate_size_map[
                site_id
            ] = len(
                members
            )


    candidate_rows = work[
        work[
            "Klassifikations_ID"
        ].isin(
            candidate_id_map.keys()
        )
    ].copy()


    candidate_rows[
        "Kandidatenverbund_ID"
    ] = (

        candidate_rows[
            "Klassifikations_ID"
        ]

        .map(
            candidate_id_map
        )
    )


    candidate_rows[
        "Anzahl_Standorte_im_Kandidatenverbund"
    ] = (

        candidate_rows[
            "Klassifikations_ID"
        ]

        .map(
            candidate_size_map
        )
    )


    candidate_rows[
        "Kandidaten_Grund"
    ] = (

        candidate_rows[
            "Klassifikations_ID"
        ]

        .map(
            lambda site_id:
            " | ".join(
                sorted(
                    reasons[
                        site_id
                    ]
                )
            )
        )
    )


    candidate_rows[
        "Adresse"
    ] = candidate_rows.apply(
        lambda row:
        make_address(
            row[
                "Strasse_Hausnummer"
            ],
            row[
                "PLZ"
            ],
            row[
                "Ort"
            ],
        ),
        axis=1,
    )


    candidate_rows[
        "Referenz_Standort"
    ] = candidate_rows.apply(
        lambda row:
        combine_reference_description(
            row[
                "Ref_CompanyName"
            ],
            row[
                "Ref_LocationName"
            ],
            row[
                "Ref_Address"
            ],
        ),
        axis=1,
    )


    review = candidate_rows[
        [
            "Kandidatenverbund_ID",

            "Anzahl_Standorte_im_Kandidatenverbund",

            "Kandidaten_Grund",

            "Klassifikations_ID",

            "Unternehmens_ID",

            "Firmenname",

            "Standortname",

            "Adresse",

            "WZ_Gruppe_Code",

            "WZ_Gruppe_Bezeichnung",

            "Standorttaetigkeit_kurz",

            "Ref_USER_UniqueID",

            "Referenz_Standort",
        ]
    ].copy()


    # ========================================================
    # MANUELLE SPALTEN
    #
    # Gruppe ist bereits vorgeschlagen.
    #
    # Wenn der gesamte Kandidatenverbund zusammengehört,
    # musst du diese Spalte NICHT ändern.
    # ========================================================

    review[
        "Standortverbund_Entscheidung"
    ] = "OFFEN"


    review[
        "Standortverbund_Gruppe"
    ] = review[
        "Kandidatenverbund_ID"
    ]


    review[
        "Standortverbund_Begruendung"
    ] = ""


    return review.sort_values(
        [
            "Kandidatenverbund_ID",
            "Firmenname",
            "Adresse",
        ],
        kind="stable",
    ).reset_index(
        drop=True
    )


# ============================================================
# 15. KOMPAKTES WZ-PRÜFUNGSBLATT
# ============================================================

def build_wz_review(
        results,
):

    manual_statuses = {
        "WIDERSPRUCH",
        "PRUEFEN_BRANCHE",
        "STANDORTFUNKTION_ABWEICHEND",
        "WZ_ZU_GROB_ODER_FEHLT",
    }


    confidence_low = (

            results[
                "Konfidenz"
            ]

            .map(
                clean_string
            )

            .str.lower()
            == "gering"
    )


    mask = (

            results[
                "Gesamt_Check"
            ].isin(
                manual_statuses
            )

            |

            confidence_low
    )


    review = results[
        mask
    ].copy()


    def determine_review_reason(row):

        reasons = []


        if row[
            "Gesamt_Check"
        ] in manual_statuses:

            reasons.append(
                row[
                    "Gesamt_Check"
                ]
            )


        if (
                clean_string(
                    row[
                        "Konfidenz"
                    ]
                ).lower()
                == "gering"
        ):

            reasons.append(
                "GERINGE_KONFIDENZ"
            )


        return " | ".join(
            reasons
        )


    review[
        "Pruefgrund"
    ] = review.apply(
        determine_review_reason,
        axis=1,
    )


    review[
        "Adresse"
    ] = review.apply(
        lambda row:
        make_address(
            row[
                "Strasse_Hausnummer"
            ],
            row[
                "PLZ"
            ],
            row[
                "Ort"
            ],
        ),
        axis=1,
    )


    review[
        "Referenz_Standort"
    ] = review.apply(
        lambda row:
        combine_reference_description(
            row[
                "Ref_CompanyName"
            ],
            row[
                "Ref_LocationName"
            ],
            row[
                "Ref_Address"
            ],
        ),
        axis=1,
    )


    # --------------------------------------------------------
    # Umbenennung nur im Prüfblatt.
    #
    # Damit ist sofort sichtbar:
    # Das ist das ursprüngliche KI-Ergebnis.
    # --------------------------------------------------------

    review[
        "WZ_AI_Code"
    ] = review[
        "WZ_Gruppe_Code"
    ]


    review[
        "WZ_AI_Bezeichnung"
    ] = review[
        "WZ_Gruppe_Bezeichnung"
    ]


    compact_columns = [

        "Klassifikations_ID",

        "Firmenname",

        "Standortname",

        "Adresse",

        "Pruefgrund",

        "WZ_AI_Code",

        "WZ_AI_Bezeichnung",

        "WZ_Sichere_Ebene",

        "WZ_Sicherer_Code",

        "Standorttaetigkeit_kurz",

        "Konfidenz",

        "WZ_Begruendung",

        "Recherchequelle_1",

        "Recherchequelle_2",

        "Alt_USER_Sector",

        "Alt_USER_Goods",

        "Sektor_Check",

        "Goods_Check",

        "Referenz_Standort",

        "Match_Methode",

        "Referenz_Mehrfach_Match",

        "Standortfunktion",
    ]


    review = review[
        compact_columns
    ].copy()


    # ========================================================
    # MANUELLE ENTSCHEIDUNGSSPALTEN
    # ========================================================

    review[
        "WZ_Entscheidung"
    ] = "OFFEN"


    review[
        "WZ_Manuell_Code"
    ] = ""


    review[
        "WZ_Manuell_Ebene"
    ] = ""


    review[
        "WZ_Manuell_Begruendung"
    ] = ""


    return review.reset_index(
        drop=True
    )


# ============================================================
# 16. BESTEHENDE MANUELLE EINTRÄGE BEWAHREN
# ============================================================

def load_existing_manual_reviews():
    """
    Falls WZ2025_Validierung_v4.xlsx bereits existiert,
    werden deine bisherigen manuellen Entscheidungen
    vor dem Überschreiben eingelesen.
    """

    old_wz = None
    old_site = None


    if not OUTPUT_FILE.exists():

        return (
            old_wz,
            old_site,
        )


    try:

        excel = pd.ExcelFile(
            OUTPUT_FILE
        )


        if (
                "WZ_Pruefung"
                in excel.sheet_names
        ):

            old_wz = pd.read_excel(
                OUTPUT_FILE,
                sheet_name="WZ_Pruefung",
                dtype=str,
                keep_default_na=False,
            )


        if (
                "Standortverbund_Pruefung"
                in excel.sheet_names
        ):

            old_site = pd.read_excel(
                OUTPUT_FILE,
                sheet_name="Standortverbund_Pruefung",
                dtype=str,
                keep_default_na=False,
            )


    except Exception as exc:

        print(
            "WARNUNG: Vorhandene manuelle Entscheidungen "
            "konnten nicht eingelesen werden:"
        )

        print(
            exc
        )


    return (
        old_wz,
        old_site,
    )


def restore_manual_columns(
        new_df,
        old_df,
        key_column,
        manual_columns,
):

    if (
            old_df is None
            or
            old_df.empty
            or
            key_column not in old_df.columns
    ):

        return new_df


    old_work = (

        old_df
        .drop_duplicates(
            subset=[
                key_column
            ],
            keep="last",
        )
        .set_index(
            key_column
        )
    )


    for column in manual_columns:

        if column not in old_work.columns:

            continue


        mapping = (
            old_work[
                column
            ]
            .to_dict()
        )


        new_df[
            column
        ] = new_df.apply(
            lambda row:
            mapping.get(
                row[
                    key_column
                ],
                row[
                    column
                ],
            ),
            axis=1,
        )


    return new_df


# ============================================================
# 17. MANUELLE EINTRÄGE AUF TYPOFEHLER PRÜFEN
# ============================================================

def validate_manual_inputs(
        wz_review,
        site_review,
):

    # ========================================================
    # WZ
    # ========================================================

    wz_decisions = (
        wz_review[
            "WZ_Entscheidung"
        ]
        .map(
            clean_string
        )
    )


    invalid_wz = (
        ~wz_decisions.isin(
            VALID_WZ_DECISIONS
        )
    )


    if invalid_wz.any():

        values = (
            wz_decisions[
                invalid_wz
            ]
            .unique()
            .tolist()
        )


        raise ValueError(
            "Ungültige Werte in WZ_Entscheidung:\n"
            f"{values}\n\n"
            f"Erlaubt: {sorted(VALID_WZ_DECISIONS)}"
        )


    wz_levels = (
        wz_review[
            "WZ_Manuell_Ebene"
        ]
        .map(
            clean_string
        )
    )


    invalid_levels = (
        ~wz_levels.isin(
            VALID_WZ_LEVELS
        )
    )


    if invalid_levels.any():

        values = (
            wz_levels[
                invalid_levels
            ]
            .unique()
            .tolist()
        )


        raise ValueError(
            "Ungültige Werte in WZ_Manuell_Ebene:\n"
            f"{values}"
        )


    # --------------------------------------------------------
    # Bei Korrektur muss Code vorhanden sein.
    # --------------------------------------------------------

    correction_mask = (
            wz_decisions
            ==
            "WZ_KORRIGIEREN"
    )


    missing_code = (
            correction_mask
            &
            wz_review[
                "WZ_Manuell_Code"
            ]
            .map(
                clean_string
            )
            .eq("")
    )


    if missing_code.any():

        raise ValueError(
            "Bei mindestens einem WZ_KORRIGIEREN-Fall "
            "fehlt WZ_Manuell_Code."
        )


    # --------------------------------------------------------
    # Bei NUR_WZ2_SICHER ebenfalls Code nötig.
    # --------------------------------------------------------

    wz2_mask = (
            wz_decisions
            ==
            "NUR_WZ2_SICHER"
    )


    missing_wz2_code = (
            wz2_mask
            &
            wz_review[
                "WZ_Manuell_Code"
            ]
            .map(
                clean_string
            )
            .eq("")
    )


    if missing_wz2_code.any():

        raise ValueError(
            "Bei NUR_WZ2_SICHER fehlt mindestens "
            "einmal WZ_Manuell_Code."
        )


    # ========================================================
    # STANDORTVERBUND
    # ========================================================

    site_decisions = (
        site_review[
            "Standortverbund_Entscheidung"
        ]
        .map(
            clean_string
        )
    )


    invalid_site = (
        ~site_decisions.isin(
            VALID_SITE_DECISIONS
        )
    )


    if invalid_site.any():

        values = (
            site_decisions[
                invalid_site
            ]
            .unique()
            .tolist()
        )


        raise ValueError(
            "Ungültige Werte in Standortverbund_Entscheidung:\n"
            f"{values}"
        )


    merge_mask = (
            site_decisions
            ==
            "ZUSAMMENFASSEN"
    )


    missing_group = (
            merge_mask
            &
            site_review[
                "Standortverbund_Gruppe"
            ]
            .map(
                clean_string
            )
            .eq("")
    )


    if missing_group.any():

        raise ValueError(
            "Bei mindestens einem ZUSAMMENFASSEN-Fall "
            "fehlt Standortverbund_Gruppe."
        )


# ============================================================
# 18. MATCHING-PRÜFBLATT
# ============================================================

def build_matching_review(
        results,
):

    review = results[
        results[
            "Match_Status"
        ]
        != "MATCH"
        ].copy()


    review[
        "Adresse"
    ] = review.apply(
        lambda row:
        make_address(
            row[
                "Strasse_Hausnummer"
            ],
            row[
                "PLZ"
            ],
            row[
                "Ort"
            ],
        ),
        axis=1,
    )


    review[
        "Kandidat_Referenz"
    ] = review.apply(
        lambda row:
        combine_reference_description(
            row[
                "Kandidat_CompanyName"
            ],
            row[
                "Kandidat_LocationName"
            ],
            row[
                "Kandidat_Address"
            ],
        ),
        axis=1,
    )


    return review[
        [
            "Klassifikations_ID",

            "Firmenname",

            "Standortname",

            "Adresse",

            "Match_Status",

            "Match_Methode",

            "Match_Score",

            "Kandidat_USER_UniqueID",

            "Kandidat_Referenz",
        ]
    ].copy()


# ============================================================
# 19. ZUSAMMENFASSUNG
# ============================================================

def build_summary(
        results,
        reference_coverage,
        batch_files,
        wz_review,
        site_review,
):

    total_sites = len(
        results
    )


    wz_group_count = int(

        results[
            "WZ_Gruppe_Code"
        ]

        .map(
            clean_string
        )

        .ne("")

        .sum()
    )


    matched_new = int(

        (
                results[
                    "Match_Status"
                ]
                == "MATCH"
        ).sum()
    )


    total_reference = len(
        reference_coverage
    )


    matched_reference = int(

        reference_coverage[
            "Sicher_wiedergefunden"
        ].sum()
    )


    rows = [

        {
            "Kennzahl":
                "Klassifizierte Batchdateien",

            "Wert":
                len(
                    batch_files
                ),
        },

        {
            "Kennzahl":
                "Standorte gesamt",

            "Wert":
                total_sites,
        },

        {
            "Kennzahl":
                "Standorte mit sicherer WZ-Gruppe",

            "Wert":
                wz_group_count,
        },

        {
            "Kennzahl":
                "Anteil Standorte mit sicherer WZ-Gruppe",

            "Wert":
                (
                    wz_group_count
                    / total_sites
                    if total_sites
                    else np.nan
                ),
        },

        {
            "Kennzahl":
                "Neue Standorte sicher mit Referenz gematcht",

            "Wert":
                matched_new,
        },

        {
            "Kennzahl":
                "Anteil neue Standorte mit Referenzmatch",

            "Wert":
                (
                    matched_new
                    / total_sites
                    if total_sites
                    else np.nan
                ),
        },

        # ====================================================
        # NEUE WICHTIGE KPI
        # ====================================================

        {
            "Kennzahl":
                "Referenzstandorte gesamt",

            "Wert":
                total_reference,
        },

        {
            "Kennzahl":
                "Referenzstandorte sicher wiedergefunden",

            "Wert":
                matched_reference,
        },

        {
            "Kennzahl":
                "Referenz-Abdeckung",

            "Wert":
                (
                    matched_reference
                    / total_reference
                    if total_reference
                    else np.nan
                ),
        },

        {
            "Kennzahl":
                "Referenzstandorte nicht sicher wiedergefunden",

            "Wert":
                (
                        total_reference
                        - matched_reference
                ),
        },

        {
            "Kennzahl":
                "WZ-Fälle zur manuellen Prüfung",

            "Wert":
                len(
                    wz_review
                ),
        },

        {
            "Kennzahl":
                "WZ-Prüfung noch offen",

            "Wert":
                int(
                    (
                        wz_review[
                            "WZ_Entscheidung"
                        ]
                        .map(
                            clean_string
                        )
                        .isin(
                            {
                                "",
                                "OFFEN",
                            }
                        )
                    ).sum()
                ),
        },

        {
            "Kennzahl":
                "Standortverbund-Kandidatenzeilen",

            "Wert":
                len(
                    site_review
                ),
        },

        {
            "Kennzahl":
                "Kandidatenverbünde",

            "Wert":
                site_review[
                    "Kandidatenverbund_ID"
                ].nunique(),
        },

        {
            "Kennzahl":
                "Standortverbund-Prüfung noch offen",

            "Wert":
                int(
                    (
                        site_review[
                            "Standortverbund_Entscheidung"
                        ]
                        .map(
                            clean_string
                        )
                        .isin(
                            {
                                "",
                                "OFFEN",
                            }
                        )
                    ).sum()
                ),
        },
    ]


    # ========================================================
    # VALIDIERUNGSSTATUS
    # ========================================================

    for status in [
        "OK_STRIKT",
        "OK_PLAUSIBEL",
        "OK_NUR_SEKTOR",
        "STANDORTFUNKTION_ABWEICHEND",
        "PRUEFEN_BRANCHE",
        "WIDERSPRUCH",
        "WZ_ZU_GROB_ODER_FEHLT",
        "NICHT_VERGLEICHBAR",
    ]:

        rows.append(
            {
                "Kennzahl":
                    f"Validierung: {status}",

                "Wert":
                    int(
                        (
                                results[
                                    "Gesamt_Check"
                                ]
                                == status
                        ).sum()
                    ),
            }
        )


    return pd.DataFrame(
        rows
    )


# ============================================================
# 20. BATCHÜBERSICHT
# ============================================================

def build_batch_summary(
        results,
):

    rows = []


    for batch_id, group in results.groupby(
            "Batch_ID",
            sort=True,
    ):

        rows.append(
            {

                "Batch_ID":
                    batch_id,

                "Standorte":
                    len(
                        group
                    ),

                "MATCH":
                    int(
                        (
                                group[
                                    "Match_Status"
                                ]
                                == "MATCH"
                        ).sum()
                    ),

                "WIDERSPRUCH":
                    int(
                        (
                                group[
                                    "Gesamt_Check"
                                ]
                                == "WIDERSPRUCH"
                        ).sum()
                    ),

                "PRUEFEN_BRANCHE":
                    int(
                        (
                                group[
                                    "Gesamt_Check"
                                ]
                                == "PRUEFEN_BRANCHE"
                        ).sum()
                    ),

                "STANDORTFUNKTION_ABWEICHEND":
                    int(
                        (
                                group[
                                    "Gesamt_Check"
                                ]
                                ==
                                "STANDORTFUNKTION_ABWEICHEND"
                        ).sum()
                    ),
            }
        )


    return pd.DataFrame(
        rows
    )


# ============================================================
# 21. ANLEITUNGSBLATT
# ============================================================

def build_instructions():

    rows = [

        {
            "Bereich":
                "WZ_Pruefung",

            "Feld":
                "WZ_Entscheidung",

            "Erlaubte_Eingabe":
                "KI_BESTAETIGT",

            "Bedeutung":
                "Die KI-WZ-Klassifikation ist trotz "
                "Validierungsauffälligkeit korrekt.",
        },

        {
            "Bereich":
                "WZ_Pruefung",

            "Feld":
                "WZ_Entscheidung",

            "Erlaubte_Eingabe":
                "WZ_KORRIGIEREN",

            "Bedeutung":
                "KI-WZ ist falsch. WZ_Manuell_Code und "
                "WZ_Manuell_Ebene ausfüllen.",
        },

        {
            "Bereich":
                "WZ_Pruefung",

            "Feld":
                "WZ_Entscheidung",

            "Erlaubte_Eingabe":
                "NUR_WZ2_SICHER",

            "Bedeutung":
                "3-Steller ist nicht belastbar. Sicheren "
                "zweistelligen Code in WZ_Manuell_Code eintragen "
                "und Ebene = Abteilung.",
        },

        {
            "Bereich":
                "WZ_Pruefung",

            "Feld":
                "WZ_Entscheidung",

            "Erlaubte_Eingabe":
                "NICHT_SICHER_KLASSIFIZIERBAR",

            "Bedeutung":
                "Auch nach manueller Prüfung keine ausreichend "
                "sichere WZ-Zuordnung.",
        },

        {
            "Bereich":
                "WZ_Pruefung",

            "Feld":
                "WZ_Entscheidung",

            "Erlaubte_Eingabe":
                "OFFEN",

            "Bedeutung":
                "Noch nicht entschieden.",
        },

        {
            "Bereich":
                "Standortverbund_Pruefung",

            "Feld":
                "Standortverbund_Entscheidung",

            "Erlaubte_Eingabe":
                "ZUSAMMENFASSEN",

            "Bedeutung":
                "Dieser Standort gehört mit allen Zeilen gleicher "
                "Standortverbund_Gruppe zu einem gemeinsamen "
                "physischen Standort.",
        },

        {
            "Bereich":
                "Standortverbund_Pruefung",

            "Feld":
                "Standortverbund_Entscheidung",

            "Erlaubte_Eingabe":
                "GETRENNT_LASSEN",

            "Bedeutung":
                "Standort bleibt eigenständig.",
        },

        {
            "Bereich":
                "Standortverbund_Pruefung",

            "Feld":
                "Standortverbund_Gruppe",

            "Erlaubte_Eingabe":
                "Vorgeschlagene KAND_...-ID",

            "Bedeutung":
                "Bei vollständiger Zusammenfassung nichts ändern. "
                "Wenn ein Kandidatenverbund geteilt werden muss, "
                "z. B. KAND_ABC_A und KAND_ABC_B verwenden.",
        },

        {
            "Bereich":
                "Allgemein",

            "Feld":
                "Begründung",

            "Erlaubte_Eingabe":
                "Freitext",

            "Bedeutung":
                "Kurze manuelle Begründung reicht, z. B. "
                "'gleiches Werksgelände' oder "
                "'Standort produziert Spezialchemikalien'.",
        },

        {
            "Bereich":
                "Allgemein",

            "Feld":
                "Wichtig",

            "Erlaubte_Eingabe":
                "",

            "Bedeutung":
                "Bestehende Daten- und ID-Spalten nicht verändern. "
                "Nur die vorgesehenen manuellen Spalten bearbeiten.",
        },
    ]


    return pd.DataFrame(
        rows
    )


# ============================================================
# 22. EXCEL SCHREIBEN
# ============================================================

def write_workbook(
        summary,
        wz_review,
        site_review,
        matching_review,
        reference_coverage,
        batch_summary,
        results,
):

    reference_missing = (
        reference_coverage[
            ~reference_coverage[
                "Sicher_wiedergefunden"
            ]
        ]
        .copy()
    )


    instructions = (
        build_instructions()
    )


    with pd.ExcelWriter(
            OUTPUT_FILE,
            engine="openpyxl",
    ) as writer:


        instructions.to_excel(
            writer,
            sheet_name="Anleitung",
            index=False,
        )


        summary.to_excel(
            writer,
            sheet_name="Zusammenfassung",
            index=False,
        )


        # ====================================================
        # DIE BEIDEN WICHTIGSTEN ARBEITSBLÄTTER
        # ====================================================

        wz_review.to_excel(
            writer,
            sheet_name="WZ_Pruefung",
            index=False,
        )


        site_review.to_excel(
            writer,
            sheet_name="Standortverbund_Pruefung",
            index=False,
        )


        # ====================================================
        # KOMPAKTE WEITERE PRÜFTABELLEN
        # ====================================================

        matching_review.to_excel(
            writer,
            sheet_name="Matching_Pruefung",
            index=False,
        )


        reference_coverage.to_excel(
            writer,
            sheet_name="Referenzabdeckung",
            index=False,
        )


        reference_missing.to_excel(
            writer,
            sheet_name="Referenz_nicht_gematcht",
            index=False,
        )


        batch_summary.to_excel(
            writer,
            sheet_name="Nach_Batch",
            index=False,
        )


        # ====================================================
        # VOLLSTÄNDIGES AUDITBLATT
        #
        # Nicht manuell bearbeiten.
        # ====================================================

        results.to_excel(
            writer,
            sheet_name="Alle_Ergebnisse",
            index=False,
        )


# ============================================================
# 23. EXCEL FORMATIERUNG UND DROPDOWNS
# ============================================================

def format_workbook():

    workbook = load_workbook(
        OUTPUT_FILE
    )


    normal_header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )


    manual_header_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )


    manual_columns_by_sheet = {

        "WZ_Pruefung": {
            "WZ_Entscheidung",
            "WZ_Manuell_Code",
            "WZ_Manuell_Ebene",
            "WZ_Manuell_Begruendung",
        },

        "Standortverbund_Pruefung": {
            "Standortverbund_Entscheidung",
            "Standortverbund_Gruppe",
            "Standortverbund_Begruendung",
        },
    }


    for worksheet in workbook.worksheets:

        worksheet.freeze_panes = "A2"

        if (
                worksheet.max_row >= 1
                and
                worksheet.max_column >= 1
        ):

            worksheet.auto_filter.ref = (
                worksheet.dimensions
            )


        headers = {}


        for cell in worksheet[1]:

            header = clean_string(
                cell.value
            )


            headers[
                header
            ] = cell.column


            cell.font = Font(
                bold=True
            )


            if (
                    worksheet.title
                    in manual_columns_by_sheet
                    and
                    header
                    in manual_columns_by_sheet[
                worksheet.title
            ]
            ):

                cell.fill = (
                    manual_header_fill
                )

            else:

                cell.fill = (
                    normal_header_fill
                )


            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


        # ====================================================
        # SPALTENBREITEN
        # ====================================================

        for column_cells in worksheet.columns:

            column_letter = (
                column_cells[0]
                .column_letter
            )


            header = clean_string(
                column_cells[0].value
            )


            # Lange Freitextspalten
            if header in {
                "WZ_Begruendung",
                "WZ_Manuell_Begruendung",
                "Standortverbund_Begruendung",
                "Bedeutung",
            }:

                width = 45

            elif header in {
                "Referenz_Standort",
                "Adresse",
                "Kandidaten_Grund",
                "Recherchequelle_1",
                "Recherchequelle_2",
            }:

                width = 35

            else:

                max_length = 0


                for cell in list(
                        column_cells
                )[:200]:

                    if cell.value is None:
                        continue

                    max_length = max(
                        max_length,
                        len(
                            str(
                                cell.value
                            )
                        ),
                    )


                width = min(
                    max(
                        max_length + 2,
                        10,
                        ),
                    28,
                )


            worksheet.column_dimensions[
                column_letter
            ].width = width


        # Alle Zellen oben ausrichten
        for row in worksheet.iter_rows(
                min_row=2
        ):

            for cell in row:

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )


    # ========================================================
    # DROPDOWN:
    # WZ_Entscheidung
    # ========================================================

    if "WZ_Pruefung" in workbook.sheetnames:

        ws = workbook[
            "WZ_Pruefung"
        ]


        headers = {
            cell.value:
                cell.column_letter
            for cell in ws[1]
        }


        if "WZ_Entscheidung" in headers:

            validation = DataValidation(
                type="list",
                formula1=(
                    '"OFFEN,KI_BESTAETIGT,'
                    'WZ_KORRIGIEREN,NUR_WZ2_SICHER,'
                    'NICHT_SICHER_KLASSIFIZIERBAR"'
                ),
                allow_blank=True,
            )


            ws.add_data_validation(
                validation
            )


            column = headers[
                "WZ_Entscheidung"
            ]


            validation.add(
                f"{column}2:{column}{max(ws.max_row, 2)}"
            )


        if "WZ_Manuell_Ebene" in headers:

            validation = DataValidation(
                type="list",
                formula1=(
                    '"Gruppe,Abteilung,Abschnitt"'
                ),
                allow_blank=True,
            )


            ws.add_data_validation(
                validation
            )


            column = headers[
                "WZ_Manuell_Ebene"
            ]


            validation.add(
                f"{column}2:{column}{max(ws.max_row, 2)}"
            )


    # ========================================================
    # DROPDOWN:
    # STANDORTVERBUND
    # ========================================================

    if (
            "Standortverbund_Pruefung"
            in workbook.sheetnames
    ):

        ws = workbook[
            "Standortverbund_Pruefung"
        ]


        headers = {
            cell.value:
                cell.column_letter
            for cell in ws[1]
        }


        if (
                "Standortverbund_Entscheidung"
                in headers
        ):

            validation = DataValidation(
                type="list",
                formula1=(
                    '"OFFEN,ZUSAMMENFASSEN,GETRENNT_LASSEN"'
                ),
                allow_blank=True,
            )


            ws.add_data_validation(
                validation
            )


            column = headers[
                "Standortverbund_Entscheidung"
            ]


            validation.add(
                f"{column}2:{column}{max(ws.max_row, 2)}"
            )


    workbook.save(
        OUTPUT_FILE
    )


# ============================================================
# 24. HAUPTPROGRAMM
# ============================================================

def main():

    print()
    print("=" * 78)
    print(
        "WZ2025 VALIDIERUNG V4"
    )
    print("=" * 78)
    print()


    # ========================================================
    # Vorhandene manuelle Eingaben sichern
    # ========================================================

    print(
        "0. Vorhandene manuelle Entscheidungen werden eingelesen ..."
    )


    (
        old_wz_review,
        old_site_review,
    ) = load_existing_manual_reviews()


    # ========================================================
    # Batches
    # ========================================================

    print()
    print(
        "1. WZ-Batches werden eingelesen ..."
    )


    (
        sites,
        batch_files,
    ) = load_all_batches()


    print(
        f"   -> {len(batch_files):,} Batchdateien"
    )


    print(
        f"   -> {len(sites):,} Standorte"
    )


    # ========================================================
    # Referenz
    # ========================================================

    print()
    print(
        "2. Referenzdatei wird eingelesen ..."
    )


    reference = (
        prepare_reference()
    )


    print(
        f"   -> {len(reference):,} Referenzstandorte"
    )


    # ========================================================
    # Matching
    # ========================================================

    print()
    print(
        "3. Matching gegen Referenz ..."
    )


    results = (
        match_sites_to_reference(
            sites,
            reference,
        )
    )


    # ========================================================
    # WZ-Validierung
    # ========================================================

    print()
    print(
        "4. WZ-Validierung ..."
    )


    results = (
        add_validation_results(
            results
        )
    )


    # ========================================================
    # Referenzabdeckung
    # ========================================================

    print()
    print(
        "5. Referenzabdeckung aus Sicht des alten "
        "Datensatzes ..."
    )


    reference_coverage = (
        build_reference_coverage(
            reference,
            results,
        )
    )


    total_reference = len(
        reference_coverage
    )


    matched_reference = int(

        reference_coverage[
            "Sicher_wiedergefunden"
        ].sum()
    )


    print(
        f"   -> {matched_reference:,} / "
        f"{total_reference:,} Referenzstandorte "
        "sicher wiedergefunden"
    )


    print(
        f"   -> {matched_reference / total_reference:.2%}"
    )


    # ========================================================
    # WZ-Prüfblatt
    # ========================================================

    print()
    print(
        "6. Kompaktes WZ-Prüfblatt wird aufgebaut ..."
    )


    wz_review = (
        build_wz_review(
            results
        )
    )


    wz_review = restore_manual_columns(

        wz_review,

        old_wz_review,

        key_column=
        "Klassifikations_ID",

        manual_columns=[
            "WZ_Entscheidung",
            "WZ_Manuell_Code",
            "WZ_Manuell_Ebene",
            "WZ_Manuell_Begruendung",
        ],
    )


    print(
        f"   -> {len(wz_review):,} zu prüfende WZ-Fälle"
    )


    # ========================================================
    # Standortverbund-Kandidaten
    # ========================================================

    print()
    print(
        "7. Standortverbund-Kandidaten werden gesucht ..."
    )


    site_review = (
        build_site_union_review(
            results
        )
    )


    site_review = restore_manual_columns(

        site_review,

        old_site_review,

        key_column=
        "Klassifikations_ID",

        manual_columns=[
            "Standortverbund_Entscheidung",
            "Standortverbund_Gruppe",
            "Standortverbund_Begruendung",
        ],
    )


    print(
        f"   -> {len(site_review):,} Standortzeilen"
    )


    print(
        f"   -> "
        f"{site_review['Kandidatenverbund_ID'].nunique():,} "
        "mögliche Standortverbünde"
    )


    # ========================================================
    # Bereits vorhandene manuelle Eingaben prüfen
    # ========================================================

    print()
    print(
        "8. Manuelle Eingaben werden auf gültige Werte geprüft ..."
    )


    validate_manual_inputs(
        wz_review,
        site_review,
    )


    # ========================================================
    # Weitere kompakte Tabellen
    # ========================================================

    matching_review = (
        build_matching_review(
            results
        )
    )


    batch_summary = (
        build_batch_summary(
            results
        )
    )


    summary = (
        build_summary(
            results,
            reference_coverage,
            batch_files,
            wz_review,
            site_review,
        )
    )


    # ========================================================
    # Speichern
    # ========================================================

    print()
    print(
        "9. Validierungs- und Arbeitsdatei wird geschrieben ..."
    )


    write_workbook(
        summary,
        wz_review,
        site_review,
        matching_review,
        reference_coverage,
        batch_summary,
        results,
    )


    format_workbook()


    # ========================================================
    # Abschluss
    # ========================================================

    print()
    print("=" * 78)
    print("FERTIG")
    print("=" * 78)

    print()

    print(
        f"Ausgabedatei:\n"
        f"{OUTPUT_FILE.resolve()}"
    )

    print()

    print(
        "Für deine manuelle Bearbeitung sind vor allem "
        "diese beiden Blätter relevant:"
    )

    print(
        "   1. WZ_Pruefung"
    )

    print(
        "   2. Standortverbund_Pruefung"
    )

    print()

    print(
        "Gelb markierte Spalten sind für deine "
        "manuellen Eingaben vorgesehen."
    )

    print()

    print(
        "Es wurden noch KEINE WZ-Codes korrigiert und "
        "KEINE Standorte zusammengeführt."
    )

    print("=" * 78)


# ============================================================
# 25. PROGRAMMSTART
# ============================================================

if __name__ == "__main__":

    main()