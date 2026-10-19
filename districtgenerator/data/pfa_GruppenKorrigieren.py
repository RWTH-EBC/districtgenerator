from pathlib import Path
import re
import shutil

import numpy as np
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


# ============================================================
# 1. KONFIGURATION
# ============================================================
#
# DIESER CODE IST EIN NEUER, SEPARATER SCHRITT.
#
# Er ersetzt NICHT die Gruppierungs- oder Exportcodes.
#
# Eingang:
#
#     Profilgruppen_Export_Waermemenge20_vorlaeufig/
#         Profilgruppe_G01.xlsx
#         Profilgruppe_G02.xlsx
#         ...
#
# Ausgang:
#
#     Profilgruppen_Pruefung_Waermemenge/
#         Pruefung_G01.xlsx
#         Pruefung_G02.xlsx
#         ...
#
# Pro Gruppe wird NUR dann eine Prüfdatei erzeugt, wenn es dort
# mindestens einen Eintrag im Blatt "Zahlen_Plausibilitaet" gibt.
#
# Ziel:
#
#   - nur auffällige Potentiale von Hand prüfen,
#   - trotzdem Vergleich mit ALLEN anderen Potentialen desselben
#     physischen Standortes ermöglichen,
#   - Originalwerte niemals überschreiben,
#   - Korrekturen / Ausschlüsse eindeutig und maschinenlesbar
#     dokumentieren,
#   - später aus:
#
#         ursprünglicher Profilgruppe_Gxx.xlsx
#         +
#         bearbeiteter Pruefung_Gxx.xlsx
#
#     die endgültigen Generator-Eingangswerte erzeugen können.
# ============================================================


BASE_DIR = Path("WZ2025_Vorbereitung")

SOURCE_DIR = (
        BASE_DIR
        / "Quellenklassifikation_final"
)

GROUP_DIR = (
        SOURCE_DIR
        / "Profilgruppen_Export_Waermemenge20_vorlaeufig"
)

OUTPUT_DIR = (
        SOURCE_DIR
        / "Profilgruppen_Pruefung_Waermemenge"
)

GROUP_PATTERN = "Profilgruppe_G*.xlsx"

CLEAN_OLD_REVIEW_FILES = True


# ============================================================
# 2. BLATTNAMEN
# ============================================================

SHEET_POTENTIALS = "Abwaermepotentiale"
SHEET_SITES = "Standortprofile"
SHEET_PLAUSIBILITY = "Zahlen_Plausibilitaet"

OUTPUT_SHEET_REVIEW = "Prueffaelle"
OUTPUT_SHEET_SITE_CONTEXT = "Standortvergleich"
OUTPUT_SHEET_GUIDE = "Anleitung"


# ============================================================
# 3. MANUELLE ENTSCHEIDUNGEN
# ============================================================
#
# Diese Werte werden später im finalen Postprocessing eindeutig
# interpretiert:
#
# OFFEN
#     Prüfung noch nicht abgeschlossen.
#
# BESTAETIGT
#     Originalwerte werden unverändert übernommen.
#
# KORRIGIERT
#     Nichtleere Korrekturfelder ersetzen später die jeweiligen
#     Originalwerte. Alle übrigen Originalwerte bleiben bestehen.
#
# AUSSCHLIESSEN
#     Dieses einzelne Abwärmepotential wird vollständig aus der
#     späteren Profil-/Generatorbildung ausgeschlossen.
#
# ============================================================

DECISION_VALUES = [
    "OFFEN",
    "BESTAETIGT",
    "KORRIGIERT",
    "AUSSCHLIESSEN",
]


# ============================================================
# 4. BEGRÜNDUNGSGRUNDLAGEN
# ============================================================
#
# Mehrfachauswahl ist in normaler Excel-Datenvalidierung nicht
# komfortabel möglich. Deshalb wird eine Hauptgrundlage gewählt;
# Details kommen in "Pruefbegruendung".
# ============================================================

BASIS_VALUES = [
    "FREITEXT_IM_DATENSATZ",
    "VERGLEICH_AM_GLEICHEN_STANDORT",
    "EINHEITEN_PLAUSIBILITAET",
    "TECHNISCHE_UNTERLAGE",
    "DUPLIKAT_PRUEFUNG",
    "SONSTIGES",
]


# ============================================================
# 5. ALLGEMEINE HILFSFUNKTIONEN
# ============================================================

def clean_string(value) -> str:
    if pd.isna(value):
        return ""

    text = str(value).strip()

    if text.lower() in {"nan", "none"}:
        return ""

    return text


def numeric_value(value):
    """
    Robuste Umwandlung deutscher/englischer Zahlenformate.
    Keine automatische Korrektur.
    """

    if pd.isna(value):
        return np.nan

    if isinstance(
            value,
            (
                    int,
                    float,
                    np.integer,
                    np.floating,
            ),
    ):
        return float(value)

    text = clean_string(value)

    if not text:
        return np.nan

    text = text.replace(" ", "")

    # Beispiel: 1.234,56
    if "," in text and "." in text:

        if text.rfind(",") > text.rfind("."):

            text = (
                text
                .replace(".", "")
                .replace(",", ".")
            )

    # Beispiel: 1234,56
    elif "," in text:

        text = text.replace(",", ".")

    try:
        return float(text)

    except ValueError:
        return np.nan


def numeric_series(series: pd.Series) -> pd.Series:
    return series.map(numeric_value)


def first_existing_column(
        df,
        candidates,
        required=False,
        label="Spalte",
):
    for column in candidates:

        if column in df.columns:
            return column

    if required:

        raise ValueError(
            f"{label} wurde nicht gefunden.\n\n"
            "Gesuchte Kandidaten:\n"
            + "\n".join(
                f"- {column}"
                for column in candidates
            )
        )

    return None


def first_existing_value(
        row,
        candidates,
):
    """
    Liest aus einer pandas-Series den ersten vorhandenen,
    nichtleeren Wert.
    """

    for column in candidates:

        if column in row.index:

            value = row[column]

            if clean_string(value):
                return value

    return ""


def safe_divide(
        numerator,
        denominator,
):
    if (
            pd.isna(numerator)
            or
            pd.isna(denominator)
            or
            denominator == 0
    ):
        return np.nan

    return numerator / denominator


def group_id_from_filename(path: Path) -> str:
    """
    Profilgruppe_G39.xlsx -> G39
    """

    match = re.search(
        r"Profilgruppe_(G[^.]+)",
        path.name,
        flags=re.IGNORECASE,
    )

    if match:
        return match.group(1).upper()

    return path.stem


# ============================================================
# 6. BENÖTIGTE ORIGINALSPALTEN FINDEN
# ============================================================

def identify_potential_columns(df):
    """
    Die Exportdateien enthalten die Originalspalten weiterhin.
    Hier werden die relevanten Felder robust gefunden.
    """

    return {
        "id":
            first_existing_column(
                df,
                [
                    "Originaldatensatz_ID",
                ],
                required=True,
                label="Originaldatensatz_ID",
            ),

        "site":
            first_existing_column(
                df,
                [
                    "Standortverbund_ID",
                ],
                required=True,
                label="Standortverbund_ID",
            ),

        "company":
            first_existing_column(
                df,
                [
                    "Firmenname",
                    "Profil_Firmenname",
                ],
            ),

        "location":
            first_existing_column(
                df,
                [
                    "Standortname",
                    "Profil_Standortname",
                ],
            ),

        "name":
            first_existing_column(
                df,
                [
                    "Name des Abwärmepotentials",
                ],
                required=True,
                label="Name des Abwärmepotentials",
            ),

        "heat":
            first_existing_column(
                df,
                [
                    "Wärmemenge pro Jahr (in kWh/a)",
                    "Waermemenge_pro_Jahr_kWh_a",
                ],
                required=True,
                label="Jahreswärmemenge",
            ),

        "power":
            first_existing_column(
                df,
                [
                    "Maximale thermische Leistung (in kW)",
                    "Maximale_thermische_Leistung_kW",
                ],
            ),

        "temperature":
            first_existing_column(
                df,
                [
                    "Durchschnittliches Temperaturniveau (in °C)",
                    "Durchschnittliches_Temperaturniveau_C",
                ],
            ),

        "temp_range":
            first_existing_column(
                df,
                [
                    "Temperaturbereich",
                ],
            ),

        "source_class":
            first_existing_column(
                df,
                [
                    "Quellklasse_Code",
                ],
            ),

        "source_class_label":
            first_existing_column(
                df,
                [
                    "Quellklasse",
                ],
            ),

        "source_reason":
            first_existing_column(
                df,
                [
                    "Quellklasse_Begruendung",
                ],
            ),

        "source_confidence":
            first_existing_column(
                df,
                [
                    "Quellklasse_Konfidenz",
                ],
            ),

        "extra":
            first_existing_column(
                df,
                [
                    "Ergänzende Informationen zum Abwärmepotential",
                ],
            ),

        "notes":
            first_existing_column(
                df,
                [
                    "Weitere Hinweise",
                ],
            ),

        "profile_wz":
            first_existing_column(
                df,
                [
                    "Gruppen_Profil_WZ_Code",
                    "Profil_WZ_Final_Code",
                ],
            ),

        "wz3":
            first_existing_column(
                df,
                [
                    "Gruppen_WZ3_Code",
                    "Profil_WZ3_Code",
                ],
            ),

        "site_qx_heat":
            first_existing_column(
                df,
                [
                    "Standort_QX_Waermengenanteil",
                    "QX_Anteil_Waermemenge_bearbeitet",
                ],
            ),
    }


# ============================================================
# 7. PLAUSIBILITÄTSFELDER FINDEN
# ============================================================

def identify_plausibility_columns(df):
    """
    Das vorherige Postprocessing kann die Spalten leicht anders
    benennen. Deshalb mehrere mögliche Namen.
    """

    return {
        "id":
            first_existing_column(
                df,
                [
                    "Originaldatensatz_ID",
                ],
                required=True,
                label="Originaldatensatz_ID im Plausibilitätsblatt",
            ),

        "priority":
            first_existing_column(
                df,
                [
                    "Pruefprioritaet",
                    "Prüfpriorität",
                    "Zahlen_Pruefprioritaet",
                    "Zahlen_Plausibilitaet_Prioritaet",
                ],
            ),

        "flags":
            first_existing_column(
                df,
                [
                    "Plausibilitaetsflags",
                    "Plausibilitaets_Flags",
                    "Plausibilitaetsflag",
                    "Plausibilitaets_Flag",
                    "Zahlen_Plausibilitaetsflags",
                ],
            ),

        "full_load_hours":
            first_existing_column(
                df,
                [
                    "Aequivalente_Vollbenutzungsstunden",
                    "Äquivalente_Vollbenutzungsstunden",
                    "Vollbenutzungsstunden",
                ],
            ),

        "min_power":
            first_existing_column(
                df,
                [
                    "Min_Leistung_physikalisch_kW",
                    "Min_Leistung_bei_8784h_kW",
                ],
            ),
    }


# ============================================================
# 8. GRUPPENDATEI EINLESEN
# ============================================================

def load_group_workbook(path: Path):
    """
    Liest nur die für diesen Prüfexport benötigten Blätter.
    """

    excel = pd.ExcelFile(path)

    required_sheets = {
        SHEET_POTENTIALS,
        SHEET_PLAUSIBILITY,
    }

    missing = (
            required_sheets
            - set(excel.sheet_names)
    )

    if SHEET_PLAUSIBILITY in missing:
        # Keine Prüffälle -> Gruppe wird einfach übersprungen.
        return None

    if SHEET_POTENTIALS in missing:

        raise ValueError(
            f"{path.name}: Blatt '{SHEET_POTENTIALS}' fehlt."
        )

    potentials = pd.read_excel(
        path,
        sheet_name=SHEET_POTENTIALS,
        dtype=str,
        keep_default_na=False,
    )

    plausibility = pd.read_excel(
        path,
        sheet_name=SHEET_PLAUSIBILITY,
        dtype=str,
        keep_default_na=False,
    )

    # Leeres Plausibilitätsblatt -> kein Prüfexport nötig.
    if plausibility.empty:
        return None

    sites = pd.DataFrame()

    if SHEET_SITES in excel.sheet_names:

        sites = pd.read_excel(
            path,
            sheet_name=SHEET_SITES,
            dtype=str,
            keep_default_na=False,
        )

    return (
        potentials,
        plausibility,
        sites,
    )


# ============================================================
# 9. PLAUSIBILITÄTSINFOS AUF ID-EBENE MAPPEN
# ============================================================

def build_plausibility_map(
        plausibility,
):
    """
    Gibt pro Originaldatensatz_ID die wichtigsten bisherigen
    Diagnosefelder zurück.

    Mehrfach vorkommende Einträge werden zu einem String
    zusammengeführt, statt still einen Eintrag zu verlieren.
    """

    cols = identify_plausibility_columns(
        plausibility
    )

    id_col = cols["id"]

    result = {}

    for _, row in plausibility.iterrows():

        original_id = clean_string(
            row[id_col]
        )

        if not original_id:
            continue

        priority = ""

        if cols["priority"]:

            priority = clean_string(
                row[cols["priority"]]
            )

        flags = ""

        if cols["flags"]:

            flags = clean_string(
                row[cols["flags"]]
            )

        hours = np.nan

        if cols["full_load_hours"]:

            hours = numeric_value(
                row[cols["full_load_hours"]]
            )

        min_power = np.nan

        if cols["min_power"]:

            min_power = numeric_value(
                row[cols["min_power"]]
            )

        if original_id not in result:

            result[original_id] = {
                "priority": priority,
                "flags": flags,
                "hours": hours,
                "min_power": min_power,
            }

        else:

            old = result[original_id]

            if priority and priority not in old["priority"]:

                old["priority"] = (
                        old["priority"]
                        + " | "
                        + priority
                ).strip(" |")

            if flags and flags not in old["flags"]:

                old["flags"] = (
                        old["flags"]
                        + " | "
                        + flags
                ).strip(" |")

            if pd.isna(old["hours"]) and pd.notna(hours):
                old["hours"] = hours

            if pd.isna(old["min_power"]) and pd.notna(min_power):
                old["min_power"] = min_power

    return result


# ============================================================
# 10. STANDORT-KENNZAHLEN BILDEN
# ============================================================

def build_site_statistics(
        potentials,
        columns,
):
    """
    Vergleichswerte werden ausschließlich aus den Potentialen
    desselben physischen Standortes gebildet.

    Sie sind nur Entscheidungshilfe und keine automatischen
    Korrekturwerte.
    """

    work = potentials.copy()

    work["_Heat_num"] = numeric_series(
        work[columns["heat"]]
    )

    if columns["power"]:

        work["_Power_num"] = numeric_series(
            work[columns["power"]]
        )

    else:

        work["_Power_num"] = np.nan

    if columns["temperature"]:

        work["_Temp_num"] = numeric_series(
            work[columns["temperature"]]
        )

    else:

        work["_Temp_num"] = np.nan

    stats = {}

    for site_id, group in work.groupby(
            columns["site"],
            sort=False,
    ):

        site_key = clean_string(site_id)

        heat_values = (
            group["_Heat_num"]
            .dropna()
        )

        positive_heat_values = (
            heat_values[
                heat_values > 0
                ]
        )

        power_values = (
            group["_Power_num"]
            .dropna()
        )

        positive_power_values = (
            power_values[
                power_values > 0
                ]
        )

        stats[site_key] = {
            "n_potentials":
                int(len(group)),

            "sum_heat":
                (
                    float(heat_values.sum())
                    if len(heat_values)
                    else np.nan
                ),

            "median_heat":
                (
                    float(positive_heat_values.median())
                    if len(positive_heat_values)
                    else np.nan
                ),

            "min_heat":
                (
                    float(positive_heat_values.min())
                    if len(positive_heat_values)
                    else np.nan
                ),

            "max_heat":
                (
                    float(positive_heat_values.max())
                    if len(positive_heat_values)
                    else np.nan
                ),

            "sum_power":
                (
                    float(power_values.sum())
                    if len(power_values)
                    else np.nan
                ),

            "median_power":
                (
                    float(positive_power_values.median())
                    if len(positive_power_values)
                    else np.nan
                ),

            "min_power":
                (
                    float(positive_power_values.min())
                    if len(positive_power_values)
                    else np.nan
                ),

            "max_power":
                (
                    float(positive_power_values.max())
                    if len(positive_power_values)
                    else np.nan
                ),
        }

    return stats


# ============================================================
# 11. PRÜFFALL-TABELLE ERSTELLEN
# ============================================================

def build_review_table(
        group_id,
        potentials,
        plausibility,
):
    """
    Eine Zeile je zu prüfendem Potential.

    Originalwerte werden bewusst separat von manuellen
    Korrekturfeldern geführt.
    """

    pcols = identify_potential_columns(
        potentials
    )

    plaus_map = build_plausibility_map(
        plausibility
    )

    flagged_ids = set(
        plaus_map.keys()
    )

    site_stats = build_site_statistics(
        potentials,
        pcols,
    )

    rows = []

    for _, source in potentials.iterrows():

        original_id = clean_string(
            source[pcols["id"]]
        )

        if original_id not in flagged_ids:
            continue

        site_id = clean_string(
            source[pcols["site"]]
        )

        company = (
            clean_string(source[pcols["company"]])
            if pcols["company"]
            else ""
        )

        location = (
            clean_string(source[pcols["location"]])
            if pcols["location"]
            else ""
        )

        name = clean_string(
            source[pcols["name"]]
        )

        heat = numeric_value(
            source[pcols["heat"]]
        )

        power = (
            numeric_value(source[pcols["power"]])
            if pcols["power"]
            else np.nan
        )

        temperature = (
            numeric_value(source[pcols["temperature"]])
            if pcols["temperature"]
            else np.nan
        )

        temp_range = (
            clean_string(source[pcols["temp_range"]])
            if pcols["temp_range"]
            else ""
        )

        source_class = (
            clean_string(source[pcols["source_class"]])
            if pcols["source_class"]
            else ""
        )

        source_class_label = (
            clean_string(source[pcols["source_class_label"]])
            if pcols["source_class_label"]
            else ""
        )

        extra = (
            clean_string(source[pcols["extra"]])
            if pcols["extra"]
            else ""
        )

        notes = (
            clean_string(source[pcols["notes"]])
            if pcols["notes"]
            else ""
        )

        profile_wz = (
            clean_string(source[pcols["profile_wz"]])
            if pcols["profile_wz"]
            else ""
        )

        wz3 = (
            clean_string(source[pcols["wz3"]])
            if pcols["wz3"]
            else ""
        )

        site_qx_heat = (
            numeric_value(source[pcols["site_qx_heat"]])
            if pcols["site_qx_heat"]
            else np.nan
        )

        diagnosis = plaus_map[
            original_id
        ]

        hours = safe_divide(
            heat,
            power,
        )

        # Falls das alte Plausibilitätsblatt bereits einen Wert
        # enthält, diesen verwenden.
        if (
                pd.isna(hours)
                and
                pd.notna(diagnosis["hours"])
        ):
            hours = diagnosis["hours"]

        min_power_8784 = (
            safe_divide(
                heat,
                8784.0,
            )
            if pd.notna(heat)
            else np.nan
        )

        # Hilfswerte für sehr typische Einheitenfehler:
        #
        # Niedrige h:
        #   Was wäre, wenn Wärmemenge eigentlich MWh war?
        #
        # Hohe h:
        #   Was wäre, wenn Wärmemenge durch 1000 kleiner wäre?
        #
        # Diese Werte sind NUR Hinweise.
        hours_if_heat_x1000 = (
            hours * 1000
            if pd.notna(hours)
            else np.nan
        )

        hours_if_heat_div1000 = (
            hours / 1000
            if pd.notna(hours)
            else np.nan
        )

        stats = site_stats.get(
            site_id,
            {}
        )

        rows.append(
            {
                # --------------------------------------------
                # Identifikation
                # --------------------------------------------
                "Experiment_Gruppe":
                    group_id,

                "Originaldatensatz_ID":
                    original_id,

                "Standortverbund_ID":
                    site_id,

                "Firmenname":
                    company,

                "Standortname":
                    location,

                "Profil_WZ_Code":
                    profile_wz,

                "WZ3_Code":
                    wz3,

                "Name_des_Abwaermepotentials":
                    name,

                "Quellklasse_Code":
                    source_class,

                "Quellklasse":
                    source_class_label,

                # --------------------------------------------
                # Originalwerte
                # --------------------------------------------
                "Waermemenge_original_kWh_a":
                    heat,

                "Leistung_original_kW":
                    power,

                "Temperatur_original_C":
                    temperature,

                "Temperaturbereich":
                    temp_range,

                # --------------------------------------------
                # Harte/naheliegende Plausibilitätsdiagnose
                # --------------------------------------------
                "Aequivalente_Vollbenutzungsstunden":
                    hours,

                "Min_Leistung_bei_8784h_kW":
                    min_power_8784,

                "Vollbenutzungsstunden_wenn_Waermemenge_x1000":
                    hours_if_heat_x1000,

                "Vollbenutzungsstunden_wenn_Waermemenge_durch1000":
                    hours_if_heat_div1000,

                "Pruefprioritaet":
                    diagnosis["priority"],

                "Plausibilitaetsflags":
                    diagnosis["flags"],

                # --------------------------------------------
                # Vergleich innerhalb desselben Standortes
                # --------------------------------------------
                "Standort_QX_Waermengenanteil_vorlaeufig":
                    site_qx_heat,

                "Standort_Anzahl_Potentiale":
                    stats.get(
                        "n_potentials",
                        np.nan,
                    ),

                "Standort_Summe_Waermemenge_kWh_a":
                    stats.get(
                        "sum_heat",
                        np.nan,
                    ),

                "Standort_Median_Waermemenge_kWh_a":
                    stats.get(
                        "median_heat",
                        np.nan,
                    ),

                "Standort_Min_Waermemenge_kWh_a":
                    stats.get(
                        "min_heat",
                        np.nan,
                    ),

                "Standort_Max_Waermemenge_kWh_a":
                    stats.get(
                        "max_heat",
                        np.nan,
                    ),

                "Standort_Median_Leistung_kW":
                    stats.get(
                        "median_power",
                        np.nan,
                    ),

                # --------------------------------------------
                # Freitext zur fachlichen Beurteilung
                # --------------------------------------------
                "Ergaenzende_Informationen":
                    extra,

                "Weitere_Hinweise":
                    notes,

                # --------------------------------------------
                # MANUELLE FELDER
                # --------------------------------------------
                "Pruefstatus":
                    "OFFEN",

                "Begruendungsgrundlage":
                    "",

                "Korrigierte_Waermemenge_kWh_a":
                    "",

                "Korrigierte_Leistung_kW":
                    "",

                "Korrigierte_Temperatur_C":
                    "",

                "Ausschlussgrund":
                    "",

                "Pruefbegruendung":
                    "",
            }
        )

    review = pd.DataFrame(
        rows
    )

    if review.empty:
        return review, pcols, flagged_ids

    # Hohe Priorität zuerst, dann Standort.
    priority_rank = {
        "hoch": 0,
        "high": 0,
        "mittel": 1,
        "medium": 1,
        "gering": 2,
        "low": 2,
        "": 3,
    }

    review["_priority_sort"] = (
        review["Pruefprioritaet"]
        .map(
            lambda value:
            priority_rank.get(
                clean_string(value).casefold(),
                3,
            )
        )
    )

    review = (
        review
        .sort_values(
            [
                "_priority_sort",
                "Standortverbund_ID",
                "Originaldatensatz_ID",
            ],
            kind="stable",
        )
        .drop(
            columns=[
                "_priority_sort",
            ]
        )
        .reset_index(drop=True)
    )

    return review, pcols, flagged_ids


# ============================================================
# 12. STANDORTVERGLEICH ERSTELLEN
# ============================================================

def build_site_context_table(
        group_id,
        potentials,
        pcols,
        flagged_ids,
        review,
):
    """
    Enthält ALLE Potentiale, aber NUR von Standorten mit mindestens
    einem Prüffall.

    Damit kann der Prüfer direkt sehen, ob z. B. andere Anlagen am
    selben Standort ähnliche Größenordnungen, Temperaturen oder
    Quellklassen haben.
    """

    affected_sites = set(
        review["Standortverbund_ID"]
        .map(clean_string)
    )

    context = potentials[
        potentials[pcols["site"]]
        .map(clean_string)
        .isin(affected_sites)
    ].copy()

    rows = []

    for _, source in context.iterrows():

        original_id = clean_string(
            source[pcols["id"]]
        )

        site_id = clean_string(
            source[pcols["site"]]
        )

        heat = numeric_value(
            source[pcols["heat"]]
        )

        power = (
            numeric_value(source[pcols["power"]])
            if pcols["power"]
            else np.nan
        )

        temperature = (
            numeric_value(source[pcols["temperature"]])
            if pcols["temperature"]
            else np.nan
        )

        rows.append(
            {
                "Experiment_Gruppe":
                    group_id,

                "IST_PRUEFFALL":
                    (
                        "JA"
                        if original_id in flagged_ids
                        else "NEIN"
                    ),

                "Originaldatensatz_ID":
                    original_id,

                "Standortverbund_ID":
                    site_id,

                "Firmenname":
                    (
                        clean_string(source[pcols["company"]])
                        if pcols["company"]
                        else ""
                    ),

                "Standortname":
                    (
                        clean_string(source[pcols["location"]])
                        if pcols["location"]
                        else ""
                    ),

                "Name_des_Abwaermepotentials":
                    clean_string(
                        source[pcols["name"]]
                    ),

                "Quellklasse_Code":
                    (
                        clean_string(source[pcols["source_class"]])
                        if pcols["source_class"]
                        else ""
                    ),

                "Waermemenge_kWh_a":
                    heat,

                "Leistung_kW":
                    power,

                "Temperatur_C":
                    temperature,

                "Temperaturbereich":
                    (
                        clean_string(source[pcols["temp_range"]])
                        if pcols["temp_range"]
                        else ""
                    ),

                "Aequivalente_Vollbenutzungsstunden":
                    safe_divide(
                        heat,
                        power,
                    ),

                "Ergaenzende_Informationen":
                    (
                        clean_string(source[pcols["extra"]])
                        if pcols["extra"]
                        else ""
                    ),

                "Weitere_Hinweise":
                    (
                        clean_string(source[pcols["notes"]])
                        if pcols["notes"]
                        else ""
                    ),
            }
        )

    result = pd.DataFrame(
        rows
    )

    if result.empty:
        return result

    # Prüffälle innerhalb eines Standortes zuerst.
    result["_review_sort"] = (
        result["IST_PRUEFFALL"]
        .map(
            {
                "JA": 0,
                "NEIN": 1,
            }
        )
    )

    result = (
        result
        .sort_values(
            [
                "Standortverbund_ID",
                "_review_sort",
                "Quellklasse_Code",
                "Name_des_Abwaermepotentials",
            ],
            kind="stable",
        )
        .drop(
            columns=[
                "_review_sort",
            ]
        )
        .reset_index(drop=True)
    )

    return result


# ============================================================
# 13. ANLEITUNGSBLATT
# ============================================================

def build_guide_table():
    """
    Kurze, eindeutige Regeln für die manuelle Prüfung.
    """

    rows = [
        [
            "1",
            "Nur Prüffälle bearbeiten",
            (
                "Im Blatt 'Prueffaelle' steht genau eine Zeile je "
                "auffälligem Potential. Andere Potentiale müssen "
                "nicht bewertet werden."
            ),
        ],
        [
            "2",
            "Standortvergleich nutzen",
            (
                "Im Blatt 'Standortvergleich' stehen alle Potentiale "
                "der betroffenen Standorte. Dort Größenordnung, "
                "Temperatur, Quellklasse und Freitexte vergleichen."
            ),
        ],
        [
            "3",
            "Nicht allein aus Vollbenutzungsstunden korrigieren",
            (
                "Ein unplausibles Verhältnis aus Jahreswärmemenge und "
                "Leistung beweist zunächst nur eine Inkonsistenz. "
                "Welcher Wert falsch ist, muss aus Freitext, "
                "Einheitenlogik oder Standortvergleich begründet werden."
            ),
        ],
        [
            "4",
            "BESTAETIGT",
            (
                "Wählen, wenn der Originaldatensatz trotz Auffälligkeit "
                "plausibel bzw. fachlich nachvollziehbar ist. "
                "Korrekturfelder leer lassen."
            ),
        ],
        [
            "5",
            "KORRIGIERT",
            (
                "Nur die tatsächlich geänderten Werte in die "
                "Korrekturfelder eintragen. Nicht geänderte "
                "Korrekturfelder leer lassen."
            ),
        ],
        [
            "6",
            "AUSSCHLIESSEN",
            (
                "Wählen, wenn das Potential nicht belastbar korrigierbar "
                "ist, ein klarer unbrauchbarer Platzhalter/Duplikat "
                "vorliegt oder das Potential für die spätere "
                "Profilbildung vollständig entfallen soll."
            ),
        ],
        [
            "7",
            "Begründung dokumentieren",
            (
                "Begruendungsgrundlage auswählen und in "
                "'Pruefbegruendung' kurz die konkrete Entscheidung "
                "begründen. So bleibt die spätere Verarbeitung "
                "nachvollziehbar."
            ),
        ],
        [
            "8",
            "Spätere Verarbeitung",
            (
                "Die ursprüngliche Profilgruppe_Gxx.xlsx bleibt "
                "unverändert. Ein späterer Processing-Code verbindet "
                "sie über Originaldatensatz_ID mit dieser bearbeiteten "
                "Prüfdatei."
            ),
        ],
    ]

    return pd.DataFrame(
        rows,
        columns=[
            "Schritt",
            "Thema",
            "Hinweis",
        ],
    )


# ============================================================
# 14. EXCEL-FORMATIERUNG
# ============================================================

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)

EDIT_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC",
)

WARNING_FILL = PatternFill(
    fill_type="solid",
    fgColor="FCE4D6",
)

CONTEXT_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7",
)

GREEN_FILL = PatternFill(
    fill_type="solid",
    fgColor="E2F0D9",
)

THIN_GRAY = Side(
    style="thin",
    color="D9D9D9",
)


def format_header(
        worksheet,
):
    for cell in worksheet[1]:

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def set_reasonable_widths(
        worksheet,
        custom_widths=None,
):
    custom_widths = custom_widths or {}

    headers = {
        clean_string(cell.value):
            cell.column
        for cell
        in worksheet[1]
        if clean_string(cell.value)
    }

    for header, column in headers.items():

        if header in custom_widths:
            width = custom_widths[header]

        elif header in {
            "Ergaenzende_Informationen",
            "Weitere_Hinweise",
            "Plausibilitaetsflags",
            "Pruefbegruendung",
            "Ausschlussgrund",
        }:
            width = 42

        elif header in {
            "Name_des_Abwaermepotentials",
            "Firmenname",
            "Standortname",
        }:
            width = 30

        elif header in {
            "Standortverbund_ID",
            "Originaldatensatz_ID",
        }:
            width = 20

        else:
            width = 18

        worksheet.column_dimensions[
            get_column_letter(column)
        ].width = width


def apply_common_formatting(
        worksheet,
):
    format_header(
        worksheet
    )

    worksheet.freeze_panes = "A2"

    if worksheet.max_row >= 2:

        worksheet.auto_filter.ref = (
            f"A1:"
            f"{get_column_letter(worksheet.max_column)}"
            f"{worksheet.max_row}"
        )

    for row in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            cell.border = Border(
                bottom=THIN_GRAY,
            )

    set_reasonable_widths(
        worksheet
    )


def format_review_sheet(
        worksheet,
):
    apply_common_formatting(
        worksheet
    )

    headers = {
        clean_string(cell.value):
            cell.column
        for cell
        in worksheet[1]
        if clean_string(cell.value)
    }

    # --------------------------------------------------------
    # Manuell editierbare Spalten gelb markieren.
    # --------------------------------------------------------

    editable_headers = [
        "Pruefstatus",
        "Begruendungsgrundlage",
        "Korrigierte_Waermemenge_kWh_a",
        "Korrigierte_Leistung_kW",
        "Korrigierte_Temperatur_C",
        "Ausschlussgrund",
        "Pruefbegruendung",
    ]

    for header in editable_headers:

        if header not in headers:
            continue

        column = headers[header]

        for row in range(
                1,
                worksheet.max_row + 1,
        ):

            worksheet.cell(
                row=row,
                column=column,
            ).fill = EDIT_FILL

        worksheet.cell(
            row=1,
            column=column,
        ).font = Font(
            bold=True,
            color="000000",
        )

    # --------------------------------------------------------
    # Prüfstatus-Dropdown
    # --------------------------------------------------------

    if (
            "Pruefstatus"
            in headers
            and
            worksheet.max_row >= 2
    ):

        column = get_column_letter(
            headers["Pruefstatus"]
        )

        validation = DataValidation(
            type="list",
            formula1=(
                    '"'
                    + ",".join(DECISION_VALUES)
                    + '"'
            ),
            allow_blank=False,
        )

        worksheet.add_data_validation(
            validation
        )

        validation.add(
            f"{column}2:{column}{worksheet.max_row}"
        )

    # --------------------------------------------------------
    # Begründungsgrundlage-Dropdown
    # --------------------------------------------------------

    if (
            "Begruendungsgrundlage"
            in headers
            and
            worksheet.max_row >= 2
    ):

        column = get_column_letter(
            headers["Begruendungsgrundlage"]
        )

        validation = DataValidation(
            type="list",
            formula1=(
                    '"'
                    + ",".join(BASIS_VALUES)
                    + '"'
            ),
            allow_blank=True,
        )

        worksheet.add_data_validation(
            validation
        )

        validation.add(
            f"{column}2:{column}{worksheet.max_row}"
        )

    # --------------------------------------------------------
    # Prozentformat
    # --------------------------------------------------------

    if "Standort_QX_Waermengenanteil_vorlaeufig" in headers:

        column = headers[
            "Standort_QX_Waermengenanteil_vorlaeufig"
        ]

        for row in range(
                2,
                worksheet.max_row + 1,
        ):

            worksheet.cell(
                row=row,
                column=column,
            ).number_format = "0.0%"

    # --------------------------------------------------------
    # Zahlenformate
    # --------------------------------------------------------

    numeric_headers = [
        "Waermemenge_original_kWh_a",
        "Leistung_original_kW",
        "Temperatur_original_C",
        "Aequivalente_Vollbenutzungsstunden",
        "Min_Leistung_bei_8784h_kW",
        "Vollbenutzungsstunden_wenn_Waermemenge_x1000",
        "Vollbenutzungsstunden_wenn_Waermemenge_durch1000",
        "Standort_Summe_Waermemenge_kWh_a",
        "Standort_Median_Waermemenge_kWh_a",
        "Standort_Min_Waermemenge_kWh_a",
        "Standort_Max_Waermemenge_kWh_a",
        "Standort_Median_Leistung_kW",
        "Korrigierte_Waermemenge_kWh_a",
        "Korrigierte_Leistung_kW",
        "Korrigierte_Temperatur_C",
    ]

    for header in numeric_headers:

        if header not in headers:
            continue

        column = headers[header]

        for row in range(
                2,
                worksheet.max_row + 1,
        ):

            worksheet.cell(
                row=row,
                column=column,
            ).number_format = '#,##0.00'

    # Prüffälle generell leicht orange markieren, aber manuelle
    # Felder bleiben gelb.
    for row in range(
            2,
            worksheet.max_row + 1,
    ):

        for column in range(
                1,
                worksheet.max_column + 1,
        ):

            header = clean_string(
                worksheet.cell(
                    row=1,
                    column=column,
                ).value
            )

            if header not in editable_headers:

                worksheet.cell(
                    row=row,
                    column=column,
                ).fill = WARNING_FILL


def format_context_sheet(
        worksheet,
):
    apply_common_formatting(
        worksheet
    )

    headers = {
        clean_string(cell.value):
            cell.column
        for cell
        in worksheet[1]
        if clean_string(cell.value)
    }

    if "IST_PRUEFFALL" in headers:

        flag_col = headers[
            "IST_PRUEFFALL"
        ]

        for row in range(
                2,
                worksheet.max_row + 1,
        ):

            is_review = (
                    clean_string(
                        worksheet.cell(
                            row=row,
                            column=flag_col,
                        ).value
                    )
                    == "JA"
            )

            if is_review:

                for column in range(
                        1,
                        worksheet.max_column + 1,
                ):

                    worksheet.cell(
                        row=row,
                        column=column,
                    ).fill = WARNING_FILL

            else:

                for column in range(
                        1,
                        worksheet.max_column + 1,
                ):

                    worksheet.cell(
                        row=row,
                        column=column,
                    ).fill = CONTEXT_FILL

    numeric_headers = [
        "Waermemenge_kWh_a",
        "Leistung_kW",
        "Temperatur_C",
        "Aequivalente_Vollbenutzungsstunden",
    ]

    for header in numeric_headers:

        if header not in headers:
            continue

        column = headers[header]

        for row in range(
                2,
                worksheet.max_row + 1,
        ):

            worksheet.cell(
                row=row,
                column=column,
            ).number_format = '#,##0.00'


def format_guide_sheet(
        worksheet,
):
    apply_common_formatting(
        worksheet
    )

    worksheet.column_dimensions["A"].width = 10
    worksheet.column_dimensions["B"].width = 32
    worksheet.column_dimensions["C"].width = 100

    for row in range(
            2,
            worksheet.max_row + 1,
    ):

        worksheet.cell(
            row=row,
            column=1,
        ).fill = GREEN_FILL


# ============================================================
# 15. EINZELNE PRÜFDATEI SCHREIBEN
# ============================================================

def write_review_workbook(
        group_id,
        source_file,
        review,
        site_context,
):
    output_file = (
            OUTPUT_DIR
            / f"Pruefung_{group_id}.xlsx"
    )

    guide = build_guide_table()

    with pd.ExcelWriter(
            output_file,
            engine="openpyxl",
    ) as writer:

        review.to_excel(
            writer,
            sheet_name=OUTPUT_SHEET_REVIEW,
            index=False,
        )

        site_context.to_excel(
            writer,
            sheet_name=OUTPUT_SHEET_SITE_CONTEXT,
            index=False,
        )

        guide.to_excel(
            writer,
            sheet_name=OUTPUT_SHEET_GUIDE,
            index=False,
        )

    workbook = load_workbook(
        output_file
    )

    format_review_sheet(
        workbook[
            OUTPUT_SHEET_REVIEW
        ]
    )

    format_context_sheet(
        workbook[
            OUTPUT_SHEET_SITE_CONTEXT
        ]
    )

    format_guide_sheet(
        workbook[
            OUTPUT_SHEET_GUIDE
        ]
    )

    # --------------------------------------------------------
    # Quelle der Prüfdatei dokumentieren.
    # --------------------------------------------------------

    guide_sheet = workbook[
        OUTPUT_SHEET_GUIDE
    ]

    last_row = guide_sheet.max_row + 2

    guide_sheet.cell(
        row=last_row,
        column=1,
        value="Quelle",
    )

    guide_sheet.cell(
        row=last_row,
        column=2,
        value="Ursprüngliche Gruppendatei",
    )

    guide_sheet.cell(
        row=last_row,
        column=3,
        value=source_file.name,
    )

    # Prüfdatei hat keine rechenintensiven Formeln.
    try:
        workbook.calculation.calcMode = "manual"
        workbook.calculation.fullCalcOnLoad = False
        workbook.calculation.forceFullCalc = False

    except Exception:
        pass

    workbook.save(
        output_file
    )

    return output_file


# ============================================================
# 16. GESAMTÜBERSICHT
# ============================================================

def write_overview(
        overview_rows,
):
    if not overview_rows:
        return

    overview = pd.DataFrame(
        overview_rows
    )

    overview = overview.sort_values(
        [
            "Anzahl_Prueffaelle",
            "Experiment_Gruppe",
        ],
        ascending=[
            False,
            True,
        ],
        kind="stable",
    )

    output_file = (
            OUTPUT_DIR
            / "_Pruefuebersicht.xlsx"
    )

    with pd.ExcelWriter(
            output_file,
            engine="openpyxl",
    ) as writer:

        overview.to_excel(
            writer,
            sheet_name="Gruppen",
            index=False,
        )

    workbook = load_workbook(
        output_file
    )

    apply_common_formatting(
        workbook["Gruppen"]
    )

    workbook.save(
        output_file
    )


# ============================================================
# 17. HAUPTPROGRAMM
# ============================================================

def main():

    print()
    print("=" * 80)
    print(
        "MANUELLE ZAHLENPRÜFUNG JE PROFILGRUPPE"
    )
    print("=" * 80)
    print()

    # --------------------------------------------------------
    # 1. Eingabeordner prüfen
    # --------------------------------------------------------

    if not GROUP_DIR.exists():

        raise FileNotFoundError(
            "Gruppenordner nicht gefunden:\n"
            f"{GROUP_DIR}"
        )

    group_files = sorted(
        [
            file
            for file
            in GROUP_DIR.glob(
            GROUP_PATTERN
        )
            if not file.name.startswith("~$")
        ]
    )

    if not group_files:

        raise FileNotFoundError(
            "Keine Profilgruppe_G*.xlsx-Dateien gefunden in:\n"
            f"{GROUP_DIR}"
        )

    print(
        f"1. Gefundene Gruppendateien: {len(group_files):,}"
    )

    # --------------------------------------------------------
    # 2. Ausgabeordner
    # --------------------------------------------------------

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if CLEAN_OLD_REVIEW_FILES:

        for file in OUTPUT_DIR.glob(
                "Pruefung_G*.xlsx"
        ):

            if not file.name.startswith("~$"):
                file.unlink()

        overview_file = (
                OUTPUT_DIR
                / "_Pruefuebersicht.xlsx"
        )

        if overview_file.exists():
            overview_file.unlink()

    # --------------------------------------------------------
    # 3. Gruppen bearbeiten
    # --------------------------------------------------------

    print()
    print(
        "2. Prüfdateien werden erzeugt ..."
    )

    overview_rows = []

    groups_without_review = 0

    total_cases = 0

    total_affected_sites = 0

    for group_file in group_files:

        group_id = group_id_from_filename(
            group_file
        )

        loaded = load_group_workbook(
            group_file
        )

        if loaded is None:

            groups_without_review += 1

            print(
                f"   {group_id}: keine Prüffälle"
            )

            continue

        (
            potentials,
            plausibility,
            sites,
        ) = loaded

        (
            review,
            pcols,
            flagged_ids,
        ) = build_review_table(
            group_id,
            potentials,
            plausibility,
        )

        if review.empty:

            groups_without_review += 1

            print(
                f"   {group_id}: Plausibilitätsblatt vorhanden, "
                "aber keine eindeutig zuordenbaren Prüffälle"
            )

            continue

        site_context = build_site_context_table(
            group_id,
            potentials,
            pcols,
            flagged_ids,
            review,
        )

        output_file = write_review_workbook(
            group_id,
            group_file,
            review,
            site_context,
        )

        affected_sites = (
            review[
                "Standortverbund_ID"
            ]
            .nunique()
        )

        total_cases += len(
            review
        )

        total_affected_sites += (
            affected_sites
        )

        overview_rows.append(
            {
                "Experiment_Gruppe":
                    group_id,

                "Anzahl_Prueffaelle":
                    int(
                        len(review)
                    ),

                "Betroffene_Standorte":
                    int(
                        affected_sites
                    ),

                "Vergleichszeilen_am_Standort":
                    int(
                        len(site_context)
                    ),

                "Pruefdatei":
                    output_file.name,

                "Urspruengliche_Gruppendatei":
                    group_file.name,
            }
        )

        print(
            f"   {group_id}: "
            f"{len(review):,} Prüffälle | "
            f"{affected_sites:,} Standorte | "
            f"{len(site_context):,} Vergleichspotentiale"
        )

    # --------------------------------------------------------
    # 4. Übersicht
    # --------------------------------------------------------

    write_overview(
        overview_rows
    )

    # --------------------------------------------------------
    # 5. Abschluss
    # --------------------------------------------------------

    print()
    print("=" * 80)
    print("FERTIG")
    print("=" * 80)
    print()

    print(
        "Ausgabeordner:"
    )

    print(
        OUTPUT_DIR.resolve()
    )

    print()

    print(
        f"Erzeugte Prüfdateien: {len(overview_rows):,}"
    )

    print(
        f"Gruppen ohne Prüffälle: {groups_without_review:,}"
    )

    print(
        f"Prüffälle insgesamt: {total_cases:,}"
    )

    print(
        f"Betroffene Standort-Gruppen-Kombinationen: "
        f"{total_affected_sites:,}"
    )

    print()
    print(
        "Jede Pruefung_Gxx.xlsx enthält:"
    )

    print(
        "   1. Prueffaelle"
    )

    print(
        "      -> nur auffällige Potentiale + manuelle "
        "Entscheidungs-/Korrekturfelder"
    )

    print(
        "   2. Standortvergleich"
    )

    print(
        "      -> alle Potentiale der betroffenen Standorte; "
        "Prüffälle sind orange markiert"
    )

    print(
        "   3. Anleitung"
    )

    print()
    print(
        "Für das spätere Processing eindeutig:"
    )

    print(
        "   BESTAETIGT   -> Originalwerte verwenden"
    )

    print(
        "   KORRIGIERT   -> ausgefüllte Korrekturwerte verwenden"
    )

    print(
        "   AUSSCHLIESSEN -> Potential vollständig entfernen"
    )

    print(
        "   OFFEN        -> noch nicht final verarbeiten"
    )

    print("=" * 80)


# ============================================================
# 18. PROGRAMMSTART
# ============================================================

if __name__ == "__main__":

    main()
