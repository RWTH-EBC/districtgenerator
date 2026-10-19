from pathlib import Path
import re

import numpy as np
import pandas as pd


# ============================================================
# 1. KONFIGURATION
# ============================================================

BASE_DIR = Path("WZ2025_Vorbereitung")
SOURCE_DIR = BASE_DIR / "Quellenklassifikation_final"

INPUT_FILE = SOURCE_DIR / "Abwaermepotentiale_mit_Quellklassen.xlsx"
INPUT_SHEET = "Abwaermepotentiale"

# Neue Datei, damit eine ältere Vortest-Auswertung erhalten bleibt.
OUTPUT_FILE = SOURCE_DIR / "Branchenanalyse_Q1_Q7_v2.xlsx"

# Für die spätere Profilbildung wird QX primär nach dem Anteil an der
# JAHRESWÄRMEMENGE beurteilt. Die gemeldete maximale thermische Leistung
# bleibt ausschließlich als Diagnosefeld erhalten und entscheidet NICHT
# mehr über die Eignung eines Standorts.
QX_HEAT_THRESHOLDS = [0.00, 0.10, 0.20, 0.30, 0.40, 0.50]
PRIMARY_QX_HEAT_THRESHOLD = 0.20

# ------------------------------------------------------------
# Plausibilitätsprüfung der gemeldeten Zahlenwerte
# ------------------------------------------------------------
#
# WICHTIG:
# Diese Regeln dienen ausschließlich zur MARKIERUNG verdächtiger
# Angaben. Positive numerische Jahreswärmemengen werden NICHT
# automatisch korrigiert oder wegen eines Plausibilitätsflags
# aus der 20-%-Berechnung entfernt.
#
# Damit bleibt der derzeitige QX-Wärmemengenanteil nachvollziehbar
# und kann nach einer späteren manuellen/KI-gestützten Prüfung
# gezielt neu berechnet werden.
# ------------------------------------------------------------

# 8.784 h decken auch ein Schaltjahr ab. Darüber kann Q/P_max
# physikalisch nicht liegen, sofern beide Angaben dieselbe
# Systemgrenze beschreiben.
MAX_PHYSICAL_HOURS_PER_YEAR = 8784.0

# Sehr kleine äquivalente Vollbenutzungsstunden sind nicht per se
# falsch, aber prüfenswert.
LOW_FULL_LOAD_HOURS_STRONG = 10.0
LOW_FULL_LOAD_HOURS_WARNING = 100.0

# Heuristik für einen möglichen MWh->kWh-Fehler:
# Ist Q/P < 10 h/a, aber (Q*1000)/P läge in einem plausiblen
# Jahresbereich, wird nur ein Kandidatenflag gesetzt.
X1000_CANDIDATE_MIN_HOURS = 100.0
X1000_CANDIDATE_MAX_HOURS = MAX_PHYSICAL_HOURS_PER_YEAR


# ============================================================
# 2. QUELLKLASSEN Q1-Q7 + QX
# ============================================================

SOURCE_CLASS_LABELS = {
    "Q1": "Flüssiger Wärmestrom",
    "Q2": "Prozessdampf / Brüden / Kondensat",
    "Q3": "Abluft / nicht verbrennungsbedingter Prozessgasstrom",
    "Q4": "Rauchgas / Verbrennungs- und Oxidationsabgas",
    "Q5": "Diffuse / feste Abwärme",
    "Q6": "Kälteanlagen- / Verflüssigerwärme",
    "Q7": "Druckluft- / Kompressorwärme",
    "QX": "Nicht eindeutig bestimmbar",
}

SOURCE_CLASS_ORDER = ["Q1", "Q2", "Q3", "Q4", "Q5", "Q6", "Q7", "QX"]
VALID_SOURCE_CLASSES = set(SOURCE_CLASS_ORDER)
ASSIGNABLE_SOURCE_CLASSES = {"Q1", "Q2", "Q3", "Q4", "Q5", "Q6", "Q7"}
FALLBACK_SOURCE_CLASSES = {"Q6", "Q7"}


# ============================================================
# 3. PFA-SPALTENNAMEN
# ============================================================

HEAT_COLUMN = "Wärmemenge pro Jahr (in kWh/a)"
POWER_COLUMN = "Maximale thermische Leistung (in kW)"
TEMPERATURE_COLUMN = "Durchschnittliches Temperaturniveau (in °C)"
TEMPERATURE_RANGE_COLUMN = "Temperaturbereich"
SOURCE_NAME_COLUMN = "Name des Abwärmepotentials"
ADDITIONAL_INFO_COLUMN = "Ergänzende Informationen zum Abwärmepotential"
NOTES_COLUMN = "Weitere Hinweise"


# ============================================================
# 4. HILFSFUNKTIONEN
# ============================================================

def clean_string(value) -> str:
    """NaN/None -> leerer String, sonst getrimmter Text."""
    if pd.isna(value):
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none"}:
        return ""
    return text


def normalize_company_name(value) -> str:
    """Nur für technische Fallback-Schlüssel."""
    text = clean_string(value).lower()
    text = (
        text.replace("ä", "ae")
        .replace("ö", "oe")
        .replace("ü", "ue")
        .replace("ß", "ss")
    )
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def first_non_empty(series) -> str:
    for value in series:
        value = clean_string(value)
        if value:
            return value
    return ""


def unique_non_empty(series):
    result = []
    seen = set()
    for value in series:
        value = clean_string(value)
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    return result


def unique_join(series, separator=" | ") -> str:
    return separator.join(unique_non_empty(series))


def safe_ratio(numerator, denominator):
    if denominator is None or pd.isna(denominator) or denominator == 0:
        return np.nan
    return numerator / denominator


def parse_number(value):
    """Robuste Interpretation deutscher/internationaler Zahlenformate."""
    text = clean_string(value)
    if not text:
        return np.nan

    text = text.replace("\u00A0", "").replace(" ", "")
    match = re.search(r"[-+]?\d[\d.,]*", text)
    if not match:
        return np.nan

    number = match.group(0)
    if "," in number and "." in number:
        if number.rfind(",") > number.rfind("."):
            number = number.replace(".", "").replace(",", ".")
        else:
            number = number.replace(",", "")
    elif "," in number:
        number = number.replace(",", ".")
    elif number.count(".") > 1:
        number = number.replace(".", "")

    try:
        return float(number)
    except ValueError:
        return np.nan


def contains_any(text_value, patterns) -> bool:
    """Case-insensitive Teilstringsuche für Freitext-Flags."""
    text = clean_string(text_value).casefold()
    return any(pattern.casefold() in text for pattern in patterns)


def join_flags(*flags) -> str:
    """Mehrere optionale Flag-Texte kompakt zusammenführen."""
    return " | ".join(flag for flag in flags if flag)


# ============================================================
# 5. WZ-CODES
# ============================================================

def normalize_wz_code(value) -> str:
    """
    Behält die vorhandene WZ-Tiefe vollständig bei.

    38.21 -> 38.21
    38.2  -> 38.2
    6.1   -> 06.1
    35    -> 35
    """
    text = clean_string(value).replace(",", ".").replace(" ", "")
    if not text:
        return ""

    match = re.match(r"^(\d{1,2})(.*)$", text)
    if not match:
        return text

    division = int(match.group(1))
    remainder = match.group(2)
    return f"{division:02d}{remainder}"


def derive_wz3(value) -> str:
    """38.21 -> 38.2; 20.3 -> 20.3; 38 -> leer."""
    code = normalize_wz_code(value)
    match = re.match(r"^(\d{2})\.(\d)", code)
    if not match:
        return ""
    return f"{match.group(1)}.{match.group(2)}"


def derive_wz2(value) -> str:
    """38.21 -> 38; 20.3 -> 20."""
    code = normalize_wz_code(value)
    match = re.match(r"^(\d{2})", code)
    return match.group(1) if match else ""


# ============================================================
# 6. DATEN EINLESEN
# ============================================================

def load_data() -> pd.DataFrame:
    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"Eingabedatei nicht gefunden:\n{INPUT_FILE}")

    df = pd.read_excel(
        INPUT_FILE,
        sheet_name=INPUT_SHEET,
        dtype=str,
        keep_default_na=False,
    )

    required_columns = {
        "Originaldatensatz_ID",
        "Klassifikations_ID",
        "Unternehmens_ID",
        "Firmenname",
    }
    missing = required_columns - set(df.columns)
    if missing:
        raise ValueError(
            "Folgende notwendige Spalten fehlen:\n\n"
            + "\n".join(sorted(missing))
        )

    # Ältere Zwischenstände sollen weiterhin lesbar bleiben. Deshalb werden
    # neue Felder notfalls leer ergänzt und später mit konservativen Fallbacks
    # behandelt.
    optional_columns = [
        "Standortname",
        "Standortverbund_ID",
        "Standortverbund_Status",
        "Profil_Klassifikations_ID",
        "Ist_Hauptfunktion",
        "WZ_Status",
        "WZ_Final_Code",
        "WZ_Final_Ebene",
        "WZ3_Code",
        "WZ2_Code",
        "WZ_Gruppe_Code",
        "WZ_Gruppe_Bezeichnung",
        "WZ_Final_Herkunft",
        "WZ_AI_Konfidenz",
        "WZ_Konfidenz",  # ältere Dateien
        "WZ_Qualitaetsflag",
        "Standorttaetigkeit_kurz",
        "Standortfunktion",
        "Referenz_Mehrfach_Match",
        "Profil_Unternehmens_ID",
        "Profil_Firmenname",
        "Profil_Standortname",
        "Profil_WZ_Status",
        "Profil_WZ_Final_Code",
        "Profil_WZ_Final_Ebene",
        "Profil_WZ3_Code",
        "Profil_WZ2_Code",
        "Profil_WZ_Gruppe_Code",
        "Profil_WZ_Gruppe_Bezeichnung",
        "Profil_Standorttaetigkeit_kurz",
        "Quellklassifikation_Status",
        "Quellklasse_Code",
        "Quellklasse",
        "Quellklasse_Begruendung",
        "Quellklasse_Konfidenz",
        HEAT_COLUMN,
        POWER_COLUMN,
        TEMPERATURE_COLUMN,
        TEMPERATURE_RANGE_COLUMN,
        SOURCE_NAME_COLUMN,
        ADDITIONAL_INFO_COLUMN,
        NOTES_COLUMN,
    ]

    for column in optional_columns:
        if column not in df.columns:
            df[column] = ""

    return df


# ============================================================
# 7. QUELLENEBENE VORBEREITEN
# ============================================================

def prepare_source_data(df: pd.DataFrame) -> pd.DataFrame:
    """
    Wesentliche Änderungen:

    - Profilebene = Standortverbund_ID statt Klassifikations_ID.
    - Branchenzuordnung = Profil_WZ_Final_Code.
    - 38.21 bleibt als exakter Profilcode erhalten.
    - Q6/Q7 gelten als zuordenbar; nur QX ist Restklasse.
    """
    data = df.copy()

    for column in [
        "Originaldatensatz_ID",
        "Klassifikations_ID",
        "Unternehmens_ID",
        "Firmenname",
        "Standortname",
        "Standortverbund_ID",
        "Profil_Klassifikations_ID",
        "Profil_Unternehmens_ID",
        "Profil_Firmenname",
        "Profil_Standortname",
    ]:
        data[column] = data[column].map(clean_string)

    missing_site_id = data["Klassifikations_ID"].eq("")
    if missing_site_id.any():
        raise ValueError(
            f"{missing_site_id.sum()} Quellen besitzen keine Klassifikations_ID."
        )

    # --------------------------------------------------------
    # Physischer Profilstandort
    # --------------------------------------------------------
    data["_Site_Key"] = data["Standortverbund_ID"]
    fallback_site = data["_Site_Key"].eq("")
    data.loc[fallback_site, "_Site_Key"] = data.loc[
        fallback_site,
        "Klassifikations_ID",
    ]

    # --------------------------------------------------------
    # Profil-Unternehmen
    # --------------------------------------------------------
    def create_profile_company_key(row):
        company_id = clean_string(row.get("Profil_Unternehmens_ID", ""))
        if not company_id:
            company_id = clean_string(row.get("Unternehmens_ID", ""))
        if company_id:
            return "ID::" + company_id

        company_name = normalize_company_name(row.get("Profil_Firmenname", ""))
        if not company_name:
            company_name = normalize_company_name(row.get("Firmenname", ""))
        if company_name:
            return "NAME::" + company_name

        return "SITE::" + clean_string(row["_Site_Key"])

    data["_Company_Key"] = data.apply(create_profile_company_key, axis=1)

    # --------------------------------------------------------
    # Profil-WZ: exakten Code bewahren, Hierarchie zusätzlich führen
    # --------------------------------------------------------
    def choose_profile_final_code(row):
        for candidate in [
            row.get("Profil_WZ_Final_Code", ""),
            row.get("WZ_Final_Code", ""),
            row.get("WZ_Gruppe_Code", ""),
        ]:
            candidate = normalize_wz_code(candidate)
            if candidate:
                return candidate
        return ""

    data["_Profile_WZ_Final_Code"] = data.apply(
        choose_profile_final_code,
        axis=1,
    )

    data["_Profile_WZ3_Code"] = data["Profil_WZ3_Code"].map(normalize_wz_code)
    missing = data["_Profile_WZ3_Code"].eq("")
    data.loc[missing, "_Profile_WZ3_Code"] = data.loc[
        missing,
        "_Profile_WZ_Final_Code",
    ].map(derive_wz3)

    data["_Profile_WZ2_Code"] = data["Profil_WZ2_Code"].map(normalize_wz_code)
    missing = data["_Profile_WZ2_Code"].eq("")
    data.loc[missing, "_Profile_WZ2_Code"] = data.loc[
        missing,
        "_Profile_WZ_Final_Code",
    ].map(derive_wz2)

    # Lokale Teilstandort-WZ nur für Diagnosezwecke.
    data["_Local_WZ_Final_Code"] = data["WZ_Final_Code"].map(normalize_wz_code)
    missing = data["_Local_WZ_Final_Code"].eq("")
    data.loc[missing, "_Local_WZ_Final_Code"] = data.loc[
        missing,
        "WZ_Gruppe_Code",
    ].map(normalize_wz_code)

    # --------------------------------------------------------
    # Quellklassifikation Q1-Q7 + QX
    # --------------------------------------------------------
    data["Quellklasse_Code"] = (
        data["Quellklasse_Code"]
        .map(clean_string)
        .str.upper()
        .str.replace(" ", "", regex=False)
    )

    invalid_code = (
            data["Quellklasse_Code"].ne("")
            & ~data["Quellklasse_Code"].isin(VALID_SOURCE_CLASSES)
    )
    if invalid_code.any():
        values = (
            data.loc[invalid_code, "Quellklasse_Code"]
            .value_counts()
            .to_dict()
        )
        raise ValueError(
            "Ungültige Quellklassen gefunden:\n"
            f"{values}\n\nZulässig sind Q1-Q7 und QX."
        )

    data["_Quelle_bearbeitet"] = data["Quellklasse_Code"].isin(VALID_SOURCE_CLASSES)
    data["_Quelle_QX"] = data["Quellklasse_Code"].eq("QX")
    data["_Quelle_zuordenbar"] = data["Quellklasse_Code"].isin(
        ASSIGNABLE_SOURCE_CLASSES
    )
    data["_Quelle_Fallback"] = data["Quellklasse_Code"].isin(
        FALLBACK_SOURCE_CLASSES
    )
    data["_Quelle_Q6"] = data["Quellklasse_Code"].eq("Q6")
    data["_Quelle_Q7"] = data["Quellklasse_Code"].eq("Q7")

    # --------------------------------------------------------
    # Wärmemenge / Leistung
    # --------------------------------------------------------
    data["_Waermemenge_raw"] = data[HEAT_COLUMN].map(clean_string)
    data["_Waermemenge"] = data[HEAT_COLUMN].map(parse_number)
    data["_Waermemenge_vorhanden"] = data["_Waermemenge"].notna()
    data["_Waermemenge_positiv"] = data["_Waermemenge"].gt(0)
    data["_Waermemenge_Text_aber_nicht_numerisch"] = (
            data["_Waermemenge_raw"].ne("") & data["_Waermemenge"].isna()
    )
    data["_Waermemenge_nicht_positiv"] = (
            data["_Waermemenge"].notna() & ~data["_Waermemenge_positiv"]
    )

    data["_Leistung_raw"] = data[POWER_COLUMN].map(clean_string)
    data["_Leistung"] = data[POWER_COLUMN].map(parse_number)
    data["_Leistung_vorhanden"] = data["_Leistung"].notna()
    data["_Leistung_positiv"] = data["_Leistung"].gt(0)
    data["_Leistung_Text_aber_nicht_numerisch"] = (
            data["_Leistung_raw"].ne("") & data["_Leistung"].isna()
    )

    # --------------------------------------------------------
    # Zahlenplausibilität – NUR MARKIEREN, NICHT FILTERN
    # --------------------------------------------------------
    #
    # Q/P_max ist ein sehr nützlicher Konsistenzcheck. Ein Wert
    # über 8.784 h/a ist ein harter Widerspruch des Wertepaares,
    # beweist aber NICHT automatisch, dass die Wärmemenge falsch
    # ist. Ebenso kann die Leistung, die Systemgrenze oder die
    # Zeitbasis falsch sein. Deshalb wird hier nichts korrigiert.
    # --------------------------------------------------------

    valid_pair = data["_Waermemenge_positiv"] & data["_Leistung_positiv"]

    data["_Aequivalente_Vollbenutzungsstunden"] = np.nan
    data.loc[valid_pair, "_Aequivalente_Vollbenutzungsstunden"] = (
            data.loc[valid_pair, "_Waermemenge"]
            / data.loc[valid_pair, "_Leistung"]
    )

    data["_Physikalischer_Widerspruch_Waerme_Leistung"] = (
        data["_Aequivalente_Vollbenutzungsstunden"].gt(
            MAX_PHYSICAL_HOURS_PER_YEAR
        )
    )

    data["_Vollbenutzungsstunden_unter_10"] = (
            data["_Aequivalente_Vollbenutzungsstunden"].notna()
            & data["_Aequivalente_Vollbenutzungsstunden"].lt(
        LOW_FULL_LOAD_HOURS_STRONG
    )
    )

    data["_Vollbenutzungsstunden_unter_100"] = (
            data["_Aequivalente_Vollbenutzungsstunden"].notna()
            & data["_Aequivalente_Vollbenutzungsstunden"].lt(
        LOW_FULL_LOAD_HOURS_WARNING
    )
    )

    # Möglicher MWh->kWh-Faktorfehler: nur Kandidatenflag.
    data["_Aequiv_Stunden_wenn_Waermemenge_x1000"] = np.nan
    data.loc[valid_pair, "_Aequiv_Stunden_wenn_Waermemenge_x1000"] = (
            data.loc[valid_pair, "_Waermemenge"] * 1000.0
            / data.loc[valid_pair, "_Leistung"]
    )

    data["_Waermemenge_x1000_Kandidat"] = (
            data["_Vollbenutzungsstunden_unter_10"]
            & data["_Aequiv_Stunden_wenn_Waermemenge_x1000"].between(
        X1000_CANDIDATE_MIN_HOURS,
        X1000_CANDIDATE_MAX_HOURS,
        inclusive="both",
    )
    )

    # Freitexthinweise helfen, harte Widersprüche einzuordnen.
    combined_text = (
            data[SOURCE_NAME_COLUMN].map(clean_string)
            + " | "
            + data[ADDITIONAL_INFO_COLUMN].map(clean_string)
            + " | "
            + data[NOTES_COLUMN].map(clean_string)
    )

    data["_Text_Hinweis_Schaetzung"] = combined_text.map(
        lambda value: contains_any(
            value,
            [
                "geschätzt",
                "geschaetzt",
                "schätzung",
                "schaetzung",
                "hochgerechnet",
                "hochrechnung",
            ],
        )
    )

    data["_Text_Hinweis_Keine_Max_Leistung"] = combined_text.map(
        lambda value: contains_any(
            value,
            [
                "keine max. thermische leistung",
                "keine max thermische leistung",
                "keine maximale thermische leistung",
                "keine max. leistung",
                "keine maximale leistung",
            ],
        )
    )

    data["_Text_Hinweis_Typenschild"] = combined_text.map(
        lambda value: contains_any(
            value,
            ["typenschild", "nennleistung"],
        )
    )

    data["_Leistung_Platzhalter_Verdacht"] = (
            data["_Leistung_positiv"]
            & data["_Leistung"].le(1.0)
            & (
                    data["_Text_Hinweis_Keine_Max_Leistung"]
                    | data["_Physikalischer_Widerspruch_Waerme_Leistung"]
            )
    )

    # Hohe Prüfpriorität: tatsächlicher Zahlen-/Systemgrenzenkonflikt oder
    # möglicher Einheitenfehler. Schätz-/Hochrechnungshinweise allein
    # gelten NICHT als Fehler.
    data["_Zahlenpruefung_hohe_Prioritaet"] = (
            data["_Waermemenge_Text_aber_nicht_numerisch"]
            | data["_Waermemenge_nicht_positiv"]
            | data["_Physikalischer_Widerspruch_Waerme_Leistung"]
            | data["_Waermemenge_x1000_Kandidat"]
            | data["_Leistung_Platzhalter_Verdacht"]
    )

    data["_Zahlenpruefung_mittlere_Prioritaet"] = (
            ~data["_Zahlenpruefung_hohe_Prioritaet"]
            & (
                    data["_Vollbenutzungsstunden_unter_100"]
                    | data["_Text_Hinweis_Schaetzung"]
                    | data["_Text_Hinweis_Typenschild"]
            )
    )

    def build_numeric_quality_flag(row):
        flags = []

        if row["_Waermemenge_Text_aber_nicht_numerisch"]:
            flags.append("WAERMEMENGE_NICHT_NUMERISCH")
        if row["_Waermemenge_nicht_positiv"]:
            flags.append("WAERMEMENGE_NICHT_POSITIV")
        if row["_Physikalischer_Widerspruch_Waerme_Leistung"]:
            flags.append("WIDERSPRUCH_Q_DURCH_P_UEBER_8784H")
        if row["_Waermemenge_x1000_Kandidat"]:
            flags.append("WAERMEMENGE_X1000_KANDIDAT")
        elif row["_Vollbenutzungsstunden_unter_10"]:
            flags.append("VOLLBENUTZUNGSSTUNDEN_UNTER_10")
        elif row["_Vollbenutzungsstunden_unter_100"]:
            flags.append("VOLLBENUTZUNGSSTUNDEN_UNTER_100")
        if row["_Leistung_Platzhalter_Verdacht"]:
            flags.append("LEISTUNG_PLATZHALTER_VERDACHT")
        if row["_Text_Hinweis_Keine_Max_Leistung"]:
            flags.append("TEXT_KEINE_MAX_LEISTUNG")
        if row["_Text_Hinweis_Typenschild"]:
            flags.append("TEXT_TYPENSCHILD_ODER_NENNLEISTUNG")
        if row["_Text_Hinweis_Schaetzung"]:
            flags.append("TEXT_SCHAETZUNG_ODER_HOCHRECHNUNG")

        return " | ".join(flags) if flags else "UNAUFFAELLIG"

    data["_Zahlen_Plausibilitaetsflags"] = data.apply(
        build_numeric_quality_flag,
        axis=1,
    )

    # --------------------------------------------------------
    # Temperatur / sonstige Datenvollständigkeit
    # --------------------------------------------------------
    avg_temp_present = data[TEMPERATURE_COLUMN].map(clean_string).ne("")
    range_temp_present = data[TEMPERATURE_RANGE_COLUMN].map(clean_string).ne("")
    data["_Temperatur_vorhanden"] = avg_temp_present | range_temp_present

    data["_Quellenname_vorhanden"] = data[SOURCE_NAME_COLUMN].map(clean_string).ne("")
    data["_Zusatzinfo_vorhanden"] = data[ADDITIONAL_INFO_COLUMN].map(clean_string).ne("")
    data["_Hinweise_vorhanden"] = data[NOTES_COLUMN].map(clean_string).ne("")

    # Für die spätere Profilbildung ist die gemeldete Maximalleistung
    # KEINE Pflichtangabe mehr. Deshalb werden zwei Vollständigkeitsmaße
    # getrennt geführt.
    data["_Profil_Kernangaben_vollstaendig"] = (
            data["_Quellenname_vorhanden"]
            & data["_Waermemenge_vorhanden"]
            & data["_Temperatur_vorhanden"]
    )

    data["_Technische_Kernangaben_mit_Leistung_vollstaendig"] = (
            data["_Profil_Kernangaben_vollstaendig"]
            & data["_Leistung_vorhanden"]
    )

    # Abwärtskompatibilität für bestehende Auswertungen.
    data["_Kernangaben_vollstaendig"] = (
        data["_Profil_Kernangaben_vollstaendig"]
    )

    return data


# ============================================================
# 8. STANDORTANALYSE AUF PHYSISCHER STANDORTEBENE
# ============================================================

def build_site_analysis(data: pd.DataFrame) -> pd.DataFrame:
    """Eine Zeile pro finalem physischen Standortverbund."""
    rows = []

    for site_key, group in data.groupby("_Site_Key", sort=False):
        total_sources = len(group)
        processed_mask = group["_Quelle_bearbeitet"]
        processed_count = int(processed_mask.sum())
        open_count = total_sources - processed_count

        qx_count = int(group["_Quelle_QX"].sum())
        assignable_count = int(group["_Quelle_zuordenbar"].sum())
        fallback_count = int(group["_Quelle_Fallback"].sum())

        if processed_count == 0:
            processing_status = "nicht_begonnen"
        elif processed_count == total_sources:
            processing_status = "vollstaendig"
        else:
            processing_status = "teilweise"

        # ----------------------------------------------------
        # Standortstruktur / Hauptfunktion
        # --------------------------------------------------------
        part_ids = unique_non_empty(group["Klassifikations_ID"])
        profile_main_id = first_non_empty(group["Profil_Klassifikations_ID"])
        if not profile_main_id:
            profile_main_id = part_ids[0] if part_ids else ""

        main_group = group[group["Klassifikations_ID"].eq(profile_main_id)]
        if main_group.empty:
            main_group = group

        # ----------------------------------------------------
        # Profil-WZ muss innerhalb eines finalen Standortverbunds
        # eindeutig sein.
        # ----------------------------------------------------
        final_codes = unique_non_empty(group["_Profile_WZ_Final_Code"])
        wz3_codes = unique_non_empty(group["_Profile_WZ3_Code"])
        wz2_codes = unique_non_empty(group["_Profile_WZ2_Code"])

        if len(final_codes) > 1:
            raise ValueError(
                f"Standortverbund '{site_key}' besitzt mehrere "
                f"Profil_WZ_Final_Codes: {final_codes}"
            )
        if len(wz3_codes) > 1:
            raise ValueError(
                f"Standortverbund '{site_key}' besitzt mehrere "
                f"Profil_WZ3_Codes: {wz3_codes}"
            )
        if len(wz2_codes) > 1:
            raise ValueError(
                f"Standortverbund '{site_key}' besitzt mehrere "
                f"Profil_WZ2_Codes: {wz2_codes}"
            )

        profile_final_code = final_codes[0] if final_codes else ""
        profile_wz3_code = wz3_codes[0] if wz3_codes else ""
        profile_wz2_code = wz2_codes[0] if wz2_codes else ""
        local_codes = unique_non_empty(group["_Local_WZ_Final_Code"])

        # ----------------------------------------------------
        # Klassenhäufigkeiten
        # ----------------------------------------------------
        class_counts = {
            code: int(group["Quellklasse_Code"].eq(code).sum())
            for code in SOURCE_CLASS_ORDER
        }
        source_confidence = (
            group["Quellklasse_Konfidenz"].map(clean_string).str.lower()
        )

        # ----------------------------------------------------
        # Wärmemengen
        # ----------------------------------------------------
        known_heat_total = group.loc[
            group["_Waermemenge_vorhanden"], "_Waermemenge"
        ].sum()
        processed_heat = group.loc[
            processed_mask & group["_Waermemenge_vorhanden"], "_Waermemenge"
        ].sum()
        qx_heat = group.loc[
            group["_Quelle_QX"] & group["_Waermemenge_vorhanden"], "_Waermemenge"
        ].sum()
        q6_heat = group.loc[
            group["_Quelle_Q6"] & group["_Waermemenge_vorhanden"], "_Waermemenge"
        ].sum()
        q7_heat = group.loc[
            group["_Quelle_Q7"] & group["_Waermemenge_vorhanden"], "_Waermemenge"
        ].sum()

        processed_with_heat = int(
            (processed_mask & group["_Waermemenge_vorhanden"]).sum()
        )
        processed_with_positive_heat = int(
            (processed_mask & group["_Waermemenge_positiv"]).sum()
        )

        class_heat = {
            code: group.loc[
                group["Quellklasse_Code"].eq(code) & group["_Waermemenge_vorhanden"],
                "_Waermemenge",
            ].sum()
            for code in SOURCE_CLASS_ORDER
        }

        # ----------------------------------------------------
        # Leistung – nur noch Diagnose / Plausibilitätskontrolle
        # ----------------------------------------------------
        known_power_total = group.loc[
            group["_Leistung_vorhanden"], "_Leistung"
        ].sum()
        processed_power = group.loc[
            processed_mask & group["_Leistung_vorhanden"], "_Leistung"
        ].sum()
        qx_power = group.loc[
            group["_Quelle_QX"] & group["_Leistung_vorhanden"], "_Leistung"
        ].sum()
        q6_power = group.loc[
            group["_Quelle_Q6"] & group["_Leistung_vorhanden"], "_Leistung"
        ].sum()
        q7_power = group.loc[
            group["_Quelle_Q7"] & group["_Leistung_vorhanden"], "_Leistung"
        ].sum()

        processed_with_power = int(
            (processed_mask & group["_Leistung_vorhanden"]).sum()
        )

        class_power = {
            code: group.loc[
                group["Quellklasse_Code"].eq(code) & group["_Leistung_vorhanden"],
                "_Leistung",
            ].sum()
            for code in SOURCE_CLASS_ORDER
        }

        # ----------------------------------------------------
        # Zahlenplausibilität auf Standortebene
        # ----------------------------------------------------
        hard_numeric_issues = int(group["_Zahlenpruefung_hohe_Prioritaet"].sum())
        medium_numeric_issues = int(
            group["_Zahlenpruefung_mittlere_Prioritaet"].sum()
        )
        physical_conflicts = int(
            group["_Physikalischer_Widerspruch_Waerme_Leistung"].sum()
        )
        x1000_candidates = int(group["_Waermemenge_x1000_Kandidat"].sum())
        placeholder_power = int(group["_Leistung_Platzhalter_Verdacht"].sum())
        estimated_sources = int(group["_Text_Hinweis_Schaetzung"].sum())
        low_hours_100 = int(group["_Vollbenutzungsstunden_unter_100"].sum())

        if hard_numeric_issues > 0:
            numeric_quality_status = "PRUEFEN_HOCH"
        elif medium_numeric_issues > 0:
            numeric_quality_status = "PRUEFEN_MITTEL"
        else:
            numeric_quality_status = "UNAUFFAELLIG"

        wz_ai_confidence = first_non_empty(main_group["WZ_AI_Konfidenz"])
        if not wz_ai_confidence:
            wz_ai_confidence = first_non_empty(main_group["WZ_Konfidenz"])

        row = {
            # Standortstruktur
            "Standortverbund_ID": site_key,
            "Standortverbund_Status": first_non_empty(group["Standortverbund_Status"]),
            "Anzahl_Teilstandorte": len(part_ids),
            "Teil_Klassifikations_IDs": " | ".join(part_ids),
            "Profil_Klassifikations_ID": profile_main_id,
            "Profil_Unternehmens_ID": (
                    first_non_empty(group["Profil_Unternehmens_ID"])
                    or first_non_empty(main_group["Unternehmens_ID"])
            ),
            "Profil_Firmenname": (
                    first_non_empty(group["Profil_Firmenname"])
                    or first_non_empty(main_group["Firmenname"])
            ),
            "Profil_Standortname": (
                    first_non_empty(group["Profil_Standortname"])
                    or first_non_empty(main_group["Standortname"])
            ),
            "_Company_Key": first_non_empty(group["_Company_Key"]),

            # Profil-WZ
            "Profil_WZ_Final_Code": profile_final_code,
            "Profil_WZ3_Code": profile_wz3_code,
            "Profil_WZ2_Code": profile_wz2_code,
            "Profil_WZ3_Bezeichnung": (
                    first_non_empty(group["Profil_WZ_Gruppe_Bezeichnung"])
                    or first_non_empty(group["WZ_Gruppe_Bezeichnung"])
            ),
            "Profil_Standorttaetigkeit_kurz": (
                    first_non_empty(group["Profil_Standorttaetigkeit_kurz"])
                    or first_non_empty(main_group["Standorttaetigkeit_kurz"])
            ),
            "Profil_WZ_Status": (
                    first_non_empty(group["Profil_WZ_Status"])
                    or first_non_empty(main_group["WZ_Status"])
            ),
            "Profil_WZ_Final_Herkunft": first_non_empty(main_group["WZ_Final_Herkunft"]),
            "Profil_WZ_AI_Konfidenz": wz_ai_confidence,
            "Profil_WZ_Qualitaetsflag": first_non_empty(main_group["WZ_Qualitaetsflag"]),
            "Profil_Standortfunktion": first_non_empty(main_group["Standortfunktion"]),
            "Teilstandort_WZ_Final_Codes": " | ".join(local_codes),
            "Anzahl_unterschiedliche_Teilstandort_WZ_Codes": len(local_codes),

            # Bearbeitungsstand
            "Anzahl_Quellen_gesamt": total_sources,
            "Anzahl_Quellen_KI_bearbeitet": processed_count,
            "Anzahl_Quellen_noch_offen": open_count,
            "Quellklassifikation_Abdeckung": safe_ratio(processed_count, total_sources),
            "Quellklassifikation_Status": processing_status,

            # Q1-Q7/QX
            **{f"Anzahl_{code}": class_counts[code] for code in SOURCE_CLASS_ORDER},
            "Anzahl_zuordenbare_Quellen": assignable_count,
            "Anzahl_Fallback_Quellen_Q6_Q7": fallback_count,
            "QX_Anteil_bearbeitet": safe_ratio(qx_count, processed_count),
            "Zuordenbar_Anteil_bearbeitet": safe_ratio(assignable_count, processed_count),
            "Fallback_Anteil_bearbeitet": safe_ratio(fallback_count, processed_count),

            # Konfidenzen
            "Quellen_Konfidenz_hoch": int(source_confidence.eq("hoch").sum()),
            "Quellen_Konfidenz_mittel": int(source_confidence.eq("mittel").sum()),
            "Quellen_Konfidenz_gering": int(source_confidence.eq("gering").sum()),
            "Quellen_Konfidenz_keine": int(source_confidence.eq("keine").sum()),

            # Wärmemenge
            "Summe_bekannte_Waermemenge_kWh": known_heat_total,
            "Waermemenge_bearbeitete_Quellen_kWh": processed_heat,
            "Waermemenge_QX_kWh": qx_heat,
            "Waermemenge_Q6_kWh": q6_heat,
            "Waermemenge_Q7_kWh": q7_heat,
            "QX_Anteil_Waermemenge_bearbeitet": safe_ratio(qx_heat, processed_heat),
            "Anteil_bearbeitete_Quellen_mit_Waermemenge": safe_ratio(
                processed_with_heat,
                processed_count,
            ),
            "Anteil_bearbeitete_Quellen_mit_positiver_Waermemenge": safe_ratio(
                processed_with_positive_heat,
                processed_count,
            ),

            # Leistung – nur Diagnose
            "Summe_bekannte_Leistung_kW": known_power_total,
            "Leistung_bearbeitete_Quellen_kW": processed_power,
            "Leistung_QX_kW": qx_power,
            "Leistung_Q6_kW": q6_power,
            "Leistung_Q7_kW": q7_power,
            "QX_Anteil_Leistung_bearbeitet": safe_ratio(qx_power, processed_power),
            "Fallback_Anteil_Leistung_bearbeitet": safe_ratio(
                q6_power + q7_power,
                processed_power,
                ),
            "Anteil_bearbeitete_Quellen_mit_Leistung": safe_ratio(
                processed_with_power,
                processed_count,
            ),

            # Datenvollständigkeit
            "Anteil_Quellen_mit_Waermemenge": safe_ratio(
                int(group["_Waermemenge_vorhanden"].sum()), total_sources
            ),
            "Anteil_Quellen_mit_Leistung": safe_ratio(
                int(group["_Leistung_vorhanden"].sum()), total_sources
            ),
            "Anteil_Quellen_mit_Temperatur": safe_ratio(
                int(group["_Temperatur_vorhanden"].sum()), total_sources
            ),
            "Anteil_Quellen_mit_Name": safe_ratio(
                int(group["_Quellenname_vorhanden"].sum()), total_sources
            ),
            "Anteil_Kernangaben_vollstaendig": safe_ratio(
                int(group["_Kernangaben_vollstaendig"].sum()), total_sources
            ),
            "Waermemenge_nicht_numerisch": int(
                group["_Waermemenge_Text_aber_nicht_numerisch"].sum()
            ),
            "Leistung_nicht_numerisch": int(
                group["_Leistung_Text_aber_nicht_numerisch"].sum()
            ),

            # Zahlenplausibilität – nur Diagnose, KEIN Filter
            "Zahlen_Qualitaetsstatus": numeric_quality_status,
            "Anzahl_Quellen_Zahlenpruefung_hoch": hard_numeric_issues,
            "Anzahl_Quellen_Zahlenpruefung_mittel": medium_numeric_issues,
            "Anzahl_Physikalische_Widersprueche_Q_durch_P": physical_conflicts,
            "Anzahl_Waermemenge_x1000_Kandidaten": x1000_candidates,
            "Anzahl_Leistung_Platzhalter_Verdacht": placeholder_power,
            "Anzahl_Quellen_Schaetzung_Hochrechnung": estimated_sources,
            "Anzahl_Quellen_Vollbenutzungsstunden_unter_100": low_hours_100,
        }

        # Direkte normalisierte Standortanteile.
        #
        # Für die eigentliche Profilbildung ist ab jetzt der Anteil an der
        # Jahreswärmemenge die primäre Größe. Leistungsanteile bleiben nur
        # zur Diagnose erhalten.
        for code in SOURCE_CLASS_ORDER:
            row[f"Quellenanteil_{code}_bearbeitet"] = safe_ratio(
                class_counts[code], processed_count
            )
            row[f"Waermengenanteil_{code}_bearbeitet"] = safe_ratio(
                class_heat[code], processed_heat
            )
            row[f"Leistungsanteil_{code}_bearbeitet"] = safe_ratio(
                class_power[code], processed_power
            )

        rows.append(row)

    result = pd.DataFrame(rows)

    if result.empty:
        return result

    # --------------------------------------------------------
    # VORLÄUFIGER Eignungsstandard für die spätere Profilbildung
    # --------------------------------------------------------
    # Ein Standort gilt rechnerisch als auswertbar, wenn:
    #   1. die Quellenklassifikation vollständig ist,
    #   2. für alle bearbeiteten Quellen eine numerische Jahreswärmemenge
    #      vorliegt,
    #   3. die Summe der Jahreswärmemengen > 0 ist.
    #
    # Vorläufig geeignet nach dem 20-%-Standard ist er zusätzlich, wenn:
    #   4. ein finaler Profil-WZ-Code vorhanden ist und
    #   5. QX höchstens 20 % der GEMELDETEN Jahreswärmemenge ausmacht.
    #
    # GANZ WICHTIG:
    # Plausibilitätsauffällige positive Jahreswärmemengen werden an dieser
    # Stelle bewusst NICHT herausgefiltert und NICHT korrigiert. Der
    # 20-%-Status ist deshalb bis zur Zahlenprüfung als VORLÄUFIG zu lesen.
    # --------------------------------------------------------
    result["Waermemengenprofil_auswertbar"] = (
            result["Quellklassifikation_Status"].eq("vollstaendig")
            & result["Waermemenge_bearbeitete_Quellen_kWh"].gt(0)
            & result["Anteil_bearbeitete_Quellen_mit_Waermemenge"].eq(1)
            & result["Anteil_bearbeitete_Quellen_mit_positiver_Waermemenge"].eq(1)
            & result["QX_Anteil_Waermemenge_bearbeitet"].notna()
    )

    result["Geeignet_QX_Waermemenge_max_20pct"] = (
            result["Profil_WZ_Final_Code"].ne("")
            & result["Waermemengenprofil_auswertbar"]
            & (
                    result["QX_Anteil_Waermemenge_bearbeitet"]
                    <= PRIMARY_QX_HEAT_THRESHOLD
            )
    )

    # Expliziter Alias: Die Eignung ist vorläufig, solange markierte
    # Zahlenwerte nicht geprüft wurden. Das bestehende Feld bleibt aus
    # Kompatibilitätsgründen erhalten.
    result["Geeignet_QX_Waermemenge_max_20pct_vorlaeufig"] = (
        result["Geeignet_QX_Waermemenge_max_20pct"]
    )

    result["QX20_Eignung_kann_sich_durch_Zahlenpruefung_aendern"] = (
        result["Anzahl_Quellen_Zahlenpruefung_hoch"].gt(0)
    )

    def suitability_reason(row):
        if clean_string(row["Profil_WZ_Final_Code"]) == "":
            return "KEIN_PROFIL_WZ"
        if row["Quellklassifikation_Status"] != "vollstaendig":
            return "QUELLKLASSIFIKATION_NICHT_VOLLSTAENDIG"
        if row["Anteil_bearbeitete_Quellen_mit_Waermemenge"] != 1:
            return "WAERMEMENGEN_NICHT_VOLLSTAENDIG"
        if row["Anteil_bearbeitete_Quellen_mit_positiver_Waermemenge"] != 1:
            return "MINDESTENS_EINE_WAERMEMENGE_NICHT_POSITIV"
        if not pd.notna(row["Waermemenge_bearbeitete_Quellen_kWh"]):
            return "KEINE_BERECHENBARE_WAERMEMENGE"
        if row["Waermemenge_bearbeitete_Quellen_kWh"] <= 0:
            return "WAERMEMENGE_NICHT_POSITIV"
        if not pd.notna(row["QX_Anteil_Waermemenge_bearbeitet"]):
            return "QX_WAERMEMENGENANTEIL_NICHT_BERECHENBAR"
        if row["QX_Anteil_Waermemenge_bearbeitet"] > PRIMARY_QX_HEAT_THRESHOLD:
            return "QX_WAERMEMENGENANTEIL_UEBER_20PCT"
        return "GEEIGNET"

    result["Eignungsgrund_QX_Waermemenge_20pct"] = result.apply(
        suitability_reason,
        axis=1,
    )

    return result


# ============================================================
# 9. BRANCHENAGGREGATION
# ============================================================

def aggregate_industries(site_df: pd.DataFrame, level: str) -> pd.DataFrame:
    """
    PROFIL = exakter finaler Profilcode, z. B. 38.21
    WZ3    = übergeordnete Gruppe, z. B. 38.2
    WZ2    = Abteilung, z. B. 38
    """
    if level == "PROFIL":
        code_column = "Profil_WZ_Final_Code"
        output_code_column = "Profil_WZ_Final_Code"
    elif level == "WZ3":
        code_column = "Profil_WZ3_Code"
        output_code_column = "WZ3_Code"
    elif level == "WZ2":
        code_column = "Profil_WZ2_Code"
        output_code_column = "WZ2_Code"
    else:
        raise ValueError("level muss 'PROFIL', 'WZ3' oder 'WZ2' sein.")

    usable = site_df[site_df[code_column].ne("")].copy()
    rows = []

    for wz_code, group in usable.groupby(code_column, sort=True):
        number_sites = len(group)
        number_companies = group["_Company_Key"].nunique()
        company_counts = group["_Company_Key"].value_counts()
        largest_company_sites = int(company_counts.max()) if len(company_counts) else 0

        complete_mask = group["Quellklassifikation_Status"].eq("vollstaendig")
        partial_mask = group["Quellklassifikation_Status"].eq("teilweise")
        unstarted_mask = group["Quellklassifikation_Status"].eq("nicht_begonnen")
        started_mask = ~unstarted_mask

        row = {
            output_code_column: wz_code,
            "Standorte_mit_WZ": number_sites,
            "Unternehmen_mit_WZ": number_companies,
            "Standorte_pro_Unternehmen": safe_ratio(number_sites, number_companies),
            "Standorte_groesstes_Unternehmen": largest_company_sites,
            "Anteil_groesstes_Unternehmen": safe_ratio(
                largest_company_sites, number_sites
            ),
            "Standorte_mit_mindestens_einer_klassifizierten_Quelle": int(
                started_mask.sum()
            ),
            "Standorte_Quellen_vollstaendig_klassifiziert": int(
                complete_mask.sum()
            ),
            "Standorte_Quellen_teilweise_klassifiziert": int(partial_mask.sum()),
            "Standorte_Quellen_noch_nicht_klassifiziert": int(
                unstarted_mask.sum()
            ),
            "Unternehmen_mit_mindestens_einer_klassifizierten_Quelle": (
                group.loc[started_mask, "_Company_Key"].nunique()
            ),
            "Quellen_gesamt": int(group["Anzahl_Quellen_gesamt"].sum()),
            "Quellen_KI_bearbeitet": int(
                group["Anzahl_Quellen_KI_bearbeitet"].sum()
            ),
            "Quellen_noch_offen": int(group["Anzahl_Quellen_noch_offen"].sum()),
        }

        # Quellklassen und Standortabdeckung der Klassen
        for code in SOURCE_CLASS_ORDER:
            row[f"Quellen_{code}"] = int(group[f"Anzahl_{code}"].sum())
            row[f"Standorte_mit_{code}"] = int(group[f"Anzahl_{code}"].gt(0).sum())

        total_processed = row["Quellen_KI_bearbeitet"]
        total_qx = row["Quellen_QX"]
        total_fallback = row["Quellen_Q6"] + row["Quellen_Q7"]

        row["QX_Anteil_bearbeitete_Quellen"] = safe_ratio(total_qx, total_processed)
        row["Fallback_Q6_Q7_Anteil_bearbeitete_Quellen"] = safe_ratio(
            total_fallback, total_processed
        )
        row["Mittlere_Quellklassifikation_Abdeckung"] = (
            group["Quellklassifikation_Abdeckung"].mean()
        )
        row["Median_Quellklassifikation_Abdeckung"] = (
            group["Quellklassifikation_Abdeckung"].median()
        )

        # Wärmemenge und Leistung – Diagnose, nicht Branchengewichtung
        total_known_heat = group["Summe_bekannte_Waermemenge_kWh"].sum()
        processed_heat = group["Waermemenge_bearbeitete_Quellen_kWh"].sum()
        qx_heat = group["Waermemenge_QX_kWh"].sum()

        total_known_power = group["Summe_bekannte_Leistung_kW"].sum()
        processed_power = group["Leistung_bearbeitete_Quellen_kW"].sum()
        qx_power = group["Leistung_QX_kW"].sum()
        fallback_power = group["Leistung_Q6_kW"].sum() + group["Leistung_Q7_kW"].sum()

        row["Summe_bekannte_Waermemenge_kWh"] = total_known_heat
        row["Waermemenge_bearbeitete_Quellen_kWh"] = processed_heat
        row["Waermemenge_QX_kWh"] = qx_heat
        row["QX_Anteil_Waermemenge_bearbeitet"] = safe_ratio(qx_heat, processed_heat)

        row["Summe_bekannte_Leistung_kW"] = total_known_power
        row["Leistung_bearbeitete_Quellen_kW"] = processed_power
        row["Leistung_QX_kW"] = qx_power
        row["QX_Anteil_Leistung_bearbeitet"] = safe_ratio(qx_power, processed_power)
        row["Fallback_Q6_Q7_Anteil_Leistung_bearbeitet"] = safe_ratio(
            fallback_power, processed_power
        )

        # ----------------------------------------------------
        # Ungewichtete Standortprofile auf Basis der Jahreswärmemenge:
        # nur vollständig klassifizierte Standorte mit 100 %
        # vorhandenen Jahreswärmemengen.
        # ----------------------------------------------------
        profile_mask = (
                complete_mask
                & group["Waermemenge_bearbeitete_Quellen_kWh"].gt(0)
                & group["Anteil_bearbeitete_Quellen_mit_Waermemenge"].eq(1)
                & group["Anteil_bearbeitete_Quellen_mit_positiver_Waermemenge"].eq(1)
        )
        row["Standorte_fuer_Waermemengenprofil_auswertbar"] = int(
            profile_mask.sum()
        )

        for code in SOURCE_CLASS_ORDER:
            share_column = f"Waermengenanteil_{code}_bearbeitet"
            row[
                f"Mittlerer_Waermemengenanteil_{code}_vollstaendige_Profile"
            ] = group.loc[profile_mask, share_column].mean()
            row[
                f"Median_Waermemengenanteil_{code}_vollstaendige_Profile"
            ] = group.loc[profile_mask, share_column].median()

        # Datenvollständigkeit
        row["Mittlerer_Anteil_Quellen_mit_Waermemenge"] = (
            group["Anteil_Quellen_mit_Waermemenge"].mean()
        )
        row["Mittlerer_Anteil_bearbeitete_Quellen_mit_Waermemenge"] = (
            group["Anteil_bearbeitete_Quellen_mit_Waermemenge"].mean()
        )
        row["Mittlerer_Anteil_bearbeitete_Quellen_mit_positiver_Waermemenge"] = (
            group["Anteil_bearbeitete_Quellen_mit_positiver_Waermemenge"].mean()
        )
        row["Mittlerer_Anteil_Quellen_mit_Leistung"] = (
            group["Anteil_Quellen_mit_Leistung"].mean()
        )
        row["Mittlerer_Anteil_Quellen_mit_Temperatur"] = (
            group["Anteil_Quellen_mit_Temperatur"].mean()
        )
        row["Standorte_Waermemenge_vollstaendig"] = int(
            group["Anteil_Quellen_mit_Waermemenge"].eq(1).sum()
        )
        row["Standorte_Leistung_vollstaendig"] = int(
            group["Anteil_Quellen_mit_Leistung"].eq(1).sum()
        )
        row["Standorte_Temperatur_vollstaendig"] = int(
            group["Anteil_Quellen_mit_Temperatur"].eq(1).sum()
        )

        # Zahlenplausibilität – diese Standorte werden NICHT aus den
        # QX-Schwellen herausgenommen. Die Kennzahlen zeigen nur, wie viel
        # spätere Prüfung noch nötig ist.
        row["Standorte_Zahlenpruefung_hoch"] = int(
            group["Anzahl_Quellen_Zahlenpruefung_hoch"].gt(0).sum()
        )
        row["Standorte_Zahlenpruefung_mittel"] = int(
            group["Anzahl_Quellen_Zahlenpruefung_mittel"].gt(0).sum()
        )
        row["Standorte_mit_Q_durch_P_Widerspruch"] = int(
            group["Anzahl_Physikalische_Widersprueche_Q_durch_P"].gt(0).sum()
        )
        row["Standorte_mit_Waermemenge_x1000_Kandidat"] = int(
            group["Anzahl_Waermemenge_x1000_Kandidaten"].gt(0).sum()
        )

        # AI-Konfidenz/alte Qualitätsflags nur als Diagnose.
        wz_conf = group["Profil_WZ_AI_Konfidenz"].map(clean_string).str.lower()
        row["WZ_AI_Konfidenz_hoch"] = int(wz_conf.eq("hoch").sum())
        row["WZ_AI_Konfidenz_mittel"] = int(wz_conf.eq("mittel").sum())
        row["WZ_AI_Konfidenz_gering"] = int(wz_conf.eq("gering").sum())
        row["WZ_AI_Konfidenz_fehlend"] = int(wz_conf.eq("").sum())

        quality = group["Profil_WZ_Qualitaetsflag"].map(clean_string)
        for flag in [
            "OK_STRIKT",
            "OK_PLAUSIBEL",
            "OK_NUR_SEKTOR",
            "STANDORTFUNKTION_ABWEICHEND",
            "PRUEFEN_BRANCHE",
            "WIDERSPRUCH",
            "WZ_ZU_GROB_ODER_FEHLT",
            "NICHT_VERGLEICHBAR",
        ]:
            row[f"WZ_Flag_{flag}"] = int(quality.eq(flag).sum())
        row["WZ_Qualitaetsflag_fehlend"] = int(quality.eq("").sum())

        # QX-Sensitivität pro Branche – nach Anteil an der Jahreswärmemenge
        for threshold in QX_HEAT_THRESHOLDS:
            threshold_pct = int(round(threshold * 100))
            suitable_mask = (
                    complete_mask
                    & group["Anteil_bearbeitete_Quellen_mit_Waermemenge"].eq(1)
                    & group["Anteil_bearbeitete_Quellen_mit_positiver_Waermemenge"].eq(1)
                    & group["Waermemenge_bearbeitete_Quellen_kWh"].gt(0)
                    & group["QX_Anteil_Waermemenge_bearbeitet"].notna()
                    & (group["QX_Anteil_Waermemenge_bearbeitet"] <= threshold)
            )
            row[
                f"Profile_vollstaendig_QX_Waermemenge_max_{threshold_pct}pct"
            ] = int(suitable_mask.sum())
            row[
                f"Unternehmen_vollstaendig_QX_Waermemenge_max_{threshold_pct}pct"
            ] = group.loc[suitable_mask, "_Company_Key"].nunique()

        # Inhalt der jeweiligen Aggregation dokumentieren
        if level == "PROFIL":
            row["WZ3_Code_zugeordnet"] = first_non_empty(group["Profil_WZ3_Code"])
            row["WZ3_Bezeichnung"] = first_non_empty(
                group["Profil_WZ3_Bezeichnung"]
            )
        elif level == "WZ3":
            row["WZ3_Bezeichnung"] = first_non_empty(
                group["Profil_WZ3_Bezeichnung"]
            )
            row["Enthaltene_Profil_WZ_Codes"] = unique_join(
                group["Profil_WZ_Final_Code"]
            )
        else:
            row["Enthaltene_Profil_WZ_Codes"] = unique_join(
                group["Profil_WZ_Final_Code"]
            )
            row["Enthaltene_WZ3_Gruppen"] = unique_join(group["Profil_WZ3_Code"])
            row["Enthaltene_WZ3_Bezeichnungen"] = unique_join(
                group["Profil_WZ3_Bezeichnung"]
            )

        rows.append(row)

    result = pd.DataFrame(rows)
    if not result.empty:
        result = result.sort_values(output_code_column, kind="stable").reset_index(
            drop=True
        )
    return result


# ============================================================
# 10. ZAHLENPLAUSIBILITÄT AUF QUELLENEBENE
# ============================================================

def build_numeric_plausibility_table(data: pd.DataFrame) -> pd.DataFrame:
    """
    Prüft auffällige Zahlenwerte, ohne sie zu korrigieren oder aus
    der Profilberechnung zu entfernen.

    Die Tabelle enthält nur Quellen mit mindestens einem Prüfhinweis.
    """

    mask = (
            data["_Zahlenpruefung_hohe_Prioritaet"]
            | data["_Zahlenpruefung_mittlere_Prioritaet"]
    )

    flagged = data.loc[mask].copy()

    if flagged.empty:
        return pd.DataFrame(
            columns=[
                "Originaldatensatz_ID",
                "Standortverbund_ID",
                "Klassifikations_ID",
                "Unternehmens_ID",
                "Firmenname",
                "Standortname",
                SOURCE_NAME_COLUMN,
                HEAT_COLUMN,
                POWER_COLUMN,
                "Aequivalente_Vollbenutzungsstunden",
                "Aequiv_Stunden_wenn_Waermemenge_x1000",
                "Quellklasse_Code",
                "Zahlen_Pruefprioritaet",
                "Zahlen_Plausibilitaetsflags",
                ADDITIONAL_INFO_COLUMN,
                NOTES_COLUMN,
            ]
        )

    flagged["Zahlen_Pruefprioritaet"] = np.where(
        flagged["_Zahlenpruefung_hohe_Prioritaet"],
        "HOCH",
        "MITTEL",
    )

    flagged["Aequivalente_Vollbenutzungsstunden"] = (
        flagged["_Aequivalente_Vollbenutzungsstunden"]
    )
    flagged["Aequiv_Stunden_wenn_Waermemenge_x1000"] = (
        flagged["_Aequiv_Stunden_wenn_Waermemenge_x1000"]
    )
    flagged["Zahlen_Plausibilitaetsflags"] = (
        flagged["_Zahlen_Plausibilitaetsflags"]
    )

    columns = [
        "Originaldatensatz_ID",
        "Standortverbund_ID",
        "Klassifikations_ID",
        "Unternehmens_ID",
        "Firmenname",
        "Standortname",
        SOURCE_NAME_COLUMN,
        HEAT_COLUMN,
        POWER_COLUMN,
        "Aequivalente_Vollbenutzungsstunden",
        "Aequiv_Stunden_wenn_Waermemenge_x1000",
        TEMPERATURE_COLUMN,
        TEMPERATURE_RANGE_COLUMN,
        "Quellklasse_Code",
        "Quellklasse",
        "Zahlen_Pruefprioritaet",
        "Zahlen_Plausibilitaetsflags",
        ADDITIONAL_INFO_COLUMN,
        NOTES_COLUMN,
    ]

    columns = [column for column in columns if column in flagged.columns]

    return (
        flagged[columns]
        .sort_values(
            ["Zahlen_Pruefprioritaet", "Firmenname", "Standortname"],
            ascending=[True, True, True],
            kind="stable",
        )
        .reset_index(drop=True)
    )


# ============================================================
# 11. ALLGEMEINE DATENQUALITÄT
# ============================================================

def build_data_quality_summary(data: pd.DataFrame, site_df: pd.DataFrame) -> pd.DataFrame:
    total_sources = len(data)
    total_sites = len(site_df)
    total_part_sites = data["Klassifikations_ID"].nunique()
    total_companies = site_df["_Company_Key"].nunique()
    merged_sites = int(site_df["Anzahl_Teilstandorte"].gt(1).sum())

    processed_sources = int(data["_Quelle_bearbeitet"].sum())
    qx_sources = int(data["_Quelle_QX"].sum())
    q6_sources = int(data["_Quelle_Q6"].sum())
    q7_sources = int(data["_Quelle_Q7"].sum())

    profile_wz_sites = int(site_df["Profil_WZ_Final_Code"].ne("").sum())
    wz3_sites = int(site_df["Profil_WZ3_Code"].ne("").sum())

    processed_heat = site_df["Waermemenge_bearbeitete_Quellen_kWh"].sum()
    qx_heat = site_df["Waermemenge_QX_kWh"].sum()
    q6_heat = site_df["Waermemenge_Q6_kWh"].sum()
    q7_heat = site_df["Waermemenge_Q7_kWh"].sum()

    processed_power = site_df["Leistung_bearbeitete_Quellen_kW"].sum()
    qx_power = site_df["Leistung_QX_kW"].sum()
    q6_power = site_df["Leistung_Q6_kW"].sum()
    q7_power = site_df["Leistung_Q7_kW"].sum()

    rows = [
        {"Kennzahl": "Abwärmequellen gesamt", "Wert": total_sources},
        {
            "Kennzahl": "Ursprüngliche Klassifikations_IDs / Teilstandorte",
            "Wert": total_part_sites,
        },
        {"Kennzahl": "Finale physische Standortprofile", "Wert": total_sites},
        {
            "Kennzahl": "Davon zusammengeführte Standortverbünde",
            "Wert": merged_sites,
        },
        {"Kennzahl": "Profil-Unternehmen gesamt", "Wert": total_companies},
        {
            "Kennzahl": "Standortprofile mit finalem Profil-WZ-Code",
            "Wert": profile_wz_sites,
        },
        {
            "Kennzahl": "Anteil Standortprofile mit finalem Profil-WZ-Code",
            "Wert": safe_ratio(profile_wz_sites, total_sites),
        },
        {"Kennzahl": "Standortprofile mit Profil-WZ3", "Wert": wz3_sites},
        {
            "Kennzahl": "Anteil Standortprofile mit Profil-WZ3",
            "Wert": safe_ratio(wz3_sites, total_sites),
        },
        {"Kennzahl": "Quellen bereits klassifiziert", "Wert": processed_sources},
        {
            "Kennzahl": "Anteil Quellen bereits klassifiziert",
            "Wert": safe_ratio(processed_sources, total_sources),
        },
        {
            "Kennzahl": "Q6 Kälteanlagen-/Verflüssigerwärme",
            "Wert": q6_sources,
        },
        {
            "Kennzahl": "Q7 Druckluft-/Kompressorwärme",
            "Wert": q7_sources,
        },
        {"Kennzahl": "Q6+Q7 Fallback-Klassen", "Wert": q6_sources + q7_sources},
        {
            "Kennzahl": "Anteil Q6+Q7 unter klassifizierten Quellen",
            "Wert": safe_ratio(q6_sources + q7_sources, processed_sources),
        },
        {"Kennzahl": "QX unter klassifizierten Quellen", "Wert": qx_sources},
        {
            "Kennzahl": "QX-Anteil unter klassifizierten Quellen",
            "Wert": safe_ratio(qx_sources, processed_sources),
        },
        {
            "Kennzahl": "QX-Anteil an klassifizierter Jahreswärmemenge",
            "Wert": safe_ratio(qx_heat, processed_heat),
        },
        {
            "Kennzahl": "Q6+Q7-Anteil an klassifizierter Jahreswärmemenge",
            "Wert": safe_ratio(q6_heat + q7_heat, processed_heat),
        },
        {
            "Kennzahl": "Standortprofile geeignet QX-Waermemenge <=20%",
            "Wert": int(site_df["Geeignet_QX_Waermemenge_max_20pct"].sum()),
        },
        {
            "Kennzahl": "QX-Anteil an klassifizierter Leistung (nur Diagnose)",
            "Wert": safe_ratio(qx_power, processed_power),
        },
        {
            "Kennzahl": "Q6+Q7-Anteil an klassifizierter Leistung (nur Diagnose)",
            "Wert": safe_ratio(q6_power + q7_power, processed_power),
        },
        {
            "Kennzahl": "Standortprofile Quellenklassifikation vollständig",
            "Wert": int(
                site_df["Quellklassifikation_Status"].eq("vollstaendig").sum()
            ),
        },
        {
            "Kennzahl": "Standortprofile Quellenklassifikation teilweise",
            "Wert": int(site_df["Quellklassifikation_Status"].eq("teilweise").sum()),
        },
        {
            "Kennzahl": "Standortprofile Quellenklassifikation nicht begonnen",
            "Wert": int(
                site_df["Quellklassifikation_Status"].eq("nicht_begonnen").sum()
            ),
        },
        {
            "Kennzahl": "Quellen mit Wärmemenge",
            "Wert": int(data["_Waermemenge_vorhanden"].sum()),
        },
        {
            "Kennzahl": "Anteil Quellen mit Wärmemenge",
            "Wert": data["_Waermemenge_vorhanden"].mean(),
        },
        {
            "Kennzahl": "Quellen mit Leistung",
            "Wert": int(data["_Leistung_vorhanden"].sum()),
        },
        {
            "Kennzahl": "Anteil Quellen mit Leistung",
            "Wert": data["_Leistung_vorhanden"].mean(),
        },
        {
            "Kennzahl": "Quellen mit Temperaturangabe",
            "Wert": int(data["_Temperatur_vorhanden"].sum()),
        },
        {
            "Kennzahl": "Anteil Quellen mit Temperaturangabe",
            "Wert": data["_Temperatur_vorhanden"].mean(),
        },
        {
            "Kennzahl": "Quellen mit vollständigen Kernangaben",
            "Wert": int(data["_Kernangaben_vollstaendig"].sum()),
        },
        {
            "Kennzahl": "Anteil Quellen mit vollständigen Kernangaben",
            "Wert": data["_Kernangaben_vollstaendig"].mean(),
        },
        {
            "Kennzahl": "Wärmemenge Text vorhanden aber nicht numerisch",
            "Wert": int(data["_Waermemenge_Text_aber_nicht_numerisch"].sum()),
        },
        {
            "Kennzahl": "Numerische Wärmemenge <= 0 (harter formaler Fehler)",
            "Wert": int(data["_Waermemenge_nicht_positiv"].sum()),
        },
        {
            "Kennzahl": "Leistung Text vorhanden aber nicht numerisch",
            "Wert": int(data["_Leistung_Text_aber_nicht_numerisch"].sum()),
        },
        {
            "Kennzahl": "Quellen mit hoher Zahlen-Prüfpriorität (nur markiert)",
            "Wert": int(data["_Zahlenpruefung_hohe_Prioritaet"].sum()),
        },
        {
            "Kennzahl": "Quellen mit mittlerer Zahlen-Prüfpriorität (nur markiert)",
            "Wert": int(data["_Zahlenpruefung_mittlere_Prioritaet"].sum()),
        },
        {
            "Kennzahl": "Quellen mit Q/P über 8.784 h/a",
            "Wert": int(
                data["_Physikalischer_Widerspruch_Waerme_Leistung"].sum()
            ),
        },
        {
            "Kennzahl": "Quellen mit Wärmemenge-x1000-Kandidatenflag",
            "Wert": int(data["_Waermemenge_x1000_Kandidat"].sum()),
        },
        {
            "Kennzahl": "Quellen mit Leistung-Platzhalterverdacht",
            "Wert": int(data["_Leistung_Platzhalter_Verdacht"].sum()),
        },
        {
            "Kennzahl": "Quellen mit Schätz-/Hochrechnungshinweis",
            "Wert": int(data["_Text_Hinweis_Schaetzung"].sum()),
        },
        {
            "Kennzahl": "Standorte, deren QX20-Status sich nach Zahlenprüfung ändern könnte",
            "Wert": int(
                site_df[
                    "QX20_Eignung_kann_sich_durch_Zahlenpruefung_aendern"
                ].sum()
            ),
        },
    ]

    for code in SOURCE_CLASS_ORDER:
        count = int(data["Quellklasse_Code"].eq(code).sum())
        rows.append(
            {
                "Kennzahl": f"{code} – {SOURCE_CLASS_LABELS[code]}",
                "Wert": count,
            }
        )
        rows.append(
            {
                "Kennzahl": f"Anteil {code} unter klassifizierten Quellen",
                "Wert": safe_ratio(count, processed_sources),
            }
        )

    return pd.DataFrame(rows)


# ============================================================
# 12. FEHLENDE INHALTE
# ============================================================

def build_missing_content_table(data: pd.DataFrame) -> pd.DataFrame:
    total = len(data)
    definitions = [
        ("Name des Abwärmepotentials", data["_Quellenname_vorhanden"]),
        ("Wärmemenge pro Jahr", data["_Waermemenge_vorhanden"]),
        ("Maximale thermische Leistung", data["_Leistung_vorhanden"]),
        ("Temperatur", data["_Temperatur_vorhanden"]),
        ("Ergänzende Informationen", data["_Zusatzinfo_vorhanden"]),
        ("Weitere Hinweise", data["_Hinweise_vorhanden"]),
    ]

    rows = []
    for field_name, present_mask in definitions:
        present = int(present_mask.sum())
        missing = total - present
        rows.append(
            {
                "Feld": field_name,
                "Vorhanden": present,
                "Fehlend": missing,
                "Anteil_vorhanden": safe_ratio(present, total),
                "Anteil_fehlend": safe_ratio(missing, total),
            }
        )
    return pd.DataFrame(rows)


# ============================================================
# 13. QX-SCHWELLENANALYSE – JAHRESWÄRMEMENGENANTEIL
# ============================================================

def build_qx_threshold_analysis(site_df: pd.DataFrame) -> pd.DataFrame:
    """
    Primäre Sensitivitätsanalyse für die spätere Profilbildung.

    Berücksichtigt werden nur Standortprofile mit finalem Profil-WZ und
    vollständig klassifizierten Quellen. Für eine belastbare Schwellenprüfung
    müssen zusätzlich für 100 % der bearbeiteten Quellen numerische
    Jahreswärmemengen vorliegen und deren Summe muss positiv sein.

    Die gemeldete Maximalleistung spielt hier ausdrücklich keine Rolle.

    Plausibilitätsauffällige, aber positive numerische Jahreswärmemengen
    werden bewusst NICHT aus dieser Schwellenanalyse entfernt. Die
    Ergebnisse sind daher bis zur Zahlenprüfung als vorläufig zu lesen.
    """
    base = site_df[
        site_df["Profil_WZ_Final_Code"].ne("")
        & site_df["Quellklassifikation_Status"].eq("vollstaendig")
        ].copy()

    computable = base[
        base["QX_Anteil_Waermemenge_bearbeitet"].notna()
    ].copy()

    complete_heat_base = computable[
        computable["Anteil_bearbeitete_Quellen_mit_Waermemenge"].eq(1)
        & computable["Anteil_bearbeitete_Quellen_mit_positiver_Waermemenge"].eq(1)
        & computable["Waermemenge_bearbeitete_Quellen_kWh"].gt(0)
        ].copy()

    rows = []
    for threshold in QX_HEAT_THRESHOLDS:
        selected = complete_heat_base[
            complete_heat_base["QX_Anteil_Waermemenge_bearbeitet"] <= threshold
            ]

        rows.append(
            {
                "Maximaler_QX_Waermemengenanteil": threshold,
                "Vollstaendige_Profile_mit_Profil_WZ": len(base),
                "Davon_QX_Waermemengenanteil_berechenbar": len(computable),
                "Davon_mit_100pct_positiver_Waermemengenangabe": len(complete_heat_base),
                "Profile_unter_Schwelle": len(selected),
                "Unternehmen_unter_Schwelle": selected["_Company_Key"].nunique(),
            }
        )

    return pd.DataFrame(rows)


# ============================================================
# 14. TAXONOMIE UND METHODISCHE HINWEISE
# ============================================================

def build_taxonomy_table() -> pd.DataFrame:
    definitions = {
        "Q1": (
            "Konkreter flüssiger Wärmestrom, z. B. Kühlwasser, Warmwasser, "
            "Öl oder Glykol. Bei Kälteanlagen/Kompressoren vorrangig, wenn "
            "der Flüssigkeitspfad bekannt ist."
        ),
        "Q2": (
            "Prozessdampf, Brüden und zugehöriges wasserbasiertes Kondensat. "
            "Kältemittel sind ausgeschlossen."
        ),
        "Q3": (
            "Nicht verbrennungsbedingter Luft-/Gasstrom, z. B. Abluft, "
            "Kühlluft oder heiße Druckluft. Bei Kälteanlagen vorrangig, "
            "wenn der luftseitige Pfad bekannt ist."
        ),
        "Q4": "Rauchgas bzw. Gasstrom aus Verbrennung oder thermischer Oxidation.",
        "Q5": "Feste, produktgebundene, Oberflächen- oder diffuse Abwärme.",
        "Q6": (
            "Kältemittelseitige Kälteanlagen-/Verflüssigerwärme oder Fallback, "
            "wenn die Kälteanlage eindeutig ist, der sekundäre Luft-/"
            "Flüssigkeitspfad aber unbekannt bleibt."
        ),
        "Q7": (
            "Fallback für eindeutig identifizierte Druckluft-/Luftkompression, "
            "wenn der konkrete Luft-/Flüssigkeitskühlpfad nicht angegeben ist."
        ),
        "QX": (
            "Technisch nicht eindeutig bestimmbar oder mehrere wesentlich "
            "verschiedene Wärmewege gemeinsam bilanziert."
        ),
    }

    return pd.DataFrame(
        [
            {
                "Quellklasse_Code": code,
                "Quellklasse": SOURCE_CLASS_LABELS[code],
                "Kurzdefinition": definitions[code],
                "Für_Profile_zuordenbar": code != "QX",
                "Fallback_Klasse": code in FALLBACK_SOURCE_CLASSES,
            }
            for code in SOURCE_CLASS_ORDER
        ]
    )


def build_notes_table() -> pd.DataFrame:
    notes = [
        (
            "Standortebene",
            "Die Analyse verwendet den finalen physischen Standortverbund als "
            "Profilebene. Mehrere Klassifikations_IDs eines bestätigten Verbunds "
            "werden gemeinsam ausgewertet."
        ),
        (
            "Profil-WZ",
            "Für die Branchenzuordnung wird Profil_WZ_Final_Code verwendet. "
            "Feinere manuelle Codes wie 38.21 bleiben exakt erhalten."
        ),
        (
            "WZ3 / WZ2",
            "WZ3 und WZ2 werden zusätzlich ausgewertet. So kann 38.21 separat "
            "analysiert und zugleich unter 38.2 bzw. 38 aggregiert werden."
        ),
        (
            "Q6",
            "Q6 ist eine gültige, zuordenbare Kälteanlagen-/Verflüssigerklasse. "
            "Wenn Luft oder Flüssigkeit bekannt war, sollte bereits Q3 bzw. Q1 "
            "klassifiziert worden sein."
        ),
        (
            "Q7",
            "Q7 ist eine gültige Fallback-Klasse für eindeutige Druckluft-/"
            "Kompressorwärme bei unbekanntem konkretem Kühlpfad."
        ),
        (
            "QX",
            "Nur QX gilt als nicht eindeutig zuordenbare Restklasse. Q6 und Q7 "
            "werden nicht zum QX-Anteil gerechnet."
        ),
        (
            "QX-Schwellenanalyse",
            "Die primäre Sensitivitätsanalyse verwendet den QX-Anteil an der "
            "Jahreswärmemenge des physischen Standorts. Für die Schwellenprüfung "
            "werden vollständige numerische Jahreswärmemengen verlangt."
        ),
        (
            "20%-Eignungsstandard – vorläufig",
            "Der aktuelle 20-%-Status basiert auf den gemeldeten positiven "
            "Jahreswärmemengen. Plausibilitätsauffällige positive Werte werden "
            "bewusst noch nicht korrigiert oder ausgeschlossen. Nur fehlende, "
            "nicht numerische oder nicht positive Jahreswärmemengen verhindern "
            "eine belastbare Profilberechnung. Der Status ist bis zur Zahlenprüfung "
            "vorläufig."
        ),
        (
            "Zahlenplausibilität",
            "Verdächtige Wärmemengen-/Leistungspaare werden nur markiert. Dazu "
            "gehören insbesondere Q/P über 8.784 h/a, mögliche x1000-"
            "Einheitenfehler, sehr niedrige Vollbenutzungsstunden sowie Freitext-"
            "Hinweise auf Schätzung, Hochrechnung, Typenschild oder fehlende "
            "Maximalleistung. Kein Flag verändert automatisch einen Originalwert."
        ),
        (
            "Profilnormalisierung",
            "Die Standorttabelle enthält Waermemengenanteile Q1-Q7/QX je Standort. "
            "Diese sind die primäre energetische Grundlage der späteren Profile. "
            "Leistungsanteile bleiben ausschließlich als Diagnose erhalten."
        ),
        (
            "Standortgewichtung",
            "Bei mittleren Waermemengenanteilen zählt jeder physische Standort gleich; "
            "große Standorte erhalten nicht automatisch höheres Gewicht."
        ),
        (
            "Branchenzusammenlegung",
            "Es erfolgt keinerlei automatische Zusammenlegung verschiedener "
            "Profil-WZ-Codes oder WZ-Gruppen."
        ),
        (
            "Unvollständige Klassifikation",
            "Noch nicht klassifizierte Quellen werden nicht als QX behandelt. "
            "Teilweise klassifizierte Standorte bleiben separat erkennbar."
        ),
    ]
    return pd.DataFrame(notes, columns=["Thema", "Hinweis"])


# ============================================================
# 15. EXCEL FORMATIEREN
# ============================================================

def format_workbook(filepath: Path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    workbook = load_workbook(filepath)
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"

        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        for column_cells in worksheet.columns:
            letter = column_cells[0].column_letter
            max_length = 0
            for cell in list(column_cells)[:250]:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            worksheet.column_dimensions[letter].width = min(
                max(max_length + 2, 10),
                45,
            )

        headers = {
            cell.column: clean_string(cell.value)
            for cell in worksheet[1]
        }
        for column_index, header in headers.items():
            if "Anteil" in header or "Abdeckung" in header:
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row=row, column=column_index).number_format = "0.0%"

    workbook.save(filepath)


# ============================================================
# 16. HAUPTPROGRAMM
# ============================================================

def main():
    print()
    print("=" * 78)
    print("BRANCHENANALYSE – Q1-Q7 + QX / FINALE STANDORTVERBÜNDE")
    print("=" * 78)
    print()

    print("1. Vollständige Tabelle wird eingelesen ...")
    raw = load_data()
    print(f"   -> {len(raw):,} Abwärmequellen")

    print("2. Analysefelder auf Quellenebene werden vorbereitet ...")
    data = prepare_source_data(raw)

    print("3. Finale physische Standortprofile werden aufgebaut ...")
    site_analysis = build_site_analysis(data)
    print(f"   -> {len(site_analysis):,} physische Standortprofile")
    print(
        f"   -> {data['Klassifikations_ID'].nunique():,} "
        "ursprüngliche Klassifikations_IDs"
    )
    print(
        f"   -> {site_analysis['_Company_Key'].nunique():,} Profil-Unternehmen"
    )

    print("4. Branchenübersicht nach exaktem Profil-WZ-Code ...")
    industry_profile = aggregate_industries(site_analysis, level="PROFIL")
    print(f"   -> {len(industry_profile):,} besetzte finale Profil-WZ-Codes")

    print("5. Branchenübersicht auf WZ3-Ebene ...")
    industry_wz3 = aggregate_industries(site_analysis, level="WZ3")
    print(f"   -> {len(industry_wz3):,} besetzte WZ3-Gruppen")

    print("6. Branchenübersicht auf WZ2-Ebene ...")
    industry_wz2 = aggregate_industries(site_analysis, level="WZ2")
    print(f"   -> {len(industry_wz2):,} besetzte WZ2-Abteilungen")

    print("7. Datenqualität und QX-Sensitivität werden analysiert ...")
    data_quality = build_data_quality_summary(data, site_analysis)
    numeric_plausibility = build_numeric_plausibility_table(data)
    missing_content = build_missing_content_table(data)
    qx_thresholds = build_qx_threshold_analysis(site_analysis)
    taxonomy = build_taxonomy_table()
    notes = build_notes_table()

    print(
        f"   -> {len(numeric_plausibility):,} Quellen mit mindestens einem "
        "Zahlen-Prüfhinweis"
    )

    print("8. Excel-Auswertung wird geschrieben ...")
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        site_analysis.to_excel(
            writer, sheet_name="Standort_Analyse", index=False
        )
        industry_profile.to_excel(
            writer, sheet_name="Branchen_ProfilWZ", index=False
        )
        industry_wz3.to_excel(
            writer, sheet_name="Branchen_WZ3", index=False
        )
        industry_wz2.to_excel(
            writer, sheet_name="Branchen_WZ2", index=False
        )
        data_quality.to_excel(
            writer, sheet_name="Datenqualitaet", index=False
        )
        numeric_plausibility.to_excel(
            writer, sheet_name="Zahlen_Plausibilitaet", index=False
        )
        missing_content.to_excel(
            writer, sheet_name="Fehlende_Inhalte", index=False
        )
        qx_thresholds.to_excel(
            writer, sheet_name="Schwellenanalyse_QX", index=False
        )
        taxonomy.to_excel(
            writer, sheet_name="Quellklassen_Taxonomie", index=False
        )
        notes.to_excel(
            writer, sheet_name="Hinweise", index=False
        )

    format_workbook(OUTPUT_FILE)

    total_sites = len(site_analysis)
    merged_sites = int(site_analysis["Anzahl_Teilstandorte"].gt(1).sum())
    profile_wz_sites = int(site_analysis["Profil_WZ_Final_Code"].ne("").sum())
    wz3_sites = int(site_analysis["Profil_WZ3_Code"].ne("").sum())

    complete_sites = int(
        site_analysis["Quellklassifikation_Status"].eq("vollstaendig").sum()
    )
    partial_sites = int(
        site_analysis["Quellklassifikation_Status"].eq("teilweise").sum()
    )
    unstarted_sites = int(
        site_analysis["Quellklassifikation_Status"].eq("nicht_begonnen").sum()
    )

    processed_sources = int(data["_Quelle_bearbeitet"].sum())
    q6_sources = int(data["_Quelle_Q6"].sum())
    q7_sources = int(data["_Quelle_Q7"].sum())
    qx_sources = int(data["_Quelle_QX"].sum())

    print()
    print("=" * 78)
    print("ERGEBNIS")
    print("=" * 78)
    print()
    print(f"Physische Standortprofile:                  {total_sites:,}")
    print(f"Davon zusammengeführte Verbünde:            {merged_sites:,}")
    print(f"Mit finalem Profil-WZ-Code:                  {profile_wz_sites:,}")
    print(f"Mit Profil-WZ3:                              {wz3_sites:,}")
    print()
    print(f"Quellenklassifikation vollständig:           {complete_sites:,}")
    print(f"Quellenklassifikation teilweise:             {partial_sites:,}")
    print(f"Quellenklassifikation noch nicht begonnen:   {unstarted_sites:,}")
    print()
    print(f"Klassifizierte Quellen gesamt:               {processed_sources:,}")
    print(f"Q6 Kälteanlagen-/Verflüssigerwärme:          {q6_sources:,}")
    print(f"Q7 Druckluft-/Kompressorwärme:                {q7_sources:,}")
    print(f"QX nicht eindeutig bestimmbar:               {qx_sources:,}")

    if processed_sources:
        print(
            f"QX-Anteil klassifizierte Quellen:            "
            f"{qx_sources / processed_sources:.2%}"
        )
        print(
            f"Q6+Q7-Anteil klassifizierte Quellen:         "
            f"{(q6_sources + q7_sources) / processed_sources:.2%}"
        )

    print()
    print(f"Finale Profil-WZ-Codes besetzt:               {len(industry_profile):,}")
    print(f"WZ3-Gruppen besetzt:                          {len(industry_wz3):,}")
    print(f"WZ2-Abteilungen besetzt:                      {len(industry_wz2):,}")
    print()
    print("Ausgabedatei:")
    print(OUTPUT_FILE.resolve())
    print()
    print("WICHTIG:")
    print("- Es wurden keine Standorte ausgeschlossen.")
    print("- Es wurden keine Branchen automatisch zusammengelegt.")
    print("- Q6 und Q7 zählen als gültige, zuordenbare Klassen.")
    print("- Nur QX zählt als nicht eindeutig bestimmbar.")
    print("- Die QX-Schwellenanalyse basiert auf dem Anteil an der Jahreswärmemenge.")
    print("- Für den 20-%-Standard werden vollständige numerische Jahreswärmemengen verlangt.")
    print("- Plausibilitätsauffällige positive Wärmemengen werden NICHT herausgefiltert.")
    print("- Der aktuelle QX20-Eignungsstatus ist deshalb ausdrücklich VORLÄUFIG.")
    print("- Verdächtige Zahlenwerte stehen separat im Blatt 'Zahlen_Plausibilitaet'.")
    print("- Leistungswerte bleiben erhalten, sind aber nur noch Diagnosefelder.")
    print("- Feinere Profil-WZ-Codes wie 38.21 bleiben separat erhalten.")
    print("=" * 78)


# ============================================================
# 17. PROGRAMMSTART
# ============================================================

if __name__ == "__main__":
    main()
