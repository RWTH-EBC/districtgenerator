from pathlib import Path
import re

import numpy as np
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# 1. KONFIGURATION
# ============================================================
#
# DIESER CODE ERSETZT DEN BISHERIGEN EXPORT-CODE
#
#     export_profilgruppen_qx20.py
#
# VOLLSTÄNDIG.
#
# Neue methodische Grundlage:
#
#   1. Manuelle Gruppierung aus:
#        - Zusammenlegung_gt2
#        - Zusammenlegung_bis2
#
#   2. Nur VORLÄUFIG geeignete physische Standorte:
#
#        Geeignet_QX_Waermemenge_max_20pct_vorlaeufig = True
#
#      also QX <= 20 % der JAHRESWÄRMEMENGE.
#
#   3. Plausibilitätsauffällige positive Wärmemengen werden
#      ausdrücklich NOCH NICHT ausgeschlossen oder korrigiert.
#
#   4. Von einem geeigneten Standort werden ALLE zugehörigen
#      Abwärmepotentiale exportiert, einschließlich einzelner
#      QX-Potentiale.
#
#   5. Pro Experimentgruppe entsteht genau EINE Exceldatei.
#
#   6. Die gemeldete maximale Leistung ist für die Auswahl der
#      Standorte NICHT relevant.
# ============================================================


BASE_DIR = Path("WZ2025_Vorbereitung")

SOURCE_DIR = (
        BASE_DIR
        / "Quellenklassifikation_final"
)


# ------------------------------------------------------------
# A) Standortanalyse aus dem neuen Wärmemengen-Postprocessing
# ------------------------------------------------------------

ANALYSIS_FILE = (
        SOURCE_DIR
        / "Branchenanalyse_Q1_Q7_v2.xlsx"
)

ANALYSIS_SHEET = "Standort_Analyse"

# Optionales Blatt mit bereits identifizierten auffälligen
# Zahlenwerten.
PLAUSIBILITY_SHEET = "Zahlen_Plausibilitaet"


# ------------------------------------------------------------
# B) Vollständige Potentialtabelle inklusive Q1-Q7/QX
# ------------------------------------------------------------

POTENTIAL_FILE = (
        SOURCE_DIR
        / "Abwaermepotentiale_mit_Quellklassen.xlsx"
)

POTENTIAL_SHEET = "Abwaermepotentiale"


# ------------------------------------------------------------
# C) Fertige manuelle Gruppierungsdatei
#
# Falls None:
# Der Code sucht automatisch die neueste passende Datei.
#
# Wenn du ganz sicher eine bestimmte Datei verwenden willst:
#
# GROUPING_FILE = SOURCE_DIR / "WZ_Gruppen_Zusammenlegung_Experiment.xlsx"
# ------------------------------------------------------------

GROUPING_FILE = None


GROUPING_FILE_CANDIDATES = [

    SOURCE_DIR
    / "WZ_Gruppen_Zusammenlegung_Experiment.xlsx",

    SOURCE_DIR
    / "WZ_Gruppen_Zusammenlegung_Experiment_GPT.xlsx",

    Path(
        "WZ_Gruppen_Zusammenlegung_Experiment.xlsx"
    ),

    Path(
        "WZ_Gruppen_Zusammenlegung_Experiment_GPT.xlsx"
    ),
    ]


GROUPING_FILE_GLOB = (
    "WZ_Gruppen_Zusammenlegung_Experiment*.xlsx"
)


# ------------------------------------------------------------
# Die manuelle Gruppierung wird DIREKT aus diesen beiden
# Eingabeblättern gelesen.
#
# Das Blatt "Gruppen_Ergebnis" wird absichtlich NICHT benötigt.
# Damit sind dessen Excel-Formeln für diesen Export irrelevant.
# ------------------------------------------------------------

GROUPING_SHEETS = [

    "Zusammenlegung_gt2",

    "Zusammenlegung_bis2",
]


# ------------------------------------------------------------
# Neue Ausgabe
#
# Eigener neuer Ordner, damit die früheren leistungsbasierten
# Exporte nicht mit den neuen Wärmemengen-Exporten vermischt
# werden.
# ------------------------------------------------------------

OUTPUT_DIR = (
        SOURCE_DIR
        / "Profilgruppen_Export_Waermemenge20_vorlaeufig"
)


# ------------------------------------------------------------
# Gruppendateien heißen deterministisch:
#
#     Profilgruppe_G01.xlsx
#     Profilgruppe_G02.xlsx
#     ...
#
# Dadurch entstehen beim erneuten Lauf nicht wegen veränderter
# Standortzahlen zusätzliche Dateivarianten.
# ------------------------------------------------------------

GROUP_FILE_PREFIX = "Profilgruppe_"


# ------------------------------------------------------------
# Falls True:
# Bereits vorhandene Profilgruppe_*.xlsx-Dateien im NEUEN
# Ausgabeordner werden vor dem Lauf gelöscht.
#
# _Export_Uebersicht.xlsx wird ebenfalls neu geschrieben.
#
# Der Code löscht NICHTS außerhalb von OUTPUT_DIR.
# ------------------------------------------------------------

CLEAN_OLD_GROUP_FILES = True


# ============================================================
# 2. METHODISCHER STANDARD
# ============================================================

QX_HEAT_THRESHOLD = 0.20


# ============================================================
# 3. ALLGEMEINE HILFSFUNKTIONEN
# ============================================================

def clean_string(
        value,
) -> str:
    """
    Vereinheitlicht leere Werte.
    """

    if pd.isna(
            value
    ):

        return ""


    text = str(
        value
    ).strip()


    if text.lower() in {

        "nan",
        "none",

    }:

        return ""


    return text


def normalize_wz_code(
        value,
) -> str:
    """
    Erhält die vorhandene WZ-Tiefe.

    Beispiele:

        38.21 -> 38.21
        38.2  -> 38.2
        6.1   -> 06.1
        05.1  -> 05.1
        0.5.1 -> 05.1
    """

    text = (

        clean_string(
            value
        )

        .replace(
            ",",
            ".",
        )

        .replace(
            " ",
            "",
        )
    )


    if not text:

        return ""


    malformed = re.fullmatch(

        r"0\.(\d)\.(\d+)",

        text,
    )


    if malformed:

        return (

            f"0"
            f"{malformed.group(1)}."
            f"{malformed.group(2)}"
        )


    match = re.match(

        r"^(\d{1,2})(.*)$",

        text,
    )


    if not match:

        return text


    return (

        f"{int(match.group(1)):02d}"
        f"{match.group(2)}"
    )


def normalize_group_id(
        value,
) -> str:
    """
    Beispiele:

        g01 -> G01
        G01 -> G01
    """

    return (

        clean_string(
            value
        )

        .upper()

        .replace(
            " ",
            "",
        )
    )


def bool_value(
        value,
) -> bool:
    """
    Robuste Umwandlung von Excel-/pandas-Werten in bool.
    """

    if isinstance(
            value,
            bool,
    ):

        return value


    if pd.isna(
            value
    ):

        return False


    if isinstance(
            value,
            (
                    int,
                    float,
                    np.integer,
                    np.floating,
            ),
    ):

        return bool(
            value
        )


    text = clean_string(
        value
    ).casefold()


    if text in {

        "true",
        "wahr",
        "1",
        "ja",
        "yes",
        "y",

    }:

        return True


    return False


def numeric_fraction(
        value,
):
    """
    Wandelt z. B.

        0.2
        20
        "20 %"
        "20,0 %"

    in 0..1 um.
    """

    text = clean_string(
        value
    )


    if not text:

        return np.nan


    has_percent = (
            "%"
            in text
    )


    text = (

        text

        .replace(
            "%",
            "",
        )

        .replace(
            " ",
            "",
        )

        .replace(
            ",",
            ".",
        )
    )


    try:

        number = float(
            text
        )


    except ValueError:

        return np.nan


    if has_percent:

        return (
                number
                / 100.0
        )


    if (
            1
            < number
            <= 100
    ):

        return (
                number
                / 100.0
        )


    return number


def numeric_series(
        series: pd.Series,
) -> pd.Series:
    """
    Robuste numerische Umwandlung für deutsche/englische
    Schreibweisen.

    Es wird NICHT korrigiert.
    """

    def convert(
            value,
    ):

        if pd.isna(
                value
        ):

            return np.nan


        if isinstance(
                value,
                (
                        int,
                        float,
                        np.integer,
                        np.floating,
                ),
        ):

            return float(
                value
            )


        text = clean_string(
            value
        )


        if not text:

            return np.nan


        text = text.replace(
            " ",
            "",
        )


        # 1.234,56
        if (
                ","
                in text
                and
                "."
                in text
        ):

            if (
                    text.rfind(
                        ","
                    )
                    >
                    text.rfind(
                        "."
                    )
            ):

                text = (

                    text

                    .replace(
                        ".",
                        "",
                    )

                    .replace(
                        ",",
                        ".",
                    )
                )


        # 1234,56
        elif (
                ","
                in text
        ):

            text = text.replace(
                ",",
                ".",
            )


        try:

            return float(
                text
            )


        except ValueError:

            return np.nan


    return series.map(
        convert
    )


def first_existing_column(
        df: pd.DataFrame,
        candidates,
        required=False,
        label="Spalte",
):
    """
    Gibt die erste tatsächlich vorhandene Spalte zurück.
    """

    for column in candidates:

        if column in df.columns:

            return column


    if required:

        raise ValueError(

            f"{label} wurde nicht gefunden.\n\n"

            "Gesuchte Kandidaten:\n"

            + "\n".join(

                f"- {column}"

                for column
                in candidates
            )
        )


    return None


def find_header_row(
        file_path,
        sheet_name,
        required_header,
        max_rows=25,
):
    """
    Die Gruppierungsdatei enthält Titelzeilen oberhalb der
    eigentlichen Kopfzeile.

    Diese Funktion sucht daher nach der Tabellenkopfzeile.
    """

    preview = pd.read_excel(

        file_path,

        sheet_name=
        sheet_name,

        header=None,

        nrows=
        max_rows,

        dtype=str,

        keep_default_na=False,
    )


    for row_index in range(
            len(
                preview
            )
    ):


        values = {

            clean_string(
                value
            )

            for value
            in preview.iloc[
                row_index
            ].tolist()
        }


        if (
                required_header
                in values
        ):

            return row_index


    return None


def safe_ratio(
        numerator,
        denominator,
):
    if (
            denominator is None
            or
            denominator == 0
            or
            pd.isna(
                denominator
            )
    ):

        return np.nan


    return (
            numerator
            / denominator
    )


# ============================================================
# 4. GRUPPIERUNGSDATEI FINDEN
# ============================================================

def find_grouping_file() -> Path:
    """
    Wenn GROUPING_FILE gesetzt ist:
        -> exakt diese Datei.

    Ansonsten:
        -> neueste passende Gruppierungsdatei.
    """

    if GROUPING_FILE is not None:


        path = Path(
            GROUPING_FILE
        )


        if not path.exists():

            raise FileNotFoundError(

                "Die explizit gesetzte Gruppierungsdatei "
                "wurde nicht gefunden:\n"

                f"{path}"
            )


        return path


    candidates = []


    for candidate in (

            GROUPING_FILE_CANDIDATES

    ):


        if (
                candidate.exists()
                and
                not candidate.name.startswith(
                    "~$"
                )
        ):

            candidates.append(
                candidate
            )


    for folder in [

        SOURCE_DIR,

        Path("."),

    ]:


        if not folder.exists():

            continue


        for file in folder.glob(

                GROUPING_FILE_GLOB
        ):


            if file.name.startswith(
                    "~$"
            ):

                continue


            if file not in candidates:

                candidates.append(
                    file
                )


    if not candidates:

        raise FileNotFoundError(

            "Keine Gruppierungsdatei gefunden.\n\n"

            "Bitte im Konfigurationsblock GROUPING_FILE "
            "direkt auf deine fertige Gruppierungsdatei setzen."
        )


    # Neueste Datei wählen.
    return max(

        candidates,

        key=
        lambda path:
        path.stat().st_mtime,
    )


# ============================================================
# 5. MANUELLE GRUPPENZUORDNUNG EINLESEN
# ============================================================

def load_group_assignments(
        grouping_file: Path,
) -> pd.DataFrame:
    """
    Liest ausschließlich die manuellen Eingabespalten aus:

        Zusammenlegung_gt2
        Zusammenlegung_bis2

    Das Formelergebnisblatt "Gruppen_Ergebnis" wird NICHT
    verwendet.

    Dadurch ist der Export unabhängig von möglichen #NAME?-
    oder anderen Excel-Formelfehlern im Ergebnisblatt.
    """

    excel = pd.ExcelFile(
        grouping_file
    )


    frames = []


    for sheet_name in (

            GROUPING_SHEETS

    ):


        if (
                sheet_name
                not in excel.sheet_names
        ):

            raise ValueError(

                f"Das benötigte Blatt '{sheet_name}' "
                "fehlt in der Gruppierungsdatei."
            )


        header_row = find_header_row(

            grouping_file,

            sheet_name,

            "Profil_WZ_Code",
        )


        if (
                header_row
                is None
        ):

            raise ValueError(

                f"Im Blatt '{sheet_name}' wurde "
                "die Kopfzeile mit 'Profil_WZ_Code' "
                "nicht gefunden."
            )


        table = pd.read_excel(

            grouping_file,

            sheet_name=
            sheet_name,

            header=
            header_row,

            dtype=str,

            keep_default_na=False,
        )


        required = {

            "Profil_WZ_Code",

            "Experiment_Gruppe",
        }


        missing = (

                required

                - set(
            table.columns
        )
        )


        if missing:

            raise ValueError(

                f"Im Blatt '{sheet_name}' fehlen:\n"

                + "\n".join(
                    sorted(
                        missing
                    )
                )
            )


        table[
            "Profil_WZ_Code"
        ] = (

            table[
                "Profil_WZ_Code"
            ]

            .map(
                normalize_wz_code
            )
        )


        table[
            "Experiment_Gruppe"
        ] = (

            table[
                "Experiment_Gruppe"
            ]

            .map(
                normalize_group_id
            )
        )


        # Nur tatsächlich zugeordnete WZ-Codes.
        table = table[

            table[
                "Profil_WZ_Code"
            ]

            .ne("")

            &

            table[
                "Experiment_Gruppe"
            ]

            .ne("")

            ].copy()


        if table.empty:

            continue


        keep = [

            column

            for column in [

                "Profil_WZ_Code",

                "Kurzname",

                "WZ3_Code",

                "WZ2_Code",

                "Standorte_gesamt",

                "Unternehmen_gesamt",

                "Standorte_QX_Waermemenge_max_20pct_vorlaeufig",

                "Unternehmen_QX_Waermemenge_max_20pct_vorlaeufig",

                "Anteil_Standorte_QX_Waermemenge_max_20pct_vorlaeufig",

                "Standorte_Zahlenpruefung_hoch_gesamt",

                "Standorte_Zahlenpruefung_hoch_unter_Geeigneten",

                "Anteil_geeignete_Standorte_mit_Zahlenpruefung_hoch",

                "Standorte_Eignung_kann_sich_aendern",

                "Experiment_Gruppe",

                "Kommentar",

            ]

            if column
               in table.columns
        ]


        table = table[
            keep
        ].copy()


        table[
            "Quellblatt_Gruppierung"
        ] = sheet_name


        frames.append(
            table
        )


    if not frames:

        raise ValueError(

            "In den beiden Zusammenlegungsblättern wurde "
            "keine Experimentgruppe vergeben."
        )


    assignments = pd.concat(

        frames,

        ignore_index=True,

        sort=False,
    )


    # ========================================================
    # KONSISTENZ:
    # Ein Profil_WZ_Code darf exakt einer Experimentgruppe
    # zugeordnet sein.
    # ========================================================

    conflict = (

        assignments

        .groupby(
            "Profil_WZ_Code"
        )[
            "Experiment_Gruppe"
        ]

        .nunique()

        .loc[
            lambda series:
            series > 1
        ]
    )


    if not conflict.empty:


        raise ValueError(

            "Mindestens ein Profil_WZ_Code ist mehreren "
            "Experimentgruppen zugeordnet:\n\n"

            + "\n".join(

                conflict
                .index
                .tolist()
            )
        )


    assignments = (

        assignments

        .sort_values(
            [
                "Experiment_Gruppe",

                "Profil_WZ_Code",
            ]
        )

        .drop_duplicates(
            subset=[
                "Profil_WZ_Code",
            ],

            keep=
            "first",
        )

        .reset_index(
            drop=True
        )
    )


    return assignments


# ============================================================
# 6. STANDORTANALYSE EINLESEN
# ============================================================

def load_site_analysis():
    """
    Liest die neue Standortanalyse.

    Hier wird ausdrücklich der neue WÄRMEMENGENSTANDARD verlangt.
    Ein alter leistungsbasierter Filter wird NICHT als Fallback
    akzeptiert.
    """

    if not ANALYSIS_FILE.exists():

        raise FileNotFoundError(

            "Branchenanalyse nicht gefunden:\n"

            f"{ANALYSIS_FILE}"
        )


    sites = pd.read_excel(

        ANALYSIS_FILE,

        sheet_name=
        ANALYSIS_SHEET,

        dtype=str,

        keep_default_na=False,
    )


    # --------------------------------------------------------
    # Physische Standort-ID:
    # Für diesen finalen Export ist Standortverbund_ID
    # verpflichtend.
    # --------------------------------------------------------

    site_id_column = first_existing_column(

        sites,

        [
            "Standortverbund_ID",
        ],

        required=True,

        label=
        "Standortverbund_ID",
    )


    profile_wz_column = first_existing_column(

        sites,

        [
            "Profil_WZ_Final_Code",

            "Profil_WZ_Code",
        ],

        required=True,

        label=
        "finaler Profil-WZ-Code",
    )


    wz3_column = first_existing_column(

        sites,

        [
            "Profil_WZ3_Code",

            "WZ3_Code",
        ],
    )


    wz2_column = first_existing_column(

        sites,

        [
            "Profil_WZ2_Code",

            "WZ2_Code",
        ],
    )


    qx_heat_column = first_existing_column(

        sites,

        [
            "QX_Anteil_Waermemenge_bearbeitet",

            "Waermengenanteil_QX_bearbeitet",
        ],

        required=True,

        label=
        "QX-Anteil an der Jahreswärmemenge",
    )


    suitable_column = first_existing_column(

        sites,

        [
            "Geeignet_QX_Waermemenge_max_20pct_vorlaeufig",

            "Geeignet_QX_Waermemenge_max_20pct",
        ],

        required=True,

        label=
        "vorläufiger Eignungsstatus nach Wärmemengenstandard",
    )


    company_column = first_existing_column(

        sites,

        [
            "_Company_Key",

            "Unternehmens_ID",

            "Profil_Unternehmens_ID",

            "Profil_Firmenname",

            "Firmenname",
        ],
    )


    quality_status_column = first_existing_column(

        sites,

        [
            "Zahlen_Qualitaetsstatus",
        ],
    )


    hard_quality_column = first_existing_column(

        sites,

        [
            "Anzahl_Quellen_Zahlenpruefung_hoch",
        ],
    )


    change_column = first_existing_column(

        sites,

        [
            "QX20_Eignung_kann_sich_durch_Zahlenpruefung_aendern",
        ],
    )


    # ========================================================
    # INTERNE SCHLÜSSEL
    # ========================================================

    sites[
        "_Standort_ID"
    ] = (

        sites[
            site_id_column
        ]

        .map(
            clean_string
        )
    )


    sites[
        "_Profil_WZ_Code"
    ] = (

        sites[
            profile_wz_column
        ]

        .map(
            normalize_wz_code
        )
    )


    if wz3_column:


        sites[
            "_Profil_WZ3_Code"
        ] = (

            sites[
                wz3_column
            ]

            .map(
                normalize_wz_code
            )
        )


    else:


        sites[
            "_Profil_WZ3_Code"
        ] = ""


    if wz2_column:


        sites[
            "_Profil_WZ2_Code"
        ] = (

            sites[
                wz2_column
            ]

            .map(
                normalize_wz_code
            )
        )


    else:


        sites[
            "_Profil_WZ2_Code"
        ] = ""


    sites[
        "_QX_Waermengenanteil"
    ] = (

        sites[
            qx_heat_column
        ]

        .map(
            numeric_fraction
        )
    )


    sites[
        "_Geeignet_vorlaeufig"
    ] = (

        sites[
            suitable_column
        ]

        .map(
            bool_value
        )
    )


    if company_column:


        sites[
            "_Unternehmen_Key"
        ] = (

            sites[
                company_column
            ]

            .map(
                clean_string
            )
        )


    else:


        sites[
            "_Unternehmen_Key"
        ] = (

                "SITE::"

                + sites[
                    "_Standort_ID"
                ]
        )


    if quality_status_column:


        sites[
            "_Zahlen_Qualitaetsstatus"
        ] = (

            sites[
                quality_status_column
            ]

            .map(
                clean_string
            )
        )


    else:


        sites[
            "_Zahlen_Qualitaetsstatus"
        ] = ""


    if hard_quality_column:


        hard_numbers = (

            numeric_series(

                sites[
                    hard_quality_column
                ]
            )

            .fillna(
                0
            )
        )


        sites[
            "_Zahlenpruefung_hoch"
        ] = (

                hard_numbers
                > 0
        )


    else:


        sites[
            "_Zahlenpruefung_hoch"
        ] = False


    if change_column:


        sites[
            "_Eignung_kann_sich_aendern"
        ] = (

            sites[
                change_column
            ]

            .map(
                bool_value
            )
        )


    else:


        # Konservativer Fallback.
        sites[
            "_Eignung_kann_sich_aendern"
        ] = (

            sites[
                "_Zahlenpruefung_hoch"
            ]
        )


    # ========================================================
    # VALIDIERUNG
    # ========================================================

    if (

            sites[
                "_Standort_ID"
            ]

                    .eq("")

                    .any()
    ):

        raise ValueError(

            "Mindestens ein Standort besitzt keine "
            "Standortverbund_ID."
        )


    if not sites[
        "_Standort_ID"
    ].is_unique:

        raise ValueError(

            "Im Blatt Standort_Analyse ist Standortverbund_ID "
            "nicht eindeutig."
        )


    # --------------------------------------------------------
    # Sicherheitsprüfung:
    # Eignungsflag und QX-Anteil sollen grundsätzlich zum
    # neuen 20-%-Standard passen.
    #
    # Auffällige positive Wärmemengen sind hier KEIN Fehler.
    # --------------------------------------------------------

    mismatch = (

            sites[
                "_Geeignet_vorlaeufig"
            ]

            &

            sites[
                "_QX_Waermengenanteil"
            ].notna()

            &

            (
                    sites[
                        "_QX_Waermengenanteil"
                    ]
                    >
                    QX_HEAT_THRESHOLD
                    + 1e-9
            )
    )


    if mismatch.any():

        print()

        print(
            "WARNUNG:"
        )

        print(

            f"{int(mismatch.sum())} Standorte sind laut "
            "Eignungsflag geeignet, obwohl ihr gespeicherter "
            "QX-Wärmemengenanteil >20 % ist."
        )

        print(

            "Der Export folgt dem expliziten Eignungsflag aus "
            "der neuen Branchenanalyse."
        )

        print()


    return sites


# ============================================================
# 7. OPTIONAL: ZAHLEN-PLAUSIBILITÄTSBLATT EINLESEN
# ============================================================

def load_plausibility_table():
    """
    Falls das Blatt Zahlen_Plausibilitaet vorhanden ist,
    wird es mit in die Gruppendateien übernommen.

    Es dient nur der späteren Prüfung und beeinflusst die
    aktuelle Standortauswahl NICHT.
    """

    try:

        excel = pd.ExcelFile(
            ANALYSIS_FILE
        )


        if (
                PLAUSIBILITY_SHEET
                not in excel.sheet_names
        ):

            return pd.DataFrame()


        table = pd.read_excel(

            ANALYSIS_FILE,

            sheet_name=
            PLAUSIBILITY_SHEET,

            dtype=str,

            keep_default_na=False,
        )


        return table


    except Exception as exc:


        print()

        print(
            "WARNUNG:"
        )

        print(

            "Das Blatt Zahlen_Plausibilitaet konnte nicht "
            "eingelesen werden."
        )

        print(
            exc
        )

        print()


        return pd.DataFrame()


# ============================================================
# 8. ABWÄRMEPOTENTIALE EINLESEN
# ============================================================

def load_potentials():
    """
    Liest alle Originalpotentiale inklusive Quellklasse.

    Für den aktuellen finalen Workflow ist Standortverbund_ID
    verpflichtend.
    """

    if not POTENTIAL_FILE.exists():

        raise FileNotFoundError(

            "Datei mit Abwärmepotentialen nicht gefunden:\n"

            f"{POTENTIAL_FILE}"
        )


    potentials = pd.read_excel(

        POTENTIAL_FILE,

        sheet_name=
        POTENTIAL_SHEET,

        dtype=str,

        keep_default_na=False,
    )


    site_id_column = first_existing_column(

        potentials,

        [
            "Standortverbund_ID",
        ],

        required=True,

        label=
        "Standortverbund_ID in der Potentialtabelle",
    )


    potentials[
        "_Standort_ID"
    ] = (

        potentials[
            site_id_column
        ]

        .map(
            clean_string
        )
    )


    return potentials


# ============================================================
# 9. MANUELLE GRUPPIERUNG AUF DIE STANDORTE ANWENDEN
# ============================================================

def build_suitable_group_sites(
        assignments,
        sites,
):
    """
    Reihenfolge:

        Profil_WZ_Code
             ↓
        manuelle Gxx-Zuordnung
             ↓
        nur Standorte mit diesem Profil-WZ
             ↓
        Geeignet_QX_Waermemenge_max_20pct_vorlaeufig
             ↓
        Export

    WICHTIG:
    Zahlenauffälligkeiten sind KEIN Ausschlusskriterium.
    """

    assignment_map = dict(

        zip(

            assignments[
                "Profil_WZ_Code"
            ],

            assignments[
                "Experiment_Gruppe"
            ],
        )
    )


    result = sites.copy()


    result[
        "_Experiment_Gruppe"
    ] = (

        result[
            "_Profil_WZ_Code"
        ]

        .map(
            assignment_map
        )

        .fillna("")
    )


    # Nur WZ-Codes, die der Benutzer tatsächlich einer
    # Experimentgruppe zugeordnet hat.
    result = result[

        result[
            "_Experiment_Gruppe"
        ]

        .ne("")

    ].copy()


    # ========================================================
    # DER EINZIGE EIGNUNGSFILTER:
    #
    # vorläufige Eignung nach JAHRESWÄRMEMENGE.
    #
    # KEIN Leistungsfilter.
    # KEIN Ausschluss wegen Zahlenplausibilitätsflag.
    # ========================================================

    result = result[

        result[
            "_Geeignet_vorlaeufig"
        ]

    ].copy()


    return result


# ============================================================
# 10. QUELLKLASSENÜBERSICHT NACH JAHRESWÄRMEMENGE
# ============================================================

def build_source_summary(
        group_potentials,
        flagged_original_ids=None,
):
    """
    Primäre Profilkennzahl ist jetzt die Jahreswärmemenge.

    Die gemeldete Leistung wird in dieser Übersicht bewusst
    NICHT aggregiert.
    """

    if (
            "Quellklasse_Code"
            not in group_potentials.columns
    ):

        return pd.DataFrame()


    heat_column = first_existing_column(

        group_potentials,

        [
            "Wärmemenge pro Jahr (in kWh/a)",

            "Waermemenge_pro_Jahr_kWh_a",

            "AnnualHeatAmount_kWh",
        ],
    )


    original_id_column = first_existing_column(

        group_potentials,

        [
            "Originaldatensatz_ID",
        ],
    )


    work = group_potentials.copy()


    if heat_column:


        work[
            "_Waermemenge_num"
        ] = numeric_series(

            work[
                heat_column
            ]
        )


    else:


        work[
            "_Waermemenge_num"
        ] = np.nan


    if flagged_original_ids is None:

        flagged_original_ids = set()


    if original_id_column:


        work[
            "_Plausibilitaetsauffaellig"
        ] = (

            work[
                original_id_column
            ]

            .map(
                clean_string
            )

            .isin(
                flagged_original_ids
            )
        )


    else:


        work[
            "_Plausibilitaetsauffaellig"
        ] = False


    rows = []


    class_order = [

        "Q1",

        "Q2",

        "Q3",

        "Q4",

        "Q5",

        "Q6",

        "Q7",

        "QX",
    ]


    for code in class_order:


        subset = work[

            work[
                "Quellklasse_Code"
            ]

            .map(
                clean_string
            )

            .str.upper()

            .eq(
                code
            )

        ]


        row = {

            "Quellklasse_Code":
                code,

            "Anzahl_Potentiale":
                int(
                    len(
                        subset
                    )
                ),

            "Anzahl_Plausibilitaetsauffaellige_Potentiale":
                int(

                    subset[
                        "_Plausibilitaetsauffaellig"
                    ]

                    .sum()
                ),

            "Waermemenge_kWh_a_gemeldet":
                subset[
                    "_Waermemenge_num"
                ].sum(
                    min_count=1
                ),
        }


        if (
                "Quellklasse"
                in subset.columns
        ):


            labels = (

                subset[
                    "Quellklasse"
                ]

                .map(
                    clean_string
                )

                .loc[
                    lambda series:
                    series.ne("")
                ]
            )


            row[
                "Quellklasse"
            ] = (

                labels.iloc[0]

                if len(
                    labels
                )

                else ""
            )


        rows.append(
            row
        )


    summary = pd.DataFrame(
        rows
    )


    total_heat = summary[
        "Waermemenge_kWh_a_gemeldet"
    ].sum(
        min_count=1
    )


    if (
            pd.notna(
                total_heat
            )
            and
            total_heat > 0
    ):


        summary[
            "Anteil_Waermemenge_gemeldet"
        ] = (

                summary[
                    "Waermemenge_kWh_a_gemeldet"
                ]

                / total_heat
        )


    else:


        summary[
            "Anteil_Waermemenge_gemeldet"
        ] = np.nan


    preferred = [

        "Quellklasse_Code",

    ]


    if (
            "Quellklasse"
            in summary.columns
    ):

        preferred.append(
            "Quellklasse"
        )


    preferred += [

        "Anzahl_Potentiale",

        "Anzahl_Plausibilitaetsauffaellige_Potentiale",

        "Waermemenge_kWh_a_gemeldet",

        "Anteil_Waermemenge_gemeldet",
    ]


    return summary[
        preferred
    ]


# ============================================================
# 11. EXCEL-FORMATIERUNG
# ============================================================

HEADER_FILL = PatternFill(

    fill_type=
    "solid",

    fgColor=
    "1F4E78",
)


HEADER_FONT = Font(

    color=
    "FFFFFF",

    bold=
    True,
)


WARNING_FILL = PatternFill(

    fill_type=
    "solid",

    fgColor=
    "FCE4D6",
)


def format_sheet(
        worksheet,
        freeze="A2",
        max_width=42,
):
    """
    Einfache, robuste Tabellenformatierung.
    """

    if (
            worksheet.max_row
            < 1
            or
            worksheet.max_column
            < 1
    ):

        return


    worksheet.freeze_panes = freeze


    # Header
    for cell in worksheet[
        1
    ]:


        cell.fill = HEADER_FILL

        cell.font = HEADER_FONT

        cell.alignment = Alignment(

            horizontal=
            "center",

            vertical=
            "center",

            wrap_text=
            True,
        )


    # Autofilter
    if (
            worksheet.max_row
            >= 2
    ):

        worksheet.auto_filter.ref = (

            f"A1:"
            f"{get_column_letter(worksheet.max_column)}"
            f"{worksheet.max_row}"
        )


    # Spaltenbreiten
    for column_index in range(

            1,

            worksheet.max_column + 1,
    ):


        letter = get_column_letter(
            column_index
        )


        max_length = 0


        for row_index in range(

                1,

                min(
                    worksheet.max_row,
                    250,
                ) + 1,
        ):


            value = worksheet.cell(

                row=
                row_index,

                column=
                column_index,

            ).value


            if value is None:

                continue


            max_length = max(

                max_length,

                len(
                    str(
                        value
                    )
                ),
            )


        width = min(

            max(

                10,

                max_length + 2,
                ),

            max_width,
        )


        worksheet.column_dimensions[
            letter
        ].width = width


    # Body
    for row in worksheet.iter_rows(

            min_row=2,

            max_row=
            worksheet.max_row,
    ):


        for cell in row:


            cell.alignment = Alignment(

                vertical=
                "top",

                wrap_text=
                True,
            )


def format_key_value_sheet(
        worksheet,
):
    """
    Gruppeninfo.
    """

    worksheet.freeze_panes = "A2"


    worksheet.column_dimensions[
        "A"
    ].width = 45


    worksheet.column_dimensions[
        "B"
    ].width = 90


    # Kopfzeile
    for cell in worksheet[
        1
    ]:


        cell.fill = HEADER_FILL

        cell.font = HEADER_FONT


    for row in range(

            2,

            worksheet.max_row + 1,
    ):


        worksheet[
            f"A{row}"
        ].font = Font(
            bold=True
        )


        worksheet[
            f"A{row}"
        ].alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )


        worksheet[
            f"B{row}"
        ].alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )


def apply_percent_formats(
        worksheet,
):
    """
    Formatiert Anteil-/QX-Spalten als Prozent.
    """

    if (
            worksheet.max_row
            < 2
    ):

        return


    headers = {

        clean_string(
            cell.value
        ):
            cell.column

        for cell
        in worksheet[
            1
        ]

        if clean_string(
            cell.value
        )
    }


    for header, column in headers.items():


        if (

                "Anteil"
                in header

                or

                "QX_Waermengenanteil"
                in header

        ):


            for row in range(

                    2,

                    worksheet.max_row + 1,
            ):


                worksheet.cell(

                    row=
                    row,

                    column=
                    column,

                ).number_format = (
                    "0.0%"
                )


# ============================================================
# 12. GRUPPENINFO ERSTELLEN
# ============================================================

def build_group_info(
        group_id,
        group_assignments,
        group_sites,
        group_potentials,
        group_plausibility,
):
    """
    Erstellt eine kompakte Dokumentation der Gruppe.
    """

    profile_codes_assigned = sorted(

        set(

            group_assignments[
                "Profil_WZ_Code"
            ]
        )
    )


    profile_codes_present = sorted(

        set(

            group_sites[
                "_Profil_WZ_Code"
            ]
        )
    )


    wz3_codes_present = sorted(

        {

            code

            for code
            in group_sites[
            "_Profil_WZ3_Code"
        ]

            if code
        }
    )


    companies = (

        group_sites[
            "_Unternehmen_Key"
        ]

        .replace(
            "",
            pd.NA,
        )

        .dropna()

        .nunique()
    )


    heat_column = first_existing_column(

        group_potentials,

        [
            "Wärmemenge pro Jahr (in kWh/a)",

            "Waermemenge_pro_Jahr_kWh_a",

            "AnnualHeatAmount_kWh",
        ],
    )


    total_heat = np.nan

    qx_heat = np.nan

    qx_share_group = np.nan


    if heat_column:


        heat_num = numeric_series(

            group_potentials[
                heat_column
            ]
        )


        total_heat = heat_num.sum(
            min_count=1
        )


        if (
                "Quellklasse_Code"
                in group_potentials.columns
        ):


            qx_mask = (

                group_potentials[
                    "Quellklasse_Code"
                ]

                .map(
                    clean_string
                )

                .str.upper()

                .eq(
                    "QX"
                )
            )


            qx_heat = (

                heat_num[
                    qx_mask
                ]

                .sum(
                    min_count=1
                )
            )


            qx_share_group = safe_ratio(

                qx_heat,

                total_heat,
            )


    high_quality_sites = int(

        group_sites[
            "_Zahlenpruefung_hoch"
        ]

        .sum()
    )


    changing_sites = int(

        group_sites[
            "_Eignung_kann_sich_aendern"
        ]

        .sum()
    )


    info = [

        [
            "Experiment_Gruppe",
            group_id,
        ],

        [
            "Eignungsstandard",
            (
                "VORLÄUFIG: QX <= 20 % der gemeldeten "
                "Jahreswärmemenge"
            ),
        ],

        [
            "Umgang_mit_Zahlenauffaelligkeiten",
            (
                "Plausibilitätsauffällige positive Wärmemengen "
                "bleiben enthalten und werden nicht automatisch "
                "korrigiert oder gefiltert."
            ),
        ],

        [
            "Anzahl_geeignete_Standorte_vorlaeufig",
            int(
                len(
                    group_sites
                )
            ),
        ],

        [
            "Anzahl_Unternehmen",
            int(
                companies
            ),
        ],

        [
            "Anzahl_Abwaermepotentiale",
            int(
                len(
                    group_potentials
                )
            ),
        ],

        [
            "Anzahl_Potentiale_in_Zahlen_Plausibilitaet",
            int(
                len(
                    group_plausibility
                )
            ),
        ],

        [
            "Anzahl_Standorte_Zahlenpruefung_hoch",
            high_quality_sites,
        ],

        [
            "Anzahl_Standorte_Eignung_kann_sich_aendern",
            changing_sites,
        ],

        [
            "Zugeordnete_Profil_WZ_Codes",
            ", ".join(
                profile_codes_assigned
            ),
        ],

        [
            "Tatsaechlich_vertretene_Profil_WZ_Codes",
            ", ".join(
                profile_codes_present
            ),
        ],

        [
            "Tatsaechlich_vertretene_WZ3_Codes",
            ", ".join(
                wz3_codes_present
            ),
        ],

        [
            "Summe_gemeldete_Waermemenge_kWh_a",
            total_heat,
        ],

        [
            "Summe_QX_Waermemenge_kWh_a",
            qx_heat,
        ],

        [
            "Aggregierter_QX_Waermengenanteil_der_Gruppe",
            qx_share_group,
        ],
    ]


    return pd.DataFrame(

        info,

        columns=[
            "Kennzahl",
            "Wert",
        ],
    )


# ============================================================
# 13. EINE GRUPPENDATEI ERSTELLEN
# ============================================================

def export_one_group(
        group_id,
        group_assignments,
        group_sites,
        potentials,
        plausibility_table,
):
    """
    Exportiert eine einzelne Experimentgruppe.
    """

    # ========================================================
    # 1. GEEIGNETE STANDORT-IDs
    # ========================================================

    site_ids = set(

        group_sites[
            "_Standort_ID"
        ]
    )


    if not site_ids:

        return None


    # ========================================================
    # 2. ALLE POTENTIALE DIESER STANDORTE
    # ========================================================

    group_potentials = potentials[

        potentials[
            "_Standort_ID"
        ]

        .isin(
            site_ids
        )

    ].copy()


    # ========================================================
    # 3. VOLLSTÄNDIGKEITSPRÜFUNG
    # ========================================================

    found_site_ids = set(

        group_potentials[
            "_Standort_ID"
        ]
    )


    missing_site_ids = (

            site_ids

            - found_site_ids
    )


    if missing_site_ids:

        raise ValueError(

            f"Gruppe {group_id}: "
            f"{len(missing_site_ids)} geeignete Standorte "
            "wurden in der Potentialtabelle nicht gefunden.\n\n"

            "Beispiele:\n"

            + "\n".join(

                sorted(
                    missing_site_ids
                )[:20]
            )
        )


    # ========================================================
    # 4. STANDORTINFORMATIONEN AUF DIE QUELLEN MAPPEN
    # ========================================================

    site_meta = (

        group_sites

        .set_index(
            "_Standort_ID"
        )
    )


    group_potentials[
        "Experiment_Gruppe"
    ] = group_id


    group_potentials[
        "Gruppen_Profil_WZ_Code"
    ] = (

        group_potentials[
            "_Standort_ID"
        ]

        .map(

            site_meta[
                "_Profil_WZ_Code"
            ]
        )
    )


    group_potentials[
        "Gruppen_WZ3_Code"
    ] = (

        group_potentials[
            "_Standort_ID"
        ]

        .map(

            site_meta[
                "_Profil_WZ3_Code"
            ]
        )
    )


    group_potentials[
        "Gruppen_WZ2_Code"
    ] = (

        group_potentials[
            "_Standort_ID"
        ]

        .map(

            site_meta[
                "_Profil_WZ2_Code"
            ]
        )
    )


    group_potentials[
        "Standort_QX_Waermengenanteil"
    ] = (

        group_potentials[
            "_Standort_ID"
        ]

        .map(

            site_meta[
                "_QX_Waermengenanteil"
            ]
        )
    )


    group_potentials[
        "Standort_Geeignet_Waermemenge20_vorlaeufig"
    ] = True


    group_potentials[
        "Standort_Zahlenpruefung_hoch"
    ] = (

        group_potentials[
            "_Standort_ID"
        ]

        .map(

            site_meta[
                "_Zahlenpruefung_hoch"
            ]
        )
    )


    group_potentials[
        "Standort_Eignung_kann_sich_durch_Zahlenpruefung_aendern"
    ] = (

        group_potentials[
            "_Standort_ID"
        ]

        .map(

            site_meta[
                "_Eignung_kann_sich_aendern"
            ]
        )
    )


    group_potentials[
        "Standort_Zahlen_Qualitaetsstatus"
    ] = (

        group_potentials[
            "_Standort_ID"
        ]

        .map(

            site_meta[
                "_Zahlen_Qualitaetsstatus"
            ]
        )
    )


    # Interne Hilfsspalte nicht exportieren.
    group_potentials = group_potentials.drop(

        columns=[
            "_Standort_ID",
        ],

        errors="ignore",
    )


    # Neue methodische Spalten ganz nach vorne.
    front_columns = [

        "Experiment_Gruppe",

        "Gruppen_Profil_WZ_Code",

        "Gruppen_WZ3_Code",

        "Gruppen_WZ2_Code",

        "Standort_QX_Waermengenanteil",

        "Standort_Geeignet_Waermemenge20_vorlaeufig",

        "Standort_Zahlenpruefung_hoch",

        "Standort_Eignung_kann_sich_durch_Zahlenpruefung_aendern",

        "Standort_Zahlen_Qualitaetsstatus",
    ]


    remaining = [

        column

        for column
        in group_potentials.columns

        if column
           not in front_columns
    ]


    group_potentials = group_potentials[

        front_columns

        + remaining
        ]


    # ========================================================
    # 5. STANDORTPROFILE
    # ========================================================

    site_export = group_sites.copy()


    site_export[
        "Experiment_Gruppe"
    ] = group_id


    # Interne technische Hilfsspalten entfernen.
    internal_columns = [

        column

        for column
        in site_export.columns

        if column.startswith(
            "_"
        )
    ]


    site_export = site_export.drop(

        columns=
        internal_columns,

        errors=
        "ignore",
    )


    # Experimentgruppe nach vorne.
    site_export = site_export[

        [
            "Experiment_Gruppe"
        ]

        + [

            column

            for column
            in site_export.columns

            if column
               != "Experiment_Gruppe"
        ]
        ]


    # ========================================================
    # 6. PLAUSIBILITÄTSTABELLE DIESER GRUPPE
    # ========================================================

    group_plausibility = pd.DataFrame()


    if not plausibility_table.empty:


        plausibility_site_column = first_existing_column(

            plausibility_table,

            [
                "Standortverbund_ID",
            ],
        )


        if plausibility_site_column:


            group_plausibility = plausibility_table[

                plausibility_table[
                    plausibility_site_column
                ]

                .map(
                    clean_string
                )

                .isin(
                    site_ids
                )

            ].copy()


            if not group_plausibility.empty:


                group_plausibility.insert(

                    0,

                    "Experiment_Gruppe",

                    group_id,
                )


    # ========================================================
    # 7. QUELLKLASSENÜBERSICHT
    # ========================================================

    flagged_original_ids = set()


    if not group_plausibility.empty:


        plausibility_original_id_column = first_existing_column(

            group_plausibility,

            [
                "Originaldatensatz_ID",
            ],
        )


        if plausibility_original_id_column:


            flagged_original_ids = set(

                group_plausibility[
                    plausibility_original_id_column
                ]

                .map(
                    clean_string
                )

                .loc[
                    lambda series:
                    series.ne("")
                ]
            )


    source_summary = build_source_summary(

        group_potentials,

        flagged_original_ids=
        flagged_original_ids,
    )


    # ========================================================
    # 8. GRUPPENINFO
    # ========================================================

    group_info = build_group_info(

        group_id,

        group_assignments,

        group_sites,

        group_potentials,

        group_plausibility,
    )


    # ========================================================
    # 9. DATEINAME
    # ========================================================

    output_file = (

            OUTPUT_DIR

            / (
                f"{GROUP_FILE_PREFIX}"
                f"{group_id}.xlsx"
            )
    )


    # ========================================================
    # 10. EXCEL SCHREIBEN
    # ========================================================

    with pd.ExcelWriter(

            output_file,

            engine=
            "openpyxl",

    ) as writer:


        group_info.to_excel(

            writer,

            sheet_name=
            "Gruppeninfo",

            index=False,
        )


        group_assignments.to_excel(

            writer,

            sheet_name=
            "WZ_Zuordnung",

            index=False,
        )


        site_export.to_excel(

            writer,

            sheet_name=
            "Standortprofile",

            index=False,
        )


        group_potentials.to_excel(

            writer,

            sheet_name=
            "Abwaermepotentiale",

            index=False,
        )


        source_summary.to_excel(

            writer,

            sheet_name=
            "Quellklassen_Uebersicht",

            index=False,
        )


        if not group_plausibility.empty:


            group_plausibility.to_excel(

                writer,

                sheet_name=
                "Zahlen_Plausibilitaet",

                index=False,
            )


    # ========================================================
    # 11. FORMATIEREN
    # ========================================================

    workbook = load_workbook(
        output_file
    )


    format_key_value_sheet(

        workbook[
            "Gruppeninfo"
        ]
    )


    for sheet_name in [

        "WZ_Zuordnung",

        "Standortprofile",

        "Abwaermepotentiale",

        "Quellklassen_Uebersicht",

        "Zahlen_Plausibilitaet",

    ]:


        if (
                sheet_name
                not in workbook.sheetnames
        ):

            continue


        worksheet = workbook[
            sheet_name
        ]


        format_sheet(
            worksheet
        )


        apply_percent_formats(
            worksheet
        )


    # Gruppeninfo:
    # aggregierter QX-Anteil als Prozent.
    group_info_sheet = workbook[
        "Gruppeninfo"
    ]


    for row in range(

            2,

            group_info_sheet.max_row + 1,
    ):


        key = clean_string(

            group_info_sheet[
                f"A{row}"
            ].value
        )


        if (
                key
                ==
                "Aggregierter_QX_Waermengenanteil_der_Gruppe"
        ):


            group_info_sheet[
                f"B{row}"
            ].number_format = "0.0%"


    workbook.save(
        output_file
    )


    return {

        "Experiment_Gruppe":
            group_id,

        "Datei":
            output_file.name,

        "Geeignete_Standorte_vorlaeufig":
            int(
                len(
                    group_sites
                )
            ),

        "Unternehmen":
            int(

                group_sites[
                    "_Unternehmen_Key"
                ]

                .replace(
                    "",
                    pd.NA,
                )

                .dropna()

                .nunique()
            ),

        "Abwaermepotentiale":
            int(
                len(
                    group_potentials
                )
            ),

        "Standorte_Zahlenpruefung_hoch":
            int(

                group_sites[
                    "_Zahlenpruefung_hoch"
                ]

                .sum()
            ),

        "Standorte_Eignung_kann_sich_aendern":
            int(

                group_sites[
                    "_Eignung_kann_sich_aendern"
                ]

                .sum()
            ),

        "Plausibilitaetsauffaellige_Potentiale":
            int(
                len(
                    group_plausibility
                )
            ),

        "Profil_WZ_Codes":
            ", ".join(

                sorted(

                    set(

                        group_sites[
                            "_Profil_WZ_Code"
                        ]
                    )
                )
            ),

        "WZ3_Codes":
            ", ".join(

                sorted(

                    {

                        value

                        for value
                        in group_sites[
                        "_Profil_WZ3_Code"
                    ]

                        if value
                    }
                )
            ),
    }


# ============================================================
# 14. EXPORTÜBERSICHT ERSTELLEN
# ============================================================

def write_export_overview(
        export_rows,
        grouping_file,
):
    """
    Kleine Kontroll-Datei im Ausgabeordner.

    Sie ist KEINE Profilgruppe und heißt deshalb nicht
    Profilgruppe_*.xlsx.
    """

    if not export_rows:

        return


    overview = pd.DataFrame(
        export_rows
    )


    overview = overview.sort_values(

        [
            "Geeignete_Standorte_vorlaeufig",

            "Experiment_Gruppe",
        ],

        ascending=[
            False,
            True,
        ],

        kind="stable",
    )


    method = pd.DataFrame(

        [

            [
                "Gruppierungsdatei",
                str(
                    grouping_file.resolve()
                ),
            ],

            [
                "Eignungsstandard",
                (
                    "Geeignet_QX_Waermemenge_max_20pct_vorlaeufig"
                ),
            ],

            [
                "QX_Grenze",
                QX_HEAT_THRESHOLD,
            ],

            [
                "Zahlenauffaellige_Waermemengen",
                (
                    "noch nicht automatisch gefiltert oder korrigiert"
                ),
            ],

            [
                "Leistung_fuer_Eignung_verwendet",
                "NEIN",
            ],

        ],

        columns=[
            "Kennzahl",
            "Wert",
        ],
    )


    output_file = (

            OUTPUT_DIR
            / "_Export_Uebersicht.xlsx"
    )


    with pd.ExcelWriter(

            output_file,

            engine=
            "openpyxl",

    ) as writer:


        overview.to_excel(

            writer,

            sheet_name=
            "Gruppen",

            index=False,
        )


        method.to_excel(

            writer,

            sheet_name=
            "Methodik",

            index=False,
        )


    workbook = load_workbook(
        output_file
    )


    format_sheet(
        workbook[
            "Gruppen"
        ]
    )


    format_key_value_sheet(
        workbook[
            "Methodik"
        ]
    )


    workbook[
        "Methodik"
    ][
        "B3"
    ].number_format = "0.0%"


    workbook.save(
        output_file
    )


# ============================================================
# 15. HAUPTPROGRAMM
# ============================================================

def main():


    print()

    print(
        "=" * 80
    )

    print(
        "EXPORT DER PROFILGRUPPEN – "
        "VORLÄUFIGER 20-%-STANDARD NACH JAHRESWÄRMEMENGE"
    )

    print(
        "=" * 80
    )

    print()


    # ========================================================
    # 1. GRUPPIERUNG
    # ========================================================

    grouping_file = (
        find_grouping_file()
    )


    print(
        "1. Verwendete Gruppierungsdatei:"
    )


    print(
        f"   {grouping_file.resolve()}"
    )


    assignments = (
        load_group_assignments(
            grouping_file
        )
    )


    print(

        f"   -> {len(assignments):,} "
        "manuell zugeordnete Profil-WZ-Codes"
    )


    print(

        f"   -> "
        f"{assignments['Experiment_Gruppe'].nunique():,} "
        "Experimentgruppen"
    )


    # ========================================================
    # 2. STANDORTANALYSE
    # ========================================================

    print()

    print(
        "2. Neue Standortanalyse wird eingelesen ..."
    )


    sites = load_site_analysis()


    print(

        f"   -> {len(sites):,} "
        "physische Standorte insgesamt"
    )


    print(

        "   -> vorläufig geeignet nach Wärmemengenstandard: "

        f"{int(sites['_Geeignet_vorlaeufig'].sum()):,}"
    )


    # ========================================================
    # 3. PLAUSIBILITÄTSTABELLE
    # ========================================================

    print()

    print(
        "3. Zahlen-Plausibilitätsinformationen werden eingelesen ..."
    )


    plausibility_table = (
        load_plausibility_table()
    )


    print(

        f"   -> {len(plausibility_table):,} "
        "auffällige Quellen in Zahlen_Plausibilitaet"
    )


    # ========================================================
    # 4. POTENTIALE
    # ========================================================

    print()

    print(
        "4. Vollständige Abwärmepotentiale werden eingelesen ..."
    )


    potentials = load_potentials()


    print(

        f"   -> {len(potentials):,} "
        "Abwärmepotentiale"
    )


    # ========================================================
    # 5. MANUELLE GRUPPEN + VORLÄUFIGE EIGNUNG
    # ========================================================

    print()

    print(
        "5. Manuelle Gxx-Zuordnung wird mit dem neuen "
        "Wärmemengen-Eignungsstatus verknüpft ..."
    )


    suitable_group_sites = (
        build_suitable_group_sites(

            assignments,

            sites,
        )
    )


    print(

        f"   -> {len(suitable_group_sites):,} "
        "vorläufig geeignete Standorte in zugeordneten Gruppen"
    )


    print(

        f"   -> "
        f"{suitable_group_sites['_Experiment_Gruppe'].nunique():,} "
        "Experimentgruppen mit mindestens einem geeigneten Standort"
    )


    print(

        "   -> davon mit hoher Zahlen-Prüfpriorität: "

        f"{int(suitable_group_sites['_Zahlenpruefung_hoch'].sum()):,}"
    )


    # ========================================================
    # 6. AUSGABEORDNER
    # ========================================================

    OUTPUT_DIR.mkdir(

        parents=True,

        exist_ok=True,
    )


    if CLEAN_OLD_GROUP_FILES:


        for file in OUTPUT_DIR.glob(

                f"{GROUP_FILE_PREFIX}*.xlsx"
        ):


            if not file.name.startswith(
                    "~$"
            ):

                file.unlink()


        overview_file = (

                OUTPUT_DIR
                / "_Export_Uebersicht.xlsx"
        )


        if overview_file.exists():

            overview_file.unlink()


    # ========================================================
    # 7. PRO EXPERIMENTGRUPPE EINE DATEI
    # ========================================================

    print()

    print(
        "6. Einzelne Gruppendateien werden erzeugt ..."
    )


    export_rows = []


    all_group_ids = sorted(

        assignments[
            "Experiment_Gruppe"
        ]

        .unique()
    )


    for group_id in all_group_ids:


        group_assignments = assignments[

            assignments[
                "Experiment_Gruppe"
            ]

            .eq(
                group_id
            )

        ].copy()


        group_sites = suitable_group_sites[

            suitable_group_sites[
                "_Experiment_Gruppe"
            ]

            .eq(
                group_id
            )

        ].copy()


        # Keine Datei, wenn in dieser manuell definierten Gruppe
        # nach dem neuen Wärmemengenstandard kein einziger
        # vorläufig geeigneter Standort verbleibt.
        if group_sites.empty:


            print(

                f"   {group_id}: "
                "0 geeignete Standorte -> keine Gruppendatei"
            )

            continue


        result = export_one_group(

            group_id=

            group_id,

            group_assignments=

            group_assignments,

            group_sites=

            group_sites,

            potentials=

            potentials,

            plausibility_table=

            plausibility_table,
        )


        export_rows.append(
            result
        )


        print(

            f"   {group_id}: "
            f"{result['Geeignete_Standorte_vorlaeufig']:,} Standorte | "
            f"{result['Abwaermepotentiale']:,} Potentiale | "
            f"{result['Plausibilitaetsauffaellige_Potentiale']:,} "
            "auffällige Potentiale"
        )


    # ========================================================
    # 8. KONTROLLÜBERSICHT
    # ========================================================

    write_export_overview(

        export_rows,

        grouping_file,
    )


    # ========================================================
    # 9. ABSCHLUSS
    # ========================================================

    print()

    print(
        "=" * 80
    )

    print(
        "FERTIG"
    )

    print(
        "=" * 80
    )

    print()


    print(
        "Ausgabeordner:"
    )


    print(
        OUTPUT_DIR.resolve()
    )


    print()

    print(

        f"Erzeugte Gruppendateien: "
        f"{len(export_rows):,}"
    )


    if export_rows:


        print()

        print(
            "Gesamtsumme über die exportierten Gruppen:"
        )


        print(

            "   vorläufig geeignete Standorte: "

            f"{sum(row['Geeignete_Standorte_vorlaeufig'] for row in export_rows):,}"
        )


        print(

            "   Abwärmepotentiale:             "

            f"{sum(row['Abwaermepotentiale'] for row in export_rows):,}"
        )


        print(

            "   auffällige Potentiale:         "

            f"{sum(row['Plausibilitaetsauffaellige_Potentiale'] for row in export_rows):,}"
        )


    print()

    print(
        "Jede Profilgruppe_*.xlsx enthält:"
    )


    print(
        "   1. Gruppeninfo"
    )


    print(
        "   2. WZ_Zuordnung"
    )


    print(
        "   3. Standortprofile"
    )


    print(
        "   4. Abwaermepotentiale"
    )


    print(
        "   5. Quellklassen_Uebersicht"
    )


    print(
        "   6. Zahlen_Plausibilitaet   (nur wenn Auffälligkeiten vorhanden)"
    )


    print()

    print(
        "Zusätzlich:"
    )


    print(
        "   _Export_Uebersicht.xlsx"
    )


    print()

    print(
        "WICHTIG:"
    )


    print(

        "   Standortauswahl ausschließlich über "
        "Geeignet_QX_Waermemenge_max_20pct_vorlaeufig."
    )


    print(

        "   Plausibilitätsauffällige positive Wärmemengen bleiben "
        "vollständig enthalten."
    )


    print(

        "   Die gemeldete Maximalleistung beeinflusst die Auswahl "
        "der Standorte nicht."
    )


    print(
        "=" * 80
    )


# ============================================================
# 16. PROGRAMMSTART
# ============================================================

if __name__ == "__main__":

    main()
