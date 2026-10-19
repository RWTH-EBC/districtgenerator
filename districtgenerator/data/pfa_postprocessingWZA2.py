from pathlib import Path
import re

import numpy as np
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation


# ============================================================
# 1. KONFIGURATION
# ============================================================

BASE_DIR = Path("WZ2025_Vorbereitung")


# ------------------------------------------------------------
# Deine vollständig manuell bearbeitete Validierungsdatei
# ------------------------------------------------------------

INPUT_FILE = (
        BASE_DIR
        / "WZ2025_Validierung_v4.xlsx"
)


# ------------------------------------------------------------
# Neue, kleinere Folge-Prüfdatei
# ------------------------------------------------------------

OUTPUT_FILE = (
        BASE_DIR
        / "WZ2025_Zusatzpruefung_v1.xlsx"
)


ALL_RESULTS_SHEET = "Alle_Ergebnisse"

WZ_REVIEW_SHEET = "WZ_Pruefung"

SITE_REVIEW_SHEET = (
    "Standortverbund_Pruefung"
)


# ============================================================
# 2. ALLGEMEINE HILFSFUNKTIONEN
# ============================================================

def clean_string(value) -> str:
    """
    Wandelt Excelwerte in saubere Strings um.

    NaN / None werden zu "".
    """

    if pd.isna(value):
        return ""

    text = str(value).strip()

    if text.lower() in {
        "nan",
        "none",
    }:
        return ""

    return text


# ============================================================
# 3. WZ-CODES
# ============================================================

def normalize_wz_code(value) -> str:
    """
    Bereinigt einen WZ-Code, BEHÄLT aber die vorhandene
    Klassifikationstiefe vollständig bei.

    Wichtig:

        38.21 -> 38.21
        38.2  -> 38.2
        20.3  -> 20.3
        5.1   -> 05.1
        35    -> 35

    38.21 wird ausdrücklich NICHT auf 38.2 reduziert.
    """

    text = (
        clean_string(value)
        .replace(",", ".")
        .replace(" ", "")
    )


    if not text:
        return ""


    # Führende Null für einstellige WZ-Abteilungen ergänzen.
    match = re.match(
        r"^(\d{1,2})(.*)$",
        text,
    )


    if not match:

        # Beispielsweise ein Abschnittscode "C".
        # Der wird unverändert zurückgegeben.
        return text


    division = int(
        match.group(1)
    )

    remainder = (
        match.group(2)
    )


    return (
        f"{division:02d}"
        f"{remainder}"
    )


def derive_wz3(value) -> str:
    """
    Leitet ZUSÄTZLICH den übergeordneten WZ-3-Steller ab.

    Der Original-/Finalcode selbst wird NICHT verändert.

    Beispiele:

        38.21 -> 38.2
        38.2  -> 38.2
        20.3  -> 20.3
        05.1  -> 05.1
        38    -> ""
        C     -> ""
    """

    text = normalize_wz_code(
        value
    )


    match = re.match(
        r"^(\d{2})\.(\d)",
        text,
    )


    if not match:
        return ""


    return (
        f"{match.group(1)}."
        f"{match.group(2)}"
    )


def derive_wz2(value) -> str:
    """
    Leitet die zweistellige WZ-Abteilung ab.

    Beispiele:

        38.21 -> 38
        38.2  -> 38
        20.3  -> 20
        05.1  -> 05
        35    -> 35
        C     -> ""
    """

    text = normalize_wz_code(
        value
    )


    match = re.match(
        r"^(\d{2})",
        text,
    )


    if not match:
        return ""


    return match.group(1)


def get_safe_wz2_from_row(row) -> str:
    """
    Gibt den bisherigen sicheren WZ-Code nur dann als
    final nutzbaren Code zurück, wenn die sichere Ebene
    tatsächlich eine WZ-Abteilung / ein 2-Steller ist.

    Ein bloßer Abschnitt wie "C" wird NICHT als finaler
    WZ-Code übernommen.
    """

    level = (
        clean_string(
            row.get(
                "WZ_Sichere_Ebene",
                "",
            )
        )
        .casefold()
    )


    safe_code = normalize_wz_code(
        row.get(
            "WZ_Sicherer_Code",
            "",
        )
    )


    if (
            "abteilung" in level
            or
            "2-steller" in level
    ):

        return safe_code


    return ""


# ============================================================
# 4. SONSTIGE HILFSFUNKTIONEN
# ============================================================

def make_address(
        street,
        postal,
        city,
) -> str:
    """
    Erstellt eine kompakte lesbare Adresse.
    """

    parts = [
        clean_string(street),
        clean_string(postal),
        clean_string(city),
    ]


    return ", ".join(
        part
        for part in parts
        if part
    )


def combine_reference_candidate(
        row,
) -> str:
    """
    Fasst einen unsicheren Referenzkandidaten kompakt zusammen.
    """

    parts = [

        clean_string(
            row.get(
                "Kandidat_CompanyName",
                "",
            )
        ),

        clean_string(
            row.get(
                "Kandidat_LocationName",
                "",
            )
        ),

        clean_string(
            row.get(
                "Kandidat_Address",
                "",
            )
        ),
    ]


    return " | ".join(
        part
        for part in parts
        if part
    )


# ============================================================
# 5. VALIDIERUNGSDATEI EINLESEN
# ============================================================

def load_input_data():
    """
    Liest die drei benötigten Blätter aus der bereits
    manuell bearbeiteten WZ2025_Validierung_v4.xlsx.
    """

    if not INPUT_FILE.exists():

        raise FileNotFoundError(
            f"Eingabedatei nicht gefunden:\n"
            f"{INPUT_FILE}"
        )


    all_results = pd.read_excel(
        INPUT_FILE,
        sheet_name=ALL_RESULTS_SHEET,
        dtype=str,
        keep_default_na=False,
    )


    wz_review = pd.read_excel(
        INPUT_FILE,
        sheet_name=WZ_REVIEW_SHEET,
        dtype=str,
        keep_default_na=False,
    )


    site_review = pd.read_excel(
        INPUT_FILE,
        sheet_name=SITE_REVIEW_SHEET,
        dtype=str,
        keep_default_na=False,
    )


    # ========================================================
    # TECHNISCHE INTEGRITÄT
    # ========================================================

    if not all_results[
        "Klassifikations_ID"
    ].is_unique:

        raise ValueError(
            "Klassifikations_ID ist in "
            "'Alle_Ergebnisse' nicht eindeutig."
        )


    return (
        all_results,
        wz_review,
        site_review,
    )


# ============================================================
# 6. MANUELLE WZ-ENTSCHEIDUNGEN ANWENDEN
#
# WICHTIG:
#
# Hier werden noch KEINE Ursprungsdateien verändert.
#
# Wir erzeugen lediglich eine aktuelle Sicht auf:
#
#   WZ_Final_Code_Aktuell
#   WZ3_Aktuell_Code
#   WZ2_Aktuell_Code
#
# ============================================================

def add_effective_wz(
        all_results,
        wz_review,
):

    result = all_results.copy()


    manual_columns = [

        "Klassifikations_ID",

        "WZ_Entscheidung",

        "WZ_Manuell_Code",

        "WZ_Manuell_Ebene",

        "WZ_Manuell_Begruendung",
    ]


    manual = (

        wz_review[
            manual_columns
        ]

        .drop_duplicates(
            subset=[
                "Klassifikations_ID"
            ],
            keep="last",
        )
    )


    result = result.merge(

        manual,

        on="Klassifikations_ID",

        how="left",

        validate="1:1",
    )


    for column in manual_columns[1:]:

        result[column] = (

            result[column]
            .fillna("")
            .map(
                clean_string
            )
        )


    final_codes = []

    wz3_codes = []

    wz2_codes = []

    sources = []


    for _, row in result.iterrows():


        decision = (
            clean_string(
                row[
                    "WZ_Entscheidung"
                ]
            )
            .upper()
        )


        manual_code = (
            normalize_wz_code(
                row[
                    "WZ_Manuell_Code"
                ]
            )
        )


        ai_code = (
            normalize_wz_code(
                row[
                    "WZ_Gruppe_Code"
                ]
            )
        )


        safe_wz2 = (
            get_safe_wz2_from_row(
                row
            )
        )


        # ====================================================
        # A) MANUELLER WZ-CODE
        #
        # Die exakte vom Nutzer eingegebene Tiefe bleibt
        # vollständig erhalten.
        #
        # Beispiel:
        #
        #   WZ_Manuell_Code = 38.21
        #
        #   Final = 38.21
        #   WZ3   = 38.2
        #   WZ2   = 38
        # ====================================================

        if (
                decision
                ==
                "WZ_KORRIGIEREN"
        ):

            final_code = (
                manual_code
            )

            source = (
                "MANUELL_KORRIGIERT"
            )


        # ====================================================
        # B) MANUELL NUR WZ2 SICHER
        # ====================================================

        elif (
                decision
                ==
                "NUR_WZ2_SICHER"
        ):

            final_code = (
                    manual_code
                    or
                    safe_wz2
            )

            source = (
                "NUR_WZ2_SICHER"
            )


        # ====================================================
        # C) NICHT KLASSIFIZIERBAR
        # ====================================================

        elif (
                decision
                ==
                "NICHT_SICHER_KLASSIFIZIERBAR"
        ):

            final_code = ""

            source = (
                "NICHT_SICHER_KLASSIFIZIERBAR"
            )


        # ====================================================
        # D) KI BESTÄTIGT / KEINE MANUELLE PRÜFUNG
        # ====================================================

        else:

            if ai_code:

                final_code = (
                    ai_code
                )

            else:

                # Nur ein tatsächlich sicherer 2-Steller
                # darf als Fallback verwendet werden.
                final_code = (
                    safe_wz2
                )


            if (
                    decision
                    ==
                    "KI_BESTAETIGT"
            ):

                source = (
                    "KI_MANUELL_BESTAETIGT"
                )

            elif ai_code:

                source = "KI"

            elif safe_wz2:

                source = (
                    "KI_FALLBACK_WZ2"
                )

            else:

                source = (
                    "NOCH_OFFEN"
                )


        # ====================================================
        # HIERARCHIE ZUSÄTZLICH ABLEITEN
        # ====================================================

        final_codes.append(
            final_code
        )


        # Wenn bewusst nur WZ2 sicher ist,
        # darf daraus KEIN künstlicher WZ3 entstehen.
        if (
                decision
                ==
                "NUR_WZ2_SICHER"
        ):

            wz3_codes.append(
                ""
            )

        else:

            wz3_codes.append(
                derive_wz3(
                    final_code
                )
            )


        wz2_codes.append(
            derive_wz2(
                final_code
            )
        )


        sources.append(
            source
        )


    result[
        "WZ_Final_Code_Aktuell"
    ] = final_codes


    result[
        "WZ3_Aktuell_Code"
    ] = wz3_codes


    result[
        "WZ2_Aktuell_Code"
    ] = wz2_codes


    result[
        "WZ_Aktuell_Quelle"
    ] = sources


    return result


# ============================================================
# 7. MANUELLE STANDORTVERBÜNDE INTERPRETIEREN
#
# Unterstützt:
#
#   ZUSAMMENFASSEN
#   GETRENNT_LASSEN
#
# sowie deine eigenen Einträge:
#
#   ID1
#   ID2
#   ID3
#   ...
#
# ============================================================

def interpret_site_groups(
        site_review,
):

    work = site_review.copy()


    interpreted_groups = []

    interpretation_types = []

    custom_labels = 0


    for _, row in work.iterrows():


        decision = (
            clean_string(
                row[
                    "Standortverbund_Entscheidung"
                ]
            )
            .upper()
        )


        candidate = clean_string(
            row[
                "Kandidatenverbund_ID"
            ]
        )


        suggested_group = clean_string(
            row[
                "Standortverbund_Gruppe"
            ]
        )


        # ====================================================
        # A) NORMAL ZUSAMMENGEFASST
        # ====================================================

        if (
                decision
                ==
                "ZUSAMMENFASSEN"
        ):

            final_group = (
                    suggested_group
                    or
                    candidate
            )


            interpretation_type = (
                "ZUSAMMENFASSEN"
            )


        # ====================================================
        # B) GETRENNT
        # ====================================================

        elif decision in {
            "",
            "OFFEN",
            "GETRENNT_LASSEN",
        }:

            final_group = ""

            interpretation_type = (
                "GETRENNT"
            )


        # ====================================================
        # C) DEINE EIGENEN ID1 / ID2 / ...
        #
        # Die Kennung gilt nur innerhalb des jeweiligen
        # ursprünglichen Kandidatenverbunds.
        #
        # Dadurch kann "ID1" in zwei verschiedenen
        # Kandidatenverbünden vorkommen, ohne dass diese
        # versehentlich zusammengeführt werden.
        # ====================================================

        else:

            custom_labels += 1


            final_group = (
                    candidate
                    + "::"
                    + decision
            )


            interpretation_type = (
                "MANUELLE_TEILGRUPPE"
            )


        interpreted_groups.append(
            final_group
        )


        interpretation_types.append(
            interpretation_type
        )


    work[
        "Interpretierter_Standortverbund"
    ] = interpreted_groups


    work[
        "Interpretationsart"
    ] = interpretation_types


    # ========================================================
    # NUR GRUPPEN MIT MINDESTENS 2 TEILSTANDORTEN
    # SIND TATSÄCHLICHE VERBÜNDE
    # ========================================================

    counts = (

        work.loc[
            work[
                "Interpretierter_Standortverbund"
            ].ne(""),
            "Interpretierter_Standortverbund",
        ]

        .value_counts()
    )


    valid_groups = set(

        counts[
            counts >= 2
            ].index
    )


    work[
        "Ist_bestaetigter_Standortverbund"
    ] = (

        work[
            "Interpretierter_Standortverbund"
        ]

        .isin(
            valid_groups
        )
    )


    return (
        work,
        custom_labels,
    )


# ============================================================
# 8. HAUPTFUNKTION DER STANDORTVERBÜNDE PRÜFEN
#
# Eine weitere Prüfung ist nur nötig, wenn innerhalb eines
# bestätigten Standortverbundes
#
#   - unterschiedliche FINALE WZ-Codes vorkommen
#     ODER
#   - mindestens ein finaler WZ-Code fehlt.
#
# WICHTIG:
#
# Es werden die EXAKTEN finalen Codes verglichen.
#
# 38.2 und 38.21 gelten hier bewusst als verschieden.
# ============================================================

def build_main_function_review(
        effective_results,
        interpreted_sites,
):

    confirmed = interpreted_sites[

        interpreted_sites[
            "Ist_bestaetigter_Standortverbund"
        ]

    ].copy()


    # ========================================================
    # Finale WZ-Informationen ergänzen
    # ========================================================

    additional = effective_results[
        [
            "Klassifikations_ID",

            "WZ_Final_Code_Aktuell",

            "WZ3_Aktuell_Code",

            "WZ2_Aktuell_Code",

            "WZ_Aktuell_Quelle",

            "WZ_Gruppe_Code",

            "WZ_Gruppe_Bezeichnung",

            "WZ_Sichere_Ebene",

            "WZ_Sicherer_Code",

            "Standortfunktion",

            "WZ_Manuell_Begruendung",
        ]
    ].copy()


    additional = additional.rename(
        columns={

            "WZ_Gruppe_Code":
                "WZ_AI_Code_Effektiv",

            "WZ_Gruppe_Bezeichnung":
                "WZ_AI_Bezeichnung_Effektiv",

            "WZ_Sichere_Ebene":
                "WZ_Sichere_Ebene_Effektiv",

            "WZ_Sicherer_Code":
                "WZ_Sicherer_Code_Effektiv",

            "Standortfunktion":
                "Standortfunktion_Effektiv",
        }
    )


    confirmed = confirmed.merge(

        additional,

        on="Klassifikations_ID",

        how="left",

        validate="1:1",
    )


    rows = []


    for (
            group_id,
            group,
    ) in confirmed.groupby(
        "Interpretierter_Standortverbund",
        sort=True,
    ):


        # ====================================================
        # EXAKTE FINALE CODES
        #
        # Beispiel:
        # 38.2 und 38.21 bleiben unterschiedlich.
        # ====================================================

        final_codes = sorted(

            set(

                clean_string(
                    code
                )

                for code
                in group[
                    "WZ_Final_Code_Aktuell"
                ]

                if clean_string(
                    code
                )
            )
        )


        missing_codes = int(

            group[
                "WZ_Final_Code_Aktuell"
            ]

            .map(
                clean_string
            )

            .eq("")

            .sum()
        )


        reasons = []


        if len(
                final_codes
        ) > 1:

            reasons.append(
                "UNTERSCHIEDLICHE_FINALE_WZ_CODES"
            )


        if (
                missing_codes
                > 0
        ):

            reasons.append(
                "WZ_CODE_FEHLT_BEI_TEILSTANDORT"
            )


        # ----------------------------------------------------
        # Gleicher finaler WZ-Code für alle?
        #
        # Dann ist keine weitere Prüfung nötig.
        # ----------------------------------------------------

        if not reasons:
            continue


        review_reason = (
            " | ".join(
                reasons
            )
        )


        codes_text = (
            " | ".join(
                final_codes
            )
        )


        for _, row in group.iterrows():


            rows.append(
                {

                    "Standortverbund_Gruppe":
                        group_id,

                    "Anzahl_Teilstandorte":
                        len(
                            group
                        ),

                    "Pruefgrund":
                        review_reason,

                    "Vorhandene_finale_WZ_Codes":
                        codes_text,

                    "Klassifikations_ID":
                        row[
                            "Klassifikations_ID"
                        ],

                    "Firmenname":
                        row[
                            "Firmenname"
                        ],

                    "Standortname":
                        row[
                            "Standortname"
                        ],

                    "Adresse":
                        row[
                            "Adresse"
                        ],

                    # ----------------------------------------
                    # EXAKTER finaler Code
                    # ----------------------------------------

                    "WZ_Final_Code_Aktuell":
                        row[
                            "WZ_Final_Code_Aktuell"
                        ],

                    # ----------------------------------------
                    # Zusätzliche Hierarchie
                    # ----------------------------------------

                    "WZ3_Aktuell_Code":
                        row[
                            "WZ3_Aktuell_Code"
                        ],

                    "WZ2_Aktuell_Code":
                        row[
                            "WZ2_Aktuell_Code"
                        ],

                    "WZ_Aktuell_Quelle":
                        row[
                            "WZ_Aktuell_Quelle"
                        ],

                    # ----------------------------------------
                    # ursprüngliche KI-Klassifikation
                    # ----------------------------------------

                    "WZ_AI_Code":
                        row[
                            "WZ_AI_Code_Effektiv"
                        ],

                    "WZ_AI_Bezeichnung":
                        row[
                            "WZ_AI_Bezeichnung_Effektiv"
                        ],

                    "Standorttaetigkeit_kurz":
                        row[
                            "Standorttaetigkeit_kurz"
                        ],

                    "Standortfunktion":
                        row[
                            "Standortfunktion_Effektiv"
                        ],

                    "WZ_Manuell_Begruendung":
                        row[
                            "WZ_Manuell_Begruendung"
                        ],

                    # ========================================
                    # MANUELLE ENTSCHEIDUNG
                    #
                    # Pro Standortverbund genau eine Zeile
                    # auf JA setzen.
                    # ========================================

                    "Hauptfunktion":
                        "NEIN",

                    "Hauptfunktion_Begruendung":
                        "",
                }
            )


    review = pd.DataFrame(
        rows
    )


    if not review.empty:

        review = review.sort_values(
            [
                "Standortverbund_Gruppe",
                "Firmenname",
                "Standortname",
            ],
            kind="stable",
        ).reset_index(
            drop=True
        )


    return review


# ============================================================
# 9. NEUE / NICHT GEMATCHTE STANDORTE OHNE WZ3
#
# Gesucht werden:
#
# - kein sicher akzeptierter Referenzmatch
# - nach deiner ersten manuellen Prüfung weiterhin kein WZ3
#
# Bereits bewusst als
#
# - NUR_WZ2_SICHER
# - NICHT_SICHER_KLASSIFIZIERBAR
# - WZ_KORRIGIEREN
#
# entschiedene Fälle werden NICHT erneut vorgelegt.
# ============================================================

def build_new_without_wz3_review(
        effective_results,
):

    work = effective_results.copy()


    already_finally_reviewed = (

        work[
            "WZ_Entscheidung"
        ]

        .map(
            clean_string
        )

        .str.upper()

        .isin(
            {
                "WZ_KORRIGIEREN",
                "NUR_WZ2_SICHER",
                "NICHT_SICHER_KLASSIFIZIERBAR",
            }
        )
    )


    mask = (

            (
                    work[
                        "Match_Status"
                    ]
                    != "MATCH"
            )

            &

            (
                work[
                    "WZ3_Aktuell_Code"
                ]

                .map(
                    clean_string
                )

                .eq("")
            )

            &

            (
                ~already_finally_reviewed
            )
    )


    review = work[
        mask
    ].copy()


    review[
        "Adresse"
    ] = review.apply(
        lambda row:
        make_address(
            row[
                "Strasse_Hausnummer"
            ],
            row[
                "PLZ"
            ],
            row[
                "Ort"
            ],
        ),
        axis=1,
    )


    review[
        "Referenz_Kandidat"
    ] = review.apply(
        combine_reference_candidate,
        axis=1,
    )


    output = review[
        [
            "Klassifikations_ID",

            "Firmenname",

            "Standortname",

            "Adresse",

            "Match_Status",

            "Match_Methode",

            "Match_Score",

            "Referenz_Kandidat",

            # --------------------------------------------
            # Aktueller Stand
            # --------------------------------------------

            "WZ_Final_Code_Aktuell",

            "WZ2_Aktuell_Code",

            # --------------------------------------------
            # bisherige sichere KI-Ebene
            # --------------------------------------------

            "WZ_Sichere_Ebene",

            "WZ_Sicherer_Code",

            "WZ_Sichere_Bezeichnung",

            # --------------------------------------------
            # fachlicher Kontext
            # --------------------------------------------

            "Standorttaetigkeit_kurz",

            "Konfidenz",

            "WZ_Begruendung",

            "Recherchequelle_1",

            "Recherchequelle_2",
        ]
    ].copy()


    # ========================================================
    # MANUELLE ENTSCHEIDUNGEN
    #
    # WZ_Manuell_Code darf auch FEINER als WZ3 sein.
    #
    # Beispiel:
    #     38.21
    #
    # wird vollständig erhalten.
    # ========================================================

    output[
        "Neue_WZ_Entscheidung"
    ] = "OFFEN"


    output[
        "WZ_Manuell_Code"
    ] = ""


    output[
        "WZ_Manuell_Begruendung"
    ] = ""


    return output.reset_index(
        drop=True
    )


# ============================================================
# 10. BEREITS BEARBEITETE ZUSATZPRÜFUNG BEWAHREN
#
# Falls du dieses Skript später erneut ausführst,
# werden bereits eingetragene Entscheidungen aus
#
# WZ2025_Zusatzpruefung_v1.xlsx
#
# wieder übernommen.
# ============================================================

def restore_existing_output(
        main_review,
        new_wz_review,
):

    if not OUTPUT_FILE.exists():

        return (
            main_review,
            new_wz_review,
        )


    try:

        excel = pd.ExcelFile(
            OUTPUT_FILE
        )


        # ====================================================
        # HAUPTFUNKTION
        # ====================================================

        if (
                "Standortverbund_Hauptfunktion"
                in excel.sheet_names
                and
                not main_review.empty
        ):

            old = pd.read_excel(
                OUTPUT_FILE,
                sheet_name=
                "Standortverbund_Hauptfunktion",
                dtype=str,
                keep_default_na=False,
            )


            needed = {
                "Klassifikations_ID",
                "Hauptfunktion",
                "Hauptfunktion_Begruendung",
            }


            if needed.issubset(
                    old.columns
            ):

                old_lookup = (

                    old[
                        list(
                            needed
                        )
                    ]

                    .drop_duplicates(
                        subset=[
                            "Klassifikations_ID"
                        ],
                        keep="last",
                    )

                    .set_index(
                        "Klassifikations_ID"
                    )
                )


                for column in [
                    "Hauptfunktion",
                    "Hauptfunktion_Begruendung",
                ]:

                    mapping = (
                        old_lookup[
                            column
                        ]
                        .to_dict()
                    )


                    main_review[
                        column
                    ] = main_review.apply(

                        lambda row:
                        mapping.get(
                            row[
                                "Klassifikations_ID"
                            ],
                            row[
                                column
                            ],
                        ),

                        axis=1,
                    )


        # ====================================================
        # NEUE WZ-FÄLLE
        # ====================================================

        if (
                "Neue_ohne_WZ3"
                in excel.sheet_names
                and
                not new_wz_review.empty
        ):

            old = pd.read_excel(
                OUTPUT_FILE,
                sheet_name=
                "Neue_ohne_WZ3",
                dtype=str,
                keep_default_na=False,
            )


            needed = {
                "Klassifikations_ID",
                "Neue_WZ_Entscheidung",
                "WZ_Manuell_Code",
                "WZ_Manuell_Begruendung",
            }


            if needed.issubset(
                    old.columns
            ):

                old_lookup = (

                    old[
                        list(
                            needed
                        )
                    ]

                    .drop_duplicates(
                        subset=[
                            "Klassifikations_ID"
                        ],
                        keep="last",
                    )

                    .set_index(
                        "Klassifikations_ID"
                    )
                )


                for column in [
                    "Neue_WZ_Entscheidung",
                    "WZ_Manuell_Code",
                    "WZ_Manuell_Begruendung",
                ]:

                    mapping = (
                        old_lookup[
                            column
                        ]
                        .to_dict()
                    )


                    new_wz_review[
                        column
                    ] = new_wz_review.apply(

                        lambda row:
                        mapping.get(
                            row[
                                "Klassifikations_ID"
                            ],
                            row[
                                column
                            ],
                        ),

                        axis=1,
                    )


    except Exception as exc:

        print()
        print(
            "WARNUNG:"
        )

        print(
            "Vorhandene Entscheidungen aus der "
            "Zusatzprüfdatei konnten nicht übernommen werden."
        )

        print(
            exc
        )


    return (
        main_review,
        new_wz_review,
    )


# ============================================================
# 11. ZUSAMMENFASSUNG
# ============================================================

def build_summary(
        effective_results,
        wz_review,
        interpreted_sites,
        main_review,
        new_wz_review,
        custom_labels,
):

    total_sites = len(
        effective_results
    )


    wz3_count = int(

        effective_results[
            "WZ3_Aktuell_Code"
        ]

        .map(
            clean_string
        )

        .ne("")

        .sum()
    )


    decisions = (

        wz_review[
            "WZ_Entscheidung"
        ]

        .map(
            clean_string
        )

        .value_counts()
    )


    confirmed_sites = interpreted_sites[

        interpreted_sites[
            "Ist_bestaetigter_Standortverbund"
        ]

    ]


    confirmed_groups = (

        confirmed_sites[
            "Interpretierter_Standortverbund"
        ]

        .nunique()
    )


    # ========================================================
    # 38.21 separat sichtbar machen
    # ========================================================

    code_3821_count = int(

        (
                effective_results[
                    "WZ_Final_Code_Aktuell"
                ]
                == "38.21"
        ).sum()
    )


    rows = [

        {
            "Kennzahl":
                "Standorte gesamt",

            "Wert":
                total_sites,
        },

        {
            "Kennzahl":
                "Standorte mit WZ3 nach erster manueller Prüfung",

            "Wert":
                wz3_count,
        },

        {
            "Kennzahl":
                "WZ3-Abdeckung nach erster manueller Prüfung",

            "Wert":
                (
                    wz3_count
                    / total_sites
                    if total_sites
                    else np.nan
                ),
        },

        {
            "Kennzahl":
                "Finaler Code 38.21",

            "Wert":
                code_3821_count,
        },

        {
            "Kennzahl":
                "WZ-Prüffälle gesamt",

            "Wert":
                len(
                    wz_review
                ),
        },

        {
            "Kennzahl":
                "KI bestätigt",

            "Wert":
                int(
                    decisions.get(
                        "KI_BESTAETIGT",
                        0,
                    )
                ),
        },

        {
            "Kennzahl":
                "WZ manuell korrigiert",

            "Wert":
                int(
                    decisions.get(
                        "WZ_KORRIGIEREN",
                        0,
                    )
                ),
        },

        {
            "Kennzahl":
                "Nur WZ2 sicher",

            "Wert":
                int(
                    decisions.get(
                        "NUR_WZ2_SICHER",
                        0,
                    )
                ),
        },

        {
            "Kennzahl":
                "Nicht sicher klassifizierbar",

            "Wert":
                int(
                    decisions.get(
                        "NICHT_SICHER_KLASSIFIZIERBAR",
                        0,
                    )
                ),
        },

        {
            "Kennzahl":
                "Bestätigte Standortverbünde",

            "Wert":
                confirmed_groups,
        },

        {
            "Kennzahl":
                "Teilstandorte in bestätigten Verbünden",

            "Wert":
                len(
                    confirmed_sites
                ),
        },

        {
            "Kennzahl":
                "Eigene ID1/ID2/... Einträge interpretiert",

            "Wert":
                custom_labels,
        },

        {
            "Kennzahl":
                "Standortverbünde mit Hauptfunktionsprüfung",

            "Wert":
                (
                    main_review[
                        "Standortverbund_Gruppe"
                    ]
                    .nunique()

                    if not main_review.empty

                    else 0
                ),
        },

        {
            "Kennzahl":
                "Zeilen in Hauptfunktionsprüfung",

            "Wert":
                len(
                    main_review
                ),
        },

        {
            "Kennzahl":
                "Neue/nicht gematchte Standorte ohne WZ3",

            "Wert":
                len(
                    new_wz_review
                ),
        },
    ]


    return pd.DataFrame(
        rows
    )


# ============================================================
# 12. ANLEITUNGSBLATT
# ============================================================

def build_instructions():

    return pd.DataFrame(
        [

            {
                "Blatt":
                    "Standortverbund_Hauptfunktion",

                "Eingabe":
                    "Hauptfunktion = JA",

                "Bedeutung":
                    "Pro Standortverbund genau eine Zeile auf JA setzen. "
                    "Dieser Teilstandort bestimmt später die "
                    "Branchenzuordnung des gesamten Standortverbunds.",
            },

            {
                "Blatt":
                    "Standortverbund_Hauptfunktion",

                "Eingabe":
                    "Hauptfunktion = NEIN",

                "Bedeutung":
                    "Teilfunktion des gemeinsamen Standortes. "
                    "Die Abwärmequellen bleiben später trotzdem "
                    "Bestandteil des gemeinsamen Standortprofils.",
            },

            {
                "Blatt":
                    "Standortverbund_Hauptfunktion",

                "Eingabe":
                    "Hauptfunktion_Begruendung",

                "Bedeutung":
                    "Bei der JA-Zeile kurze Begründung eintragen, "
                    "z.B. 'Krankenhaus ist Hauptfunktion; "
                    "BHKW ist interne Energieversorgung'.",
            },

            {
                "Blatt":
                    "Neue_ohne_WZ3",

                "Eingabe":
                    "WZ_CODE_FESTLEGEN",

                "Bedeutung":
                    "Einen finalen WZ-Code in WZ_Manuell_Code eintragen. "
                    "Der Code darf auch feiner als WZ3 sein, z.B. 38.21.",
            },

            {
                "Blatt":
                    "Neue_ohne_WZ3",

                "Eingabe":
                    "NUR_WZ2_SICHER",

                "Bedeutung":
                    "Eine feinere Zuordnung ist nicht sicher. "
                    "Der vorhandene sichere 2-Steller bleibt bestehen.",
            },

            {
                "Blatt":
                    "Neue_ohne_WZ3",

                "Eingabe":
                    "NICHT_SICHER_KLASSIFIZIERBAR",

                "Bedeutung":
                    "Auch nach manueller Prüfung keine belastbare "
                    "WZ-Zuordnung möglich.",
            },

            {
                "Blatt":
                    "Allgemein",

                "Eingabe":
                    "Wichtig",

                "Bedeutung":
                    "IDs und Informationsspalten nicht verändern. "
                    "Nur gelb markierte Spalten bearbeiten.",
            },

            {
                "Blatt":
                    "Allgemein",

                "Eingabe":
                    "WZ-Code 38.21",

                "Bedeutung":
                    "38.21 bleibt als finaler Profilcode vollständig "
                    "erhalten. Zusätzlich wird lediglich WZ3 = 38.2 "
                    "und WZ2 = 38 abgeleitet.",
            },
        ]
    )


# ============================================================
# 13. EXCELDATEI SCHREIBEN
# ============================================================

def write_output(
        summary,
        instructions,
        main_review,
        new_wz_review,
):

    with pd.ExcelWriter(
            OUTPUT_FILE,
            engine="openpyxl",
    ) as writer:


        instructions.to_excel(
            writer,
            sheet_name="Anleitung",
            index=False,
        )


        summary.to_excel(
            writer,
            sheet_name="Zusammenfassung",
            index=False,
        )


        main_review.to_excel(
            writer,
            sheet_name=
            "Standortverbund_Hauptfunktion",
            index=False,
        )


        new_wz_review.to_excel(
            writer,
            sheet_name="Neue_ohne_WZ3",
            index=False,
        )


# ============================================================
# 14. FORMATIERUNG UND DROPDOWNS
# ============================================================

def format_output():

    workbook = load_workbook(
        OUTPUT_FILE
    )


    normal_header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )


    manual_header_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )


    manual_columns = {

        "Standortverbund_Hauptfunktion": {
            "Hauptfunktion",
            "Hauptfunktion_Begruendung",
        },

        "Neue_ohne_WZ3": {
            "Neue_WZ_Entscheidung",
            "WZ_Manuell_Code",
            "WZ_Manuell_Begruendung",
        },
    }


    for worksheet in workbook.worksheets:


        worksheet.freeze_panes = "A2"


        if (
                worksheet.max_row >= 1
                and
                worksheet.max_column >= 1
        ):

            worksheet.auto_filter.ref = (
                worksheet.dimensions
            )


        # ====================================================
        # KOPFZEILEN
        # ====================================================

        for cell in worksheet[1]:

            header = clean_string(
                cell.value
            )


            cell.font = Font(
                bold=True
            )


            if (
                    worksheet.title
                    in manual_columns
                    and
                    header
                    in manual_columns[
                worksheet.title
            ]
            ):

                cell.fill = (
                    manual_header_fill
                )

            else:

                cell.fill = (
                    normal_header_fill
                )


            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


        # ====================================================
        # SPALTENBREITEN
        # ====================================================

        for column_cells in worksheet.columns:


            column_letter = (
                column_cells[0]
                .column_letter
            )


            header = clean_string(
                column_cells[0].value
            )


            if header in {
                "WZ_Begruendung",
                "WZ_Manuell_Begruendung",
                "Hauptfunktion_Begruendung",
                "Bedeutung",
            }:

                width = 45


            elif header in {
                "Adresse",
                "Referenz_Kandidat",
                "Standorttaetigkeit_kurz",
                "Vorhandene_finale_WZ_Codes",
            }:

                width = 35


            else:

                max_length = 0


                for cell in list(
                        column_cells
                )[:200]:

                    if cell.value is None:
                        continue


                    max_length = max(
                        max_length,
                        len(
                            str(
                                cell.value
                            )
                        ),
                    )


                width = min(
                    max(
                        max_length + 2,
                        10,
                        ),
                    28,
                )


            worksheet.column_dimensions[
                column_letter
            ].width = width


        # ====================================================
        # TEXTUMBRUCH
        # ====================================================

        for row in worksheet.iter_rows(
                min_row=2
        ):

            for cell in row:

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )


    # ========================================================
    # DROPDOWN:
    # HAUPTFUNKTION
    # ========================================================

    sheet_name = (
        "Standortverbund_Hauptfunktion"
    )


    if sheet_name in workbook.sheetnames:

        ws = workbook[
            sheet_name
        ]


        headers = {
            cell.value:
                cell.column_letter
            for cell in ws[1]
        }


        if "Hauptfunktion" in headers:

            validation = DataValidation(
                type="list",
                formula1='"JA,NEIN"',
                allow_blank=False,
            )


            ws.add_data_validation(
                validation
            )


            column = headers[
                "Hauptfunktion"
            ]


            validation.add(
                f"{column}2:"
                f"{column}{max(ws.max_row, 2)}"
            )


    # ========================================================
    # DROPDOWN:
    # NEUE WZ-ENTSCHEIDUNG
    # ========================================================

    sheet_name = (
        "Neue_ohne_WZ3"
    )


    if sheet_name in workbook.sheetnames:

        ws = workbook[
            sheet_name
        ]


        headers = {
            cell.value:
                cell.column_letter
            for cell in ws[1]
        }


        if (
                "Neue_WZ_Entscheidung"
                in headers
        ):

            validation = DataValidation(
                type="list",

                formula1=(
                    '"OFFEN,'
                    'WZ_CODE_FESTLEGEN,'
                    'NUR_WZ2_SICHER,'
                    'NICHT_SICHER_KLASSIFIZIERBAR"'
                ),

                allow_blank=False,
            )


            ws.add_data_validation(
                validation
            )


            column = headers[
                "Neue_WZ_Entscheidung"
            ]


            validation.add(
                f"{column}2:"
                f"{column}{max(ws.max_row, 2)}"
            )


    workbook.save(
        OUTPUT_FILE
    )


# ============================================================
# 15. HAUPTPROGRAMM
# ============================================================

def main():

    print()
    print("=" * 78)
    print(
        "ZUSATZPRÜFUNG NACH MANUELLER "
        "WZ- UND STANDORTVALIDIERUNG"
    )
    print("=" * 78)
    print()


    # ========================================================
    # 1. DATEI EINLESEN
    # ========================================================

    print(
        "1. Bearbeitete Validierungsdatei wird eingelesen ..."
    )


    (
        all_results,
        wz_review,
        site_review,
    ) = load_input_data()


    print(
        f"   -> {len(all_results):,} Standorte"
    )


    print(
        f"   -> {len(wz_review):,} WZ-Prüffälle"
    )


    print(
        f"   -> {len(site_review):,} "
        "Standortverbund-Prüfzeilen"
    )


    # ========================================================
    # 2. WZ-ENTSCHEIDUNGEN INTERPRETIEREN
    # ========================================================

    print()
    print(
        "2. Manuelle WZ-Entscheidungen werden interpretiert ..."
    )


    effective_results = (
        add_effective_wz(
            all_results,
            wz_review,
        )
    )


    wz3_count = int(

        effective_results[
            "WZ3_Aktuell_Code"
        ]

        .map(
            clean_string
        )

        .ne("")

        .sum()
    )


    code_3821 = int(

        (
                effective_results[
                    "WZ_Final_Code_Aktuell"
                ]
                == "38.21"
        ).sum()
    )


    print(
        f"   -> {wz3_count:,} / "
        f"{len(effective_results):,} "
        "Standorte mit WZ3"
    )


    print(
        f"   -> "
        f"{wz3_count / len(effective_results):.2%}"
    )


    print(
        f"   -> {code_3821:,} Standorte "
        "mit finalem Code 38.21"
    )


    # ========================================================
    # 3. STANDORTVERBÜNDE INTERPRETIEREN
    # ========================================================

    print()
    print(
        "3. Manuelle Standortverbund-Entscheidungen "
        "werden interpretiert ..."
    )


    (
        interpreted_sites,
        custom_labels,
    ) = interpret_site_groups(
        site_review
    )


    confirmed_sites = interpreted_sites[

        interpreted_sites[
            "Ist_bestaetigter_Standortverbund"
        ]

    ]


    confirmed_groups = (

        confirmed_sites[
            "Interpretierter_Standortverbund"
        ]

        .nunique()
    )


    print(
        f"   -> {confirmed_groups:,} bestätigte "
        "Standortverbünde"
    )


    print(
        f"   -> {len(confirmed_sites):,} Teilstandorte "
        "in bestätigten Verbünden"
    )


    print(
        f"   -> {custom_labels:,} eigene "
        "ID1/ID2/... Einträge interpretiert"
    )


    # ========================================================
    # 4. HAUPTFUNKTIONSPRÜFUNG
    # ========================================================

    print()
    print(
        "4. Standortverbünde mit unterschiedlichen "
        "finalen WZ-Codes werden gesucht ..."
    )


    main_review = (
        build_main_function_review(
            effective_results,
            interpreted_sites,
        )
    )


    if not main_review.empty:

        main_group_count = (

            main_review[
                "Standortverbund_Gruppe"
            ]

            .nunique()
        )

    else:

        main_group_count = 0


    print(
        f"   -> {main_group_count:,} Verbünde "
        "benötigen eine Hauptfunktionsentscheidung"
    )


    print(
        f"   -> {len(main_review):,} Prüfzeilen"
    )


    # ========================================================
    # 5. NICHT GEMATCHTE STANDORTE OHNE WZ3
    # ========================================================

    print()
    print(
        "5. Neue/nicht sicher gematchte Standorte "
        "ohne WZ3 werden gesucht ..."
    )


    new_wz_review = (
        build_new_without_wz3_review(
            effective_results
        )
    )


    print(
        f"   -> {len(new_wz_review):,} Fälle"
    )


    if not new_wz_review.empty:

        print(
            "      davon UNSICHER: "
            f"{(new_wz_review['Match_Status'] == 'UNSICHER').sum():,}"
        )


        print(
            "      davon KEIN_MATCH: "
            f"{(new_wz_review['Match_Status'] == 'KEIN_MATCH').sum():,}"
        )


    # ========================================================
    # 6. EVENTUELLE FRÜHERE ENTSCHEIDUNGEN WIEDERHERSTELLEN
    # ========================================================

    (
        main_review,
        new_wz_review,
    ) = restore_existing_output(
        main_review,
        new_wz_review,
    )


    # ========================================================
    # 7. ZUSAMMENFASSUNG
    # ========================================================

    summary = (
        build_summary(
            effective_results,
            wz_review,
            interpreted_sites,
            main_review,
            new_wz_review,
            custom_labels,
        )
    )


    instructions = (
        build_instructions()
    )


    # ========================================================
    # 8. EXCEL SCHREIBEN
    # ========================================================

    print()
    print(
        "6. Zusatzprüfdatei wird geschrieben ..."
    )


    write_output(
        summary,
        instructions,
        main_review,
        new_wz_review,
    )


    format_output()


    # ========================================================
    # ABSCHLUSS
    # ========================================================

    print()
    print("=" * 78)
    print("FERTIG")
    print("=" * 78)
    print()


    print(
        f"Ausgabedatei:\n"
        f"{OUTPUT_FILE.resolve()}"
    )


    print()
    print(
        "Zu bearbeiten:"
    )


    print()
    print(
        "1. Standortverbund_Hauptfunktion"
    )


    print(
        "   Pro Verbund genau EINE Zeile auf JA setzen."
    )


    print()
    print(
        "2. Neue_ohne_WZ3"
    )


    print(
        "   Für jeden Fall die Entscheidung über "
        "das Dropdown treffen."
    )


    print()
    print(
        "Wichtig:"
    )


    print(
        "Ein manueller Code wie 38.21 bleibt "
        "vollständig als 38.21 erhalten."
    )


    print("=" * 78)


# ============================================================
# 16. PROGRAMMSTART
# ============================================================

if __name__ == "__main__":

    main()