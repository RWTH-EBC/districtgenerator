from pathlib import Path
import re
import unicodedata
import warnings

import numpy as np
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


# =============================================================================
# PFA – EXPORT FÜR EXTERNE KI-RECHERCHE ZUR STANDORTGRÖSSE – VERSION 2
# =============================================================================
#
# ZIEL
# ----
# Erzeugt eine Excel-Arbeitsdatei mit den Standorten, die nach der fertigen
# Gruppenaufbereitung noch in der Generatorbasis enthalten sind.
#
# WICHTIGE ÄNDERUNG GEGENÜBER VERSION 1
# --------------------------------------
# Die Ausgabedatei enthält ABSICHTLICH KEINE Angaben zur Abwärmemenge und
# keine technischen Abwärmeinformationen (Q-Typen, Temperaturen,
# Betriebsprofile, Quellenanteile usw.).
#
# Hintergrund:
# Die externe KI soll die reale Standortgröße ausschließlich aus unabhängigen
# Produktions-, Kapazitäts-, Beschäftigten-, Umsatz- oder anderen geeigneten
# Standortkennzahlen ableiten. Die PFA-Abwärmedaten dürfen die Einordnung der
# Standortgröße nicht beeinflussen.
#
# PRO STANDORT KÖNNEN BIS ZU 5 GRÖSSENKENNZAHLEN RECHERCHIERT WERDEN.
# Jede Kennzahl erhält u. a.:
#   - Typ / Bezeichnung
#   - Wert
#   - Einheit
#   - Bezugsjahr
#   - Standort-/Unternehmensbezug
#   - Datentyp (Istwert, Kapazität, Schätzung ...)
#   - Quelle
#   - Vertrauensniveau
#   - Vergleichbarkeit innerhalb der Gruppe
#   - Kennzeichnung, ob sie in den Größenscore einfließen soll
#
# Erst anschließend soll die KI aus den belastbaren Kennzahlen einen relativen
# Größenscore von 0 bis 100 INNERHALB DER JEWEILIGEN PROFILGRUPPE ableiten.
# =============================================================================


# =============================================================================
# 1. KONFIGURATION
# =============================================================================

BASE_DIR = Path("WZ2025_Vorbereitung")
SOURCE_DIR = BASE_DIR / "Quellenklassifikation_final"

# Nur Standorte aus diesen fertigen Gruppenordnern werden exportiert.
GROUP_PACKAGE_DIR = SOURCE_DIR / "Generator_Datenpakete_Iteration4"

# Diese Datei wird NUR für Firmenname, Standortname, Adresse und WZ-Kontext
# verwendet. Es werden KEINE Abwärmedaten daraus in die Ausgabe übernommen.
MASTER_POTENTIAL_FILE = SOURCE_DIR / "Abwaermepotentiale_mit_Quellklassen.xlsx"
MASTER_POTENTIAL_SHEET = "Abwaermepotentiale"

OUTPUT_DIR = SOURCE_DIR / "KI_Standortgroesse_Recherche_v2"
OUTPUT_FILE = OUTPUT_DIR / "KI_Standortgroesse_Recherche_v2.xlsx"

# None = alle fertigen Gruppen.
# Beispiele:
# TARGET_GROUPS = ["G62"]
# TARGET_GROUPS = ["G06", "G39", "G62"]
TARGET_GROUPS = ["G16"]

GROUP_FOLDER_PATTERN = "G*"
WRITE_ONE_FILE_PER_GROUP = True

FREEZE_PANES = "A2"
MAX_COLUMN_WIDTH = 55


# =============================================================================
# 2. ZULÄSSIGE KENNZAHLTYPEN
# =============================================================================

METRIC_TYPES = [
    "PRODUKTIONSMENGE",
    "PRODUKTIONSKAPAZITAET",
    "DURCHSATZ",
    "SEKTORALE_PHYSISCHE_KENNZAHL",
    "MITARBEITER_STANDORT",
    "UMSATZ_STANDORT",
    "ANLAGENKAPAZITAET",
    "SONSTIGE_STANDORTKENNZAHL",
    "UNTERNEHMENSWERT_FALLBACK",
]

METRIC_REFERENCES = [
    "STANDORT",
    "WERK",
    "BETRIEBSSTAETTE",
    "UNTERNEHMEN",
    "KONZERN",
]

METRIC_DATA_TYPES = [
    "IST_WERT",
    "KAPAZITAET",
    "PLANWERT",
    "GESCHAETZT",
    "ABGELEITET",
]

QUALITY_LEVELS = [
    "HOCH",
    "MITTEL",
    "NIEDRIG",
]

COMPARABILITY_LEVELS = [
    "HOCH",
    "MITTEL",
    "NIEDRIG",
    "NICHT_VERGLEICHBAR",
]

SCORE_USE_VALUES = [
    "JA",
    "NEIN",
]

SIZE_CLASSES = [
    "SEHR_KLEIN",
    "KLEIN",
    "MITTEL",
    "GROSS",
    "SEHR_GROSS",
]


# =============================================================================
# 3. KI-SPALTEN – BIS ZU 5 GRÖSSENKENNZAHLEN
# =============================================================================

BASE_KI_COLUMNS = [
    "KI_Bearbeitungsstatus",
    "KI_Standortidentifikation_Sicherheit",
    "KI_Recherchierter_Standortname",
    "KI_Standort_Webseite",
    "KI_Allgemeine_Quellen_URLs",
    "KI_Quellen_Abrufdatum",
]


def metric_columns(metric_number: int):
    prefix = f"KI_Kennzahl_{metric_number}"
    return [
        f"{prefix}_Typ",
        f"{prefix}_Bezeichnung",
        f"{prefix}_Wert",
        f"{prefix}_Einheit",
        f"{prefix}_Jahr",
        f"{prefix}_Bezug",
        f"{prefix}_Datentyp",
        f"{prefix}_Quelle_URL",
        f"{prefix}_Vertrauen",
        f"{prefix}_Vergleichbarkeit_innerhalb_Gruppe",
        f"{prefix}_Fuer_Score_verwenden",
        f"{prefix}_Kommentar",
    ]


METRIC_COLUMNS = []
for number in range(1, 6):
    METRIC_COLUMNS.extend(metric_columns(number))

SCORE_COLUMNS = [
    "KI_Anzahl_recherchierter_Kennzahlen",
    "KI_Anzahl_fuer_Score_verwendeter_Kennzahlen",
    "KI_Fuer_Score_verwendete_Kennzahlen",
    "KI_Score_Methodik",
    "KI_Relativer_Groessenscore_0_100",
    "KI_Groessenklasse_innerhalb_Gruppe",
    "KI_Vertrauen_Groessenschaetzung",
    "KI_Kurzbegruendung_Score",
    "KI_Mehrstandortunternehmen_Hinweis",
    "KI_Recherchehinweise",
]

KI_COLUMNS = BASE_KI_COLUMNS + METRIC_COLUMNS + SCORE_COLUMNS


# =============================================================================
# 4. HILFSFUNKTIONEN
# =============================================================================


def clean_string(value) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none", "<na>"}:
        return ""
    return text



def normalize_text(value) -> str:
    text = clean_string(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.casefold().replace("ß", "ss")
    text = re.sub(r"\s+", " ", text)
    return text.strip()



def first_nonempty(series):
    for value in series:
        text = clean_string(value)
        if text:
            return text
    return ""



def join_unique(values, separator=" | ", max_items=None):
    result = []
    seen = set()
    for value in values:
        text = clean_string(value)
        if not text or text in seen:
            continue
        seen.add(text)
        result.append(text)

    if max_items is not None and len(result) > max_items:
        remaining = len(result) - max_items
        result = result[:max_items] + [f"... (+{remaining} weitere)"]

    return separator.join(result)



def first_existing_column(df, candidates):
    for candidate in candidates:
        if candidate in df.columns:
            return candidate

    normalized_columns = {normalize_text(column): column for column in df.columns}
    for candidate in candidates:
        key = normalize_text(candidate)
        if key in normalized_columns:
            return normalized_columns[key]
    return None


# =============================================================================
# 5. FERTIGE GRUPPEN EINLESEN
# =============================================================================


def read_table_auto(folder: Path, basename: str, required=True):
    parquet_path = folder / f"{basename}.parquet"
    csv_path = folder / f"{basename}.csv"

    if parquet_path.exists():
        try:
            return pd.read_parquet(parquet_path)
        except Exception as exc:
            warnings.warn(
                f"{parquet_path.name} konnte nicht gelesen werden: {exc}. "
                "CSV-Fallback wird versucht."
            )

    if csv_path.exists():
        return pd.read_csv(
            csv_path,
            sep=";",
            encoding="utf-8-sig",
            low_memory=False,
        )

    if required:
        raise FileNotFoundError(
            f"Weder {parquet_path} noch {csv_path} gefunden."
        )
    return pd.DataFrame()



def find_group_folders():
    if not GROUP_PACKAGE_DIR.exists():
        raise FileNotFoundError(
            "Ordner mit fertigen Generator-Datenpaketen nicht gefunden:\n"
            f"{GROUP_PACKAGE_DIR}"
        )

    folders = sorted(
        [path for path in GROUP_PACKAGE_DIR.glob(GROUP_FOLDER_PATTERN) if path.is_dir()],
        key=lambda path: path.name,
    )

    if TARGET_GROUPS is None:
        return folders

    requested = {clean_string(group).upper() for group in TARGET_GROUPS}
    selected = [path for path in folders if path.name.upper() in requested]

    missing = sorted(requested - {path.name.upper() for path in selected})
    if missing:
        warnings.warn(
            "Folgende angeforderte Gruppen wurden nicht gefunden: "
            + ", ".join(missing)
        )
    return selected



def load_final_site_ids():
    rows = []

    for folder in find_group_folders():
        group_id = folder.name
        sites = read_table_auto(folder, "sites", required=True)

        site_col = first_existing_column(sites, ["Standortverbund_ID"])
        if site_col is None:
            raise ValueError(f"{group_id}: sites enthält keine Standortverbund_ID.")

        for site_id in sites[site_col].map(clean_string):
            if site_id:
                rows.append(
                    {
                        "Experiment_Gruppe": group_id,
                        "Standortverbund_ID": site_id,
                    }
                )

    result = pd.DataFrame(rows).drop_duplicates()
    if result.empty:
        raise RuntimeError("Keine verbliebenen Standorte gefunden.")
    return result


# =============================================================================
# 6. STANDORTMETADATEN AUS DER MASTERDATEI
# =============================================================================


def detect_excel_header(path: Path, sheet_name, expected_candidates, max_header_row=6):
    for header_row in range(max_header_row):
        try:
            test = pd.read_excel(
                path,
                sheet_name=sheet_name,
                header=header_row,
                nrows=5,
            )
        except Exception:
            continue

        normalized = {normalize_text(column) for column in test.columns}
        for candidate in expected_candidates:
            if candidate in test.columns or normalize_text(candidate) in normalized:
                return header_row
    return 0



def read_master_potentials():
    if not MASTER_POTENTIAL_FILE.exists():
        raise FileNotFoundError(
            "Zentrale Potentialdatei nicht gefunden:\n"
            f"{MASTER_POTENTIAL_FILE}"
        )

    header = detect_excel_header(
        MASTER_POTENTIAL_FILE,
        MASTER_POTENTIAL_SHEET,
        ["Standortverbund_ID", "Standortverbund_ID_final"],
    )

    return pd.read_excel(
        MASTER_POTENTIAL_FILE,
        sheet_name=MASTER_POTENTIAL_SHEET,
        header=header,
    )



def build_master_site_metadata(master):
    # WICHTIG:
    # Hier werden ausschließlich Identifikations- und Brancheninformationen
    # extrahiert. KEINE Abwärme-, Q-, Temperatur-, Leistungs- oder
    # Verfügbarkeitsinformationen gelangen in die KI-Datei.

    site_col = first_existing_column(
        master,
        ["Standortverbund_ID", "Standortverbund_ID_final", "Profil_Standortverbund_ID"],
    )
    if site_col is None:
        raise ValueError("Keine Standortverbund_ID in der Masterdatei gefunden.")

    company_col = first_existing_column(master, ["Firmenname", "Profil_Firmenname"])
    location_col = first_existing_column(master, ["Standortname", "Profil_Standortname"])
    street_col = first_existing_column(
        master,
        [
            "Straße und Hausnummer",
            "Strasse und Hausnummer",
            "Profil_Straße und Hausnummer",
            "Profil_Strasse und Hausnummer",
        ],
    )
    plz_col = first_existing_column(master, ["PLZ", "Profil_PLZ"])
    city_col = first_existing_column(master, ["Ort", "Profil_Ort"])

    wz_col = first_existing_column(
        master,
        [
            "Profil_WZ_Final_Code",
            "WZ_Final_Code",
            "Gruppen_Profil_WZ_Code",
            "Profil_WZ_Code_final",
            "Profil_WZ_Code",
        ],
    )
    wz3_col = first_existing_column(
        master,
        ["Profil_WZ3_Code", "WZ3_Code", "Gruppen_WZ3_Code", "WZ3_Code_final"],
    )
    wz_label_col = first_existing_column(
        master,
        ["WZ_Bezeichnung", "Profil_WZ_Bezeichnung", "WZ2025_Bezeichnung"],
    )

    temp = master.copy()
    temp["_site"] = temp[site_col].map(clean_string)
    temp = temp[temp["_site"].ne("")].copy()

    def values(frame, col):
        if col is None:
            return pd.Series("", index=frame.index, dtype=object)
        return frame[col]

    rows = []

    for site_id, group in temp.groupby("_site", sort=False):
        company = first_nonempty(values(group, company_col))
        location = first_nonempty(values(group, location_col))
        street = first_nonempty(values(group, street_col))
        plz = first_nonempty(values(group, plz_col))
        city = first_nonempty(values(group, city_col))

        city_line = " ".join(part for part in [plz, city] if part)
        full_address = ", ".join(part for part in [street, city_line] if part)

        wz = join_unique(values(group, wz_col), separator=", ")
        wz3 = join_unique(values(group, wz3_col), separator=", ")
        wz_label = join_unique(values(group, wz_label_col), separator=" | ", max_items=5)

        search_string = " ".join(
            part for part in [company, location, full_address] if part
        )

        production_search = " ".join(
            part
            for part in [
                company,
                city,
                "Werk Produktion Kapazität Durchsatz Mitarbeiter Umsatz",
            ]
            if part
        )

        rows.append(
            {
                "Standortverbund_ID": site_id,
                "Firmenname": company,
                "Standortname": location,
                "Strasse_und_Hausnummer": street,
                "PLZ": plz,
                "Ort": city,
                "Adresse_vollstaendig": full_address,
                "WZ_Final_Code": wz,
                "WZ3_Code": wz3,
                "WZ_Bezeichnung": wz_label,
                "Recherche_Suchstring": search_string,
                "Recherche_Suchstring_Produktion": production_search,
            }
        )

    return pd.DataFrame(rows)


# =============================================================================
# 7. KI-RECHERCHETABELLE AUFBAUEN
# =============================================================================


def build_research_table(final_sites, master_metadata):
    result = final_sites.merge(
        master_metadata,
        on="Standortverbund_ID",
        how="left",
        validate="many_to_one",
    )

    result = result.sort_values(
        ["Experiment_Gruppe", "Firmenname", "Ort", "Standortverbund_ID"],
        kind="stable",
        na_position="last",
    ).reset_index(drop=True)

    result.insert(
        0,
        "Recherche_ID",
        [f"R_{number:05d}" for number in range(1, len(result) + 1)],
    )

    result["Recherche_Ziel"] = (
        "Reale Standortgröße ausschließlich aus unabhängigen externen "
        "Produktions-, Kapazitäts-, Durchsatz-, Beschäftigten-, Umsatz- oder "
        "anderen geeigneten Standortkennzahlen bestimmen. Keine PFA-Abwärmedaten "
        "zur Größenbewertung verwenden."
    )

    for column in KI_COLUMNS:
        result[column] = ""

    identity_columns = [
        "Recherche_ID",
        "Experiment_Gruppe",
        "Standortverbund_ID",
        "Firmenname",
        "Standortname",
        "Strasse_und_Hausnummer",
        "PLZ",
        "Ort",
        "Adresse_vollstaendig",
        "WZ_Final_Code",
        "WZ3_Code",
        "WZ_Bezeichnung",
        "Recherche_Suchstring",
        "Recherche_Suchstring_Produktion",
        "Recherche_Ziel",
    ]

    ordered = [column for column in identity_columns + KI_COLUMNS if column in result.columns]
    remaining = [column for column in result.columns if column not in ordered]

    # Sicherheitsprüfung: Keine Abwärme-/PFA-Technikspalten dürfen in die
    # Ausgabedatei gelangen.
    forbidden_patterns = [
        "waermemenge",
        "wärmemenge",
        "abwaerme",
        "abwärme",
        "quellklasse",
        "q_typ",
        "q1",
        "q2",
        "q3",
        "q4",
        "q5",
        "q6",
        "q7",
        "temperatur",
        "betriebsprofil",
        "vollast",
        "leistung",
    ]

    safe_remaining = []
    for column in remaining:
        normalized = normalize_text(column)
        if any(pattern in normalized for pattern in forbidden_patterns):
            continue
        safe_remaining.append(column)

    result = result[ordered + safe_remaining]

    # Zweite harte Sicherheitsprüfung.
    for column in result.columns:
        normalized = normalize_text(column)
        if any(pattern in normalized for pattern in forbidden_patterns):
            raise RuntimeError(
                "Abwärmebezogene Spalte ist unerwartet in der KI-Ausgabe verblieben: "
                f"{column}"
            )

    return result


# =============================================================================
# 8. ANLEITUNG UND KENNZAHLENKATALOG
# =============================================================================


def build_instruction_table():
    rows = [
        [
            "Ziel",
            "Die reale Größe des konkreten Standorts recherchieren und relativ zu den anderen Standorten derselben Profilgruppe einordnen.",
        ],
        [
            "Keine Abwärmedaten",
            "Die Arbeitsdatei enthält absichtlich keine Abwärmemengen oder technischen Abwärmeinformationen. Solche Informationen dürfen auch nicht extern zur Herleitung des Größenscores herangezogen werden.",
        ],
        [
            "Bis zu 5 Kennzahlen",
            "Pro Standort können bis zu fünf unabhängige Größenkennzahlen recherchiert werden. Es müssen nicht alle fünf befüllt werden.",
        ],
        [
            "Bevorzugte Kennzahlen",
            "Produktionsmenge, Produktionskapazität und Durchsatz sind möglichst zu bevorzugen. Danach folgen branchenspezifische physische Kennzahlen und Beschäftigte am Standort.",
        ],
        [
            "Standortbezug",
            "Standort- bzw. werksspezifische Kennzahlen haben Vorrang. Unternehmens- oder Konzernwerte nur als klar gekennzeichneten Fallback verwenden.",
        ],
        [
            "Mehrere Kennzahlen kombinieren",
            "Der Größenscore soll nicht automatisch aus nur einer Zahl entstehen. Nutze alle fachlich geeigneten und ausreichend vergleichbaren Kennzahlen, die mit 'Fuer_Score_verwenden = JA' markiert sind.",
        ],
        [
            "Vergleichbarkeit",
            "Kennzahlen mit unterschiedlichen Einheiten nicht direkt numerisch gegeneinander verrechnen. Erst innerhalb vergleichbarer Kennzahltypen bzw. fachlich vergleichbarer Standorte relativ einordnen und anschließend qualitativ zusammenführen.",
        ],
        [
            "Score 0-100",
            "0 bedeutet sehr klein, 100 sehr groß INNERHALB DER JEWEILIGEN PROFILGRUPPE. Der Score ist kein absoluter branchenübergreifender Größenindex.",
        ],
        [
            "Quellenqualität",
            "Primärquellen bevorzugen: Unternehmens-/Werksseiten, Geschäftsberichte, Nachhaltigkeitsberichte, Genehmigungsunterlagen, Behörden oder Branchenverbände.",
        ],
        [
            "Unsicherheit",
            "Keine Zahlen erfinden. Bei schwacher Quellenlage oder schwer vergleichbaren Kennzahlen das Vertrauen reduzieren oder keinen Score vergeben.",
        ],
    ]

    return pd.DataFrame(rows, columns=["Thema", "Regel"])



def build_metric_catalog():
    rows = [
        [1, "PRODUKTIONSMENGE", "Tatsächlich erzeugte Produktmenge", "t/a, Stück/a, m³/a, Fahrzeuge/a ...", "SEHR_HOCH"],
        [2, "PRODUKTIONSKAPAZITAET", "Maximale / installierte Produktionskapazität", "t/a, Stück/a, m³/a ...", "SEHR_HOCH"],
        [3, "DURCHSATZ", "Verarbeitete Rohstoff- oder Materialmenge", "t/a, m³/a, Stück/a ...", "HOCH"],
        [4, "SEKTORALE_PHYSISCHE_KENNZAHL", "Branchenspezifische reale Größenkennzahl", "Betten, Fälle/a, Zimmer, Anlagenleistung ...", "HOCH"],
        [5, "MITARBEITER_STANDORT", "Beschäftigte am konkreten Standort", "Personen", "MITTEL"],
        [6, "UMSATZ_STANDORT", "Umsatz des konkreten Standorts / Werks", "EUR/a", "MITTEL"],
        [7, "ANLAGENKAPAZITAET", "Installierte Anlagen- oder Linienkapazität", "branchenspezifisch", "MITTEL_BIS_HOCH"],
        [8, "SONSTIGE_STANDORTKENNZAHL", "Andere nachvollziehbare standortbezogene Größenkennzahl", "branchenspezifisch", "ABHAENGIG_VOM_FALL"],
        [9, "UNTERNEHMENSWERT_FALLBACK", "Unternehmens-/Konzernwert ohne eindeutige Standortzuordnung", "EUR/a, Personen, Produktion ...", "NIEDRIG"],
    ]

    return pd.DataFrame(
        rows,
        columns=["Prioritaet", "Kennzahltyp", "Beschreibung", "Typische_Einheit", "Eignung_fuer_Standortgroesse"],
    )



def build_group_summary(research):
    return (
        research.groupby("Experiment_Gruppe", dropna=False)
        .agg(Anzahl_Standorte=("Standortverbund_ID", "nunique"))
        .reset_index()
        .sort_values("Experiment_Gruppe")
    )



def build_duplicate_site_table(research):
    counts = research.groupby("Standortverbund_ID")["Experiment_Gruppe"].nunique()
    duplicate_ids = counts[counts > 1].index

    columns = [
        "Standortverbund_ID",
        "Experiment_Gruppe",
        "Firmenname",
        "Standortname",
        "Ort",
    ]

    if len(duplicate_ids) == 0:
        return pd.DataFrame(columns=columns)

    return (
        research[research["Standortverbund_ID"].isin(duplicate_ids)][columns]
        .sort_values(["Standortverbund_ID", "Experiment_Gruppe"])
        .reset_index(drop=True)
    )


# =============================================================================
# 9. EXCEL FORMATIEREN
# =============================================================================


def autosize_worksheet(ws, max_width=MAX_COLUMN_WIDTH):
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(1, min(ws.max_row, 250) + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), max_width)



def add_list_validation(ws, header_name, values):
    header_map = {cell.value: cell.column for cell in ws[1]}
    if header_name not in header_map:
        return

    letter = get_column_letter(header_map[header_name])
    formula = '"{}"'.format(",".join(values))
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "Bitte einen vorgesehenen Wert auswählen."
    validation.errorTitle = "Ungültiger Wert"
    ws.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{max(ws.max_row, 2)}")



def format_workbook(path):
    wb = load_workbook(path)

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    identity_fill = PatternFill(fill_type="solid", fgColor="E7E6E6")
    search_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    metric_fills = [
        PatternFill(fill_type="solid", fgColor="E2F0D9"),
        PatternFill(fill_type="solid", fgColor="DDEBF7"),
        PatternFill(fill_type="solid", fgColor="FCE4D6"),
        PatternFill(fill_type="solid", fgColor="E4DFEC"),
        PatternFill(fill_type="solid", fgColor="FFF2CC"),
    ]
    score_fill = PatternFill(fill_type="solid", fgColor="C6E0B4")
    thin_grey = Side(style="thin", color="D9D9D9")

    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        if ws.title == "Standorte_KI":
            header_map = {cell.value: cell.column for cell in ws[1]}

            for header, col_idx in header_map.items():
                if not header:
                    continue

                fill = None
                header_str = str(header)

                if header_str.startswith("KI_Kennzahl_"):
                    match = re.match(r"KI_Kennzahl_(\d+)_", header_str)
                    if match:
                        number = int(match.group(1))
                        if 1 <= number <= 5:
                            fill = metric_fills[number - 1]
                elif header_str in SCORE_COLUMNS:
                    fill = score_fill
                elif header_str.startswith("KI_"):
                    fill = metric_fills[0]
                elif header_str.startswith("Recherche_") and header_str != "Recherche_ID":
                    fill = search_fill
                else:
                    fill = identity_fill

                for row_idx in range(2, ws.max_row + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = fill

            add_list_validation(ws, "KI_Bearbeitungsstatus", ["OFFEN", "RECHERCHIERT", "NICHT_GEFUNDEN", "MEHRDEUTIG"])
            add_list_validation(ws, "KI_Standortidentifikation_Sicherheit", QUALITY_LEVELS)

            for number in range(1, 6):
                prefix = f"KI_Kennzahl_{number}"
                add_list_validation(ws, f"{prefix}_Typ", METRIC_TYPES)
                add_list_validation(ws, f"{prefix}_Bezug", METRIC_REFERENCES)
                add_list_validation(ws, f"{prefix}_Datentyp", METRIC_DATA_TYPES)
                add_list_validation(ws, f"{prefix}_Vertrauen", QUALITY_LEVELS)
                add_list_validation(ws, f"{prefix}_Vergleichbarkeit_innerhalb_Gruppe", COMPARABILITY_LEVELS)
                add_list_validation(ws, f"{prefix}_Fuer_Score_verwenden", SCORE_USE_VALUES)

            add_list_validation(ws, "KI_Groessenklasse_innerhalb_Gruppe", SIZE_CLASSES)
            add_list_validation(ws, "KI_Vertrauen_Groessenschaetzung", QUALITY_LEVELS)

        ws.freeze_panes = FREEZE_PANES
        if ws.max_row >= 2:
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin_grey)

        autosize_worksheet(ws)

    wb.save(path)


# =============================================================================
# 10. EXCEL SCHREIBEN
# =============================================================================


def write_research_workbook(research, output_file):
    output_file.parent.mkdir(parents=True, exist_ok=True)

    instructions = build_instruction_table()
    metric_catalog = build_metric_catalog()
    group_summary = build_group_summary(research)
    duplicates = build_duplicate_site_table(research)

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        research.to_excel(writer, sheet_name="Standorte_KI", index=False)
        instructions.to_excel(writer, sheet_name="Anleitung_KI", index=False)
        metric_catalog.to_excel(writer, sheet_name="Kennzahlenkatalog", index=False)
        group_summary.to_excel(writer, sheet_name="Gruppen_Uebersicht", index=False)
        duplicates.to_excel(writer, sheet_name="Doppelte_Standorte", index=False)

    format_workbook(output_file)


# =============================================================================
# 11. HAUPTLAUF
# =============================================================================


def main():
    print()
    print("=" * 100)
    print("PFA – KI-RECHERCHEEXPORT ZUR STANDORTGRÖSSE – VERSION 2")
    print("=" * 100)

    print()
    print("1. Finale verbliebene Standorte einlesen ...")
    final_sites = load_final_site_ids()
    print(f"   Standort-Gruppen-Zuordnungen: {len(final_sites):,}")

    print()
    print("2. Nur Firmen-, Adress- und WZ-Informationen ergänzen ...")
    master = read_master_potentials()
    master_metadata = build_master_site_metadata(master)
    print(f"   Standortmetadaten: {len(master_metadata):,}")

    print()
    print("3. KI-Recherchetabelle ohne Abwärmeinformationen aufbauen ...")
    research = build_research_table(final_sites, master_metadata)

    print(f"   Standorte gesamt: {len(research):,}")
    print(f"   Spalten gesamt:   {len(research.columns):,}")

    print()
    print("4. Masterdatei schreiben ...")
    write_research_workbook(research, OUTPUT_FILE)
    print(f"   {OUTPUT_FILE}")

    if WRITE_ONE_FILE_PER_GROUP:
        group_dir = OUTPUT_DIR / "Einzelgruppen"
        group_dir.mkdir(parents=True, exist_ok=True)

        print()
        print("5. Einzeldateien je Profilgruppe schreiben ...")

        for group_id, subset in research.groupby("Experiment_Gruppe", sort=True):
            group_file = group_dir / f"KI_Standortgroesse_{group_id}_v2.xlsx"
            write_research_workbook(subset.reset_index(drop=True), group_file)
            print(f"   {group_id}: {len(subset):,} Standorte -> {group_file.name}")

    print()
    print("=" * 100)
    print("EXPORT ABGESCHLOSSEN")
    print("=" * 100)
    print()
    print("SICHERHEITSPRINZIP:")
    print("- In der Excel-Ausgabe befinden sich keine PFA-Abwärmemengen.")
    print("- Es befinden sich auch keine Q-Typen, Temperaturen, Leistungen oder Betriebsprofile darin.")
    print("- Bis zu fünf externe Größenkennzahlen können pro Standort recherchiert werden.")
    print("- Nur ausdrücklich markierte Kennzahlen sollen in den Größenscore eingehen.")
    print("- Der Score gilt relativ innerhalb der jeweiligen Profilgruppe.")
    print()


if __name__ == "__main__":
    main()
