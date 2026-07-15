import argparse
import json
import re
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.text import Text
from openpyxl import load_workbook
from patsy import build_design_matrices, dmatrix
from scipy.optimize import curve_fit
from sklearn.ensemble import RandomForestRegressor


RESULT_RE = re.compile(
    r"KPIs_district_(?P<district>[A-Z])_seed_(?P<seed>\d+)_buildings_30_"
    r"(?P<system>central|decentral)_cost_(?P<cost>min|mean|max)\.xlsx$"
)

COST_CASES = ("min", "mean", "max")
REGRESSION_TARGETS = {
    "lower": "min_lcoh_ct_per_kwh",
    "mean": "mean_lcoh_ct_per_kwh",
    "upper": "max_lcoh_ct_per_kwh",
}
MODEL_SPECS = [
    ("linear_density_only", "linear", False, False, False),
    ("log_density_only", "log", False, False, False),
    ("log_density_plus_lhd_tertile", "log", False, True, False),
    ("log_density_times_lhd_tertile", "log", False, True, True),
    ("log_density_plus_district", "log", True, False, False),
]
VALIDATION_REPEATS = 100
VALIDATION_TEST_FRACTION = 0.20
VALIDATION_RANDOM_SEED = 42
PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RESULTS_DIR = PROJECT_ROOT / "districtgenerator" / "results" / "results_paper_2"
DEFAULT_OUTPUT_PREFIX = DEFAULT_RESULTS_DIR / "central_lcoh_vs_heat_density"
DEFAULT_CONFIG_PATH = PROJECT_ROOT / "districtgenerator" / "data" / ".env.CONFIG.PAPER"

DISTRICT_COLORS = {
    "A": "#0072B2",
    "B": "#D55E00",
    "C": "#009E73",
    "D": "#CC79A7",
    "E": "#E69F00",
    "F": "#56B4E9",
    "H": "#7F3C8D",
    "I": "#6B6B6B",
}

DISPLAY_DISTRICT_LABELS = {
    "H": "G",
    "I": "H",
}
TERTILE_COLORS = {"low": "#0072B2", "medium": "#E69F00", "high": "#009E73"}
FIGURE_SIZE = (10.5, 7.0)
BAND_ALPHA = 0.16
MODEL_DISPLAY_NAMES = {
    "linear_density_only": "Linear",
    "log_density_only": "Logarithmic",
    "log_density_plus_lhd_tertile": "Log. + density class",
    "log_density_times_lhd_tertile": "Log. x density class",
    "log_density_plus_district": "Log. + district type",
}
CANDIDATE_MODEL_DISPLAY_NAMES = {
    "candidate_linear": "Linear",
    "candidate_log": "Logarithmic",
    "candidate_shifted_log": "Shifted logarithmic",
    "candidate_exp_saturation": "Exponential saturation",
    "candidate_power_law": "Power law",
    "candidate_gam_lhd": "GAM / spline",
    "candidate_gam_lhd_district": "GAM / spline + district type",
    "candidate_random_forest": "Random forest",
}


def apply_publication_style():
    plt.rcParams.update(
        {
            "figure.dpi": 120,
            "savefig.dpi": 600,
            "font.size": 13,
            "axes.labelsize": 15,
            "axes.titlesize": 15,
            "axes.titleweight": "bold",
            "axes.linewidth": 1.1,
            "xtick.labelsize": 12,
            "ytick.labelsize": 12,
            "legend.fontsize": 11,
            "legend.title_fontsize": 12,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "svg.fonttype": "path",
        }
    )


def district_color_map(districts):
    return {
        district: DISTRICT_COLORS.get(district, plt.get_cmap("tab10")(i))
        for i, district in enumerate(districts)
    }


def display_district(district):
    return DISPLAY_DISTRICT_LABELS.get(district, district)


def save_presentation_svg(fig, path):
    for axis in fig.axes:
        axis.set_title("")
        axis.grid(False, which="both")
        axis.xaxis.grid(False, which="both")
        axis.yaxis.grid(False, which="both")
    for text in fig.findobj(match=Text):
        text.set_fontsize(text.get_fontsize() * 1.25)
    fig.savefig(path, transparent=True)


def row_values_by_kpi(ws):
    values = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] is not None:
            values[str(row[0]).strip()] = list(row[1:])
    return values


def find_kpi_row(values, row_name):
    if row_name in values:
        return values[row_name]
    for key, row in values.items():
        if key.startswith(row_name):
            return row
    raise KeyError(f"Could not find KPI row starting with '{row_name}'.")


def numeric_values(values):
    return [float(value) for value in values if isinstance(value, (int, float))]


def mean_kpi_row(values, row_name):
    numbers = numeric_values(find_kpi_row(values, row_name))
    if not numbers:
        raise ValueError(f"KPI row has no numeric values: {row_name}")
    return sum(numbers) / len(numbers)


def parse_config_value(raw_value):
    value = raw_value.split("#", 1)[0].strip()
    if value.startswith("[") and value.endswith("]"):
        return [float(item.strip()) for item in value[1:-1].split(",") if item.strip()]
    if "," in value:
        return [float(item.strip()) for item in value.split(",") if item.strip()]
    return float(value)


def read_config_values(config_path):
    values = {}
    for line in config_path.read_text(encoding="utf-8").splitlines():
        if not line.strip() or line.lstrip().startswith("#") or "=" not in line:
            continue
        key, raw_value = line.split("=", 1)
        key = key.strip()
        if key in {
            "PRICE_SUPPLY_EL",
            "PRICE_SUPPLY_EL_EH",
            "REVENUE_FEED_IN_EL",
            "REVENUE_FEED_IN_EL_EH",
        }:
            values[key] = parse_config_value(raw_value)
    return values


def value_for_year(config_values, key, index):
    value = config_values[key]
    if isinstance(value, list):
        return value[index]
    return value


def sum_column(ws, column_name):
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    if column_name in header:
        idx = header.index(column_name)
    else:
        idx = next(
            (
                i for i, value in enumerate(header)
                if isinstance(value, str) and value.startswith(column_name)
            ),
            None,
        )
    if idx is None:
        return 0.0
    total = 0.0
    for row in ws.iter_rows(min_row=2, values_only=True):
        value = row[idx]
        if isinstance(value, (int, float)):
            total += float(value)
    return total


def central_heat_net_annual_cost(kpi_path, config_values):
    wb = load_workbook(kpi_path, data_only=True, read_only=True)
    yearly = row_values_by_kpi(wb["Yearly KPIs"])

    electricity_costs = numeric_values(find_kpi_row(yearly, "Electricity Costs"))
    total_feed_in_revenue = numeric_values(find_kpi_row(yearly, "Revenue from Electricity Feed-in"))
    building_el_demand = numeric_values(find_kpi_row(yearly, "Electricity Demand within District"))
    building_el_feed_in = numeric_values(find_kpi_row(yearly, "Electricity Injection within District"))

    fuel_row_names = [
        "Gas Costs",
        "Biomethane Costs",
        "Oil Costs",
        "Waste Costs",
        "Biomass Costs",
        "District Heat Costs",
        "Hydrogen Costs",
    ]
    fuel_costs_by_year = []
    for row_name in fuel_row_names:
        try:
            fuel_costs_by_year.append(numeric_values(find_kpi_row(yearly, row_name)))
        except KeyError:
            fuel_costs_by_year.append([0.0] * len(electricity_costs))

    annual_heat_operation_costs = []
    for i, electricity_cost in enumerate(electricity_costs):
        building_import_cost = building_el_demand[i] * value_for_year(config_values, "PRICE_SUPPLY_EL", i)
        building_feed_in_revenue = -building_el_feed_in[i] * value_for_year(config_values, "REVENUE_FEED_IN_EL", i)
        energy_hub_electricity_cost = electricity_cost - building_import_cost
        energy_hub_feed_in_revenue = total_feed_in_revenue[i] - building_feed_in_revenue
        central_fuel_costs = sum(fuel_costs[i] for fuel_costs in fuel_costs_by_year)
        annual_heat_operation_costs.append(
            energy_hub_electricity_cost + energy_hub_feed_in_revenue + central_fuel_costs
        )
    mean_heat_operation_cost = sum(annual_heat_operation_costs) / len(annual_heat_operation_costs)

    central_costs = 0.0
    if "Central Devices Costs" in wb.sheetnames:
        central_costs = sum_column(wb["Central Devices Costs"], "Annualized Cost Subsidized")

    return mean_heat_operation_cost + central_costs


def scalar_json_value(data, key):
    value = data[key]
    if isinstance(value, dict) and "value" in value:
        return float(value["value"])
    return float(value)


def central_heat_grid_values(results_dir, district, seed):
    network_dir = results_dir / "network"
    candidates = sorted(
        network_dir.glob(
            f"district_{district}_seed_{seed}_buildings_30_central_cost_mean_*"
            "/heat_grid_parameters_outputs.json"
        )
    )
    if not candidates:
        candidates = sorted(
            network_dir.glob(
                f"district_{district}_seed_{seed}_buildings_30_central_cost_*_*"
                "/heat_grid_parameters_outputs.json"
            )
        )
    if not candidates:
        return None

    data = json.loads(candidates[0].read_text(encoding="utf-8"))
    heat_kwh = scalar_json_value(data, "total_net_heat_demand")
    pipe_m = scalar_json_value(data, "total_pipe_length")
    if heat_kwh <= 0 or pipe_m <= 0:
        return None
    return heat_kwh, pipe_m, heat_kwh / pipe_m


def collect_kpis(results_dir):
    records = {}
    for path in results_dir.glob("KPIs_district_*_seed_*_buildings_30_*_cost_*.xlsx"):
        match = RESULT_RE.match(path.name)
        if not match:
            continue
        key = (
            match.group("district"),
            int(match.group("seed")),
            match.group("system"),
            match.group("cost"),
        )
        records[key] = path
    return records


def make_ranges(results_dir, config_values, require_complete_decentral=False):
    kpis = collect_kpis(results_dir)
    district_seeds = sorted({(district, seed) for district, seed, _, _ in kpis})
    ranges = []
    cases = []

    for district, seed in district_seeds:
        heat_grid_values = central_heat_grid_values(results_dir, district, seed)
        if heat_grid_values is None:
            continue
        total_heat_kwh, pipe_length_m, density = heat_grid_values

        central_paths = {
            cost: kpis.get((district, seed, "central", cost)) for cost in COST_CASES
        }
        if any(path is None for path in central_paths.values()):
            continue
        if require_complete_decentral:
            decentral_paths = {
                cost: kpis.get((district, seed, "decentral", cost)) for cost in COST_CASES
            }
            if any(path is None for path in decentral_paths.values()):
                continue

        lcoh_by_case = {}
        net_cost_by_case = {}
        for cost_case, path in central_paths.items():
            net_cost = central_heat_net_annual_cost(path, config_values)
            lcoh = 100.0 * net_cost / total_heat_kwh
            net_cost_by_case[cost_case] = net_cost
            lcoh_by_case[cost_case] = lcoh
            cases.append(
                {
                    "district": district,
                    "seed": seed,
                    "cost_case": cost_case,
                    "linear_heat_density_kwh_per_m_a": density,
                    "total_net_heat_demand_kwh_per_a": total_heat_kwh,
                    "total_pipe_length_m": pipe_length_m,
                    "net_annual_cost_eur_per_a": net_cost,
                    "lcoh_ct_per_kwh": lcoh,
                }
            )

        lcoh_values = list(lcoh_by_case.values())
        ranges.append(
            {
                "district": district,
                "seed": seed,
                "linear_heat_density_kwh_per_m_a": density,
                "total_net_heat_demand_kwh_per_a": total_heat_kwh,
                "total_pipe_length_m": pipe_length_m,
                "min_lcoh_ct_per_kwh": min(lcoh_values),
                "mean_lcoh_ct_per_kwh": lcoh_by_case["mean"],
                "max_lcoh_ct_per_kwh": max(lcoh_values),
                "central_min_lcoh_ct_per_kwh": lcoh_by_case["min"],
                "central_mean_lcoh_ct_per_kwh": lcoh_by_case["mean"],
                "central_max_lcoh_ct_per_kwh": lcoh_by_case["max"],
                "central_mean_net_annual_cost_eur_per_a": net_cost_by_case["mean"],
            }
        )

    return pd.DataFrame(ranges), pd.DataFrame(cases)


def add_lhd_tertiles(df):
    df = df.copy()
    df["lhd_tertile"] = pd.qcut(
        df["linear_heat_density_kwh_per_m_a"],
        q=3,
        labels=["low", "medium", "high"],
        duplicates="drop",
    )
    return df


def set_scientific_axes(ax):
    ax.set_xlabel(r"Annual linear heat density, $q_\mathrm{L}$ (kWh m$^{-1}$ a$^{-1}$)")
    ax.set_ylabel(r"LCOH (ct kWh$^{-1}$)")
    ax.grid(True, which="major", color="#D0D0D0", linewidth=0.8, alpha=0.75)
    ax.grid(True, which="minor", color="#E8E8E8", linewidth=0.5, alpha=0.55)
    ax.minorticks_on()
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(axis="both", which="major", length=5, width=1.0)
    ax.tick_params(axis="both", which="minor", length=3, width=0.8)


def add_model_label(ax, text):
    ax.text(
        0.025,
        0.975,
        text,
        transform=ax.transAxes,
        va="top",
        ha="left",
        fontsize=11,
        bbox={
            "boxstyle": "round,pad=0.35",
            "facecolor": "white",
            "edgecolor": "#B8B8B8",
            "alpha": 0.94,
        },
    )


def density_term(df, transform):
    if transform == "linear":
        return df["linear_heat_density_kwh_per_m_a"] / 1000.0
    if transform == "log":
        return np.log(df["linear_heat_density_kwh_per_m_a"])
    raise ValueError("transform must be 'linear' or 'log'.")


def density_term_from_values(density_kwh_per_m, transform):
    if transform == "linear":
        return density_kwh_per_m / 1000.0
    if transform == "log":
        return np.log(density_kwh_per_m)
    raise ValueError("transform must be 'linear' or 'log'.")


def fit_ols(
    df,
    model_name,
    response_name="mean",
    response_column="mean_lcoh_ct_per_kwh",
    transform=None,
    include_district=False,
    include_lhd_tertile=False,
    include_lhd_tertile_interaction=False,
):
    model_df = df.copy()

    x_parts = [pd.Series(1.0, index=model_df.index, name="intercept")]
    if transform is not None:
        density_column = "density_1000_kwh_per_m_a" if transform == "linear" else "ln_density_kwh_per_m_a"
        model_df[density_column] = density_term(model_df, transform)
        x_parts.append(model_df[density_column])
    if include_lhd_tertile or include_lhd_tertile_interaction:
        tertile_dummies = pd.get_dummies(model_df["lhd_tertile"], prefix="lhd", drop_first=True, dtype=float)
        x_parts.append(tertile_dummies)
        if include_lhd_tertile_interaction:
            if transform is None:
                raise ValueError("LHD tertile interaction requires a density transform.")
            interaction_terms = tertile_dummies.mul(model_df[density_column], axis=0)
            interaction_terms = interaction_terms.rename(
                columns={column: f"{density_column}_x_{column}" for column in interaction_terms.columns}
            )
            x_parts.append(interaction_terms)
    if include_district:
        dummies = pd.get_dummies(model_df["district"], prefix="district", drop_first=True, dtype=float)
        x_parts.append(dummies)

    x = pd.concat(x_parts, axis=1).astype(float)
    y = model_df[response_column].astype(float)

    results = sm.OLS(y, x).fit()
    n_obs = int(results.nobs)
    rss = float(results.ssr)
    rmse = float(np.sqrt(rss / n_obs))
    summary = {
        "model": model_name,
        "response": response_name,
        "response_column": response_column,
        "density_transform": transform if transform is not None else "none",
        "includes_district_type": include_district,
        "includes_lhd_tertile": include_lhd_tertile,
        "includes_lhd_tertile_interaction": include_lhd_tertile_interaction,
        "n_observations": n_obs,
        "n_parameters": int(results.df_model + 1),
        "df_model": float(results.df_model),
        "df_resid": float(results.df_resid),
        "r2": float(results.rsquared),
        "adjusted_r2": float(results.rsquared_adj),
        "rmse_ct_per_kwh": rmse,
        "residual_standard_error": float(np.sqrt(results.mse_resid)),
        "rss": rss,
        "aic": float(results.aic),
        "bic": float(results.bic),
        "f_statistic": float(results.fvalue) if results.fvalue is not None else np.nan,
        "f_pvalue": float(results.f_pvalue) if results.f_pvalue is not None else np.nan,
        "condition_number": float(results.condition_number),
    }
    conf_int = results.conf_int()
    coefficients = pd.DataFrame(
        {
            "model": model_name,
            "response": response_name,
            "term": x.columns,
            "coefficient": results.params.reindex(x.columns).to_numpy(),
            "std_error": results.bse.reindex(x.columns).to_numpy(),
            "t_value": results.tvalues.reindex(x.columns).to_numpy(),
            "p_value": results.pvalues.reindex(x.columns).to_numpy(),
            "ci_lower_95": conf_int.loc[x.columns, 0].to_numpy(),
            "ci_upper_95": conf_int.loc[x.columns, 1].to_numpy(),
        }
    )
    return summary, coefficients, results.params.to_dict(), results.summary().as_text()


def regression_design_matrix(
    df,
    transform=None,
    include_district=False,
    include_lhd_tertile=False,
    include_lhd_tertile_interaction=False,
    columns=None,
):
    model_df = df.copy()
    x_parts = [pd.Series(1.0, index=model_df.index, name="intercept")]
    density_column = None
    if transform is not None:
        density_column = "density_1000_kwh_per_m_a" if transform == "linear" else "ln_density_kwh_per_m_a"
        model_df[density_column] = density_term(model_df, transform)
        x_parts.append(model_df[density_column])
    if include_lhd_tertile or include_lhd_tertile_interaction:
        tertile_dummies = pd.get_dummies(model_df["lhd_tertile"], prefix="lhd", drop_first=True, dtype=float)
        x_parts.append(tertile_dummies)
        if include_lhd_tertile_interaction:
            if transform is None:
                raise ValueError("LHD tertile interaction requires a density transform.")
            interaction_terms = tertile_dummies.mul(model_df[density_column], axis=0)
            interaction_terms = interaction_terms.rename(
                columns={column: f"{density_column}_x_{column}" for column in interaction_terms.columns}
            )
            x_parts.append(interaction_terms)
    if include_district:
        x_parts.append(pd.get_dummies(model_df["district"], prefix="district", drop_first=True, dtype=float))
    x = pd.concat(x_parts, axis=1).astype(float)
    if columns is not None:
        x = x.reindex(columns=columns, fill_value=0.0)
    return x


def repeated_train_test_validation(
    df,
    repeats=VALIDATION_REPEATS,
    test_fraction=VALIDATION_TEST_FRACTION,
    random_seed=VALIDATION_RANDOM_SEED,
):
    rng = np.random.default_rng(random_seed)
    rows = []
    n_obs = len(df)
    n_test = max(1, int(round(n_obs * test_fraction)))

    for repeat in range(repeats):
        shuffled = rng.permutation(df.index.to_numpy())
        test_index = shuffled[:n_test]
        train_index = shuffled[n_test:]
        train = df.loc[train_index].copy()
        test = df.loc[test_index].copy()

        for response_name, response_column in REGRESSION_TARGETS.items():
            y_train = train[response_column].astype(float)
            y_test = test[response_column].astype(float)
            for model_name, transform, include_district, include_lhd_tertile, include_lhd_tertile_interaction in MODEL_SPECS:
                x_train = regression_design_matrix(
                    train,
                    transform=transform,
                    include_district=include_district,
                    include_lhd_tertile=include_lhd_tertile,
                    include_lhd_tertile_interaction=include_lhd_tertile_interaction,
                )
                x_test = regression_design_matrix(
                    test,
                    transform=transform,
                    include_district=include_district,
                    include_lhd_tertile=include_lhd_tertile,
                    include_lhd_tertile_interaction=include_lhd_tertile_interaction,
                    columns=x_train.columns,
                )
                result = sm.OLS(y_train, x_train).fit()
                y_pred = result.predict(x_test)
                residuals = y_test - y_pred
                sse = float(np.sum(residuals**2))
                sst = float(np.sum((y_test - y_test.mean()) ** 2))
                rows.append(
                    {
                        "repeat": repeat,
                        "model": model_name,
                        "response": response_name,
                        "response_column": response_column,
                        "n_train": len(train),
                        "n_test": len(test),
                        "train_r2": float(result.rsquared),
                        "test_r2": np.nan if sst == 0 else 1.0 - sse / sst,
                        "test_mae_ct_per_kwh": float(np.mean(np.abs(residuals))),
                        "test_rmse_ct_per_kwh": float(np.sqrt(np.mean(residuals**2))),
                    }
                )
    raw = pd.DataFrame(rows)
    summary = (
        raw.groupby(["model", "response"], as_index=False)
        .agg(
            repeats=("repeat", "nunique"),
            train_r2_mean=("train_r2", "mean"),
            train_r2_std=("train_r2", "std"),
            test_r2_mean=("test_r2", "mean"),
            test_r2_std=("test_r2", "std"),
            test_mae_mean_ct_per_kwh=("test_mae_ct_per_kwh", "mean"),
            test_mae_std_ct_per_kwh=("test_mae_ct_per_kwh", "std"),
            test_rmse_mean_ct_per_kwh=("test_rmse_ct_per_kwh", "mean"),
            test_rmse_std_ct_per_kwh=("test_rmse_ct_per_kwh", "std"),
        )
    )
    return raw, summary


def leave_one_seed_out_validation(df):
    rows = []
    for test_index in df.index:
        train = df.drop(index=test_index).copy()
        test = df.loc[[test_index]].copy()
        test_case = f"{test['district'].iloc[0]}_seed_{int(test['seed'].iloc[0])}"

        for response_name, response_column in REGRESSION_TARGETS.items():
            y_train = train[response_column].astype(float)
            y_test = test[response_column].astype(float)
            for model_name, transform, include_district, include_lhd_tertile, include_lhd_tertile_interaction in MODEL_SPECS:
                x_train = regression_design_matrix(
                    train,
                    transform=transform,
                    include_district=include_district,
                    include_lhd_tertile=include_lhd_tertile,
                    include_lhd_tertile_interaction=include_lhd_tertile_interaction,
                )
                x_test = regression_design_matrix(
                    test,
                    transform=transform,
                    include_district=include_district,
                    include_lhd_tertile=include_lhd_tertile,
                    include_lhd_tertile_interaction=include_lhd_tertile_interaction,
                    columns=x_train.columns,
                )
                result = sm.OLS(y_train, x_train).fit()
                y_pred = result.predict(x_test)
                error = float(y_test.iloc[0] - y_pred.iloc[0])
                rows.append(
                    {
                        "test_case": test_case,
                        "model": model_name,
                        "response": response_name,
                        "response_column": response_column,
                        "n_train": len(train),
                        "n_test": 1,
                        "train_r2": float(result.rsquared),
                        "actual": float(y_test.iloc[0]),
                        "predicted": float(y_pred.iloc[0]),
                        "error": error,
                        "absolute_error": abs(error),
                        "squared_error": error**2,
                    }
                )
    raw = pd.DataFrame(rows)
    summary = (
        raw.groupby(["model", "response"], as_index=False)
        .agg(
            folds=("test_case", "nunique"),
            train_r2_mean=("train_r2", "mean"),
            train_r2_std=("train_r2", "std"),
            mae_ct_per_kwh=("absolute_error", "mean"),
            rmse_ct_per_kwh=("squared_error", lambda values: float(np.sqrt(np.mean(values)))),
            error_std_ct_per_kwh=("error", "std"),
        )
    )
    return raw, summary


def candidate_density(df):
    return df["linear_heat_density_kwh_per_m_a"].astype(float).to_numpy()


def shifted_log_function(x, intercept, slope, shift):
    return intercept + slope * np.log(x + shift)


def exponential_saturation_function(x, asymptote, amplitude, rate):
    return asymptote + amplitude * np.exp(-rate * x)


def power_law_function(x, intercept, scale, exponent):
    return intercept + scale * np.power(x, exponent)


def fit_candidate_model(train, response_column, model_name, density_bounds=None):
    x = candidate_density(train)
    y = train[response_column].astype(float).to_numpy()
    x_scaled = x / 1000.0
    if density_bounds is None:
        density_bounds = (float(np.min(x)), float(np.max(x)))

    if model_name == "candidate_linear":
        design = sm.add_constant(x_scaled, has_constant="add")
        result = sm.OLS(y, design).fit()

        def predict(test):
            test_x = candidate_density(test) / 1000.0
            return result.predict(sm.add_constant(test_x, has_constant="add"))

        return result.fittedvalues, predict, int(result.df_model + 1)

    if model_name == "candidate_log":
        design = sm.add_constant(np.log(x), has_constant="add")
        result = sm.OLS(y, design).fit()

        def predict(test):
            test_x = np.log(candidate_density(test))
            return result.predict(sm.add_constant(test_x, has_constant="add"))

        return result.fittedvalues, predict, int(result.df_model + 1)

    if model_name == "candidate_shifted_log":
        min_shift = max(1e-6, -float(np.min(x)) + 1e-6)
        max_shift = max(float(np.max(x)) * 5.0, min_shift * 10.0)
        initial = [float(np.mean(y)), 1.0, max(float(np.min(x)) * 0.1, min_shift)]
        params, _ = curve_fit(
            shifted_log_function,
            x,
            y,
            p0=initial,
            bounds=([-np.inf, -np.inf, min_shift], [np.inf, np.inf, max_shift]),
            maxfev=20000,
        )

        def predict(test):
            return shifted_log_function(candidate_density(test), *params)

        return shifted_log_function(x, *params), predict, 3

    if model_name == "candidate_exp_saturation":
        initial = [float(np.mean(y)), float(y[0] - np.mean(y)), 1.0 / max(float(np.mean(x)), 1.0)]
        params, _ = curve_fit(
            exponential_saturation_function,
            x,
            y,
            p0=initial,
            bounds=([-np.inf, -np.inf, 0.0], [np.inf, np.inf, np.inf]),
            maxfev=20000,
        )

        def predict(test):
            return exponential_saturation_function(candidate_density(test), *params)

        return exponential_saturation_function(x, *params), predict, 3

    if model_name == "candidate_power_law":
        initial = [float(np.mean(y)), 1.0, -0.5]
        params, _ = curve_fit(
            power_law_function,
            x,
            y,
            p0=initial,
            bounds=([-np.inf, -np.inf, -5.0], [np.inf, np.inf, 5.0]),
            maxfev=20000,
        )

        def predict(test):
            return power_law_function(candidate_density(test), *params)

        return power_law_function(x, *params), predict, 3

    if model_name in {"candidate_gam_lhd", "candidate_gam_lhd_district"}:
        train_design = dmatrix(
            "bs(x, df=5, degree=3, include_intercept=False, lower_bound=lower_bound, upper_bound=upper_bound)",
            {"x": x, "lower_bound": density_bounds[0], "upper_bound": density_bounds[1]},
            return_type="dataframe",
        )
        spline_design_info = train_design.design_info
        if model_name == "candidate_gam_lhd_district":
            district_dummies = pd.get_dummies(train["district"], prefix="district", drop_first=True, dtype=float)
            train_design = pd.concat([train_design.reset_index(drop=True), district_dummies.reset_index(drop=True)], axis=1)
        result = sm.OLS(y, train_design.astype(float)).fit()
        columns = train_design.columns

        def predict(test):
            test_design = build_design_matrices(
                [spline_design_info],
                {
                    "x": candidate_density(test),
                    "lower_bound": density_bounds[0],
                    "upper_bound": density_bounds[1],
                },
                return_type="dataframe",
            )[0]
            if model_name == "candidate_gam_lhd_district":
                test_dummies = pd.get_dummies(test["district"], prefix="district", drop_first=True, dtype=float)
                test_design = pd.concat([test_design.reset_index(drop=True), test_dummies.reset_index(drop=True)], axis=1)
            test_design = test_design.reindex(columns=columns, fill_value=0.0)
            return result.predict(test_design.astype(float))

        return result.fittedvalues, predict, int(result.df_model + 1)

    if model_name == "candidate_random_forest":
        if RandomForestRegressor is None:
            raise ImportError("scikit-learn is required for candidate_random_forest")
        train_features = pd.DataFrame({"linear_heat_density_kwh_per_m_a": x})
        district_dummies = pd.get_dummies(train["district"], prefix="district", dtype=float)
        train_features = pd.concat([train_features.reset_index(drop=True), district_dummies.reset_index(drop=True)], axis=1)
        model = RandomForestRegressor(
            n_estimators=500,
            min_samples_leaf=3,
            random_state=VALIDATION_RANDOM_SEED,
        )
        model.fit(train_features, y)
        columns = train_features.columns

        def predict(test):
            test_features = pd.DataFrame({"linear_heat_density_kwh_per_m_a": candidate_density(test)})
            test_dummies = pd.get_dummies(test["district"], prefix="district", dtype=float)
            test_features = pd.concat([test_features.reset_index(drop=True), test_dummies.reset_index(drop=True)], axis=1)
            test_features = test_features.reindex(columns=columns, fill_value=0.0)
            return model.predict(test_features)

        return model.predict(train_features), predict, len(columns)

    raise ValueError(f"Unknown candidate model: {model_name}")


def candidate_model_comparison(df):
    candidate_models = list(CANDIDATE_MODEL_DISPLAY_NAMES)
    fit_rows = []
    loo_rows = []
    density_bounds = (
        float(df["linear_heat_density_kwh_per_m_a"].min()),
        float(df["linear_heat_density_kwh_per_m_a"].max()),
    )

    for response_name, response_column in REGRESSION_TARGETS.items():
        y_all = df[response_column].astype(float).to_numpy()
        for model_name in candidate_models:
            try:
                fitted, _, n_parameters = fit_candidate_model(df, response_column, model_name, density_bounds)
                residuals = y_all - np.asarray(fitted, dtype=float)
                rss = float(np.sum(residuals**2))
                tss = float(np.sum((y_all - np.mean(y_all)) ** 2))
                n_obs = len(y_all)
                fit_rows.append(
                    {
                        "model": model_name,
                        "model_label": CANDIDATE_MODEL_DISPLAY_NAMES[model_name],
                        "response": response_name,
                        "response_column": response_column,
                        "n_observations": n_obs,
                        "n_parameters": n_parameters,
                        "r2": np.nan if tss == 0 else 1.0 - rss / tss,
                        "rmse_ct_per_kwh": float(np.sqrt(np.mean(residuals**2))),
                        "mae_ct_per_kwh": float(np.mean(np.abs(residuals))),
                        "status": "ok",
                    }
                )
            except Exception as exc:
                fit_rows.append(
                    {
                        "model": model_name,
                        "model_label": CANDIDATE_MODEL_DISPLAY_NAMES[model_name],
                        "response": response_name,
                        "response_column": response_column,
                        "n_observations": len(y_all),
                        "n_parameters": np.nan,
                        "r2": np.nan,
                        "rmse_ct_per_kwh": np.nan,
                        "mae_ct_per_kwh": np.nan,
                        "status": f"failed: {exc}",
                    }
                )

            for test_index in df.index:
                train = df.drop(index=test_index).copy()
                test = df.loc[[test_index]].copy()
                try:
                    _, predictor, _ = fit_candidate_model(train, response_column, model_name, density_bounds)
                    predicted = float(np.asarray(predictor(test), dtype=float)[0])
                    actual = float(test[response_column].iloc[0])
                    error = actual - predicted
                    status = "ok"
                except Exception as exc:
                    actual = float(test[response_column].iloc[0])
                    predicted = np.nan
                    error = np.nan
                    status = f"failed: {exc}"
                loo_rows.append(
                    {
                        "test_case": f"{test['district'].iloc[0]}_seed_{int(test['seed'].iloc[0])}",
                        "model": model_name,
                        "model_label": CANDIDATE_MODEL_DISPLAY_NAMES[model_name],
                        "response": response_name,
                        "response_column": response_column,
                        "actual": actual,
                        "predicted": predicted,
                        "error": error,
                        "absolute_error": abs(error) if np.isfinite(error) else np.nan,
                        "squared_error": error**2 if np.isfinite(error) else np.nan,
                        "status": status,
                    }
                )

    fit_summary = pd.DataFrame(fit_rows)
    loo_raw = pd.DataFrame(loo_rows)
    loo_summary = (
        loo_raw[loo_raw["status"] == "ok"]
        .groupby(["model", "model_label", "response"], as_index=False)
        .agg(
            folds=("test_case", "nunique"),
            mae_ct_per_kwh=("absolute_error", "mean"),
            rmse_ct_per_kwh=("squared_error", lambda values: float(np.sqrt(np.mean(values)))),
            error_std_ct_per_kwh=("error", "std"),
        )
    )
    return fit_summary, loo_raw, loo_summary


def candidate_model_comparison_figure(loo_summary, loo_raw, ylabel, title):
    mean_summary = loo_summary[loo_summary["response"] == "mean"].copy()
    order = list(CANDIDATE_MODEL_DISPLAY_NAMES)
    mean_summary["model"] = pd.Categorical(mean_summary["model"], categories=order, ordered=True)
    mean_summary = mean_summary.sort_values("model")
    mean_raw = loo_raw[(loo_raw["response"] == "mean") & (loo_raw["status"] == "ok")].copy()

    fig, ax = plt.subplots(figsize=(10.8, 5.8), constrained_layout=True)
    x_positions = np.arange(len(mean_summary))
    ax.bar(
        x_positions,
        mean_summary["rmse_ct_per_kwh"],
        color="#5277A3",
        edgecolor="#222222",
        linewidth=0.7,
        alpha=0.88,
    )
    rng = np.random.default_rng(VALIDATION_RANDOM_SEED)
    for x_position, model_name in zip(x_positions, mean_summary["model"].astype(str)):
        errors = mean_raw.loc[mean_raw["model"] == model_name, "absolute_error"].astype(float).to_numpy()
        jitter = rng.normal(0.0, 0.055, size=len(errors))
        ax.scatter(
            np.full(len(errors), x_position) + jitter,
            errors,
            s=18,
            color="#222222",
            alpha=0.45,
            linewidths=0,
            zorder=3,
        )
    ax.set_ylabel(ylabel)
    ax.set_title(title)
    ax.set_xticks(x_positions)
    ax.set_xticklabels(mean_summary["model_label"], rotation=30, ha="right")
    ax.grid(axis="y", color="#D0D0D0", linewidth=0.8, alpha=0.75)
    ax.grid(axis="x", visible=False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    return fig


def candidate_model_r2(candidate_fit_summary, model_name, response_name):
    row = candidate_fit_summary[
        (candidate_fit_summary["model"] == model_name)
        & (candidate_fit_summary["response"] == response_name)
        & (candidate_fit_summary["status"] == "ok")
    ]
    if row.empty:
        return np.nan
    return float(row["r2"].iloc[0])


def candidate_model_rmse(candidate_fit_summary, model_name, response_name):
    row = candidate_fit_summary[
        (candidate_fit_summary["model"] == model_name)
        & (candidate_fit_summary["response"] == response_name)
        & (candidate_fit_summary["status"] == "ok")
    ]
    if row.empty:
        return np.nan
    return float(row["rmse_ct_per_kwh"].iloc[0])


def candidate_model_r2_label(candidate_fit_summary, model_name):
    return (
        "$R^2$ min/mean/max = "
        f"{candidate_model_r2(candidate_fit_summary, model_name, 'lower'):.3f} / "
        f"{candidate_model_r2(candidate_fit_summary, model_name, 'mean'):.3f} / "
        f"{candidate_model_r2(candidate_fit_summary, model_name, 'upper'):.3f}\n"
        "RMSE min/mean/max = "
        f"{candidate_model_rmse(candidate_fit_summary, model_name, 'lower'):.2f} / "
        f"{candidate_model_rmse(candidate_fit_summary, model_name, 'mean'):.2f} / "
        f"{candidate_model_rmse(candidate_fit_summary, model_name, 'upper'):.2f} ct/kWh"
    )


def candidate_predictions_for_responses(df, model_name, prediction_df, density_bounds):
    predictions = {}
    for response_name, response_column in REGRESSION_TARGETS.items():
        _, predictor, _ = fit_candidate_model(
            df,
            response_column,
            model_name,
            density_bounds=density_bounds,
        )
        predictions[response_name] = np.asarray(predictor(prediction_df), dtype=float)
    return predictions


def candidate_model_curve_figure(df, model_name, title, candidate_fit_summary):
    color_by_district = district_color_map(sorted(df["district"].unique()))
    fig, ax = plt.subplots(figsize=FIGURE_SIZE, constrained_layout=True)
    plot_observations_by_district(ax, df, color_by_district, alpha=0.72)

    x_min = float(df["linear_heat_density_kwh_per_m_a"].min())
    x_max = float(df["linear_heat_density_kwh_per_m_a"].max())
    density_bounds = (x_min, x_max)
    if model_name in {"candidate_gam_lhd_district", "candidate_random_forest"}:
        for district in sorted(df["district"].unique()):
            sub = df[df["district"] == district]
            x_values = np.linspace(
                float(sub["linear_heat_density_kwh_per_m_a"].min()),
                float(sub["linear_heat_density_kwh_per_m_a"].max()),
                120,
            )
            prediction_df = pd.DataFrame(
                {
                    "linear_heat_density_kwh_per_m_a": x_values,
                    "district": district,
                }
            )
            predictions = candidate_predictions_for_responses(df, model_name, prediction_df, density_bounds)
            plot_response_band(
                ax,
                x_values,
                predictions,
                color=color_by_district[district],
                linewidth=2.0,
                label=f"District {display_district(district)}",
            )
    else:
        x_values = np.linspace(x_min, x_max, 220)
        prediction_df = pd.DataFrame(
            {
                "linear_heat_density_kwh_per_m_a": x_values,
                "district": df["district"].iloc[0],
            }
        )
        predictions = candidate_predictions_for_responses(df, model_name, prediction_df, density_bounds)
        plot_response_band(ax, x_values, predictions, "#111111", label="Fitted curve", linewidth=2.5)

    ax.set_title(title)
    add_model_label(ax, candidate_model_r2_label(candidate_fit_summary, model_name))
    set_scientific_axes(ax)
    ax.legend(ncols=2, frameon=False, loc="best")
    return fig


def plot_candidate_model_comparison(loo_summary, loo_raw, output_prefix):
    fig = candidate_model_comparison_figure(
        loo_summary,
        loo_raw,
        ylabel=r"RMSE (ct kWh$^{-1}$)",
        title="Candidate function comparison, leave-one-district-seed-out",
    )
    fig.savefig(f"{output_prefix}_candidate_model_comparison.pdf")
    save_presentation_svg(fig, f"{output_prefix}_candidate_model_comparison.svg")
    plt.close(fig)


def regression_comparison(df):
    fitted_models = []
    for response_name, response_column in REGRESSION_TARGETS.items():
        for model_name, transform, include_district, include_lhd_tertile, include_lhd_tertile_interaction in MODEL_SPECS:
            fitted_models.append(
                fit_ols(
                    df,
                    model_name=model_name,
                    response_name=response_name,
                    response_column=response_column,
                    transform=transform,
                    include_district=include_district,
                    include_lhd_tertile=include_lhd_tertile,
                    include_lhd_tertile_interaction=include_lhd_tertile_interaction,
                )
            )

    summary = pd.DataFrame([item[0] for item in fitted_models])
    coefficients = pd.concat([item[1] for item in fitted_models], ignore_index=True)
    beta_by_model = {
        (item[0]["model"], item[0]["response"]): item[2]
        for item in fitted_models
    }
    statsmodels_summaries = {
        f"{item[0]['model']}__{item[0]['response']}": item[3]
        for item in fitted_models
    }
    return summary, coefficients, beta_by_model, statsmodels_summaries


def beta_for(beta_by_model, model_name, response_name="mean"):
    return beta_by_model[(model_name, response_name)]


def model_r2(summary, model_name, response_name):
    row = summary[(summary["model"] == model_name) & (summary["response"] == response_name)]
    return row["r2"].iloc[0]


def model_rmse(summary, model_name, response_name):
    row = summary[(summary["model"] == model_name) & (summary["response"] == response_name)]
    return row["rmse_ct_per_kwh"].iloc[0]


def model_r2_label(summary, model_name):
    return (
        "$R^2$ min/mean/max = "
        f"{model_r2(summary, model_name, 'lower'):.3f} / "
        f"{model_r2(summary, model_name, 'mean'):.3f} / "
        f"{model_r2(summary, model_name, 'upper'):.3f}\n"
        "RMSE min/mean/max = "
        f"{model_rmse(summary, model_name, 'lower'):.2f} / "
        f"{model_rmse(summary, model_name, 'mean'):.2f} / "
        f"{model_rmse(summary, model_name, 'upper'):.2f} ct/kWh"
    )


def predict_density_only(beta, density_kwh_per_m, transform):
    term = "density_1000_kwh_per_m_a" if transform == "linear" else "ln_density_kwh_per_m_a"
    return beta["intercept"] + beta[term] * density_term_from_values(density_kwh_per_m, transform)


def predict_density_plus_district(beta, district, density_kwh_per_m, transform):
    term = "density_1000_kwh_per_m_a" if transform == "linear" else "ln_density_kwh_per_m_a"
    district_shift = beta.get(f"district_{district}", 0.0)
    return beta["intercept"] + district_shift + beta[term] * density_term_from_values(density_kwh_per_m, transform)


def predict_log_plus_tertile(beta, tertile, density_kwh_per_m):
    tertile_shift = beta.get(f"lhd_{tertile}", 0.0)
    return beta["intercept"] + tertile_shift + beta["ln_density_kwh_per_m_a"] * np.log(density_kwh_per_m)


def predict_log_times_tertile(beta, tertile, density_kwh_per_m):
    ln_density = np.log(density_kwh_per_m)
    tertile_shift = beta.get(f"lhd_{tertile}", 0.0)
    slope_shift = beta.get(f"ln_density_kwh_per_m_a_x_lhd_{tertile}", 0.0)
    return beta["intercept"] + tertile_shift + (beta["ln_density_kwh_per_m_a"] + slope_shift) * ln_density


def predictions_for_responses(predictor, beta_by_model, model_name, *args):
    return {
        response_name: predictor(beta_for(beta_by_model, model_name, response_name), *args)
        for response_name in REGRESSION_TARGETS
    }


def plot_response_band(ax, x, predictions, color, label=None, linestyle="-", linewidth=2.0, alpha=BAND_ALPHA):
    ax.fill_between(
        x,
        predictions["lower"],
        predictions["upper"],
        color=color,
        alpha=alpha,
        linewidth=0,
    )
    ax.plot(
        x,
        predictions["mean"],
        color=color,
        linestyle=linestyle,
        linewidth=linewidth,
        label=label,
    )


def plot_observations_by_district(ax, df, color_by_district, alpha=0.78):
    for district in sorted(df["district"].unique()):
        sub = df[df["district"] == district]
        yerr_lower = sub["mean_lcoh_ct_per_kwh"] - sub["min_lcoh_ct_per_kwh"]
        yerr_upper = sub["max_lcoh_ct_per_kwh"] - sub["mean_lcoh_ct_per_kwh"]
        ax.errorbar(
            sub["linear_heat_density_kwh_per_m_a"],
            sub["mean_lcoh_ct_per_kwh"],
            yerr=[yerr_lower, yerr_upper],
            fmt="o",
            markersize=6.2,
            capsize=3,
            elinewidth=1.0,
            alpha=alpha,
            color=color_by_district[district],
            markeredgecolor="white",
            markeredgewidth=0.55,
            label=f"District {display_district(district)}",
        )


def plot_observations_by_tertile(ax, df, color_by_tertile):
    for tertile in ("low", "medium", "high"):
        sub = df[df["lhd_tertile"] == tertile]
        if sub.empty:
            continue
        yerr_lower = sub["mean_lcoh_ct_per_kwh"] - sub["min_lcoh_ct_per_kwh"]
        yerr_upper = sub["max_lcoh_ct_per_kwh"] - sub["mean_lcoh_ct_per_kwh"]
        ax.errorbar(
            sub["linear_heat_density_kwh_per_m_a"],
            sub["mean_lcoh_ct_per_kwh"],
            yerr=[yerr_lower, yerr_upper],
            fmt="o",
            markersize=6.2,
            capsize=3,
            elinewidth=1.0,
            alpha=0.72,
            color=color_by_tertile[tertile],
            markeredgecolor="white",
            markeredgewidth=0.55,
            label=None,
        )


def plot_ranges(df, output_prefix):
    if df.empty:
        raise ValueError("No complete central seed ranges found.")

    districts = sorted(df["district"].unique())
    color_by_district = district_color_map(districts)
    fig, ax = plt.subplots(figsize=FIGURE_SIZE, constrained_layout=True)
    plot_observations_by_district(ax, df, color_by_district)
    set_scientific_axes(ax)
    ax.set_title("Central LCOH versus annual linear heat density")
    ax.legend(ncols=2, frameon=False, loc="best")
    fig.savefig(f"{output_prefix}.pdf")
    plt.close(fig)


def format_signed_term(coefficient, label):
    sign = "+" if coefficient >= 0 else "-"
    return f" {sign} {abs(coefficient):.2f} {label}"


def format_compact(value):
    return f"{value:.4g}"


def metric_lines(summary, model_name):
    return [
        (
            "  R2 min/mean/max = "
            f"{model_r2(summary, model_name, 'lower'):.3f} / "
            f"{model_r2(summary, model_name, 'mean'):.3f} / "
            f"{model_r2(summary, model_name, 'upper'):.3f}"
        ),
        (
            "  RMSE min/mean/max = "
            f"{model_rmse(summary, model_name, 'lower'):.2f} / "
            f"{model_rmse(summary, model_name, 'mean'):.2f} / "
            f"{model_rmse(summary, model_name, 'upper'):.2f} ct/kWh"
        ),
    ]


def candidate_metric_lines(candidate_fit_summary, model_name):
    return [
        (
            "  R2 min/mean/max = "
            f"{candidate_model_r2(candidate_fit_summary, model_name, 'lower'):.3f} / "
            f"{candidate_model_r2(candidate_fit_summary, model_name, 'mean'):.3f} / "
            f"{candidate_model_r2(candidate_fit_summary, model_name, 'upper'):.3f}"
        ),
        (
            "  RMSE min/mean/max = "
            f"{candidate_model_rmse(candidate_fit_summary, model_name, 'lower'):.2f} / "
            f"{candidate_model_rmse(candidate_fit_summary, model_name, 'mean'):.2f} / "
            f"{candidate_model_rmse(candidate_fit_summary, model_name, 'upper'):.2f} ct/kWh"
        ),
    ]


def fitted_candidate_equation_lines(df, candidate_fit_summary=None):
    x = candidate_density(df)
    y = df[REGRESSION_TARGETS["mean"]].astype(float).to_numpy()
    density_bounds = (float(np.min(x)), float(np.max(x)))
    lines = []

    shifted_params, _ = curve_fit(
        shifted_log_function,
        x,
        y,
        p0=[float(np.mean(y)), 1.0, max(float(np.min(x)) * 0.1, 1e-6)],
        bounds=([-np.inf, -np.inf, 1e-6], [np.inf, np.inf, max(float(np.max(x)) * 5.0, 1e-5)]),
        maxfev=20000,
    )
    lines.extend(
        [
            "Model 6: shifted logarithmic model",
            "  General: LCOH(q_L) = a + b ln(q_L + c)",
            "  Meaning: c shifts the logarithmic curve along the q_L axis.",
            (
                "  Mean fit: LCOH(q_L) = "
                f"{format_compact(shifted_params[0])}"
                f"{format_signed_term(shifted_params[1], f'ln(q_L + {format_compact(shifted_params[2])})')}"
            ),
            *(candidate_metric_lines(candidate_fit_summary, "candidate_shifted_log") if candidate_fit_summary is not None else []),
            "",
        ]
    )

    exp_params, _ = curve_fit(
        exponential_saturation_function,
        x,
        y,
        p0=[float(np.mean(y)), float(y[0] - np.mean(y)), 1.0 / max(float(np.mean(x)), 1.0)],
        bounds=([-np.inf, -np.inf, 0.0], [np.inf, np.inf, np.inf]),
        maxfev=20000,
    )
    lines.extend(
        [
            "Model 7: exponential saturation model",
            "  General: LCOH(q_L) = a + b exp(-c q_L)",
            "  Meaning: the curve can approach an asymptotic value for high q_L.",
            (
                "  Mean fit: LCOH(q_L) = "
                f"{format_compact(exp_params[0])}"
                f"{format_signed_term(exp_params[1], f'exp(-{format_compact(exp_params[2])} q_L)')}"
            ),
            *(candidate_metric_lines(candidate_fit_summary, "candidate_exp_saturation") if candidate_fit_summary is not None else []),
            "",
        ]
    )

    power_params, _ = curve_fit(
        power_law_function,
        x,
        y,
        p0=[float(np.mean(y)), 1.0, -0.5],
        bounds=([-np.inf, -np.inf, -5.0], [np.inf, np.inf, 5.0]),
        maxfev=20000,
    )
    lines.extend(
        [
            "Model 8: power-law model",
            "  General: LCOH(q_L) = a + b q_L^c",
            "  Meaning: c controls the curvature.",
            (
                "  Mean fit: LCOH(q_L) = "
                f"{format_compact(power_params[0])}"
                f"{format_signed_term(power_params[1], f'q_L^{format_compact(power_params[2])}')}"
            ),
            *(candidate_metric_lines(candidate_fit_summary, "candidate_power_law") if candidate_fit_summary is not None else []),
            "",
        ]
    )

    spline_design = dmatrix(
        "bs(x, df=5, degree=3, include_intercept=False, lower_bound=lower_bound, upper_bound=upper_bound)",
        {"x": x, "lower_bound": density_bounds[0], "upper_bound": density_bounds[1]},
        return_type="dataframe",
    )
    spline_result = sm.OLS(y, spline_design.astype(float)).fit()
    spline_coeffs = ", ".join(format_compact(v) for v in spline_result.params.to_numpy())
    lines.extend(
        [
            "Model 9: spline / GAM-like model",
            "  General: LCOH(q_L) = sum_j beta_j B_j(q_L)",
            "  Meaning: B_j are cubic B-spline basis functions.",
            f"  Mean fit: beta = [{spline_coeffs}]",
            *(candidate_metric_lines(candidate_fit_summary, "candidate_gam_lhd") if candidate_fit_summary is not None else []),
            "",
        ]
    )

    district_dummies = pd.get_dummies(df["district"], prefix="district", drop_first=True, dtype=float)
    spline_district_design = pd.concat(
        [spline_design.reset_index(drop=True), district_dummies.reset_index(drop=True)],
        axis=1,
    )
    spline_district_result = sm.OLS(y, spline_district_design.astype(float)).fit()
    spline_part = spline_district_result.params.loc[spline_design.columns].to_numpy()
    district_part = spline_district_result.params.drop(index=spline_design.columns, errors="ignore")
    spline_part_text = ", ".join(format_compact(v) for v in spline_part)
    district_part_text = ", ".join(
        f"{term.replace('district_', '')}={format_compact(value)}" for term, value in district_part.items()
    )
    lines.extend(
        [
            "Model 10: spline / GAM-like model + district type",
            "  General: LCOH(q_L,d) = sum_j beta_j B_j(q_L) + gamma_d",
            "  Meaning: gamma_d is the district-type-specific shift.",
            f"  Mean fit: beta = [{spline_part_text}]",
            f"  Mean fit district shifts: {district_part_text}",
            *(candidate_metric_lines(candidate_fit_summary, "candidate_gam_lhd_district") if candidate_fit_summary is not None else []),
            "",
            "Model 11: random forest",
            "  General: LCOH(q_L,d) = RF(q_L,d)",
            "  Meaning: non-parametric benchmark using an ensemble of decision trees.",
            "  Mean fit: no compact scalar equation; prediction is the average of 500 fitted trees.",
            *(candidate_metric_lines(candidate_fit_summary, "candidate_random_forest") if candidate_fit_summary is not None else []),
        ]
    )
    return lines


def equation(beta, transform):
    if transform == "linear":
        return (
            f"{beta['intercept']:.2f}"
            + format_signed_term(beta["density_1000_kwh_per_m_a"], "q_L/1000")
        )
    if transform == "log":
        return (
            f"{beta['intercept']:.2f}"
            + format_signed_term(beta["ln_density_kwh_per_m_a"], "ln(q_L)")
        )
    raise ValueError("transform must be 'linear' or 'log'.")


def draw_equations_page(fig, beta_by_model, summary, page, candidate_fit_summary=None, df=None):
    ax = fig.add_subplot(111)
    ax.axis("off")
    ax.text(
        0.03,
        0.96,
        (
            "Central LCOH regression equations"
            if page == 1
            else "Categorical regression equations"
            if page == 2
            else "Model 6-11 equations"
        ),
        fontsize=15,
        weight="bold",
        transform=ax.transAxes,
    )

    if page == 1:
        linear_betas = {name: beta_for(beta_by_model, "linear_density_only", name) for name in REGRESSION_TARGETS}
        log_betas = {name: beta_for(beta_by_model, "log_density_only", name) for name in REGRESSION_TARGETS}
        lines = [
            "Definitions",
            "  q_L is annual linear heat density in kWh m^-1 a^-1.",
            "  LCOH = 100 * (C_EH,operation + C_central) / Q_heat",
            "  C_EH,operation excludes building-side electricity costs and revenues.",
            "  Decentral device investments, e.g. decentral PV, are not included.",
            "  Lower, mean, and upper curves are fitted to central min, mean, and max cost cases.",
            "",
            "Model 1: linear heat-density model",
            "  LCOH_min(q_L)  = " + equation(linear_betas["lower"], "linear"),
            "  LCOH_mean(q_L) = " + equation(linear_betas["mean"], "linear"),
            "  LCOH_max(q_L)  = " + equation(linear_betas["upper"], "linear"),
            *metric_lines(summary, "linear_density_only"),
            "",
            "Model 2: logarithmic heat-density model",
            "  LCOH_min(q_L)  = " + equation(log_betas["lower"], "log"),
            "  LCOH_mean(q_L) = " + equation(log_betas["mean"], "log"),
            "  LCOH_max(q_L)  = " + equation(log_betas["upper"], "log"),
            *metric_lines(summary, "log_density_only"),
        ]
    elif page == 2:
        tertile_mean = beta_for(beta_by_model, "log_density_plus_lhd_tertile", "mean")
        district_mean = beta_for(beta_by_model, "log_density_plus_district", "mean")
        lines = [
            "Model 3: ln(q_L) + heat-density class",
            "  Reference class: low heat-density tertile",
            "  LCOH_mean,low(q_L)    = " + equation(tertile_mean, "log"),
            (
                "  LCOH_mean,medium(q_L) = "
                + equation(tertile_mean, "log")
                + format_signed_term(tertile_mean.get("lhd_medium", 0.0), "")
            ),
            (
                "  LCOH_mean,high(q_L)   = "
                + equation(tertile_mean, "log")
                + format_signed_term(tertile_mean.get("lhd_high", 0.0), "")
            ),
            *metric_lines(summary, "log_density_plus_lhd_tertile"),
            "",
            "Model 4: ln(q_L) x heat-density class",
            "  This model allows both the intercept and ln(q_L) slope to vary by class.",
            *metric_lines(summary, "log_density_times_lhd_tertile"),
            "",
            "Model 5: ln(q_L) + district type",
            "  Reference district type: A",
            "  LCOH_mean,A(q_L) = " + equation(district_mean, "log"),
        ]
        for district in ("B", "C", "D", "E", "F", "H", "I"):
            lines.append(
                f"  LCOH_mean,{display_district(district)}(q_L) = "
                + equation(district_mean, "log")
                + format_signed_term(district_mean.get(f"district_{district}", 0.0), "")
            )
        lines.extend(metric_lines(summary, "log_density_plus_district"))
    elif page == 3:
        lines = [
            "Definitions",
            "  q_L: annual linear heat density in kWh m^-1 a^-1",
            "  Lower, mean, and upper curves are fitted separately; numbers below show the mean fit.",
            "",
        ]
        if df is not None:
            lines.extend(fitted_candidate_equation_lines(df, candidate_fit_summary))
    else:
        raise ValueError("page must be 1, 2, or 3")

    ax.text(
        0.03,
        0.89,
        "\n".join(lines),
        va="top",
        ha="left",
        family="monospace",
        fontsize=7.4 if page == 3 else 8.7,
        transform=ax.transAxes,
    )


def save_regression_model_report(
    df,
    output_prefix,
    beta_by_model,
    summary,
    candidate_fit_summary=None,
    candidate_loo_summary=None,
    candidate_loo_raw=None,
):
    report_prefix = f"{output_prefix}_regression_models"
    pdf_path = f"{report_prefix}.pdf"
    for old_file in Path(output_prefix).parent.glob(f"{Path(report_prefix).name}_*.svg"):
        old_file.unlink()

    districts = sorted(df["district"].unique())
    color_by_district = district_color_map(districts)
    color_by_tertile = TERTILE_COLORS
    x_global = np.linspace(
        df["linear_heat_density_kwh_per_m_a"].min(),
        df["linear_heat_density_kwh_per_m_a"].max(),
        200,
    )

    figures = []

    fig, ax = plt.subplots(figsize=FIGURE_SIZE, constrained_layout=True)
    plot_observations_by_district(ax, df, color_by_district)
    linear_predictions = predictions_for_responses(
        predict_density_only, beta_by_model, "linear_density_only", x_global, "linear"
    )
    plot_response_band(ax, x_global, linear_predictions, "#111111", "Linear fit", linewidth=2.4)
    ax.set_title(r"Model 1: linear heat density")
    add_model_label(ax, model_r2_label(summary, "linear_density_only"))
    set_scientific_axes(ax)
    ax.legend(ncols=2, frameon=False, loc="best")
    figures.append(("01_lhd", fig))

    fig, ax = plt.subplots(figsize=FIGURE_SIZE, constrained_layout=True)
    plot_observations_by_district(ax, df, color_by_district)
    log_predictions = predictions_for_responses(
        predict_density_only, beta_by_model, "log_density_only", x_global, "log"
    )
    plot_response_band(ax, x_global, log_predictions, "#111111", "Logarithmic fit", linestyle="--", linewidth=2.4)
    ax.set_title(r"Model 2: logarithmic heat density")
    add_model_label(ax, model_r2_label(summary, "log_density_only"))
    set_scientific_axes(ax)
    ax.legend(ncols=2, frameon=False, loc="best")
    figures.append(("02_ln_lhd", fig))

    fig, ax = plt.subplots(figsize=FIGURE_SIZE, constrained_layout=True)
    plot_observations_by_tertile(ax, df, color_by_tertile)
    for tertile in ("low", "medium", "high"):
        sub = df[df["lhd_tertile"] == tertile]
        if sub.empty:
            continue
        x_tertile = np.linspace(
            sub["linear_heat_density_kwh_per_m_a"].min(),
            sub["linear_heat_density_kwh_per_m_a"].max(),
            80,
        )
        tertile_predictions = predictions_for_responses(
            predict_log_plus_tertile, beta_by_model, "log_density_plus_lhd_tertile", tertile, x_tertile
        )
        plot_response_band(
            ax,
            x_tertile,
            tertile_predictions,
            color_by_tertile[tertile],
            "Logarithmic fit" if tertile == "low" else None,
        )
    ax.set_title(r"Model 3: logarithmic heat density with density class")
    add_model_label(ax, model_r2_label(summary, "log_density_plus_lhd_tertile"))
    set_scientific_axes(ax)
    ax.legend(ncols=2, frameon=False, loc="best")
    figures.append(("03_ln_lhd_plus_lhd_tertile", fig))

    fig, ax = plt.subplots(figsize=FIGURE_SIZE, constrained_layout=True)
    plot_observations_by_tertile(ax, df, color_by_tertile)
    for tertile in ("low", "medium", "high"):
        sub = df[df["lhd_tertile"] == tertile]
        if sub.empty:
            continue
        x_tertile = np.linspace(
            sub["linear_heat_density_kwh_per_m_a"].min(),
            sub["linear_heat_density_kwh_per_m_a"].max(),
            80,
        )
        tertile_predictions = predictions_for_responses(
            predict_log_times_tertile, beta_by_model, "log_density_times_lhd_tertile", tertile, x_tertile
        )
        plot_response_band(
            ax,
            x_tertile,
            tertile_predictions,
            color_by_tertile[tertile],
            "Logarithmic fit" if tertile == "low" else None,
        )
    ax.set_title(r"Model 4: logarithmic heat density with class-specific slopes")
    add_model_label(ax, model_r2_label(summary, "log_density_times_lhd_tertile"))
    set_scientific_axes(ax)
    ax.legend(ncols=2, frameon=False, loc="best")
    figures.append(("04_ln_lhd_times_lhd_tertile", fig))

    fig, ax = plt.subplots(figsize=FIGURE_SIZE, constrained_layout=True)
    plot_observations_by_district(ax, df, color_by_district)
    for district in districts:
        sub = df[df["district"] == district]
        x_district = np.linspace(
            sub["linear_heat_density_kwh_per_m_a"].min(),
            sub["linear_heat_density_kwh_per_m_a"].max(),
            80,
        )
        district_predictions = predictions_for_responses(
            predict_density_plus_district,
            beta_by_model,
            "log_density_plus_district",
            district,
            x_district,
            "log",
        )
        plot_response_band(ax, x_district, district_predictions, color_by_district[district], linewidth=1.8)
    ax.set_title(r"Model 5: logarithmic heat density with district type")
    add_model_label(ax, model_r2_label(summary, "log_density_plus_district"))
    set_scientific_axes(ax)
    ax.legend(ncols=2, frameon=False, loc="best")
    figures.append(("05_ln_lhd_plus_district_type", fig))

    candidate_curve_pages = [
        ("candidate_shifted_log", "06_model_6_shifted_log", "Model 6: shifted logarithmic"),
        ("candidate_exp_saturation", "07_model_7_exponential_saturation", "Model 7: exponential saturation"),
        ("candidate_power_law", "08_model_8_power_law", "Model 8: power law"),
        ("candidate_gam_lhd", "09_model_9_spline", "Model 9: spline / GAM-like"),
        (
            "candidate_gam_lhd_district",
            "10_model_10_spline_plus_district",
            "Model 10: spline / GAM-like with district type",
        ),
        ("candidate_random_forest", "11_model_11_random_forest", "Model 11: random forest"),
    ]
    for model_name, page_name, title in candidate_curve_pages:
        try:
            fig = candidate_model_curve_figure(df, model_name, title, candidate_fit_summary)
        except Exception:
            continue
        figures.append((page_name, fig))

    fig = plt.figure(figsize=FIGURE_SIZE, constrained_layout=True)
    draw_equations_page(fig, beta_by_model, summary, page=1, candidate_fit_summary=candidate_fit_summary, df=df)
    figures.append(("12_equations_basic_models", fig))

    fig = plt.figure(figsize=FIGURE_SIZE, constrained_layout=True)
    draw_equations_page(fig, beta_by_model, summary, page=2, candidate_fit_summary=candidate_fit_summary, df=df)
    figures.append(("13_equations_categorical_models", fig))

    fig = plt.figure(figsize=FIGURE_SIZE, constrained_layout=True)
    draw_equations_page(fig, beta_by_model, summary, page=3, candidate_fit_summary=candidate_fit_summary, df=df)
    figures.append(("14_equations_models_6_to_11", fig))

    if candidate_loo_summary is not None and candidate_loo_raw is not None:
        fig = candidate_model_comparison_figure(
            candidate_loo_summary,
            candidate_loo_raw,
            ylabel=r"RMSE (ct kWh$^{-1}$)",
            title="Model comparison, leave-one-district-seed-out",
        )
        figures.append(("15_model_comparison", fig))

    with PdfPages(pdf_path) as pdf:
        for name, fig in figures:
            pdf.savefig(fig)
            save_presentation_svg(fig, f"{report_prefix}_{name}.svg")
            plt.close(fig)


def plot_regression_comparison(df, output_prefix, beta_by_model, summary):
    if df.empty:
        raise ValueError("No complete central seed ranges found.")

    districts = sorted(df["district"].unique())
    color_by_district = district_color_map(districts)
    fig, ax = plt.subplots(figsize=FIGURE_SIZE, constrained_layout=True)

    for district in districts:
        sub = df[df["district"] == district].sort_values("linear_heat_density_kwh_per_m_a")
        color = color_by_district[district]
        yerr_lower = sub["mean_lcoh_ct_per_kwh"] - sub["min_lcoh_ct_per_kwh"]
        yerr_upper = sub["max_lcoh_ct_per_kwh"] - sub["mean_lcoh_ct_per_kwh"]
        ax.errorbar(
            sub["linear_heat_density_kwh_per_m_a"],
            sub["mean_lcoh_ct_per_kwh"],
            yerr=[yerr_lower, yerr_upper],
            fmt="o",
            markersize=6.2,
            capsize=3,
            elinewidth=1.0,
            alpha=0.55,
            color=color,
            markeredgecolor="white",
            markeredgewidth=0.55,
            label=f"District {display_district(district)}",
        )
        x_district = np.linspace(
            sub["linear_heat_density_kwh_per_m_a"].min(),
            sub["linear_heat_density_kwh_per_m_a"].max(),
            50,
        )
        district_predictions = predictions_for_responses(
            predict_density_plus_district,
            beta_by_model,
            "log_density_plus_district",
            district,
            x_district,
            "log",
        )
        plot_response_band(ax, x_district, district_predictions, color=color, linewidth=1.8, alpha=0.10)

    x_global = np.linspace(
        df["linear_heat_density_kwh_per_m_a"].min(),
        df["linear_heat_density_kwh_per_m_a"].max(),
        200,
    )
    log_predictions = predictions_for_responses(
        predict_density_only,
        beta_by_model,
        "log_density_only",
        x_global,
        "log",
    )
    plot_response_band(
        ax,
        x_global,
        log_predictions,
        "#111111",
        "Logarithmic fit",
        linestyle="--",
        linewidth=2.4,
        alpha=0.10,
    )

    def r2_triplet(model_name):
        return (
            f"{model_r2(summary, model_name, 'lower'):.2f}/"
            f"{model_r2(summary, model_name, 'mean'):.2f}/"
            f"{model_r2(summary, model_name, 'upper'):.2f}"
        )

    def rmse_triplet(model_name):
        return (
            f"{model_rmse(summary, model_name, 'lower'):.2f}/"
            f"{model_rmse(summary, model_name, 'mean'):.2f}/"
            f"{model_rmse(summary, model_name, 'upper'):.2f}"
        )

    text = (
        r"$R^2$ and RMSE shown as min/mean/max"
        "\n"
        rf"Linear model: R2 {r2_triplet('linear_density_only')}, RMSE {rmse_triplet('linear_density_only')} ct/kWh"
        "\n"
        rf"Logarithmic model: R2 {r2_triplet('log_density_only')}, RMSE {rmse_triplet('log_density_only')} ct/kWh"
        "\n"
        rf"Logarithmic model + density class: R2 {r2_triplet('log_density_plus_lhd_tertile')}, RMSE {rmse_triplet('log_density_plus_lhd_tertile')} ct/kWh"
        "\n"
        rf"Logarithmic model x density class: R2 {r2_triplet('log_density_times_lhd_tertile')}, RMSE {rmse_triplet('log_density_times_lhd_tertile')} ct/kWh"
        "\n"
        rf"Logarithmic model + district type: R2 {r2_triplet('log_density_plus_district')}, RMSE {rmse_triplet('log_density_plus_district')} ct/kWh"
    )
    add_model_label(ax, text)
    ax.set_title("Central LCOH: effect of annual linear heat density and district type")
    set_scientific_axes(ax)
    ax.legend(ncols=2, frameon=False, loc="best", handlelength=2.3)

    regression_prefix = f"{output_prefix}_regression"
    fig.savefig(f"{regression_prefix}.pdf")
    plt.close(fig)


def plot_train_test_validation(validation_summary, validation_raw, output_prefix):
    mean_summary = validation_summary[validation_summary["response"] == "mean"].copy()
    mean_summary["model_label"] = mean_summary["model"].map(MODEL_DISPLAY_NAMES)
    mean_summary = mean_summary.set_index("model").loc[list(MODEL_DISPLAY_NAMES)].reset_index()
    mean_raw = validation_raw[validation_raw["response"] == "mean"].copy()

    x = np.arange(len(mean_summary))
    colors = ["#4C78A8", "#F58518", "#54A24B", "#B279A2", "#E45756"]

    fig, ax = plt.subplots(figsize=(8.6, 5.2), constrained_layout=True)
    ax.bar(
        x,
        mean_summary["rmse_ct_per_kwh"],
        color=colors,
        alpha=0.88,
    )
    rng = np.random.default_rng(7)
    for i, model in enumerate(mean_summary["model"]):
        errors = mean_raw.loc[mean_raw["model"] == model, "absolute_error"].astype(float).to_numpy()
        jitter = rng.uniform(-0.16, 0.16, size=len(errors))
        ax.scatter(
            np.full(len(errors), x[i]) + jitter,
            errors,
            s=24,
            color="#1F1F1F",
            alpha=0.62,
            edgecolors="white",
            linewidths=0.35,
            zorder=3,
        )
    ax.set_ylabel(r"RMSE (ct kWh$^{-1}$)")
    ax.set_title("Leave-one-district-seed-out validation")
    ax.set_xticks(x)
    ax.set_xticklabels(mean_summary["model_label"], rotation=28, ha="right")
    ax.grid(True, axis="y", color="#D0D0D0", linewidth=0.8, alpha=0.75)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(axis="both", which="major", length=5, width=1.0)

    validation_prefix = f"{output_prefix}_leave_one_seed_out_validation"
    fig.savefig(f"{validation_prefix}.pdf")
    save_presentation_svg(fig, f"{validation_prefix}.svg")
    plt.close(fig)


def main():
    apply_publication_style()
    parser = argparse.ArgumentParser()
    parser.add_argument("--results-dir", type=Path, default=DEFAULT_RESULTS_DIR)
    parser.add_argument("--output-prefix", type=Path, default=DEFAULT_OUTPUT_PREFIX)
    parser.add_argument("--config-path", type=Path, default=DEFAULT_CONFIG_PATH)
    parser.add_argument(
        "--require-complete-decentral-pair",
        action="store_true",
        help="Restrict LCOH analysis to seeds that also have complete decentral results.",
    )
    args = parser.parse_args()

    if not args.results_dir.exists():
        raise FileNotFoundError(f"Results directory does not exist: {args.results_dir}")
    if not args.config_path.exists():
        raise FileNotFoundError(f"Config file does not exist: {args.config_path}")

    args.output_prefix.parent.mkdir(parents=True, exist_ok=True)
    config_values = read_config_values(args.config_path)
    ranges, cases = make_ranges(
        args.results_dir,
        config_values,
        require_complete_decentral=args.require_complete_decentral_pair,
    )
    if ranges.empty:
        raise ValueError("No complete central scenarios found.")
    ranges = add_lhd_tertiles(ranges)

    regression_summary, regression_coefficients, beta_by_model, statsmodels_summaries = regression_comparison(ranges)
    validation_raw, validation_summary = repeated_train_test_validation(ranges)
    loo_raw, loo_summary = leave_one_seed_out_validation(ranges)
    candidate_fit_summary, candidate_loo_raw, candidate_loo_summary = candidate_model_comparison(ranges)
    ranges.to_csv(f"{args.output_prefix}_ranges.csv", index=False)
    cases.to_csv(f"{args.output_prefix}_central_cases.csv", index=False)
    regression_summary.to_csv(f"{args.output_prefix}_regression_summary.csv", index=False)
    regression_coefficients.to_csv(f"{args.output_prefix}_regression_coefficients.csv", index=False)
    validation_raw.to_csv(f"{args.output_prefix}_train_test_validation_raw.csv", index=False)
    validation_summary.to_csv(f"{args.output_prefix}_train_test_validation_summary.csv", index=False)
    loo_raw.to_csv(f"{args.output_prefix}_leave_one_seed_out_validation_raw.csv", index=False)
    loo_summary.to_csv(f"{args.output_prefix}_leave_one_seed_out_validation_summary.csv", index=False)
    candidate_fit_summary.to_csv(f"{args.output_prefix}_candidate_model_fit_summary.csv", index=False)
    candidate_loo_raw.to_csv(f"{args.output_prefix}_candidate_model_leave_one_seed_out_raw.csv", index=False)
    candidate_loo_summary.to_csv(f"{args.output_prefix}_candidate_model_leave_one_seed_out_summary.csv", index=False)
    with open(f"{args.output_prefix}_statsmodels_summaries.txt", "w", encoding="utf-8") as file:
        for model_name, model_summary in statsmodels_summaries.items():
            file.write(f"{'=' * 100}\n")
            file.write(f"{model_name}\n")
            file.write(f"{'=' * 100}\n")
            file.write(model_summary)
            file.write("\n\n")

    plot_ranges(ranges, args.output_prefix)
    plot_regression_comparison(ranges, args.output_prefix, beta_by_model, regression_summary)
    plot_train_test_validation(loo_summary, loo_raw, args.output_prefix)
    save_regression_model_report(
        ranges,
        args.output_prefix,
        beta_by_model,
        regression_summary,
        candidate_fit_summary=candidate_fit_summary,
        candidate_loo_summary=candidate_loo_summary,
        candidate_loo_raw=candidate_loo_raw,
    )

    print(f"Wrote {len(ranges)} central seed ranges")
    print(f"Wrote {len(cases)} central cost cases")
    print(f"Wrote {args.output_prefix}.pdf")
    print(f"Wrote {args.output_prefix}_ranges.csv")
    print(f"Wrote {args.output_prefix}_central_cases.csv")
    print(f"Wrote {args.output_prefix}_regression.pdf")
    print(f"Wrote {args.output_prefix}_regression_summary.csv")
    print(f"Wrote {args.output_prefix}_regression_coefficients.csv")
    print(f"Wrote {args.output_prefix}_train_test_validation_raw.csv")
    print(f"Wrote {args.output_prefix}_train_test_validation_summary.csv")
    print(f"Wrote {args.output_prefix}_leave_one_seed_out_validation_raw.csv")
    print(f"Wrote {args.output_prefix}_leave_one_seed_out_validation_summary.csv")
    print(f"Wrote {args.output_prefix}_leave_one_seed_out_validation.pdf")
    print(f"Wrote {args.output_prefix}_leave_one_seed_out_validation.svg")
    print(f"Wrote {args.output_prefix}_candidate_model_fit_summary.csv")
    print(f"Wrote {args.output_prefix}_candidate_model_leave_one_seed_out_raw.csv")
    print(f"Wrote {args.output_prefix}_candidate_model_leave_one_seed_out_summary.csv")
    print(f"Wrote {args.output_prefix}_statsmodels_summaries.txt")
    print(f"Wrote {args.output_prefix}_regression_models.pdf")
    print(f"Wrote {args.output_prefix}_regression_models_*.svg")


if __name__ == "__main__":
    main()
